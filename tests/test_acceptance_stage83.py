import json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from product_tool import jobs,storage,worker,discovery_trace,attribute_projection
from product_tool.adapters.common import RawAttribute,SourceDocument,PhotoCandidate
from product_tool.normalization import normalize_facts
from product_tool.product_input import resolve_product_input

WIRELESS='MagSafe wireless charging up to 25W; Qi2 wireless charging up to 25W; Magnet array; Alignment magnet; Accessory Identification NFC; Magnetometer'
CHARGING='Fast-charge capable: Up to 50% charge in around 15 minutes when using 60W-capable adapter or higher with Adjustable Voltage Supply (available separately) paired with USB-C charging cable Up to 50% charge in 30 minutes with 35W adapter or higher paired with MagSafe Charger (both available separately)'

class SemanticDecomposition(unittest.TestCase):
 def test_independent_standards_and_sensors_preserve_complete_raw(self):
  for section in ('MagSafe and Wireless Charging','Samsung accessory specifications'):
   rows=normalize_facts([RawAttribute('Wireless charging',WIRELESS,section,True)])
   self.assertEqual(len(rows),6)
   self.assertEqual(len({r.normalized_name for r in rows}),6)
   self.assertTrue(all(r.raw_value==WIRELESS and r.raw_name=='Wireless charging' for r in rows))
   values={r.normalized_name:(r.normalized_value,r.unit) for r in rows}
   self.assertEqual(values['wireless_charging_power_qi2'],('up_to=25','W'))
   self.assertEqual(values['accessory_identification_nfc'],('true','bool'))
 def test_independent_fast_charge_modes_keep_limits_and_conditions(self):
  rows=normalize_facts([RawAttribute('Fast charging',CHARGING,'Charging',True)])
  values={r.normalized_name:(r.normalized_value,r.unit) for r in rows}
  self.assertEqual(values['fast_charge_usb_c_time_minutes'],('approx=15','min'))
  self.assertEqual(values['fast_charge_magsafe_time_minutes'],('30','min'))
  self.assertEqual(values['fast_charge_usb_c_adapter_power'],('at_least=60','W'))
  self.assertEqual(values['fast_charge_magsafe_adapter_power'],('at_least=35','W'))
  self.assertTrue(all(r.raw_value==CHARGING for r in rows))
  self.assertIn('adjustable voltage supply',values['fast_charge_usb_c_conditions'][0])
  from product_tool.value_localization import localize
  for mode in ('usb_c','magsafe'):
   shown=localize(values['fast_charge_'+mode+'_conditions'][0])
   self.assertNotIn('paired',shown)
   self.assertIn('совместно с',shown)
 def test_ambiguous_and_unknown_clauses_stay_whole(self):
  for value in (WIRELESS+'; dependent on unspecified accessory', 'MagSafe wireless charging up to 15W; MagSafe wireless charging up to 25W',CHARGING.replace('USB-C charging cable','USB-C charging cable with MagSafe support')):
   rows=normalize_facts([RawAttribute('Unresolved feature',value,'Features',True)])
   self.assertEqual(len(rows),1)
   self.assertEqual(rows[0].raw_value,value)
 def test_output_uses_children_localized_values_and_compatibility_group(self):
  with tempfile.TemporaryDirectory() as td:
   db=Path(td)/'batches.sqlite3';jobs.initialize(db)
   seed(db,'Apple','iPhone 18 Pro','A3714','')
   doc=SourceDocument('apple_model','Apple','https://www.apple.com/iphone-18-pro/specs/',match_level='model_confirmed',attributes=[RawAttribute('MagSafe и беспроводная зарядка',WIRELESS,'MagSafe and Wireless Charging',True),RawAttribute('Системные требования','Apple Account (required for some features); Internet access','System Requirements',True),RawAttribute('Время воспроизведения видео','Video playback Up to 36 hours','Power and Battery',True)])
   jobs.save_source_document(db,1,doc);jobs.resolve_product(db,1)
   rows=attribute_projection.final_attribute_rows(db,1)
   by_key={r['normalized_name']:r for r in rows}
   self.assertEqual(by_key['wireless_charging_power_magsafe']['resolved']['display_value'],'До 25 Вт')
   self.assertEqual(by_key['magnetometer']['resolved']['display_value'],'Да')
   compat=next(r for r in rows if r['display_name']=='Системные требования')
   self.assertEqual(compat['parent_group'],'Совместимость и системные требования')
   self.assertIn('Доступ к интернету',compat['resolved']['display_value'])
   self.assertNotIn('Video playback',' '.join(r['resolved']['display_value'] for r in rows))
   self.assertEqual(len({r['category_attribute_key'] for r in rows}),len(rows))
   from io import BytesIO
   from openpyxl import load_workbook
   from product_tool.exporter import export_batch
   book=load_workbook(BytesIO(export_batch(db,'b')),data_only=True)
   raw=book['Семантика и исходные значения']
   self.assertEqual(raw.max_row,7)
   self.assertTrue(all(row[4]==WIRELESS for row in raw.iter_rows(min_row=2,values_only=True)))
   self.assertIn('Мощность беспроводной зарядки Qi2, Вт',[c.value for c in book.worksheets[0][1]])
 def test_brightness_modes_remain_independent_and_cellular_5g_is_not_mass(self):
  from product_tool.value_localization import localize
  from product_tool.canonical_attributes import project
  raw='1000 nits max brightness (typical); 1600 nits peak brightness (HDR); 3000 nits peak brightness (outdoor); 1 nit minimum brightness'
  rows=normalize_facts([RawAttribute('Brightness',raw,'Display',True)])
  self.assertEqual({r.normalized_name:r.normalized_value for r in rows},{'display_brightness_typical_max':'1000','display_brightness_hdr_peak':'1600','display_brightness_outdoor_peak':'3000','display_brightness_min':'1'})
  projected=project([dict(normalized_name='calls',display_name='Calls',sources={'official':dict(raw_value='FaceTime video over 5G or Wi-Fi')},resolved={'selected_unit':'','selected_source':'official'})])
  self.assertIn('5G',projected[0]['resolved']['display_value'])
  self.assertEqual(localize('Video playback Up to 36 hours'),'Воспроизведение видео до 36 ч')

def seed(db,brand,name,article,alternate):
 with storage._connection(db) as c:
  c.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('b','supplier.xlsx','Sheet','{}',storage._now()))
  c.execute('INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (?,?,?,?,?,?,?,?,?,?)',('b',2,name,brand,article,alternate,'Паровой шкаф',0,'[]',json.dumps([alternate,name,brand])))

class ProductionParity(unittest.TestCase):
 def test_descriptive_primary_fallback_is_not_identity_or_dataset_mutation(self):
  raw=dict(brand='LG',name='LG Styler S3WER',search_code='Паровой шкаф LG Styler S3WER',alternate_code='S3WER.ALWPCOM')
  p=resolve_product_input(raw)
  self.assertEqual(p['search_code'],'S3WER.ALWPCOM')
  self.assertFalse(p['input_resolution']['identity_claim'])
  self.assertEqual(raw['search_code'],'Паровой шкаф LG Styler S3WER')
  for primary in ('S3WER.ALWPCOM','S40T','P12ED.NSAR + P12ED.USAR'):
   self.assertEqual(resolve_product_input({**raw,'search_code':primary})['search_code'],primary)
 def test_ui_batch_and_direct_run_once_have_same_dispatch_evidence_and_readiness(self):
  from fastapi.testclient import TestClient
  from product_tool.web import create_app
  from product_tool.readiness import card_readiness
  calls=[]
  class LG:
   source_key='lg_kz';site_name='LG Казахстан'
   def find_source(self,article,**kw):
    calls.append(('lg_kz',article))
    return SourceDocument('lg_kz','LG Казахстан','https://www.lg.com/kz/laundry/styler/s3wer/',found_model=article,match_level='full_sku',evidence='Structured exact sales code',html='<div data-pim-sku="S3WER.ALWPCOM.EEAK.KZ.C"></div>',attributes=[RawAttribute('Вес товара','86 kg','Specifications',True)],photo_candidates=[PhotoCandidate('https://www.lg.com/image.jpg','img')])
  class Dealer:
   def find_source(self,*a,**kw):return SourceDocument('sulpak','Sulpak','')
  results=[]
  with tempfile.TemporaryDirectory() as td:
   for mode in ('ui','direct'):
    directory=Path(td)/mode;directory.mkdir();db=directory/'batches.sqlite3'
    app=create_app(directory,start_worker=False)
    seed(db,'LG','Паровой шкаф LG Styler S3WER','Паровой шкаф LG Styler S3WER','S3WER.ALWPCOM')
    with TestClient(app) as client:
     if mode=='ui':
      response=client.post('/batches/b/search',data={'product_ids':['1'],'stages':['1','2','3','4']})
      self.assertEqual(response.status_code,200)
     else:jobs.enqueue(db,1,[1,2,3,4])
     with patch('product_tool.worker.default_lg_adapters',return_value=(LG(),None,Dealer(),None)):
      worker.run_once(db)
    j=jobs.list_jobs(db,1)[0]
    trace=discovery_trace.for_job(db,j['id'])
    results.append(dict(dispatch=[(e['event'],e.get('effective_article'),e.get('provider')) for e in trace if e['event'] in {'production_dispatch','input_search_resolution','lg_adapter_selected','lg_discovery_start','lg_discovery_end'}],sources=[(s['source_key'],s['url'],s['match_level']) for s in jobs.get_source_pages(db,1)],facts=[(f['normalized_name'],f['normalized_value']) for f in jobs.get_facts(db,1)],readiness=card_readiness(db,1)['verdict']))
    with storage._connection(db) as c:self.assertEqual(c.execute('SELECT search_code FROM products').fetchone()[0],'Паровой шкаф LG Styler S3WER')
  self.assertEqual(results[0],results[1])
  self.assertEqual(calls,[('lg_kz','S3WER.ALWPCOM')]*2)
