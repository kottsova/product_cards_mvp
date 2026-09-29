"""Schema, status, fingerprint and safe-probe regressions for source census."""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from product_tool.census.catalog import load_catalog_coverage
from product_tool.census.fingerprint import fingerprint_platform
from product_tool.census.models import AccessStatus, FingerprintStatus
from product_tool.census.endpoint_probe import blocked_hosts_from_fetch_log
from product_tool.census.probe import AccessProbe, ProbePolicy
from product_tool.census.registry import load_source_catalog
from product_tool.sources import SourceDefinition, SourceRole, default_source_registry

from _source_policy import unapproved_source_ids


class SourceCatalogTests(unittest.TestCase):
    def test_schema_loads_and_only_approved_sources_are_registered(self):
        # Generic rule: a source outside the approved list is not in the
        # catalog (active or not) and not in the executable registry.
        catalog = load_source_catalog()
        self.assertGreater(len(catalog.records), 5)
        self.assertEqual(unapproved_source_ids(item.source_id for item in catalog.records), set())
        self.assertEqual(unapproved_source_ids(item.source_id for item in catalog.active()), set())
        self.assertEqual(unapproved_source_ids(default_source_registry().definitions), set())

    def test_host_verification_defaults_to_not_checked(self):
        source = SourceDefinition(
            "candidate", "candidate", ("Brand",), (), (), SourceRole.MANUFACTURER,
            frozenset(), "model", 100, ("example.test",),
        )
        self.assertEqual(source.host_verification, "not_checked")

    def test_access_status_contract_is_complete(self):
        self.assertEqual(
            {item.value for item in AccessStatus},
            {"not_checked", "direct_access", "structured_api", "javascript_required", "browser_assisted", "search_only", "rate_limited", "captcha_or_blocked", "regional_redirect", "support_only", "legacy_missing", "unavailable"},
        )


class FingerprintingTests(unittest.TestCase):
    def assert_engine(self, engine, html="", headers=None, cookies=None, content_type=""):
        results = fingerprint_platform(html, headers=headers, cookies=cookies, content_type=content_type)
        found = next((item for item in results if item.engine == engine), None)
        self.assertIsNotNone(found, engine)
        self.assertEqual(found.status, FingerprintStatus.CONFIRMED)

    def test_known_platform_signatures(self):
        fixtures = {
            "shopify": '<script src="https://cdn.shopify.com/a.js"></script>',
            "magento": '<script src="/static/version123/frontend.js"></script>',
            "salesforce_commerce_cloud": '<form action="/on/demandware.store/Sites-X/default">',
            "woocommerce": '<script src="/wp-content/plugins/woocommerce/a.js"></script>',
            "adobe_experience_manager": '<script src="/etc.clientlibs/site/clientlib.js"></script>',
            "sitecore": '<img src="/-/media/catalog/item.jpg">',
            "next_js": '<script id="__NEXT_DATA__" type="application/json">{}</script>',
            "nuxt": '<script src="/_nuxt/app.js"></script>',
        }
        for engine, html in fixtures.items():
            with self.subTest(engine=engine):
                self.assert_engine(engine, html)

    def test_json_ld_product_and_sitemap(self):
        self.assert_engine("generic_json_ld_product", '<script type="application/ld+json">{"@type":"Product","name":"P"}</script>')
        self.assert_engine("sitemap_catalog_feed", '<urlset><url><loc>https://example.test/p</loc></url></urlset>', content_type="application/xml")

    def test_one_weak_signal_is_not_confirmed(self):
        result = next(item for item in fingerprint_platform("", cookies={"_shopify_s": "1"}) if item.engine == "shopify")
        self.assertEqual(result.status, FingerprintStatus.PROBABLE)
        self.assertLess(result.confidence, .7)


class FakeCookies(dict):
    def get_dict(self):
        return dict(self)


class FakeResponse:
    def __init__(self, status=200, url="https://example.test/", body=b"<html></html>", history=()):
        self.status_code = status
        self.url = url
        self.history = history
        self.headers = {"content-type": "text/html; charset=utf-8"}
        self.cookies = FakeCookies()
        self.encoding = "utf-8"
        self.body = body

    def iter_content(self, chunk_size=16384):
        for index in range(0, len(self.body), chunk_size):
            yield self.body[index:index + chunk_size]


class FakeSession:
    def __init__(self, responses):
        self.responses = list(responses)
        self.calls = []

    def get(self, url, **kwargs):
        self.calls.append((url, kwargs))
        return self.responses.pop(0)


class SafeProbeTests(unittest.TestCase):
    def test_403_stops_repeated_request_to_endpoint(self):
        session = FakeSession([FakeResponse(status=403)])
        probe = AccessProbe(session, policy=ProbePolicy(min_interval_seconds=0))
        first = probe.probe("https://example.test/", allowed_hosts=("example.test",))
        second = probe.probe("https://example.test/", allowed_hosts=("example.test",))
        self.assertEqual(first.access_status, AccessStatus.CAPTCHA_OR_BLOCKED)
        self.assertEqual(second.access_status, AccessStatus.CAPTCHA_OR_BLOCKED)
        self.assertEqual(len(session.calls), 1)

    def test_seeding_a_fresh_probe_with_a_persisted_block_refuses_all_contact(self):
        # Stage 13: proves the stop survives a fresh AccessProbe instance --
        # i.e. a script re-run -- not just the lifetime of one instance.
        # blocked_hosts_from_fetch_log() derives the stopped-host set from a
        # PAST run's own fetch log; a brand-new AccessProbe seeded with that
        # set must make zero requests to that host, even for a URL it has
        # never itself seen before.
        past_run_log = [
            {"url": "https://example.test/robots.txt", "status_code": 403},
            {"url": "https://example.test/", "status_code": 200},
        ]
        stopped = blocked_hosts_from_fetch_log(past_run_log)
        self.assertEqual(stopped, frozenset({"example.test"}))

        session = FakeSession([])  # would raise IndexError on any .get() call
        probe = AccessProbe(session, policy=ProbePolicy(min_interval_seconds=0), initial_stopped_hosts=stopped)
        result = probe.probe("https://example.test/some/other/never-before-seen/path", allowed_hosts=("example.test",))
        self.assertEqual(result.access_status, AccessStatus.CAPTCHA_OR_BLOCKED)
        self.assertEqual(session.calls, [])

    def test_blocked_hosts_from_fetch_log_ignores_ordinary_responses(self):
        log = [{"url": "https://ok.test/", "status_code": 200}, {"url": "https://ok.test/other", "status_code": 404}]
        self.assertEqual(blocked_hosts_from_fetch_log(log), frozenset())

    def test_redirect_to_unknown_host_is_not_followed_as_valid(self):
        redirect = FakeResponse(status=302, url="https://example.test/start")
        session = FakeSession([FakeResponse(url="https://other.test/final", history=(redirect,))])
        result = AccessProbe(session, policy=ProbePolicy(min_interval_seconds=0)).probe(
            "https://example.test/start", allowed_hosts=("example.test",),
        )
        self.assertEqual(result.access_status, AccessStatus.REGIONAL_REDIRECT)

    def test_probe_rejects_non_allowlisted_start_url(self):
        with self.assertRaises(ValueError):
            AccessProbe(FakeSession([])).probe("https://other.test/", allowed_hosts=("example.test",))


class CatalogCoverageTests(unittest.TestCase):
    def test_coverage_keeps_unresolved_brand_rows_and_reconciles(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        path = Path(temporary.name) / "catalog.xlsx"
        workbook = Workbook()
        context = workbook.active
        context.title = "Контекст"
        context.append(["Параметр", "Значение"])
        for key, value in (
            ("Уникальных товаров после фильтра", 3), ("Рабочих строк после фильтра", 4),
            ("Брендов после фильтра", 2), ("Категорий после фильтра", 2),
            ("Связок бренд + категория после фильтра", 3),
        ):
            context.append([key, value])
        workbook.create_sheet("Товары")
        sheet = workbook.create_sheet("Бренд_Категория")
        sheet.append(["Бренд", "Категория", "Уникальных товаров", "Строк в исходных выгрузках", "Примеры артикулов продавца"])
        sheet.append(["A", "Смартфоны", 1, 1, "A1"])
        sheet.append(["B", "Ноутбуки", 1, 1, "B1"])
        sheet.append([None, "Смартфоны", 1, 2, "U1"])
        workbook.save(path)
        workbook.close()
        census = load_catalog_coverage(path)
        self.assertEqual(len(census.brands), 2)
        self.assertEqual(len(census.unresolved_brand_coverages), 1)
        self.assertEqual(census.observed_totals, census.expected_totals)


if __name__ == "__main__":
    unittest.main()
