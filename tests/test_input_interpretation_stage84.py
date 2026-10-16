import hashlib,json,tempfile,unittest
from types import SimpleNamespace
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook
from fastapi.testclient import TestClient
from product_tool import storage,jobs,worker,discovery_trace
from product_tool.input_interpretation import interpret,adapter_name
from product_tool.importer import ColumnMapping,preview_row,preview_xlsx
from product_tool.web import create_app

class TitleInterpretation(unittest.TestCase):
 def test_resolutions_and_technical_quantities_are_not_model_candidates(self):
  p=interpret('24″ Монитор S3 S30GD, 100 Гц, IPS, 1920x1080','Samsung','LS24D300GAIXCI')
  self.assertEqual(p['model'],'S30GD');self.assertEqual(p['category'],'Монитор')
  p=interpret('LG S3WER 500W 50Hz 1920x1080 DDR5','LG','SUP')
  self.assertEqual(p['model_candidates'],['S3WER'])
 def test_model_category_and_supplier_are_independent(self):
  cases=[('Паровой шкаф LG Styler S3WER','LG','S3WER.ALWPCOM','S3WER','Паровой шкаф'),('Телевизор LG 86NANO81A6A','LG','SUP-TV','86NANO81A6A','Телевизор'),('Стиральная машина Samsung WW90T554CAT','Samsung','SUP-WASH','WW90T554CAT','Стиральная машина'),('Apple iPhone 18 Pro 256GB Silver','Apple','A3714','iPhone 18 Pro','Смартфон'),('JBL Tune 520BT White','JBL','SUP-JBL','Tune 520BT','Наушники'),('Dreame H14 Pro Wet and Dry Vacuum','Dreame','SUP-H14','H14 Pro','Моющий пылесос'),('Робот-пылесос Dreame X50 Ultra Complete','Dreame','SUP-X50','X50 Ultra','Робот-пылесос')]
  for name,brand,article,model,category in cases:
   with self.subTest(name=name):
    p=interpret(name,brand,article)
    self.assertEqual((p['model'],p['category'],p['supplier_article']),(model,category,article))
    self.assertEqual(p['fields']['model']['source'],'product_name')
    self.assertEqual(p['search_queries'][0]['query'],model)
    self.assertEqual(p['search_queries'][-1]['purpose'],'original_name')
 def test_configuration_scopes_and_no_guess_from_article(self):
  p=interpret('Xiaomi 15 12/256GB Black Global','Xiaomi','OPAQUE-RU-RED-2TB')
  self.assertEqual(p['configuration_hints'],dict(ram='12GB',storage='256GB',color='Black',region='Global'))
  self.assertEqual(p['model'],'Xiaomi 15')
  self.assertTrue(adapter_name(p).startswith('Xiaomi 15 ['))
  p=interpret('HyperX Alloy Origins 65 Black RU Red','HyperX','4P5D6AX#ACB')
  self.assertEqual(p['model'],'Alloy Origins 65')
  self.assertEqual(p['configuration_hints'],dict(color='Black',layout='RU',switch='Red'))
  self.assertNotIn('region',p['configuration_hints'])
  p=interpret('Razer Barracuda X 2022 Quartz Wireless','Razer','SUP')
  self.assertEqual(p['model'],'Barracuda X 2022')
  self.assertEqual(p['configuration_hints']['generation'],'2022')
  self.assertEqual(p['configuration_hints']['color'],'Quartz')
 def test_ambiguity_and_brand_column_priority(self):
  p=interpret('LG A123 B456 Black White','LG','SKU123')
  self.assertEqual(p['model'],'')
  self.assertEqual(p['fields']['model']['confidence'],'ambiguous')
  self.assertEqual(p['model_candidates'],['A123','B456'])
  self.assertNotIn('color',p['configuration_hints'])
  self.assertEqual(p['configuration']['color']['confidence'],'ambiguous')
  p=interpret('Apple iPhone 16','Samsung','SUP')
  self.assertEqual(p['brand'],'Samsung');self.assertTrue(p['warnings'])
  self.assertEqual(interpret('iPhone 16 256GB Black','Apple','SUP')['model'],'iPhone 16')
  self.assertEqual(interpret('iPhone 18 Pro 256GB','Apple','A3714',explicit_model='A3714')['model'],'iPhone 18 Pro')
 def test_size_layout_bundle_and_connectivity_are_explicit_hints(self):
  p=interpret('Apple Watch Series 10 46mm Silver Global','Apple','SUP')
  self.assertEqual(p['configuration_hints']['size'],'46mm')
  p=interpret('Dreame X50 Ultra Complete with auto-empty station','Dreame','SUP')
  self.assertEqual(p['configuration_hints']['bundle'],'Complete')
  self.assertIn('station',p['configuration_hints'])
  p=interpret('Xiaomi Smart Air Purifier 4 Lite White','Xiaomi','SUP')
  self.assertEqual(p['category'],'Очиститель воздуха')
  self.assertEqual(p['model'],'Xiaomi Smart Air Purifier 4 Lite')
 def test_memory_labels_are_owned_by_one_capacity(self):
  for suffix in ('16GB RAM 512GB SSD','RAM 16GB SSD 512GB','16GB ОЗУ 512GB SSD','RAM 16GB 512GB SSD'):
   p=interpret('Apple MacBook Pro 14 '+suffix+' Silver','Apple','SUP')
   self.assertEqual(p['configuration_hints']['ram'],'16GB',suffix)
   self.assertEqual(p['configuration_hints']['storage'],'512GB',suffix)
  self.assertEqual(interpret('Xiaomi 15 2026 256GB Black','Xiaomi','SUP')['model'],'Xiaomi 15')
 def test_recognized_commercial_families_have_categories_without_generic_nouns(self):
  for name,brand,category in [('JBL Bar 1000','JBL','Саундбар'),('HyperX QuadCast 2','HyperX','Микрофон'),('HyperX Pulsefire Haste 2','HyperX','Мышь'),('Razer Huntsman Mini','Razer','Клавиатура'),('Apple AirPods Pro 2','Apple','Наушники')]:
   self.assertEqual(interpret(name,brand,'SUP')['category'],category)

class ThreeColumnImport(unittest.TestCase):
 def test_inferred_brand_provenance_survives_unedited_confirmation(self):
  with tempfile.TemporaryDirectory() as td:
   root=Path(td);wb=Workbook();s=wb.active;s.append(['Наименование','Артикул']);s.append(['Apple iPhone 18 Pro 256GB Silver','A3714']);body=BytesIO();wb.save(body);wb.close()
   with TestClient(create_app(root,start_worker=False)) as client:
    r=client.post('/upload',files={'file':('input.xlsx',body.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},follow_redirects=False)
    draft=r.headers['location'].rsplit('/',1)[1]
    client.post(f'/drafts/{draft}/confirm',data={'brand_2':'Apple','search_code_2':'A3714','category_2':'Смартфон'})
    p=storage.get_batch(root/'batches.sqlite3',draft)['products'][0]
    self.assertEqual(p['input_interpretation']['fields']['brand']['source'],'product_name')
 def test_header_synonyms_no_required_model_or_category_column(self):
  for headers in [('Наименование','Артикул','Бренд'),('Название','Код','Производитель'),('Товар','SKU','Brand'),('Product Name','Part Number','Бренд'),('Name','Supplier SKU','Brand')]:
   with self.subTest(headers=headers),tempfile.TemporaryDirectory() as td:
    path=Path(td)/'input.xlsx';wb=Workbook();s=wb.active;s.append(headers);s.append(['Apple iPhone 18 Pro 256GB Silver','A3714','Apple']);wb.save(path);wb.close()
    before=hashlib.sha256(path.read_bytes()).hexdigest();p=preview_xlsx(path).products[0]
    self.assertEqual(p.search_code,'A3714');self.assertEqual(p.category,'Смартфон')
    self.assertEqual(p.interpretation['model'],'iPhone 18 Pro')
    self.assertEqual(p.variant_attributes,dict(storage='256GB',color='Silver'))
    self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
 def test_real_preview_confirm_batch_to_worker_preserves_raw(self):
  with tempfile.TemporaryDirectory() as td:
   root=Path(td);db=root/'batches.sqlite3';wb=Workbook();s=wb.active
   s.append(['Наименование','Артикул','Бренд']);s.append(['Xiaomi 15 12/256GB Black Global','OPAQUE-SUPPLIER-15','Xiaomi'])
   body=BytesIO();wb.save(body);wb.close();raw=body.getvalue()
   with TestClient(create_app(root,start_worker=False)) as client:
    r=client.post('/upload',files={'file':('ordinary.xlsx',raw,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},follow_redirects=False)
    draft=r.headers['location'].split('/')[-1];page=client.get(r.headers['location'])
    self.assertIn('Распознанная модель',page.text);self.assertIn('Xiaomi 15',page.text);self.assertIn('256GB',page.text)
    r=client.post(f'/drafts/{draft}/confirm',data={},follow_redirects=False);self.assertEqual(r.status_code,303)
    p=storage.get_batch(db,draft)['products'][0];pid=p['id']
    self.assertEqual((p['name'],p['supplier_article'],p['model']),('Xiaomi 15 12/256GB Black Global','OPAQUE-SUPPLIER-15','Xiaomi 15'))
    self.assertEqual(json.loads(p['original_values_json']),['Xiaomi 15 12/256GB Black Global','OPAQUE-SUPPLIER-15','Xiaomi'])
    client.post(f'/batches/{draft}/search',data={'product_ids':[str(pid)],'stages':['1','3','4']})
    observed={}
    def run(*args,**kwargs):observed.update(product=args[3],factory=kwargs['adapter_factory']);jobs.finish(db,args[1],'needs_review','Offline dispatch proof only')
    with patch('product_tool.xiaomi_pipeline.run_job',side_effect=run):self.assertTrue(worker.run_once(db))
    self.assertEqual(observed['product']['search_queries'][0]['query'],'Xiaomi 15')
    self.assertEqual(observed['product']['discovery_name'],'Xiaomi 15 [RAM=12GB;storage=256GB;color=Black;region=Global]')
    self.assertIsNone(observed['factory'])
    trace=discovery_trace.for_job(db,jobs.list_jobs(db,pid)[0]['id'])
    self.assertTrue(any(t.get('event')=='input_interpretation' for t in trace))
    self.assertEqual(hashlib.sha256((root/'uploads'/f'{draft}.xlsx').read_bytes()).hexdigest(),hashlib.sha256(raw).hexdigest())

class JBLModelHintBoundaries(unittest.TestCase):
 def test_unknown_model_with_opaque_article_attempts_title_discovery(self):
  from product_tool.adapters.jbl import JBLAdapter
  calls=[];trace=[]
  def get(url,**kw):
   calls.append(url);return SimpleNamespace(ok=False,status_code=404,marker='',text='',url=url)
  with tempfile.TemporaryDirectory() as td:
   a=JBLAdapter(session=SimpleNamespace(get=get),capture_dir=Path(td),clock=lambda:0,trace_callback=trace.append,search_provider=lambda *args,**kw:[])
   a.input_interpreted=True;a.input_queries=[{'query':'Unknown Launch White'}]
   article='../opaque / supplier';doc=a.find_source(article,name='Unknown Launch White',deadline=65)
   self.assertTrue(any('Unknown%20Launch' in u for u in calls))
   self.assertEqual(a._queries(article,'ignored'),['site:jbl.com "Unknown Launch White"'])
   self.assertNotIn(doc.match_level,{'full_sku','model_confirmed'})
   self.assertEqual(a.find_documents(doc,article.upper()),[])
 def test_opaque_supplier_article_cannot_become_document_cache_path(self):
  from product_tool.adapters.jbl import JBLAdapter
  from product_tool.adapters.common import SourceDocument
  with tempfile.TemporaryDirectory() as td:
   a=JBLAdapter(session=object(),capture_dir=Path(td)/'captures')
   article='../opaque / supplier';a.reports[article]={}
   self.assertEqual(a.find_documents(SourceDocument('jbl','JBL',''),article),[])
   files=list(Path(td).rglob('*'));self.assertEqual(len([p for p in files if p.is_file()]),1)
   file=next(p for p in files if p.is_file());self.assertEqual(file.parent,a.capture_dir)
   self.assertTrue(file.name.startswith('supplier_'));self.assertTrue(file.name.endswith('.documents.json'))
 def test_master_inventory_follows_only_published_matching_variant_and_revalidates(self):
  from product_tool.adapters.jbl import JBLAdapter
  payload=dict(product_sku=['JBLT520BTBLK','JBLT520BTWHT'],product_variant=['Black-GLOBAL-Current','White-GLOBAL-Current'],product_name=['JBL Tune 520BT']*2)
  master='<script>var utag_data_dynamic = '+json.dumps(payload)+';</script>'
  variant='''<div class="product-info product-detail"><meta itemprop="sku" content="JBLT520BTWHT"><meta itemprop="mpn" content="JBLT520BTWHT"><meta itemprop="name" content="JBL Tune 520BT"><div class="product-variations" data-current='{"color":{"displayValue":"White"}}'></div></div><div class="hero-gallery -main"><img src="/White.png"></div>'''
  seen=[]
  def get(url,**kw):
   seen.append(url);return SimpleNamespace(ok=True,url=url,text=variant)
  adapter=JBLAdapter(session=SimpleNamespace(get=get),clock=lambda:0)
  adapter.input_model='Tune 520BT';adapter.input_configuration={'color':'White'}
  found=adapter.published_variants(master,'https://id.jbl.com/TUNE520BT.html','opaque / supplier',deadline=10)
  self.assertIsNotNone(found);self.assertEqual(seen,['https://id.jbl.com/JBLT520BTWHT.html'])
  self.assertTrue(found[1]['identity']['model_confirmed']);self.assertFalse(found[1]['identity']['supplier_sku_confirmed'])
  self.assertEqual(len(found[1]['exact_photo_assets']),1)
  seen.clear();adapter.input_model='Tune 720BT'
  self.assertIsNone(adapter.published_variants(master,'https://id.jbl.com/TUNE520BT.html','opaque',deadline=10));self.assertFalse(seen)
  adapter.input_model='Tune 520BT';adapter.input_configuration={'color':'Blue'}
  self.assertIsNone(adapter.published_variants(master,'https://id.jbl.com/TUNE520BT.html','opaque',deadline=10))
 def test_ambiguous_input_stays_not_ready_and_photos_are_candidates(self):
  from product_tool import readiness,photo_inventory
  p=dict(brand='LG',input_interpretation=interpret('LG S3WER / S5BB','LG','SUP'))
  with patch('product_tool.readiness.jobs.get_product',return_value=p),patch('product_tool.readiness._resolved_card_readiness',return_value=dict(verdict='export_ready',blocking_gaps=[],advisory_gaps=[],gaps=[])):
   r=readiness.card_readiness(Path('unused'),1,manuals_required=False)
  self.assertEqual(r['verdict'],'not_ready');self.assertIn('input_model_ambiguous',r['blocking_gaps'])
  self.assertFalse(photo_inventory.identity_confirmed({},p,[],{}))
 def test_opaque_supplier_is_not_promoted_to_retail_sku(self):
  from product_tool.adapters.jbl import parse_pdp
  html='''<div class="product-wrapper" data-pid="JBLT520BTWHT"><meta itemprop="sku" content="JBLT520BTWHT"><meta itemprop="mpn" content="JBLT520BTWHT"><meta itemprop="name" content="JBL Tune 520BT"><div class="product-variations" data-current='{"color":{"displayValue":"White"}}'></div></div>'''
  # The single official Product ties its exact model name to coherent SKU/MPN.
  # This published relation works even when the commercial code abbreviates Tune.
  doc,ev=parse_pdp(html,'https://id.jbl.com/en/JBLT520BTWHT.html','SUP-15',input_model='Tune 520BT',input_configuration={'color':'White'})
  self.assertFalse(ev['identity']['exact_sku'])
  self.assertTrue(ev['identity']['model_confirmed'])
  html=html.replace('JBLT520BTWHT','JBLTUNE520BTWHT')
  doc,ev=parse_pdp(html,'https://id.jbl.com/en/item.html','SUP-15',input_model='Tune 520BT',input_configuration={'color':'White'})
  self.assertTrue(ev['identity']['model_confirmed']);self.assertFalse(ev['identity']['exact_sku'])
  doc,ev=parse_pdp(html.replace('Tune 520BT','Tune 720BT'),'https://id.jbl.com/en/item.html','SUP-15',input_model='Tune 520BT',input_configuration={'color':'White'})
  self.assertFalse(ev['identity']['model_confirmed'])
