"""Adversarial model/component contracts; synthetic HTML is never live evidence."""
import hashlib,json,sqlite3,tempfile,unittest
from pathlib import Path
from io import BytesIO
from unittest.mock import patch
from product_tool.dreame_page import parse_page,matches,component,SECTIONS
from product_tool.normalization import normalize_facts
from product_tool.adapters.dreame import DreameAdapter
from product_tool.adapters.common import SourceDocument
from product_tool import jobs,worker,readiness,card_evidence
R=Path(__file__).resolve().parents[1]/'reports/dreame_stage76'

def page(rows=(),model='Audit R1',color=''):
    product={'@type':'Product','name':'Dreame '+model,'image':'https://cdn.shopify.com/s/files/1/model.png','additionalProperty':[{'name':n,'value':v,'propertyID':s} for n,v,s in rows]}
    if color:product['color']=color
    return '<title>Dreame '+model+' | Dreame Global</title><main><script type="application/ld+json">'+json.dumps(product)+'</script></main>'

def parsed(rows=(),name='Dreame Audit R1 [region=RU]',article='SUPPLIER-001',url='https://global.dreametech.com/products/audit-r1',color=''):
    return parse_page(page(rows,color=color),url,article,name,'Роботы-пылесосы')

ROWS=[('Robot Vacuum Weight','4 kg','Robot'),('Base Station Weight','9 kg','Base Station'),('Dust Cup','300 mL','Robot'),('Clean Water Tank','4 L','Base Station'),('Dirty Water Tank','3.5 L','Base Station'),('Battery Capacity','6400 mAh','Robot')]

class Stage76Tests(unittest.TestCase):
    def test_frozen_dataset_and_trace(self):
        raw=(R/'dataset.json').read_bytes();self.assertEqual(hashlib.sha256(raw).hexdigest(),(R/'dataset.sha256').read_text().strip());rows=json.loads(raw)['rows'];self.assertEqual(len(rows),10);self.assertEqual(len({x['model'] for x in rows}),10);self.assertFalse(any('url' in x for x in rows));self.assertGreaterEqual(sum('робот' in x['category'].casefold() for x in rows),4)
        baseline=json.loads((R/'baseline/results.json').read_text(encoding='utf8'));self.assertEqual(baseline['job']['status'],'error');self.assertFalse(baseline['specs']);self.assertFalse(baseline['trace'])
    def test_supplier_sku_independent_same_model(self):
        for sku in ('SUPPLIER-001','SUPPLIER-002'):
            doc,ev=parsed(ROWS,article=sku);self.assertEqual(doc.match_level,'model_confirmed');self.assertEqual(len(doc.attributes),6);self.assertEqual(ev['input_configuration']['supplier_part_number'],sku);self.assertIsNone(ev['identity']['retail_sku'])
    def test_other_region_and_absent_region(self):
        for name,url in [('Dreame Audit R1 [region=RU]','https://de.dreametech.com/products/audit-r1'),('Dreame Audit R1','https://global.dreametech.com/products/audit-r1')]:
            doc,ev=parsed(ROWS,name=name,url=url);self.assertEqual(len(doc.attributes),6);self.assertEqual(ev['identity']['model_relation'],'model_confirmed')
    def test_near_models_do_not_match(self):
        for other in ('Audit R1 Pro','Audit R2','Audit R1 Gen 2','Audit R1 Dual','Audit R1 Complete'):self.assertFalse(matches('Dreame Audit R1',other))
        self.assertTrue(matches('Dreame Audit R1','Audit R1 Complete',{'bundle':'Complete'}));self.assertTrue(matches('Dreame Audit R1 Dual','Dreame Audit R1 Dual Wet and Dry Vacuum'))
    def test_catalog_and_other_host_not_identity(self):
        for url in ('https://global.dreametech.com/collections/all','https://dreametech.com.evil.test/products/audit-r1','http://global.dreametech.com/products/audit-r1'):
            d,e=parsed(ROWS,url=url);self.assertNotEqual(e['identity']['model_relation'],'model_confirmed');self.assertFalse(d.attributes)
    def test_station_and_robot_canonical_keys_never_collide(self):
        d,e=parsed([('Weight','4 kg','Robot'),('Weight','9 kg','Base Station'),('Clean Water Tank','100 mL','Robot on-board'),('Clean Water Tank','4 L','Base Station')]);facts=normalize_facts(d.attributes)
        self.assertEqual(len({x.normalized_name for x in facts}),4);self.assertEqual([x.section for x in facts],['Робот','База / станция','Робот','База / станция']);self.assertFalse(e['rejected_specs'])
    def test_ambiguous_robot_tank_and_bundle_candidate(self):
        d,e=parsed([('Clean Water Tank','4 L','Specifications'),('Package Contents','Robot + 3 accessories','Specifications'),('Rated Power','75 W','Robot')]);self.assertEqual([x.name for x in d.attributes],['Rated Power']);self.assertEqual(len(e['rejected_specs']),2);self.assertEqual(component('Package List'),'bundle')
    def test_regional_variation_and_commercial_fields_candidate(self):
        d,e=parsed([('Rated Voltage','Varies by region','Robot'),('Color','Black','Specifications'),('Supplier SKU','ABC-1','Specifications')]);self.assertFalse(d.attributes);self.assertEqual(len(e['rejected_specs']),3)
    def test_dimensions_keep_unordered_triple(self):
        d,e=parsed([('Robot Vacuum Dimensions','351 x 351 x 89 mm','Robot')]);self.assertEqual(len(d.attributes),1);f=normalize_facts(d.attributes)[0];self.assertNotIn('w=',f.normalized_value);self.assertEqual(e['original_raw_specs'][0]['value'],'351 x 351 x 89 mm')
    def test_explicit_axes_derived_and_scoped(self):
        d,e=parsed([('Dimensions (W x D x H)','30 x 45 x 60 cm','Base Station')]);self.assertEqual([(a.name,a.value,a.section) for a in d.attributes],[('Width','300 mm','База / станция'),('Depth','450 mm','База / станция'),('Height','600 mm','База / станция')])
    def test_runtime_modes_and_component_separation(self):
        d,e=parsed([('Maximum Runtime','Up to 35 min (Wet and Dry Vacuum)\nUp to 60 min (Handheld Vacuum)','Specifications')],name='Dreame Audit R1 [bundle=Dual]');self.assertEqual(len(d.attributes),2);self.assertTrue(all('Up to' in a.value for a in d.attributes));self.assertEqual(d.attributes[-1].section,'Ручной пылесос');self.assertEqual(e['original_raw_specs'][0]['value'].count('Up to'),2)
    def test_continuation_tank_does_not_lose_lie_flat(self):
        html=page().replace('</main>','<section id="Specifications"><table><tr><td>Used Water Tank</td><td>Standing: 650mL</td></tr><tr><td></td><td>Lie-Flat: 450mL</td></tr></table></section></main>')
        d,e=parse_page(html,'https://global.dreametech.com/products/audit-r1','Audit R1','Dreame Audit R1','Вертикальные моющие пылесосы');self.assertEqual(len(d.attributes),2);self.assertEqual([a.value for a in d.attributes],['650mL','450mL']);self.assertIn('Lie-Flat: 450mL',e['original_raw_specs'][0]['value'])
    def test_runtime_eco_turbo_distinct(self):
        d,e=parsed([('Maximum Runtime','Eco: Up to 60 min; Turbo: Up to 10 min','Robot')]);self.assertEqual(len(d.attributes),2);self.assertNotEqual(*[f.normalized_name for f in normalize_facts(d.attributes)])
    def test_dreame_image_bytes_full_decode(self):
        from PIL import Image
        from product_tool.photo_metadata import inspect_saved_photo
        from product_tool.adapters.policy_session import PolicyResponse
        image=Image.new('RGB',(500,400),'white');image.putpixel((0,0),(0,0,0));stream=BytesIO();image.save(stream,format='PNG');raw=stream.getvalue()
        class Session:
            def get(self,*args,**kwargs):return PolicyResponse(args[0],200,raw.decode('latin1'),'application/octet-stream')
        with patch('product_tool.photo_metadata.PolicyAwareSession',return_value=Session()):
            m=inspect_saved_photo('https://cdn.shopify.com/s/files/1/model.jpg','dreame_model',Path('unused.json'));self.assertEqual(m['sha256'],hashlib.sha256(raw).hexdigest());self.assertEqual(m['effective_content_type'],'image/png');self.assertEqual((m['width'],m['height']),(500,400))
        with patch('product_tool.photo_metadata.PolicyAwareSession',return_value=Session()):
            with self.assertRaises(ValueError):inspect_saved_photo('https://evil.test/image.png','dreame_model',Path('unused.json'))
    def test_manual_index_detail_relation_not_other_generation(self):
        from product_tool.adapters.dreame import manual_label_matches
        self.assertTrue(manual_label_matches('Dreame X50 Ultra','X50 Ultra/X50 Ultra Complete User Manual'))
        self.assertTrue(manual_label_matches('Dreame X40 Master','X40 Master User Manual-1'))
        self.assertFalse(manual_label_matches('Dreame H15 Pro Heat','H15 Heat User Manual'))
        self.assertFalse(manual_label_matches('Dreame H12 Dual','H12 Dual FlexReach User Manual'))
    def test_discovery_uses_published_href_and_title_not_guessed_slug(self):
        from product_tool.adapters.policy_session import PolicyResponse
        class Session:
            def __init__(self):self.urls=[]
            def get(self,url,**kwargs):
                self.urls.append(url)
                html='<a href="/products/commercial-name-99">Audit R1</a>' if url=='https://global.dreametech.com/' else page(ROWS)
                return PolicyResponse(url,200,html,'text/html')
        with tempfile.TemporaryDirectory() as tmp:
            session=Session();adapter=DreameAdapter(session=session,fetch_log_path=Path(tmp)/'fetch.json');doc=adapter.find_source('SUPPLIER-001',name='Dreame Audit R1 [region=RU]');self.assertEqual(doc.url,'https://global.dreametech.com/products/commercial-name-99');self.assertEqual(doc.match_level,'model_confirmed');self.assertEqual(session.urls,['https://global.dreametech.com/','https://global.dreametech.com/products/commercial-name-99']);self.assertTrue(adapter.reports['SUPPLIER-001']['live_success'])
    def test_gallery_model_scope_and_explicit_other_color(self):
        d,e=parsed(ROWS,name='Dreame Audit R1 [color=Black]');self.assertEqual(e['model_gallery_candidates'][0]['model_relation'],'model_confirmed');self.assertEqual(e['model_gallery_candidates'][0]['supplier_color_relation'],'unproven')
        d,e=parsed(ROWS,name='Dreame Audit R1 [color=Black]',color='White');self.assertEqual(d.photo_candidates[0].kind,'candidate');self.assertEqual(e['model_gallery_candidates'][0]['supplier_color_relation'],'different')
    def test_support_help_never_description(self):
        html=page(ROWS)+'<footer>Pairing reset Dreamehome warranty firmware troubleshooting</footer>';d,e=parse_page(html,'https://global.dreametech.com/products/audit-r1','Audit R1','Dreame Audit R1');self.assertEqual(d.description,'');self.assertEqual(len(d.attributes),6)

    def test_recommended_product_cannot_override_primary_model(self):
        html=page(ROWS).replace('<main>','<main><h1>Dreame Audit R2</h1>')
        d,e=parse_page(html,'https://global.dreametech.com/products/audit-r2','Audit R1','Dreame Audit R1')
        self.assertEqual(e['identity']['model_relation'],'unproven');self.assertFalse(d.attributes);self.assertFalse(d.photo_candidates)
        d,e=parse_page(page(ROWS).replace('<title>Dreame Audit R1','<title>Dreame Audit R2'),'https://global.dreametech.com/products/audit-r2','Audit R1','Dreame Audit R1')
        self.assertEqual(e['identity']['model_relation'],'unproven');self.assertFalse(d.attributes)

    def test_real_ru_manual_and_safety_are_distinct(self):
        import pymupdf
        from product_tool.adapters.dreame import manual_assessment
        with pymupdf.open(R/'manual_research/H12_Dual_RU.pdf') as pdf:pages=[p.get_text() for p in pdf]
        check,headings=manual_assessment(pages,['Dreame H12 Dual','H12 Dual'])
        self.assertTrue(check['accepted']);self.assertIn('ru',check['guide_languages']['present'])
        self.assertFalse(manual_assessment(pages,['Dreame H12 Pro','H12 Pro'])[0]['accepted'])
        self.assertFalse(manual_assessment(['Dreame H12 Dual\nSafety and Warranty\npress connect install clean'],['Dreame H12 Dual'])[0]['accepted'])
    def test_shared_worker_readiness_ui_and_excel(self):
        from fastapi.testclient import TestClient
        from product_tool.web import create_app
        from product_tool.exporter import export_batch
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp);db=path/'batches.sqlite3';jobs.initialize(db)
            with sqlite3.connect(db) as c:
                c.execute("INSERT INTO batches VALUES ('dr','synthetic.xlsx','Products','{}','2026-10-11')");c.execute("INSERT INTO products (id,batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (1,'dr',2,'Dreame Audit R1 [region=RU;color=Black]','Dreame','SUPPLIER-001','','Роботы-пылесосы',0,'[]','[]')")
            c.close()
            class Adapter:
                extra_sources=[]
                def find_source(self,article,**kw):
                    d,e=parsed(ROWS,name=kw['name'],article=article);self.reports={article:e};return d
                def find_documents(self,*args):return []
            with patch('product_tool.xiaomi_pipeline.photo_metadata.inspect_saved_photo',return_value={'width':800,'height':800,'size_bytes':2000,'format':'PNG'}),patch('product_tool.xiaomi_pipeline.DnsAdapter') as dealer:
                jobs.enqueue(db,1,[1,2,3,4,6]);worker.run_once(db,dreame_adapter_factory=Adapter);dealer.assert_not_called()
            r=readiness.card_readiness(db,1);self.assertEqual(r['verdict'],'export_ready');self.assertEqual(r['confirmed_specs'],6);self.assertEqual(r['blocking_gaps'],[]);self.assertEqual(r['manual_status'],'Не проверена');self.assertEqual(r['input_configuration']['supplier_part_number'],'SUPPLIER-001')
            self.assertTrue(all(not f['full_sku_confirmed'] for f in jobs.get_resolved(db,1)))
            with TestClient(create_app(path,start_worker=False)) as client:
                response=client.get('/products/1');self.assertEqual(response.status_code,200);self.assertIn('Модель и конфигурация Dreame',response.text);self.assertIn('База / станция',response.text);self.assertIn('Ёмкость аккумулятора',response.text);self.assertNotIn('Модель и конфигурация Xiaomi',response.text)
            b=load_workbook(BytesIO(export_batch(db,'dr')));self.assertIn('Dreame готовность',b.sheetnames);self.assertEqual(b['Dreame готовность']['I2'].value,'Готова');self.assertEqual(b['Фотографии'].max_row,2);self.assertEqual(b['Фотографии']['E1'].value,'Область подтверждения');b.close()

if __name__=='__main__':unittest.main()
