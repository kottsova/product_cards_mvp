"""Semantic cleanup regression over scoped official fields; offline only."""
import unittest,json,tempfile,sqlite3
from contextlib import closing
from pathlib import Path
from product_tool import normalization,jobs,storage,card_evidence,exporter,canonical_attributes
from product_tool.adapters.common import RawAttribute,SourceDocument
from product_tool.attribute_lineage import linked_facts,text_key
from product_tool.apple_identity import parse_model
from product_tool.razer_page import parse_page

ROOT=Path(__file__).resolve().parents[1]
S=ROOT/'reports/cross_brand_stage81/qualified_live'

class Stage82Regression(unittest.TestCase):
 def test_hardware_optical_continuations_and_laser_components_are_separate(self):
  from product_tool.playstation_specs import parse_hardware_specs
  text='''CFI-1216A PlayStation 5
  Specifications
  Optical drive (read-only)
  BD-ROM (66 GB/100 GB) ~10xCAV
  DVD ~3.2xCLV
  Laser
  Type: Semiconductor, continuous
  BD Wavelength: 395-415 nm
  Power: Max. 1 mW
  DVD Wavelength: 640-675 nm
  Power: Max. 1 mW
  Operating temperature 5 C to 35 C
  Type: Not an optical component
  Wireless controller'''
  text='\n'.join(line.strip() for line in text.splitlines())
  doc,raw,candidates=parse_hardware_specs(text,'https://www.playstation.com/content/dam/global_pdc/en-gb/corporate/support/manuals/ps5-docs/1200a/CFI-12XXA_PS5_Safety_Guide$en-gb.pdf','CFI-1216A')
  attrs={a.name:a.value for a in doc.attributes}
  self.assertIn('BD-ROM',attrs['Форматы чтения оптического привода'])
  self.assertNotIn('Laser',attrs['Форматы чтения оптического привода'])
  self.assertEqual('395-415 nm',attrs['Длина волны лазера BD'])
  self.assertEqual('640-675 nm',attrs['Длина волны лазера DVD'])
  self.assertEqual('Max. 1 mW',attrs['Мощность лазера BD'])
  self.assertEqual('Semiconductor, continuous',attrs['Тип лазера оптического привода'])
  self.assertTrue(all(a.get('hardware_code')=='CFI-1216A' for a in raw))
 def test_exact_model_display_and_scalar_origins_survive_without_supplier_sku(self):
  html=(S/'40/raw_sources/apple_model.html').read_text(encoding='utf8')
  doc,ev=parse_model(html,'https://support.apple.com/en-us/121029','iphone16',{})
  attrs={a.name:a.value for a in doc.attributes}
  self.assertEqual('2556 x 1179 пикселей',attrs['Разрешение экрана'])
  self.assertEqual('460',attrs['Плотность пикселей, ppi'])
  self.assertIn('Dynamic Island',attrs)
  self.assertIn('Wi-Fi',attrs);self.assertIn('Bluetooth',attrs)
  self.assertNotIn('Объём накопителя',attrs)
  ledger=next(a for a in ev['admitted'] if a['name']=='Плотность пикселей, ppi')
  self.assertTrue(ledger['raw_rows']);self.assertIn('ppi',ledger['raw_rows'][0]['raw_value'])
 def test_other_generation_is_not_admitted(self):
  html=(S/'40/raw_sources/apple_model.html').read_text(encoding='utf8')
  doc,_=parse_model(html,'https://support.apple.com/en-us/121029','iphone15pro',{})
  self.assertTrue(doc.error);self.assertEqual([],doc.attributes)
 def test_noise_test_configuration_is_not_a_model_wide_quantity(self):
  html=(S/'43/raw_sources/apple_model.html').read_text(encoding='utf8')
  doc,ev=parse_model(html,'https://support.apple.com/en-us/122209','air13m4',{})
  self.assertFalse(any(a.section=='Acoustic Performance' for a in doc.attributes))
  self.assertTrue(any('acoustic test' in r['reason'] for r in ev['candidates']))
 def test_watch_area_and_fit_remain_case_size_scoped(self):
  html=(S/'45/raw_sources/apple_model.html').read_text(encoding='utf8')
  doc,ev=parse_model(html,'https://support.apple.com/en-us/121202','watch10',{'case_size':'46','material':'aluminum','connectivity':'GPS'})
  areas=[a for a in ev['admitted'] if a['name']=='Площадь экрана, мм²']
  self.assertTrue(areas);self.assertTrue(all(a['scope']=='configuration-sensitive' for a in areas))
  self.assertFalse(any('989' in a.value for a in doc.attributes if a.name=='Площадь экрана, мм²'))
 def test_headphone_and_microphone_multi_field_ssr_cells_do_not_mix(self):
  html='''<html><h1>Razer BlackShark V2 Pro (2023) | RZ04-0453</h1><div id="at-a-glance"><table>
  <tr><td>Headphones</td><td>Frequency response: 12 Hz - 28 kHz Impedance: 32 Ω Drivers: 50 mm</td></tr>
  <tr><td>Microphone</td><td>Frequency response: 100 Hz - 10 kHz Signal-to-noise ratio: ≥ 60 dB Pick-up pattern: Unidirectional</td></tr>
  </table></div></html>'''
  doc,ev=parse_page(html,'https://mysupport.razer.com/app/answers/detail/a_id/13508','RZ04-04530100-R3M1','Razer BlackShark V2 Pro (2023)','Наушники')
  attrs={a.name:a.value for a in doc.attributes}
  self.assertEqual('12 Hz - 28 kHz',attrs['Наушники: Частотный диапазон'])
  self.assertEqual('100 Hz - 10 kHz',attrs['Микрофон: Частотный диапазон'])
  self.assertEqual('32 Ω',attrs['Наушники: Сопротивление'])
  self.assertTrue(all(a.get('derived_fields') for a in ev['accepted_specs']))
 def test_normalized_lineage_does_not_cross_scope_or_accept_different_value(self):
  def fact(name,value,section=''):
   n=normalization.normalize_fact(RawAttribute(name,value,section,value_cell=True))
   return {'id':1,'raw_name':name,'raw_value':value,'section':section,'normalized_name':n.normalized_name,'normalized_value':n.normalized_value,'unit':n.unit}
  raw={'raw_label':'Net mass','value':'1000 g','section':'Dimensions'}
  ledger={**raw,'derived_fields':[{'name':'Вес товара','value':'1000 g'}]}
  self.assertTrue(linked_facts(raw,[ledger],[fact('Вес товара','1 kg')]))
  self.assertFalse(linked_facts(raw,[ledger],[fact('Вес товара','2 kg')]))
  mic={'raw_label':'Frequency response','value':'100 Hz','section':'Microphone'}
  self.assertFalse(linked_facts(mic,[],[fact('Frequency response','100 Hz','Headphones')]))
 def test_boolean_support_preserves_negative_features_and_unknown_quantities(self):
  self.assertEqual(('true','bool'),normalization.normalize_value('Feature','Supported',value_cell=True))
  self.assertEqual(('false','bool'),normalization.normalize_value('Feature','Not supported',value_cell=True))
  absent=normalization.normalize_fact(RawAttribute('Вес','—',value_cell=True))
  self.assertEqual(('unknown','unknown'),(absent.normalized_value,absent.unit))
 def test_annotation_key_does_not_erase_technical_digits(self):
  self.assertEqual(text_key('Bluetooth 5.3¹⁰'),text_key('Bluetooth 5.3'))
  self.assertNotEqual(text_key('Bluetooth 5.3'),text_key('Bluetooth 5.4'))
 def test_frequency_qualified_sensitivity_keeps_component_excel_identity(self):
  rows=[]
  for component,value in [('Наушники','100 dBSPL/mW'),('Микрофон','-42 dBV/Pa')]:
   rows.append({'normalized_name':component+'_sensitivity','display_name':component+': Sensitivity (@1 kHz)','sources':{'razer_model':{'raw_name':component+': Sensitivity (@1 kHz)','raw_value':value,'section':'Основные характеристики'}},'resolved':{'selected_source':'razer_model','selected_unit':'','selected_value':value,'display_value':value,'status':'model_confirmed_official'}})
  projected=canonical_attributes.project(rows)
  self.assertEqual({'Наушники','Микрофон'},{r['parent_group'] for r in projected})
  self.assertEqual(2,len({r['display_name'] for r in projected}))
  self.assertTrue(any('микрофона' in r['display_name'] for r in projected))
  self.assertTrue(any('наушников' in r['display_name'] for r in projected))
 def test_raw_excluded_evidence_exports_separately_and_cannot_create_product_fact(self):
  from io import BytesIO
  from openpyxl import load_workbook
  with tempfile.TemporaryDirectory() as folder:
   db=Path(folder)/'test.sqlite3';jobs.initialize(db)
   with closing(sqlite3.connect(db)) as c:
    c.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('s82','Stage82','Products','{}',storage._now()))
    c.execute("INSERT INTO products (id,batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json,identity_json) VALUES (1,'s82',1,'Test','Apple','TEST','','Смартфоны',0,'[]','{}','{}')");c.commit()
   row={'source':'https://www.apple.com/','section':'Privacy','subsection':'','raw_label':'Privacy','raw_value':'=1+1','region':'US','component_scope':'device','identity_relation':{},'class':'D','action':'excluded_marketing','decision_reason':'No independent technical value','canonical_targets':[]}
   card_evidence.save(db,1,'canonical_mapping_audit',{'rows':[row]})
   workbook=load_workbook(BytesIO(exporter.export_batch(db,'s82')))
   self.assertEqual(2,workbook['Исходные поля и mapping'].max_row)
   self.assertNotEqual('f',workbook['Исходные поля и mapping']['H2'].data_type)
   self.assertEqual(1,workbook['Атрибуты'].max_row)
   self.assertEqual([],jobs.get_facts(db,1));workbook.close()
