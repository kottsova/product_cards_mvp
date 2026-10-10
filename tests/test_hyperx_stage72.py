"""Stage 72 offline negative controls against newly captured official page structures."""
import json,tempfile,unittest
from pathlib import Path
from product_tool.adapters.hyperx import HyperXAdapter
from product_tool.hyperx_page import refine_page
from product_tool.hyperx_presentation import label
from product_tool.adapters.hyperx_discovery import HyperXDiscovery
from product_tool import jobs,card_evidence,readiness,exporter,attribute_projection,card_presentation
from tests.test_hyperx_worker_integration import seed_product

R=Path(__file__).resolve().parents[1]/'reports/hyperx_stage72'

class HyperXStage72Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows=json.loads((R/'frozen_inputs.json').read_text(encoding='utf8'))['rows']
        cls.adapter=HyperXAdapter(urls={})
        cls.html={i:(R/'raw'/f'live_final_{i}.html').read_text(encoding='utf8') for i in range(1,11)}
    def parse(self,i,code=None,name=None):
        row=self.rows[i-1];url='https://hyperx.com/products/test'
        return refine_page(self.adapter.parse_page(self.html[i],url,catalog_code=code or row['article']),self.html[i],url,code or row['article'],name or row['name'])
    def test_all_ten_raw_facts_partition(self):
        for i in range(1,11):
            _,ev=self.parse(i)
            self.assertEqual(len(ev['raw_specs']),len(ev['accepted_specs'])+len(ev['rejected_specs']))
    def test_nine_exact_configurations(self):
        for i in range(1,10):self.assertEqual(self.parse(i)[1]['configuration_relation'],'exact_variant')
    def test_regional_ru_is_not_us(self):
        doc,ev=self.parse(10)
        self.assertEqual(doc.match_level,'model_confirmed');self.assertEqual(ev['configuration_relation'],'unverified')
        self.assertFalse(ev['exact_photo_assets'])
        self.assertFalse(any(a.name in {'Layout','SKU','Switch'} for a in doc.attributes))
    def test_wrong_color_is_not_confirmed(self):
        doc,ev=self.parse(1,name='HyperX Cloud Alpha 2 Wireless [color=White]')
        self.assertEqual(doc.match_level,'model_confirmed');self.assertFalse(ev['exact_photo_assets'])
        self.assertFalse(any(a.name=='Color' for a in doc.attributes))
    def test_wrong_layout_is_not_confirmed(self):
        _,ev=self.parse(4,name='HyperX Alloy Rise 75 [layout=UK Layout]')
        self.assertEqual(ev['configuration_relation'],'unverified');self.assertFalse(ev['exact_photo_assets'])
    def test_wrong_switch_is_not_confirmed(self):
        doc,ev=self.parse(4,name='HyperX Alloy Rise 75 [switch=HyperX Aqua - Tactile]')
        self.assertEqual(ev['configuration_relation'],'unverified')
        self.assertFalse(any(a.name=='Switch' for a in doc.attributes))
    def test_wired_wireless_are_different_models(self):
        doc,ev=self.parse(7,name='HyperX Pulsefire Haste 2 Wired')
        self.assertEqual(ev['model_relation'],'unverified');self.assertFalse(doc.attributes)
    def test_wrong_generation_is_not_confirmed(self):
        doc,ev=self.parse(8,name='HyperX QuadCast 2 S')
        self.assertEqual(doc.match_level,'mismatch');self.assertEqual(ev['model_relation'],'unverified')
    def test_close_part_number_is_not_exact(self):
        _,ev=self.parse(1,code='AJ5C8AA')
        self.assertEqual(ev['model_relation'],'model_confirmed');self.assertEqual(ev['configuration_relation'],'unverified')
    def test_shopify_product_id_does_not_become_hp_part_number(self):
        _,ev=self.parse(1,code='8683488018589')
        self.assertEqual(ev['configuration_relation'],'unverified')
        self.assertEqual(ev['identity_relations'][0]['type'],'storefront_product_id')
    def test_suffix_is_published_not_decoded(self):
        _,ev=self.parse(4)
        rel=next(r for r in ev['identity_relations'] if r.get('accepted'))
        self.assertEqual(rel['regional_suffix'],'ABA');self.assertEqual(rel['options']['Layout'],'US Layout')
    def test_microphone_sensitivity_is_separate(self):
        doc,_=self.parse(1)
        names=[a.name for a in doc.attributes if 'Sensitivity' in a.name]
        self.assertIn('Sensitivity',names);self.assertIn('Sensitivity (microphone)',names)
    def test_two_actuation_meanings_are_separate(self):
        doc,_=self.parse(4)
        names=[a.name for a in doc.attributes if 'Actuation Point' in a.name]
        self.assertEqual(len(set(names)),2)
    def test_aqua_linear_published_conflict_is_rejected(self):
        _,ev=self.parse(5)
        self.assertTrue(any(f['reason']=='official_switch_style_conflict' for f in ev['rejected_specs']))
        self.assertFalse(any(f['raw_label']=='Operation Style' for f in ev['accepted_specs']))
    def test_selected_other_variant_physical_values_are_candidates(self):
        doc,ev=self.parse(10)
        self.assertFalse(any(a.name=='Weight' for a in doc.attributes))
        self.assertTrue(any(f['raw_label']=='Weight' for f in ev['rejected_specs']))
    def test_battery_maximum_preserved_without_invented_mode(self):
        _,ev=self.parse(1)
        value=next(f['value'] for f in ev['accepted_specs'] if f['raw_label']=='Battery Life')
        self.assertEqual(value,'Up to 250 hours')
        self.assertNotIn('Bluetooth',value)
    def test_price_widget_is_not_a_spec(self):
        self.assertFalse(any(f['raw_label']=='Unit price' for f in self.parse(1)[1]['accepted_specs']))
    def test_russian_labels_keep_context(self):
        self.assertEqual(label('Sensitivity (microphone)'),'Чувствительность микрофона')
        self.assertEqual(label('Cable Length (imperial) and type'),'Длина и тип кабеля')
        self.assertEqual(label('Self-noise (RMS)'),'Собственный шум (RMS)')
        self.assertIn('Переключатели',label('Actuation Point (Switch Specifications #2)'))
    def test_frozen_dataset_has_no_seed_urls(self):
        self.assertEqual(len(self.rows),10)
        self.assertTrue(all('url' not in r for r in self.rows))
    def test_active_cooldown_is_access_failure_without_network(self):
        from product_tool.adapters.access_stop import stop_event
        from tests.test_hyperx_worker_integration import FakeSession
        with tempfile.TemporaryDirectory() as tmp:
            log=Path(tmp)/'fetch.json';original=json.dumps([stop_event('hyperx.com','rate_limit')])
            log.write_text(original,encoding='utf8');session=FakeSession({})
            adapter=HyperXAdapter(session,urls={},fetch_log_path=log,discovery_enabled=True,clock=lambda:0)
            doc=adapter.find_source('AJ5C7AA',name='HyperX Cloud Alpha 2 Wireless',deadline=100)
            self.assertEqual(doc.match_level,'blocked');self.assertFalse(session.calls)
            self.assertEqual(log.read_text(encoding='utf8'),original)
    def test_discovery_uses_published_form_links(self):
        root='<form action="/search"><input name="q"></form>'
        link='<a href="/products/hyperx-cloud-alpha-2-wireless-gaming-headset">AJ5C7AA HyperX Cloud Alpha 2 Wireless</a>'
        calls=[];events=[]
        def fetch(url,deadline):
            calls.append(url);return root if url=='https://hyperx.com/' else link if '/search?' in url else ''
        found=next(HyperXDiscovery(fetch,events.append).candidates('AJ5C7AA','HyperX Cloud Alpha 2 Wireless','Гарнитуры',999))
        self.assertEqual(found[0],'https://hyperx.com/products/hyperx-cloud-alpha-2-wireless-gaming-headset')
        self.assertTrue(any('/search?' in u for u in calls))
    def test_manual_optional_but_identity_metadata_is_not_spec_coverage(self):
        with tempfile.TemporaryDirectory() as tmp:
            db=Path(tmp)/'db.sqlite3';pid=seed_product(db,sku=self.rows[0]['article'],name=self.rows[0]['name'])
            doc,ev=self.parse(1);jobs.save_source_document(db,pid,doc);jobs.resolve_product(db,pid)
            card_evidence.save(db,pid,'hyperx',ev);jobs.set_photo_selection(db,pid,ev['exact_photo_assets'])
            r=readiness.card_readiness(db,pid)
            self.assertEqual(r['verdict'],'export_ready');self.assertEqual(r['manual_status'],'Не проверена')
            doc.attributes=[a for a in doc.attributes if a.name in {'Color','SKU'}]
            jobs.save_source_document(db,pid,doc);jobs.resolve_product(db,pid)
            self.assertEqual(readiness.card_readiness(db,pid)['verdict'],'not_ready')
    def test_common_ui_and_excel_accept_scoped_model_facts(self):
        from io import BytesIO
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as tmp:
            db=Path(tmp)/'batches.sqlite3';pid=seed_product(db,sku=self.rows[0]['article'],name=self.rows[0]['name'])
            doc,ev=self.parse(1);jobs.save_source_document(db,pid,doc);jobs.resolve_product(db,pid);card_evidence.save(db,pid,'hyperx',ev)
            rows=attribute_projection.final_attribute_rows(db,pid)
            shown=card_presentation.present_card_rows(rows,jobs.get_facts(db,pid),jobs.get_source_pages(db,pid),'Гарнитуры')
            self.assertTrue(shown)
            self.assertTrue(any(r['display_name']=='Динамик' for r in rows))
            book=load_workbook(BytesIO(exporter.export_batch(db,'b1')))
            self.assertIn('Варианты HyperX',book.sheetnames);self.assertIn('Готовность HyperX',book.sheetnames)
            self.assertIn('Фото HyperX',book.sheetnames);book.close()
            from product_tool.web import create_app
            from fastapi.testclient import TestClient
            with TestClient(create_app(Path(tmp),start_worker=False)) as client:
                response=client.get('/products/'+str(pid))
                self.assertEqual(response.status_code,200)
                self.assertNotIn('Только у базовой модели LG',response.text)
                self.assertIn('Полный SKU HyperX',response.text)
