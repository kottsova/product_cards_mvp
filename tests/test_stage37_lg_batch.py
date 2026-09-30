"""Offline batch UI acceptance using the saved LG pilot pages and files."""
from __future__ import annotations

from io import BytesIO
import gzip
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from product_tool import jobs, lg_batch, storage, worker
from product_tool.adapters.lg_policy import default_lg_adapters
from product_tool.web import create_app, _preview
from tests.coverage_controls import NoDealer

ROOT = Path(__file__).resolve().parents[1]
SAVED = [
    ROOT / "reports/source_census_2026-09-24_stage20/pilot/responses",
    ROOT / "reports/source_census_2026-09-24_stage21/probe/responses",
    ROOT / "reports/source_census_2026-09-24_stage21/verify/responses",
    ROOT / "reports/source_census_2026-09-24_stage22/docs_probe/responses",
    ROOT / "reports/source_census_2026-09-24_stage22/pilot2/responses",
]


def saved_pages():
    pages = {}
    for directory in SAVED:
        for line in (directory / "index.jsonl").read_text(encoding="utf-8").splitlines():
            item = json.loads(line)
            if item.get("saved_as"):
                with gzip.open(directory / item["saved_as"], "rt", encoding="utf-8") as handle:
                    pages[item["url"]] = (item["status"], handle.read())
    return pages


class SavedResponse:
    def __init__(self, url, status, text):
        self.url, self.status_code = url, status
        self._data = text.encode("latin-1") if "lge.com" in url else text.encode("utf-8")
        self.headers = {"Content-Type": "application/pdf" if self._data.startswith(b"%PDF-") else "text/html;charset=UTF-8"}
        self.history, self.encoding = (), "utf-8"
        self.cookies = type("Cookies", (), {"get_dict": lambda self: {}})()

    def iter_content(self, chunk_size):
        return iter((self._data,))

    def close(self):
        pass


class Replay:
    headers = {"User-Agent": "stage37-offline-replay"}

    def __init__(self):
        self.pages = saved_pages()
        self.calls = []

    def get(self, url, **kwargs):
        self.calls.append(url)
        if url not in self.pages:
            raise AssertionError(f"No saved LG response for {url}")
        return SavedResponse(url, *self.pages[url])


def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "Товары"
    sheet.append(["Категория", "Бренд", "Название", "Артикул"])
    sheet.append(["Колонки", "LG", "LG XL7S", "XL7S"])
    sheet.append(["Микроволновые печи", "LG", "LG MS2032GAS", "MS2032GAS"])
    sheet.append(["Музыкальные проигрыватели", "LG", "LG CK43", "CK43"])
    sheet.append(["Колонки", "LG", "Другой вариант колонки", "XL7S.OTHER"])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


class LgBatchTests(unittest.TestCase):
    def setUp(self):
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.database = self.root / "batches.sqlite3"
        self.client = TestClient(create_app(self.root))
        self.addCleanup(self.client.close)

    def upload(self):
        response = self.client.post("/upload", files={"file": ("small-lg.xlsx", workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        draft_id = response.headers["location"].rsplit("/", 1)[-1]
        draft = storage.get_draft(self.database, draft_id)
        preview = _preview(self.root / "uploads" / f"{draft_id}.xlsx", draft)
        form = {}
        for item in preview.products:
            for key in ("brand", "search_code", "alternate_code", "category"):
                form[f"{key}_{item.row_number}"] = getattr(item, key)
        response = self.client.post(f"/drafts/{draft_id}/confirm", data=form, follow_redirects=False)
        self.assertEqual(response.status_code, 303, response.text[:500])
        batch = storage.get_batch(self.database, draft_id)
        return batch, {p["search_code"]: p["id"] for p in batch["products"]}

    def test_batch_selection_queue_replay_cards_and_excel(self):
        batch, ids = self.upload()
        response = self.client.get(f"/batches/{batch['id']}")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Проверка LG одной кнопкой", response.text)
        self.assertEqual(lg_batch.default_ids(batch["products"]), {ids["XL7S"], ids["MS2032GAS"], ids["CK43"]})
        self.assertEqual(response.text.count('name="stages"'), 5)
        selected = [ids["XL7S"], ids["MS2032GAS"]]
        form = {"product_ids": [str(pid) for pid in selected], "stages": [str(n) for n in lg_batch.DEFAULT_STAGES]}
        response = self.client.post(f"/batches/{batch['id']}/lg-search", data=form, follow_redirects=False)
        self.assertEqual(response.status_code, 303, response.text[:500])
        self.assertEqual(lg_batch.selected_ids(self.database, batch["id"]), set(selected))
        self.assertEqual(jobs.list_jobs(self.database, ids["CK43"]), [])
        self.assertEqual(jobs.list_jobs(self.database, ids["XL7S.OTHER"]), [])
        response = self.client.get(f"/batches/{batch['id']}")
        self.assertIn("В очереди", response.text)
        self.assertIn("Инструкции", response.text)
        # Repeated click must leave exactly one active job per selected row.
        response = self.client.post(f"/batches/{batch['id']}/lg-search", data=form, follow_redirects=False)
        self.assertEqual(response.status_code, 303)
        self.assertEqual(len(jobs.list_jobs(self.database, selected[0])), 1)
        self.assertEqual(len(jobs.list_jobs(self.database, selected[1])), 1)
        replay = Replay()
        # browser_search=False: this offline test replays saved fixtures; it must never launch a real
        # browser subprocess even for a row whose sitemap lookup misses full_sku (Stage 43 fallback).
        adapters = lambda: default_lg_adapters(self.root, underlying_lg=replay, min_interval_seconds=0, browser_search=False)
        self.assertTrue(worker.run_once(self.database, adapters, dns_adapter_factory=NoDealer))
        self.assertTrue(worker.run_once(self.database, adapters, dns_adapter_factory=NoDealer))
        self.assertFalse(worker.run_once(self.database, adapters, dns_adapter_factory=NoDealer))
        self.assertTrue(replay.calls)
        self.assertFalse(any("/ck43/" in url.casefold() for url in replay.calls))
        self.assertEqual(lg_batch.batch_rows(self.database, batch["products"], set(selected))[1],
                         {"selected": 2, "queued": 0, "running": 0, "completed": 2, "review": 2})
        self.assertEqual(jobs.list_jobs(self.database, ids["XL7S"])[0]["status"], "done")
        self.assertEqual(jobs.list_jobs(self.database, ids["MS2032GAS"])[0]["status"], "needs_review")
        self.assertEqual(len([p for p in jobs.get_photo_candidates(self.database, ids["XL7S"]) if p["selected"]]), 42)
        self.assertEqual(len([p for p in jobs.get_photo_candidates(self.database, ids["MS2032GAS"]) if p["selected"]]), 35)
        self.assertIn("конфликт", " ".join(lg_batch.card_summary(self.database, ids["MS2032GAS"])["reasons"]))
        self.assertIn("PDF", " ".join(lg_batch.card_summary(self.database, ids["XL7S"])["reasons"]))
        self.assertEqual(len(jobs.list_jobs(self.database, ids["CK43"])), 0)
        for pid in selected:
            self.assertEqual(jobs.list_jobs(self.database, pid)[0]["status"] in ("done", "needs_review"), True)
            self.assertEqual(self.client.get(f"/products/{pid}").status_code, 200)
        response = self.client.get(f"/batches/{batch['id']}/export.xlsx")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(BytesIO(response.content), read_only=True)
        try:
            self.assertIn("Готовность LG", book.sheetnames)
            rows = list(book["Готовность LG"].values)
            self.assertEqual(rows[0][2:4], ("Статус задания", "Готовность карточки"))
            self.assertEqual(len(rows), 5)
            self.assertIn("Инструкции", rows[1][4])
        finally:
            book.close()

    def test_batch_progress_get_tracks_queue_without_starting_worker(self):
        import re

        batch, ids = self.upload()
        selected = [ids["XL7S"], ids["MS2032GAS"]]
        lg_batch.queue_selected(self.database, batch["id"], selected, [1])

        def count(html, key):
            match = re.search(r'<strong id="lg-' + key + r'-count">(\d+)</strong>', html)
            self.assertIsNotNone(match, key)
            return int(match.group(1))

        url = f"/batches/{batch['id']}"
        queued = self.client.get(url).text
        self.assertEqual((count(queued, "queued"), count(queued, "running")), (2, 0))
        self.assertIn("fetch(window.location.pathname, {cache: 'no-store'})", queued)
        self.assertIn(f'data-lg-job="{selected[0]}"', queued)
        self.assertIn(f'data-lg-card="{selected[1]}"', queued)
        self.assertEqual(len(jobs.list_jobs(self.database, selected[0])), 1)
        self.assertEqual(len(jobs.list_jobs(self.database, selected[1])), 1)

        first = jobs.claim_next(self.database)
        self.assertIsNotNone(first)
        running = self.client.get(url).text
        self.assertEqual((count(running, "queued"), count(running, "running")), (1, 1))
        self.assertIsNone(jobs.claim_next(self.database))
        jobs.finish(self.database, first["id"], "done", "Completed.")
        second = jobs.claim_next(self.database)
        self.assertIsNotNone(second)
        self.assertNotEqual(first["id"], second["id"])
        next_page = self.client.get(url).text
        self.assertEqual((count(next_page, "queued"), count(next_page, "running"), count(next_page, "completed")), (0, 1, 1))
        jobs.finish(self.database, second["id"], "done", "Completed.")
        finished = self.client.get(url).text
        self.assertEqual((count(finished, "queued"), count(finished, "running"), count(finished, "completed")), (0, 0, 2))
        self.assertEqual(len(jobs.list_jobs(self.database, selected[0])), 1)
        self.assertEqual(len(jobs.list_jobs(self.database, selected[1])), 1)

    def test_global_worker_claim_is_sequential(self):
        batch, ids = self.upload()
        lg_batch.queue_selected(self.database, batch["id"], [ids["XL7S"], ids["MS2032GAS"]], [1])
        first = jobs.claim_next(self.database)
        self.assertIsNotNone(first)
        self.assertIsNone(jobs.claim_next(self.database))
        jobs.finish(self.database, first["id"], "done", "Completed.")
        self.assertIsNotNone(jobs.claim_next(self.database))

    def test_second_worker_process_cannot_recover_active_job(self):
        with worker.single_worker(self.root):
            with self.assertRaisesRegex(RuntimeError, "already running"):
                with worker.single_worker(self.root):
                    pass

    def test_rejects_two_in_same_category_without_partial_queue(self):
        batch, ids = self.upload()
        form = {"product_ids": [str(ids["XL7S"]), str(ids["XL7S.OTHER"])], "stages": ["1"]}
        response = self.client.post(f"/batches/{batch['id']}/lg-search", data=form, follow_redirects=False)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(jobs.list_jobs(self.database, ids["XL7S"]), [])

    def test_unselected_data_stages_are_not_reported_as_failed_searches(self):
        batch, ids = self.upload()
        pid = ids["CK43"]
        lg_batch.queue_selected(self.database, batch["id"], [pid], [1, 6])
        reasons = " ".join(lg_batch.card_summary(self.database, pid)["reasons"])
        self.assertIn("Описание не проверялось", reasons)
        self.assertIn("Характеристики не проверялись", reasons)
        self.assertIn("Фото не проверялись", reasons)

    def test_large_document_reason_is_visible(self):
        batch, ids = self.upload()
        pid = ids["XL7S"]
        lg_batch.queue_selected(self.database, batch["id"], [pid], [1, 6])
        job = jobs.list_jobs(self.database, pid)[0]
        jobs.progress(self.database, job["id"], 6, "PDF файл превышает лимит 25 МБ.")
        reasons = lg_batch.card_summary(self.database, pid)["reasons"]
        self.assertTrue(any("слишком большой" in reason for reason in reasons))

    def test_deferred_instruction_has_explicit_reason(self):
        batch, ids = self.upload()
        pid = ids["CK43"]
        lg_batch.queue_selected(self.database, batch["id"], [pid], [1, 2, 3, 4])
        self.assertIn("отложен", " ".join(lg_batch.card_summary(self.database, pid)["reasons"]))
