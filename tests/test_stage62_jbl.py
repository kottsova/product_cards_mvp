import json,tempfile,unittest,hashlib
from pathlib import Path
from unittest.mock import patch
from product_tool.adapters.jbl import parse_pdp,JBLAdapter,doc_type
from product_tool.adapters.jbl_documents import assess_guide
from product_tool import jobs,storage,worker,card_evidence,jbl_pipeline,attribute_projection,exporter

def page(sku='JBLTESTBLKEU',mpn=None,pid=None,specs='',images=None):
 payload={'@type':'Product','name':'JBL Test','sku':sku,'mpn':sku if mpn is None else mpn,'image':images or ['https://uk.jbl.com/test_black.png']}
 return '<div class="product-wrapper" data-pid="'+(sku if pid is None else pid)+'"></div><script type="application/ld+json">'+json.dumps(payload)+'</script><div id="pdp-specs--specifications-content"><h3>Battery</h3><dl>'+specs+'</dl></div>'
def spec(k,v):return '<dt>'+k+'</dt><dd>'+v+'</dd>'
class JBLTests(unittest.TestCase):
 def parse(self,html,code='JBLTESTBLKEU'):return parse_pdp(html,'https://uk.jbl.com/JBLTESTBLKEU.html',code)
 def test_exact_identity(self):self.assertEqual(self.parse(page())[0].match_level,'full_sku')
 def test_other_color_not_exact(self):
  d,e=self.parse(page('JBLTESTBLUEU',specs=spec('Battery Life','40')));self.assertNotEqual(d.match_level,'full_sku');self.assertFalse(d.attributes);self.assertFalse(e['exact_photo_assets'])
 def test_other_region_not_exact(self):self.assertNotEqual(self.parse(page('JBLTESTBLKAM'))[0].match_level,'full_sku')
 def test_base_name_not_sku(self):self.assertNotEqual(self.parse(page('JBLTEST'))[0].match_level,'full_sku')
 def test_mpn_disagreement(self):self.assertNotEqual(self.parse(page(mpn='JBLTESTBLKAM'))[0].match_level,'full_sku')
 def test_pid_disagreement(self):self.assertNotEqual(self.parse(page(pid='OTHER'))[0].match_level,'full_sku')
 def test_foreign_host(self):self.assertNotEqual(parse_pdp(page(),'https://fake-jbl.com/x','JBLTESTBLKEU')[0].match_level,'full_sku')
 def test_multiple_products_bundle(self):self.assertNotEqual(self.parse(page()+page('JBLSMARTTX'))[0].match_level,'full_sku')
 def test_raw_battery_modes_separate(self):
  d,e=self.parse(page(specs=spec('Maximum music play time with ANC off (hours)','65')+spec('Maximum music play time with ANC on (hours)','50')+spec('Maximum talk time (hours)','33')));self.assertEqual([a.value for a in d.attributes],['65','50','33']);self.assertTrue(all(a.section=='Battery' for a in d.attributes))
 def test_conditional_inputs_candidate(self):
  d,e=self.parse(page(specs=spec('Audio inputs','USB playback available in US version, other versions Service only')));self.assertFalse(d.attributes);self.assertEqual(e['rejected_specs'][0]['reason'],'conditional_or_regional_option')
 def test_invalid_microphone_count(self):self.assertFalse(self.parse(page(specs=spec('Number of Microphones','5.3')))[0].attributes)
 def test_footer_faq_not_description(self):
  d,e=self.parse(page()+'<section id="product--support">Reset firmware FAQ</section>');self.assertEqual(d.description,'');self.assertFalse(d.attributes)
 def test_gallery_foreign_asset_candidate(self):
  d,e=self.parse(page(images=['https://evil.test/x.png']));self.assertFalse(e['exact_photo_assets'])
 def test_no_other_swatches_imported(self):
  d,e=self.parse(page()+'<img src="https://uk.jbl.com/blue.png">');self.assertEqual(len(d.photo_candidates),1)
 def test_typed_documents(self):
  self.assertEqual(doc_type('Quick Start Guide (App)'),'App Guide');self.assertEqual(doc_type('Safety Sheet'),'Safety Sheet');self.assertEqual(doc_type('Quick Start Guide (Multilingual)'),'Quick Start Guide')
 def test_safety_cannot_be_guide(self):self.assertFalse(assess_guide(['JBL Test RU руководство зарядка подключение'],'JBL Test','Safety Sheet')['verified'])
 def test_qsg_operational_ru(self):
  with patch('product_tool.adapters.jbl_documents.document_languages',return_value={'russian_present':True,'russian_instruction':False}):
   self.assertTrue(assess_guide(['JBL Test QUICK START GUIDE RU включение зарядка подключение'],'JBL Test','Quick Start Guide')['verified'])
 def test_qsg_legal_ru_rejected(self):
  with patch('product_tool.adapters.jbl_documents.document_languages',return_value={'russian_present':True,'russian_instruction':False}):self.assertFalse(assess_guide(['JBL Test QUICK START GUIDE RU гарантия безопасность'],'JBL Test','Quick Start Guide')['verified'])
 def test_qsg_spaced_cover(self):
  with patch('product_tool.adapters.jbl_documents.document_languages',return_value={'russian_present':True,'russian_instruction':False}):self.assertTrue(assess_guide(['JBL Test QUIC K ST ART GUIDE RU зарядка подключение'],'JBL Test','Quick Start Guide')['verified'])
 def test_real_official_templates_and_raw_sections(self):
  root=Path(__file__).resolve().parents[1]
  html=(root/'reports/jbl_stage62/attended/baseline.html').read_text(encoding='utf-8')
  d,e=parse_pdp(html,'https://de.jbl.com/JBLFLIP6BLKEU.html','JBLFLIP6BLKEU');self.assertEqual(d.match_level,'full_sku');self.assertEqual(len(d.attributes),21)
  html=(root/'reports/jbl_stage62/other_2.html').read_text(encoding='utf-8')
  d,e=parse_pdp(html,'https://id.jbl.com/en/bluetooth-portables/JBLCHARGE5BLK.html','JBLCHARGE5BLK');self.assertEqual(d.match_level,'full_sku');self.assertTrue(d.photo_candidates);self.assertTrue(all(a.section for a in d.attributes))
 def test_qsg_wrong_model_rejected(self):
  with patch('product_tool.adapters.jbl_documents.document_languages',return_value={'russian_present':True,'russian_instruction':False}):self.assertFalse(assess_guide(['JBL Other QUICK START GUIDE RU зарядка подключение'],'JBL Test','Quick Start Guide')['verified'])
 def test_capture_tampering(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td);(p/'x.html').write_text(page());(p/'JBLTESTBLKEU.manifest.json').write_text(json.dumps({'article':'JBLTESTBLKEU','records':[{'file':'x.html','sha256':'bad','status':200,'challenge':False,'final_url':'https://uk.jbl.com/x'}]}));self.assertIsNone(JBLAdapter(capture_dir=p)._capture('JBLTESTBLKEU'))
 def test_worker_ui_excel_projection_and_selection(self):
  with tempfile.TemporaryDirectory() as td:
   p=Path(td)/'db.sqlite3';jobs.initialize(p)
   with storage._connection(p) as db:
    db.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('jbltest','test','Products','{}',storage._now()))
    pid=db.execute('INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (?,?,?,?,?,?,?,?,?,?)',('jbltest',1,'JBL Test','JBL','JBLTESTBLKEU','','headphones',0,'[]','{}')).lastrowid
   html=page(specs=spec('Maximum music play time with ANC off (hours)','65')+spec('Maximum music play time with ANC on (hours)','50')+spec('Dimensions (Width x Height x Depth) (cm)','17.8 x 6.8 x 7.2')+spec('Charging Case Weight (g)','50')+spec('Earpiece Weight (g)','5'))
   class Fake(JBLAdapter):
    def find_source(self,article,**kwargs):
     d,r=parse_pdp(html,'https://uk.jbl.com/JBLTESTBLKEU.html',article);self.reports[article]=r;return d
    def find_documents(self,doc,article):return []
   for n in range(2):
    jobs.enqueue(p,pid,[1,2,3,4]);worker.run_once(p,jbl_adapter_factory=lambda:Fake(capture_dir=Path(td)))
    if n==0:jobs.set_photo_selection(p,pid,[],mode='exact')
   self.assertEqual(jbl_pipeline.card_readiness(p,pid)['verdict'],'not_ready')
   self.assertTrue(all('JBL' in r['reason'] for r in jobs.get_resolved(p,pid)))
   rows=attribute_projection.final_attribute_rows(p,pid);labels=[r['display_name'] for r in rows]
   self.assertIn('Ширина товара, мм',labels);self.assertIn('Вес зарядного кейса, г',labels);self.assertIn('Вес наушников (один или пара не уточнено), г',labels)
   self.assertTrue(any('без ANC' in x for x in labels));self.assertTrue(any('с ANC' in x for x in labels))
   from openpyxl import load_workbook
   from io import BytesIO
   book=load_workbook(BytesIO(exporter.export_batch(p,'jbltest')),read_only=True);self.assertIn('Готовность JBL',book.sheetnames);book.close()
if __name__=='__main__':unittest.main()
