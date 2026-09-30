"""Stage 53: saved evidence, presentation and measured photo regressions."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from product_tool import exporter, jobs, lg_batch
from product_tool.adapters.common import PhotoCandidate, ProductDocument, RawAttribute, SourceDocument
from product_tool.display import display_access_error, display_name_ru
from product_tool.fetch_history import record_fetch_attempt, save_source_snapshot, latest_source_snapshot
from product_tool.photo_metadata import image_dimensions, inspect_saved_photo
from product_tool.product_description import product_description
from product_tool.web import create_app


class Stage53PresentationTests(unittest.TestCase):
    def setUp(self):
        tmp = TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.root = Path(tmp.name)
        self.db = self.root / "batches.sqlite3"
        jobs.initialize(self.db)
        with sqlite3.connect(self.db) as con:
            con.execute("INSERT INTO batches VALUES ('b','fixture.xlsx','Sheet1','{}','now')")
            con.execute(
                "INSERT INTO products (id,batch_id,row_number,name,brand,search_code,"
                "alternate_code,category,needs_confirmation,issues_json,original_values_json) "
                "VALUES (1,'b',2,'Product','LG','TEST.USAR','','Washer',0,'[]','{}')"
            )
        con.close()

    def _book(self):
        return load_workbook(BytesIO(exporter.export_batch(self.db, "b")), read_only=True)

    def test_russian_evidence_survives_sqlite_json_ui_and_xlsx(self):
        evidence = "Полный артикул подтверждён на странице."
        jobs.save_source_document(self.db, 1, SourceDocument(
            "lg_kz", "LG Казахстан", "https://www.lg.com/kz/test",
            found_model="TEST.USAR", match_level="full_sku", evidence=evidence,
            html="<h1>TEST.USAR</h1>",
        ))
        self.assertEqual(jobs.get_source_pages(self.db, 1)[0]["evidence"], evidence)
        attempt = record_fetch_attempt(self.db, 1, "lg_kz", status="success")
        save_source_snapshot(self.db, 1, "lg_kz", attempt,
                             source_url="https://www.lg.com/kz/test",
                             content="<h1>TEST.USAR</h1>", extracted={"evidence": evidence})
        self.assertEqual(latest_source_snapshot(self.db, 1, "lg_kz")["extracted"]["evidence"], evidence)
        with TestClient(create_app(self.root)) as client:
            self.assertIn(evidence, client.get("/products/1").text)
        book = self._book()
        try:
            self.assertIn(evidence, [row[5] for row in book["Источники"].values][1:])
            self.assertNotIn("?????", str(list(book["Источники"].values)))
        finally:
            book.close()

    def test_latin_feature_has_semantic_header_and_all_conflict_sides(self):
        self.assertEqual(display_name_ru("turbowash", ["TurboWash"]), "TurboWash")
        for key,site,url,marker,section in (
            ("lg_kz", "LG Казахстан", "https://www.lg.com/kz/test", "—", "Функции"),
            ("lg_ru", "LG Россия", "https://www.lg.com/ru/test", "●", "Опции"),
        ):
            jobs.save_source_document(self.db, 1, SourceDocument(
                key, site, url, match_level="full_sku",
                attributes=[RawAttribute("TurboWash", marker, section, True)],
                html="<h1>TEST.USAR</h1>",
            ))
        jobs.resolve_product(self.db, 1)
        job_id = jobs.enqueue(self.db, 1, [1, 3])
        summary = lg_batch.card_summary(self.db, 1, jobs.list_jobs(self.db, 1)[0])
        self.assertTrue(summary["conflicts"])
        values = summary["conflicts"][0]["values"]
        self.assertTrue(any("LG Казахстан" in item and "Функции" in item for item in values))
        self.assertTrue(any("LG Россия" in item and "Опции" in item for item in values))
        with TestClient(create_app(self.root)) as client:
            html = client.get("/products/1").text
            self.assertIn("Функции", html)
            self.assertIn("Опции", html)
            self.assertIn("TurboWash", html)
        book = self._book()
        try:
            headers = next(book["Washer"].values)
            self.assertEqual(headers.count("TurboWash"), 1)
            self.assertNotIn("Дополнительная характеристика", headers)
        finally:
            book.close()

    def test_manual_status_and_original_official_urls(self):
        support = "https://www.lg.com/kz/support/product-support/cs-TEST.USAR/"
        pdf = "https://gscs-b2c.lge.com/open/downloadFile?fileId=abc"
        jobs.save_source_document(self.db, 1, SourceDocument(
            "lg_kz_support", "LG KZ support", support, match_level="full_sku",
            found_model="TEST.USAR", html="<h1>TEST.USAR</h1>",
        ))
        doc = ProductDocument("Russian manual", "Русский", "", "", pdf,
                              support, "TEST", "TEST.USAR", support, True)
        jobs.save_documents(self.db, 1, "lg_kz_support", [doc])
        jobs.save_documents(self.db, 1, "lg_kz_support", [doc])
        self.assertEqual(len(jobs.get_documents(self.db, 1)), 1)
        with TestClient(create_app(self.root)) as client:
            html = client.get("/products/1").text
            self.assertIn("Проверена", html)
            self.assertIn(pdf.replace("&", "&amp;"), html)
            self.assertIn(support, html)
        book = self._book()
        try:
            rows=list(book["Инструкции"].values)
            self.assertEqual(rows[1][-1], "Проверена")
            self.assertEqual(rows[1][7:9], (support,pdf))
        finally:
            book.close()

    def test_candidate_sheet_is_conditional_and_photo_measurement_is_truthful(self):
        exact=SourceDocument("lg_kz", "LG KZ", "https://www.lg.com/kz/test",
            match_level="full_sku", html="<h1>TEST.USAR</h1>",
            photo_candidates=[PhotoCandidate("https://www.lg.com/photo.png","exact","product_gallery",320)])
        jobs.save_source_document(self.db, 1, exact)
        book=self._book()
        try:
            self.assertNotIn("Фото-кандидаты",book.sheetnames)
            photo=list(book["Фотографии"].values)
            self.assertEqual(photo[1][5:9], ("не определено",)*4)
        finally: book.close()
        candidate=SourceDocument("lg_ru", "LG RU", "https://www.lg.com/ru/other",
            match_level="base_model", html="<h1>TEST</h1>",
            photo_candidates=[PhotoCandidate("https://www.lg.com/candidate.webp","candidate","product_gallery",100)])
        jobs.save_source_document(self.db, 1, candidate)
        jobs.set_photo_selection(self.db, 1, ["candidate"])
        photo=next(x for x in jobs.get_photo_candidates(self.db,1) if x["asset_key"]=="candidate")
        self.assertTrue(jobs.save_photo_metadata(self.db,1,photo["id"],photo["url"],
            width=2000,height=1000,size_bytes=1800000,image_format="WEBP"))
        jobs.save_source_document(self.db, 1, candidate)
        photo=next(x for x in jobs.get_photo_candidates(self.db,1) if x["asset_key"]=="candidate")
        self.assertEqual((photo["verified_width"],photo["verified_height"]), (2000,1000))
        with TestClient(create_app(self.root)) as client:
            html=client.get("/products/1").text
            self.assertIn("2000 × 1000 px",html)
            self.assertIn("WEBP",html)
        book=self._book()
        try:
            rows=list(book["Фото-кандидаты"].values)
            self.assertEqual(rows[0][-4:],("Ширина, px","Высота, px","Размер файла","Формат"))
            self.assertEqual(rows[1][-4],2000)
            self.assertEqual(rows[1][-3],1000)
            self.assertEqual(rows[1][-1],"WEBP")
        finally: book.close()

    def test_internal_challenge_does_not_look_like_site_http_403(self):
        shown=display_access_error("HTTP-ошибка: HTTP 403 [policy_challenge_confirmed]")
        self.assertIn("HTTP 200",shown)
        self.assertNotIn("HTTP 403",shown)
        self.assertEqual(display_access_error("HTTP-ошибка: HTTP 403"),
                         "HTTP-ошибка: HTTP 403")

    def test_image_header_and_help_filter(self):
        png=bytes.fromhex("89504e470d0a1a0a0000000d49484452")+(2000).to_bytes(4,"big")+(1000).to_bytes(4,"big")+b"extra"
        self.assertEqual(image_dimensions(png),(2000,1000,"PNG"))
        jpeg = bytes.fromhex("ffd8ffc0001108") + (1000).to_bytes(2,"big") + (2000).to_bytes(2,"big") + b"0"*12
        self.assertEqual(image_dimensions(jpeg),(2000,1000,"JPG"))
        self.assertIsNone(image_dimensions(b"not an image"))

        class Cookies:
            def get_dict(self): return {}
        class Response:
            status_code=200
            url="https://www.lg.com/photo.png"
            headers={"content-type":"image/png"}
            cookies=Cookies()
            history=()
            encoding=None
            def iter_content(self, chunk_size): yield png
            def close(self): pass
        class Session:
            headers={}
            def get(self,url,**kwargs): return Response()
        measured=inspect_saved_photo("https://www.lg.com/photo.png","lg_kz",
                                     self.root/"lg_fetch_log.json",underlying=Session())
        self.assertEqual(measured,{"width":2000,"height":1000,"size_bytes":len(png),"format":"PNG"})
        description=("Умная стирка\n"
                     "LG ThinQ помогает управлять режимами.\n"
                     "Вопросы и ответы\n"
                     "Как зарегистрировать технику?\n"
                     "1. Установите приложение.")
        shown=product_description(description)
        self.assertIn("LG ThinQ",shown)
        self.assertNotIn("Вопросы и ответы",shown)
        self.assertNotIn("1.",shown)


if __name__ == "__main__":
    unittest.main()
