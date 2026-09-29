"""Stage 36: two saved Bosch KZ pages through upload, queue, run_once, and export."""
from __future__ import annotations
import gzip
import io
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from product_tool import bosch_readiness, jobs, storage, worker
from product_tool.adapters.bosch_home import BoschHomeAdapter, _asset_key, parse_page, verified_pages
from product_tool.adapters.policy_session import PolicyResponse
from product_tool.coverage.catalog_units import load_catalog
from product_tool.coverage.planner import build_plan
from product_tool.web import create_app

ROOT=Path(__file__).resolve().parents[1]
SAVED=ROOT/'reports/source_census_2026-09-26_stage35/route_check/responses'
CATALOG=ROOT/'data/catalog_2026-09-21_filtered.xlsx'


class Replay:
    def __init__(self):
        index=[json.loads(x) for x in (SAVED/'index.jsonl').read_text(encoding='utf8').splitlines()]
        self.pages={}
        for entry in index:
            if '/product/' in entry['url']:
                with gzip.open(SAVED/entry['saved_as'],'rt',encoding='utf8') as f:
                    self.pages[entry['url']]=f.read()
        self.calls=[]
    def get(self,url,**kwargs):
        self.calls.append(url)
        if url not in self.pages:
            raise AssertionError('No saved response for '+url)
        return PolicyResponse(url,200,self.pages[url],'text/html')


def workbook(rows):
    book=Workbook();sheet=book.active
    sheet.title='Products'
    sheet.append(['Category','Brand','Name','Seller code','Model'])
    for u in rows:sheet.append([u.category,u.brand,u.title,u.seller_sku,''])
    output=io.BytesIO();book.save(output);book.close();return output.getvalue()


class BoschOrdinaryPath(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=TemporaryDirectory();cls.workdir=Path(cls.tmp.name)
        cls.db=cls.workdir/'batches.sqlite3'
        cls.replay=Replay();cls.manifest=verified_pages()
        all_units=load_catalog(CATALOG).units
        cls.units={code:next(u for u in all_units if u.seller_sku==code and u.brand.upper()=='BOSCH') for code in cls.manifest}
        cls.other=next(u for u in all_units if u.brand.upper()=='BOSCH' and u.category==cls.units['TWK7203'].category and u.seller_sku!='TWK7203')
        with TestClient(create_app(cls.workdir)) as client:
            upload=client.post('/upload',files={'file':('bosch.xlsx',workbook([*cls.units.values(),cls.other]),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},follow_redirects=False)
            assert upload.status_code==303,upload.text
            form={}
            for number,unit in enumerate([*cls.units.values(),cls.other],start=2):
                form.update({f'brand_{number}':unit.brand,f'search_code_{number}':unit.seller_sku,f'alternate_code_{number}':'',f'category_{number}':unit.category})
            confirm=client.post(upload.headers['location']+'/confirm',data=form,follow_redirects=False)
            assert confirm.status_code==303,confirm.text
            cls.batch_id=confirm.headers['location'].rsplit('/',1)[-1]
            cls.products={p['search_code']:p for p in storage.get_batch(cls.db,cls.batch_id)['products']}
            for code in cls.manifest:
                started=client.post(f"/products/{cls.products[code]['id']}/search",data={'stages':['1','2','3','4','6']},follow_redirects=False)
                assert started.status_code==303,started.text
            while worker.run_once(cls.db,clock=lambda:0.0,bosch_adapter_factory=lambda:BoschHomeAdapter(http=cls.replay,clock=lambda:0.0)):
                pass
            exported=client.get(f'/batches/{cls.batch_id}/export.xlsx')
            assert exported.status_code==200,exported.text[:200]
            cls.xlsx=exported.content
            cls.product_html=client.get(f"/products/{cls.products['TWK7203']['id']}").text

    @classmethod
    def tearDownClass(cls):cls.tmp.cleanup()

    def test_plan_and_enqueue_guard(self):
        plan=build_plan()
        bosch=[u for u in plan['units'] if u['family']=='bosch_home']
        self.assertEqual({u['seller_sku'] for u in bosch if u['status']=='ready_to_run'},set(self.manifest))
        self.assertEqual(sum(u['status']=='ready_to_run' for u in bosch),2)
        self.assertTrue(all(u['status']!='ready_to_run' for u in bosch if u['seller_sku'] not in self.manifest))
        with self.assertRaises(ValueError):jobs.enqueue(self.db,self.products[self.other.seller_sku]['id'],[1,2,3,4,6])
        self.assertEqual(jobs.list_jobs(self.db,self.products[self.other.seller_sku]['id']),[])
        self.assertEqual(self.replay.calls,[self.manifest[code]['page_url'] for code in self.manifest])

    def test_saved_kz_pages_become_jobs_and_cards(self):
        for code,expected_specs,expected_photos in [('TWK7203',20,4),('MMB2111M',19,28)]:
            with self.subTest(code):
                pid=self.products[code]['id']
                job=jobs.list_jobs(self.db,pid)[0]
                self.assertEqual(job['status'],'done',job['message']+' '+str(bosch_readiness.card_readiness(self.db,pid)))
                page=next(s for s in jobs.get_source_pages(self.db,pid) if s['source_key']=='bosch_home')
                self.assertEqual((page['match_level'],page['found_model']),('full_sku',code))
                readiness=bosch_readiness.card_readiness(self.db,pid)
                self.assertEqual(readiness['verdict'],'export_ready')
                self.assertEqual(readiness['revision_status'],'unknown')
                self.assertIsNone(readiness['page_enr'])
                self.assertEqual(readiness['official_facts'],expected_specs)
                self.assertEqual(readiness['official_photos_selected'],expected_photos)
                self.assertEqual(readiness['gtin'],self.manifest[code]['gtin'])
                self.assertEqual(readiness['manual_family'],self.manifest[code]['manual']['family'])
                self.assertTrue(readiness['manual_russian_by_text'])
                self.assertFalse(readiness['manual_exact_code_in_pdf'])
                docs=[d for d in jobs.get_documents(self.db,pid) if d['source_key']=='bosch_home']
                self.assertEqual(len(docs),1)
                self.assertEqual(docs[0]['direct_url'],self.manifest[code]['manual']['url'])
                self.assertEqual(docs[0]['support_model'],self.manifest[code]['manual']['family'])
                self.assertIn(code,docs[0]['title'])
                if code=='MMB2111M':self.assertEqual(len(readiness['other_manuals_unverified']),3)

    def test_kettle_photos_exclude_other_variant_and_size_duplicates(self):
        pid=self.products['TWK7203']['id']
        photos=[p for p in jobs.get_photo_candidates(self.db,pid) if p['source_key']=='bosch_home']
        self.assertEqual(len(photos),4)
        self.assertEqual(len({p['asset_key'] for p in photos}),4)
        self.assertEqual(_asset_key('https://media3.bsh-group.com/Product_Shots/TWK7203_def.webp'),
                         _asset_key('https://media3.bsh-group.com/Product_Shots/TWK7203_1200x1200.webp'))
        self.assertFalse(any('TWK7203GB' in p['url'] or 'TAT7203GB' in p['url'] for p in photos))
        html=self.replay.pages[self.manifest['TWK7203']['page_url']]
        document,report=parse_page(html,self.manifest['TWK7203']['page_url'],'TWK7203',self.manifest['TWK7203']['category'])
        self.assertEqual(len(report['photo_excluded']),4)
        self.assertEqual(len(document.photos),4)

    def test_product_page_shows_bosch_facts_and_readiness(self):
        self.assertIn('id="bosch-readiness"',self.product_html)
        self.assertIn('E-Nr',self.product_html)
        self.assertIn('<th>Bosch Home</th>',self.product_html)
        self.assertIn('TWK720.',self.product_html)

    def test_wrong_page_model_or_revision_does_not_confirm(self):
        code='TWK7203';url=self.manifest[code]['page_url'];html=self.replay.pages[url]
        changed=html.replace('"mpn":"TWK7203"','"mpn":"TWK7203GB"',1)
        self.assertNotEqual(html,changed)
        doc,_=parse_page(changed,url,code,self.manifest[code]['category'])
        self.assertEqual(doc.match_level,'mismatch')
        self.assertFalse(doc.attributes)
        doc,_=parse_page(html,url,code+'/01',self.manifest[code]['category'])
        self.assertTrue(doc.error)

    def test_z_old_queued_unselected_row_stops_before_client(self):
        from uuid import uuid4
        job_id=uuid4().hex
        product_id=self.products[self.other.seller_sku]['id']
        with storage._connection(self.db) as connection:
            connection.execute("INSERT INTO search_jobs (id,product_id,stages_json,status,message,created_at,updated_at) VALUES (?,?,?,'queued',?,?,?)",
                               (job_id,product_id,'[1,2,3,4,6]','old queued job',storage._now(),storage._now()))
        before=list(self.replay.calls)
        def forbidden():raise AssertionError('unselected Bosch row constructed a source client')
        self.assertTrue(worker.run_once(self.db,bosch_adapter_factory=forbidden,clock=lambda:0.0))
        self.assertEqual(self.replay.calls,before)
        self.assertEqual(jobs.list_jobs(self.db,product_id)[0]['status'],'needs_review')
        self.assertEqual(jobs.get_source_pages(self.db,product_id),[])

    def test_excel_export_includes_bosch_provenance_and_status(self):
        book=load_workbook(io.BytesIO(self.xlsx),read_only=True,data_only=True)
        try:
            sheet=next(s for s in book.worksheets if 'Bosch Home' in s.title)
            rows=list(sheet.values)
            self.assertEqual(len(rows),3)
            for row in rows[1:]:
                self.assertEqual(row[2],'done')
                self.assertEqual(row[3],'export_ready')
                self.assertEqual(row[6],'unknown')
                self.assertEqual(row[10],'\u041d\u0435\u0442')
            self.assertTrue(any('Bosch Home' in list(sheet.values)[0] for sheet in book.worksheets))
        finally:book.close()
