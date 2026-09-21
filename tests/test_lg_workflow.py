"""Offline acceptance tests for LG multi-source identification and comparison."""

from pathlib import Path
from tempfile import TemporaryDirectory
import sqlite3
import unittest
from xml.sax.saxutils import escape

from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
import requests

from product_tool import jobs, worker
from product_tool.adapters.common import RawAttribute, SourceDocument, fetch_with_retry
from product_tool.adapters.lg import LGAdapter, lg_base_model, normalize_lg_sku
from product_tool.adapters.mechta import MechtaAdapter
from product_tool.adapters.sulpak import SulpakAdapter
from product_tool.exporter import export_batch
from product_tool.normalization import normalize_fact
from product_tool.web import create_app


LG_BASE_URL = "https://www.lg.com/kz/laundry/styler/s3wer/"
LG_FULL_URL = "https://www.lg.com/kz/laundry/styler/s3wer-alwpcom/"
SULPAK_URL = "https://www.sulpak.kz/g/parovoj_shkaf_lg_styler_s3wer_alwpcom"
MECHTA_URL = "https://www.mechta.kz/product/parovoy-shkaf-lg-s3-wer/"


def lg_html(model: str, *, color: str = "Белый") -> str:
    return f'''<html><head><meta name="description" content="Паровой шкаф LG"></head><body>
    <h1>LG Styler {model}</h1><section id="pdp-overview-section"><p>Уход за одеждой.</p></section>
    <section id="pdp-specs-section"><div class="c-compare-selling--all">
      <div class="c-compare-selling__item"><span class="c-compare-selling__spec-name">Размеры (Ш×В×Г)</span><span class="c-compare-selling__spec-desc">445 × 1850 × 585 мм</span></div>
      <div class="c-compare-selling__item"><span class="c-compare-selling__spec-name">Вес</span><span class="c-compare-selling__spec-desc">83 кг</span></div>
      <div class="c-compare-selling__item"><span class="c-compare-selling__spec-name">Цвет</span><span class="c-compare-selling__spec-desc">{color}</span></div>
      <div class="c-compare-selling__item"><span class="c-compare-selling__spec-name">True Steam</span><span class="c-compare-selling__spec-desc">●</span></div>
    </div></section></body></html>'''


def supplier_html(site: str, *, color: str = "Серый", weight: str = "83 кг") -> str:
    return f'''<html><head><meta name="description" content="{site} LG Styler"></head><body>
    <h1>Паровой шкаф LG S3WER.ALWPCOM</h1><p>Артикул: S3WER.ALWPCOM</p>
    <table class="characteristics">
      <tr><th>Габариты (В×Ш×Г)</th><td>1850 × 445 × 585 мм</td></tr>
      <tr><th>Масса товара</th><td>{weight}</td></tr>
      <tr><th>Вес с упаковкой</th><td>90 кг</td></tr>
      <tr><th>Цвет</th><td>{color}</td></tr>
      <tr><th>True Steam</th><td>Есть</td></tr>
    </table></body></html>'''


class FakeResponse:
    def __init__(self, url: str, body: str, status: int = 200):
        self.url, self.text, self.status_code = url, body, status
        self.content = body.encode("utf-8")

    def raise_for_status(self) -> None:
        if self.status_code >= 400:
            raise requests.HTTPError(f"HTTP {self.status_code}")


class FakeSession:
    def __init__(self, pages: dict[str, str], *, sitemap_urls: list[str] | None = None):
        self.pages = pages
        self.sitemap_urls = sitemap_urls if sitemap_urls is not None else list(pages)
        self.calls: list[tuple[str, float]] = []
        self.headers: dict[str, str] = {}

    def get(self, url: str, timeout: float) -> FakeResponse:
        self.calls.append((url, timeout))
        if url.endswith("/sitemap.xml"):
            locations = "".join(f"<url><loc>{escape(item)}</loc></url>" for item in self.sitemap_urls)
            return FakeResponse(url, f"<urlset>{locations}</urlset>")
        return FakeResponse(url, self.pages.get(url, ""), 200 if url in self.pages else 404)


class StaticAdapter:
    def __init__(self, document: SourceDocument, clock=None, advance: float = 0):
        self.document, self.calls, self.clock, self.advance = document, [], clock, advance
        self.source_key, self.site_name = document.source_key, document.site_name

    def find_source(self, sku: str, *, deadline: float) -> SourceDocument:
        self.calls.append((sku, deadline))
        if self.clock and self.advance:
            self.clock.value += self.advance
        return self.document


class FakeClock:
    def __init__(self): self.value = 0.0
    def __call__(self) -> float: return self.value


def source(key: str, match: str, facts: list[tuple[str, str]], *, error: str = "") -> SourceDocument:
    names = {"lg": "LG Казахстан", "sulpak": "Sulpak", "mechta": "Mechta"}
    urls = {"lg": LG_BASE_URL, "sulpak": SULPAK_URL, "mechta": MECHTA_URL}
    return SourceDocument(
        key, names[key], urls[key], found_model="S3WER" if key == "lg" else "S3WER.ALWPCOM",
        match_level=match, evidence="fixture evidence", error=error,
        attributes=[RawAttribute(name, value) for name, value in facts],
    )


def seed_product(database: Path, brand: str = "LG", sku: str = "S3WER.ALWPCOM") -> int:
    jobs.initialize(database)
    connection = sqlite3.connect(database)
    try:
        connection.execute(
            "INSERT INTO batches (id, filename, sheet_name, mapping_json, confirmed_at) VALUES ('b1','fixture.xlsx','Товары','{}','2026-01-01')"
        )
        cursor = connection.execute(
            "INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES ('b1',2,'LG Styler',?,?, '', 'Паровые шкафы',0,'[]','{}')",
            (brand, sku),
        )
        connection.commit()
        return int(cursor.lastrowid)
    finally:
        connection.close()


class LGAdapterFixtureTests(unittest.TestCase):
    def test_lg_base_model_and_sulpak_full_sku(self) -> None:
        lg_session = FakeSession({LG_BASE_URL: lg_html("S3WER")})
        lg = LGAdapter(lg_session, clock=lambda: 0)
        official = lg.find_source("S3WER.ALWPCOM", deadline=20)
        self.assertEqual((official.found_model, official.match_level), ("S3WER", "base_model"))
        self.assertEqual(len(lg_session.calls), 2)

        supplier = SulpakAdapter(
            FakeSession({SULPAK_URL: supplier_html("Sulpak")}),
            clock=lambda: 0,
        ).find_source("S3WER.ALWPCOM", deadline=30)
        self.assertEqual((supplier.found_model, supplier.match_level), ("S3WER.ALWPCOM", "full_sku"))
        self.assertIn("S3WER.ALWPCOM", supplier.evidence)

    def test_exact_lg_sku_stops_search_immediately(self) -> None:
        session = FakeSession(
            {LG_FULL_URL: lg_html("S3WER.ALWPCOM"), LG_BASE_URL: lg_html("S3WER")},
            sitemap_urls=[LG_FULL_URL, LG_BASE_URL],
        )
        document = LGAdapter(session, clock=lambda: 0).find_source("S3WER.ALWPCOM", deadline=20)
        self.assertEqual(document.match_level, "full_sku")
        self.assertEqual([call[0] for call in session.calls], ["https://www.lg.com/kz/sitemap.xml", LG_FULL_URL])

    def test_lg_normalizer_is_brand_specific(self) -> None:
        self.assertEqual(normalize_lg_sku(" s3wer.alwpcom "), "S3WER.ALWPCOM")
        self.assertEqual(lg_base_model("S3WER.ALWPCOM"), "S3WER")

    def test_request_retry_is_bounded(self) -> None:
        class TimeoutSession:
            def __init__(self): self.calls = 0
            def get(self, url, timeout):
                self.calls += 1
                self.last_timeout = timeout
                raise requests.Timeout("blocked")
        session = TimeoutSession()
        with self.assertRaises(Exception):
            fetch_with_retry(session, "https://example.test", deadline=60, clock=lambda: 0)
        self.assertEqual(session.calls, 2)
        self.assertLessEqual(session.last_timeout, 10)


class NormalizationAndResolutionTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.database = Path(temporary.name) / "batches.sqlite3"
        self.product_id = seed_product(self.database)

    def save(self, *documents: SourceDocument) -> None:
        for document in documents:
            jobs.save_source_document(self.database, self.product_id, document)
        jobs.resolve_product(self.database, self.product_id)

    def resolved(self, name: str) -> dict:
        return next(item for item in jobs.get_resolved(self.database, self.product_id) if item["normalized_name"] == name)

    def test_suppliers_match_and_override_base_model(self) -> None:
        self.save(
            source("lg", "base_model", [("Цвет", "Белый")]),
            source("sulpak", "full_sku", [("Цвет", "Серый")]),
            source("mechta", "full_sku", [("Цвет", " серый ")]),
        )
        item = self.resolved("color")
        self.assertEqual((item["selected_value"], item["status"], item["conflict"]), ("серый", "confirmed_two_suppliers", 0))

    def test_supplier_conflict_requires_review(self) -> None:
        self.save(
            source("sulpak", "full_sku", [("Цвет", "Серый")]),
            source("mechta", "full_sku", [("Цвет", "Белый")]),
        )
        item = self.resolved("color")
        self.assertEqual(item["status"], "needs_review")
        self.assertEqual(item["selected_value"], "")
        self.assertEqual(item["conflict"], 1)

    def test_product_and_package_weight_are_separate(self) -> None:
        product = normalize_fact(RawAttribute("Масса товара", "83000 г"))
        package = normalize_fact(RawAttribute("Вес с упаковкой", "90 кг"))
        self.assertEqual((product.normalized_name, product.normalized_value, product.unit), ("product_weight", "83", "kg"))
        self.assertEqual((package.normalized_name, package.normalized_value, package.unit), ("package_weight", "90", "kg"))
        bare = normalize_fact(RawAttribute("Вес с упаковкой (кг)", "86"))
        display = normalize_fact(RawAttribute("Дисплей", "Сенсорный"))
        self.assertEqual((bare.normalized_value, bare.unit), ("86", "kg"))
        self.assertEqual(display.normalized_name, "display_type")

    def test_dimension_axis_order_normalizes_to_same_value(self) -> None:
        whd = normalize_fact(RawAttribute("Габариты (Ш×В×Г)", "445 × 1850 × 585 мм"))
        hwd = normalize_fact(RawAttribute("Размеры (В×Ш×Г)", "1850 x 445 x 585 мм"))
        self.assertEqual(whd.normalized_name, "product_dimensions")
        self.assertEqual(whd.normalized_value, hwd.normalized_value)
        self.assertEqual(whd.unit, "mm")

    def test_unselected_stage_does_not_erase_existing_facts(self) -> None:
        original = source("lg", "base_model", [("Цвет", "Белый")])
        jobs.save_source_document(self.database, self.product_id, original)
        refreshed = source("lg", "base_model", [])
        jobs.save_source_document(
            self.database, self.product_id, refreshed,
            update_description=False, update_attributes=False, update_photos=False,
        )
        self.assertEqual(len(jobs.get_facts(self.database, self.product_id)), 1)
    def test_manual_decision_survives_rerun_and_reopen(self) -> None:
        self.save(
            source("sulpak", "full_sku", [("Цвет", "Серый")]),
            source("mechta", "full_sku", [("Цвет", "Белый")]),
        )
        jobs.save_manual_decision(self.database, self.product_id, "color", "графит", "", "Проверено сотрудником")
        jobs.resolve_product(self.database, self.product_id)
        jobs.initialize(self.database)
        jobs.resolve_product(self.database, self.product_id)
        item = self.resolved("color")
        self.assertEqual((item["selected_value"], item["status"], item["selected_source"]), ("графит", "manual", "manual"))

    def test_export_contains_three_required_views(self) -> None:
        self.save(source("lg", "base_model", [("Цвет", "Белый")]))
        from openpyxl import load_workbook
        from io import BytesIO
        book = load_workbook(BytesIO(export_batch(self.database, "b1")), read_only=True)
        try:
            self.assertIn("Паровые шкафы", book.sheetnames)
            self.assertIn("Проверка источников", book.sheetnames)
            self.assertIn("Источники", book.sheetnames)
        finally:
            book.close()


class WorkerBudgetTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.database = Path(temporary.name) / "batches.sqlite3"

    def test_base_model_then_full_sku_fallback_and_stop(self) -> None:
        product = seed_product(self.database)
        jobs.enqueue(self.database, product, [1, 3])
        lg = StaticAdapter(source("lg", "base_model", [("Цвет", "Белый")]))
        sulpak = StaticAdapter(source("sulpak", "full_sku", [("Цвет", "Серый")]))
        mechta = StaticAdapter(source("mechta", "full_sku", [("Цвет", "Серый")]))
        self.assertTrue(worker.run_once(self.database, lambda: (lg, sulpak, mechta), clock=lambda: 0))
        self.assertEqual(sulpak.calls[0][0], "S3WER.ALWPCOM")
        self.assertEqual(mechta.calls[0][0], "S3WER.ALWPCOM")
        self.assertEqual(jobs.list_jobs(self.database, product)[0]["status"], "done")
        self.assertIn("двумя поставщиками", jobs.identification_status(jobs.get_source_pages(self.database, product)))

    def test_unavailable_source_does_not_exceed_fallback_budget(self) -> None:
        product = seed_product(self.database)
        jobs.enqueue(self.database, product, [1])
        clock = FakeClock()
        lg = StaticAdapter(source("lg", "base_model", []))
        sulpak = StaticAdapter(source("sulpak", "unknown", [], error="таймаут"), clock, advance=31)
        mechta = StaticAdapter(source("mechta", "full_sku", []))
        worker.run_once(self.database, lambda: (lg, sulpak, mechta), clock=clock)
        self.assertEqual(len(mechta.calls), 0)
        pages = {item["source_key"]: item for item in jobs.get_source_pages(self.database, product)}
        self.assertIn("бюджет", pages["mechta"]["error"])
        self.assertLessEqual(clock.value, 60)

    def test_fallback_factory_not_started_for_other_brands(self) -> None:
        for brand in ("Samsung", "Apple"):
            database = self.database.with_name(f"{brand}.sqlite3")
            product = seed_product(database, brand=brand, sku="MODEL.123")
            connection = sqlite3.connect(database)
            try:
                connection.execute(
                    "INSERT INTO search_jobs (id,product_id,stages_json,status,message,created_at,updated_at) VALUES (?,?,?,?,?,?,?)",
                    (f"job-{brand}", product, "[1]", "queued", "", "2026-01-01", "2026-01-01"),
                )
                connection.commit()
            finally:
                connection.close()
            called = False
            def factory():
                nonlocal called
                called = True
                raise AssertionError("fallback must not start")
            worker.run_once(database, factory, clock=lambda: 0)
            self.assertFalse(called)
            self.assertEqual(jobs.list_jobs(database, product)[0]["status"], "error")


class MultiSourceWebTests(unittest.TestCase):
    def test_product_page_manual_choice_and_export(self) -> None:
        temporary = TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        root = Path(temporary.name)
        database = root / "batches.sqlite3"
        product = seed_product(database)
        for document in (
            source("lg", "base_model", [("Цвет", "Белый")]),
            source("sulpak", "full_sku", [("Цвет", "Серый")]),
            source("mechta", "full_sku", [("Цвет", "Белый")]),
        ):
            jobs.save_source_document(database, product, document)
        jobs.resolve_product(database, product)
        with TestClient(create_app(root)) as client:
            page = client.get(f"/products/{product}")
            self.assertEqual(page.status_code, 200)
            self.assertIn("S3WER.ALWPCOM", page.text)
            self.assertIn("S3WER", page.text)
            self.assertIn("Нужна проверка", page.text)
            self.assertIn(SULPAK_URL, page.text)
            response = client.post(
                f"/products/{product}/attributes/color/decision",
                data={"value": "графит", "unit": "", "reason": "Проверено вручную"},
                follow_redirects=False,
            )
            self.assertEqual(response.status_code, 303)
            exported = client.get("/batches/b1/export.xlsx")
            self.assertEqual(exported.status_code, 200)
            self.assertTrue(exported.content.startswith(b"PK"))
        self.assertEqual(jobs.get_manual_decisions(database, product)["color"]["selected_value"], "графит")

if __name__ == "__main__":
    unittest.main()
