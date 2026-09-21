"""Offline fixtures for regional LG, manuals, Russian display and photo candidates."""
from io import BytesIO
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory
import unittest
from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from product_tool import jobs
from product_tool.adapters.common import PhotoCandidate, ProductDocument, RawAttribute, SourceDocument
from product_tool.adapters.lg import LGRUAdapter, extract_lg_ru_attributes, extract_lg_ru_description
from product_tool.adapters.supplier import extract_photo_candidates
from product_tool.display import display_name_ru, display_value
from product_tool.exporter import export_batch
from product_tool.normalization import normalize_fact
from product_tool.web import create_app

RU_URL="https://www.lg.com/ru/laundry/lg-S3WER"
SUPPORT="https://www.lg.com/ru/support/manuals?csSalesCode=S3RERB.ALBPCOM"
PDF_RU="https://gscs-b2c.lge.com/open/downloadFile?fileId=ru123"

RU_HTML=f'''<html><body><h1>LG Styler S3WER</h1>
<div id="pdp_spec"><div class="tech-spacs"><div class="tech-spacs-title">Основные</div><div class="tech-spacs-contents">
<dl><dt>Тип дисплея</dt><dd>Сенсорный</dd></dl><dl><dt>Функция Smart Diagnosis</dt><dd>●</dd></dl><dl><dt>Декоративная линия</dt><dd>-</dd></dl>
</div></div></div>
<div id="overview"><div class="text-block"><div class="title"><h2>Освежает одежду</h2></div><div class="copy"><p>Пар удаляет запахи и помогает ухаживать за тканями.</p></div></div>
<div class="text-block"><h2>Купить в интернет-магазине</h2><p>Где купить и отзывы покупателей.</p></div></div>
<a href="/ru/support/manuals?csSalesCode=S3RERB.ALBPCOM">Руководства S3RERB.ALBPCOM</a>
</body></html>'''
SUPPORT_HTML=f'''<html><body><div class="manual"><span>Русский</span><span>2025-02-03</span><span>12.5 MB</span><a title="Руководство пользователя" href="{PDF_RU}">Скачать</a></div>
<div class="manual"><span>English</span><a href="https://gscs-b2c.lge.com/open/downloadFile?fileId=en123">Download</a></div></body></html>'''

class Response:
    def __init__(self,url,text,status=200): self.url,self.text,self.status_code=url,text,status; self.content=text.encode()
    def raise_for_status(self):
        if self.status_code>=400: raise RuntimeError(self.status_code)
class Session:
    def __init__(self,pages): self.pages=pages; self.headers={}; self.calls=[]
    def get(self,url,timeout): self.calls.append(url); return Response(url,self.pages.get(url,""),200 if url in self.pages else 404)

def seed(db: Path):
    jobs.initialize(db); con=sqlite3.connect(db)
    try:
        con.execute("INSERT INTO batches VALUES ('b1','fixture.xlsx','Товары','{}','2026-01-01')")
        cur=con.execute("INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES ('b1',2,'LG Styler','LG','S3WER.ALWPCOM','','Паровые шкафы',0,'[]','{}')")
        con.commit(); return cur.lastrowid
    finally: con.close()

def doc(key,facts,match="base_model",photos=None):
    names={"lg_kz":"LG Казахстан","lg_ru":"LG Россия","sulpak":"Sulpak","mechta":"Mechta"}
    return SourceDocument(key,names[key],f"https://example.test/{key}",found_model="S3WER",match_level=match,evidence="fixture",attributes=[RawAttribute(*x) for x in facts],photo_candidates=photos or [])

class ClassificationTests(unittest.TestCase):
    def setUp(self):
        tmp=TemporaryDirectory(); self.addCleanup(tmp.cleanup); self.root=Path(tmp.name); self.db=self.root/"batches.sqlite3"; self.product=seed(self.db)
    def save(self,*docs):
        for item in docs: jobs.save_source_document(self.db,self.product,item)
        jobs.resolve_product(self.db,self.product)
    def test_base_only_is_not_conflict_and_counted_separately(self):
        self.save(doc("lg_kz",[("Drip Tray (Qty)","2")]))
        result=jobs.get_resolved(self.db,self.product)[0]; counts=jobs.result_counts(self.db,self.product)
        self.assertEqual(result["status"],"official_base_only"); self.assertFalse(result["conflict"]); self.assertFalse(result["full_sku_confirmed"])
        self.assertEqual(counts,{"conflicts":0,"official_base_only":1,"supplier_confirmed":0})
    def test_two_different_values_are_real_conflict(self):
        self.save(doc("lg_kz",[("Цвет","Белый")]),doc("lg_ru",[("Цвет","Серый")]))
        result=jobs.get_resolved(self.db,self.product)[0]
        self.assertTrue(result["conflict"]); self.assertEqual(result["status"],"official_regions_conflict"); self.assertEqual(jobs.result_counts(self.db,self.product)["conflicts"],1)
    def test_regions_match_but_do_not_confirm_full_sku(self):
        self.save(doc("lg_kz",[("Цвет","Белый")]),doc("lg_ru",[("Цвет корпуса","Белый")]))
        result=jobs.get_resolved(self.db,self.product)[0]
        self.assertEqual(result["status"],"official_base_only"); self.assertFalse(result["full_sku_confirmed"])
    def test_supplier_full_sku_count_and_manual_priority(self):
        self.save(doc("lg_kz",[("Цвет","Белый")]),doc("sulpak",[("Цвет","Серый")],"full_sku"))
        self.assertEqual(jobs.result_counts(self.db,self.product)["supplier_confirmed"],1)
        jobs.save_manual_decision(self.db,self.product,"color","графит",reason="Проверено")
        jobs.resolve_product(self.db,self.product)
        self.assertEqual(jobs.get_resolved(self.db,self.product)[0]["status"],"manual")

class DisplayAndRussiaTests(unittest.TestCase):
    def test_russian_names_and_values(self):
        self.assertEqual(display_name_ru("pants_hanger_qty"),"Количество вешалок для брюк")
        self.assertEqual(display_name_ru("тип_управления"),"Тип управления")
        self.assertEqual(display_value("true","bool"),"Да")
        self.assertEqual(display_value("w=445;h=1850;d=585","mm"),"445 × 1850 × 585 мм (Ш × В × Г)")
    def test_ru_hidden_dt_dd_and_contextual_boolean(self):
        facts=extract_lg_ru_attributes(BeautifulSoup(RU_HTML,"html.parser")); by={x.name:x.value for x in facts}
        self.assertEqual(by["Тип дисплея"],"Сенсорный"); self.assertEqual(normalize_fact(RawAttribute("Функция Smart Diagnosis","●")).normalized_value,"true")
        self.assertEqual(normalize_fact(RawAttribute("Smart Diagnosis™","● версия 3.0")).normalized_value,"true")
        ordinary=normalize_fact(RawAttribute("Декоративная линия","-")); self.assertEqual(ordinary.normalized_value,"-"); self.assertNotEqual(ordinary.unit,"bool")
    def test_overview_keeps_content_and_removes_commerce(self):
        text=extract_lg_ru_description(BeautifulSoup(RU_HTML,"html.parser"))
        self.assertIn("Освежает одежду",text); self.assertIn("Пар удаляет запахи",text); self.assertNotIn("Купить",text); self.assertNotIn("Где купить",text)

class DocumentTests(unittest.TestCase):
    def setUp(self):
        tmp=TemporaryDirectory(); self.addCleanup(tmp.cleanup); self.db=Path(tmp.name)/"batches.sqlite3"; self.product=seed(self.db)
    def test_official_relation_allows_only_russian_document(self):
        session=Session({SUPPORT:SUPPORT_HTML}); adapter=LGRUAdapter(session,clock=lambda:0)
        page=SourceDocument("lg_ru","LG Россия",RU_URL,html=RU_HTML)
        documents,error=adapter.find_documents(page,"S3WER",deadline=20)
        self.assertFalse(error); self.assertEqual(len(documents),1); self.assertEqual(documents[0].support_model,"S3RERB.AL BPCOM".replace(" ","")); self.assertEqual(documents[0].direct_url,PDF_RU); self.assertEqual(documents[0].language,"Русский")
    def test_related_model_without_official_evidence_is_rejected(self):
        adapter=LGRUAdapter(Session({}),clock=lambda:0); documents,error=adapter.find_documents(SourceDocument("lg_ru","LG Россия",RU_URL,html="<html>S3WER</html>"),"S3WER",deadline=20)
        self.assertEqual(documents,[]); self.assertIn("связь",error)
    def test_repeated_document_save_has_no_duplicates(self):
        jobs.save_source_document(self.db,self.product,doc("lg_ru",[]))
        d=ProductDocument("Руководство","Русский","2025-01-01","1 МБ",PDF_RU,SUPPORT,"S3WER","S3RERB",RU_URL,True)
        jobs.save_documents(self.db,self.product,"lg_ru",[d]); jobs.save_documents(self.db,self.product,"lg_ru",[d])
        self.assertEqual(len(jobs.get_documents(self.db,self.product)),1)

class PhotoAndExportTests(unittest.TestCase):
    def setUp(self):
        tmp=TemporaryDirectory(); self.addCleanup(tmp.cleanup); self.root=Path(tmp.name); self.db=self.root/"batches.sqlite3"; self.product=seed(self.db)
    def test_sulpak_garbage_excluded_and_sizes_deduplicated(self):
        html='''<img src="/Banners/cashback.webp"><img src="/wwwroot/img/city.webp"><img src="/cms/cms/Photo/img_a_160.webp"><img src="/cms/cms/Photo/img_a_320.webp"><img src="/cms/cms/Photo/img_a.webp"><img src="/cms/cms/Photo/sulpak_logo.svg">'''
        items=extract_photo_candidates(BeautifulSoup(html,"html.parser"),"https://www.sulpak.kz/p",source_key="sulpak")
        gallery=[x for x in items if x.kind=="product_gallery"]; excluded=[x for x in items if x.kind=="excluded"]
        self.assertEqual(len(gallery),1); self.assertTrue(gallery[0].url.endswith("img_a.webp")); self.assertGreaterEqual(len(excluded),3)
    def test_manual_selection_survives_refresh(self):
        photos=[PhotoCandidate("https://img/one.jpg","one","product_gallery"),PhotoCandidate("https://img/two.jpg","two","product_gallery")]
        jobs.save_source_document(self.db,self.product,doc("sulpak",[],"full_sku",photos))
        jobs.set_photo_selection(self.db,self.product,["two"])
        jobs.save_source_document(self.db,self.product,doc("sulpak",[],"full_sku",photos))
        selected=[x["asset_key"] for x in jobs.get_photo_candidates(self.db,self.product) if x["selected"]]
        self.assertEqual(selected,["two"])
    def test_ui_hides_technical_keys_and_export_has_new_sheets(self):
        jobs.save_source_document(self.db,self.product,doc("lg_kz",[("Drip Tray (Qty)","2")],"base_model"))
        # selected official photo and Russian manual
        photos=[PhotoCandidate("https://www.lg.com/photo.jpg","official","product_gallery")]
        jobs.save_source_document(self.db,self.product,doc("lg_ru",[("Тип дисплея","Сенсорный")],photos=photos))
        jobs.save_documents(self.db,self.product,"lg_ru",[ProductDocument("Руководство","Русский","2025-01-01","1 МБ",PDF_RU,SUPPORT,"S3WER","S3RERB",RU_URL,True)])
        jobs.resolve_product(self.db,self.product)
        with TestClient(create_app(self.root)) as client:
            page=client.get(f"/products/{self.product}").text
            self.assertIn("Количество поддонов для сбора воды",page); self.assertNotIn("drip_tray_qty",page); self.assertIn("Изменить итог",page)
        book=load_workbook(BytesIO(export_batch(self.db,"b1")),read_only=True)
        try:
            self.assertIn("Инструкции",book.sheetnames); self.assertIn("Фотографии",book.sheetnames)
            headers=[c.value for c in next(book["Паровые шкафы"].iter_rows())]; self.assertIn("Количество поддонов для сбора воды",headers); self.assertNotIn("drip_tray_qty",headers)
            photo_rows=list(book["Фотографии"].iter_rows(values_only=True)); self.assertEqual(len(photo_rows),2)
        finally: book.close()

if __name__=="__main__": unittest.main()