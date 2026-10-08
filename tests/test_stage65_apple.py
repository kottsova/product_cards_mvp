"""Offline regressions on real official bodies and identity/scoping counterexamples."""
import unittest,json,tempfile
from pathlib import Path
from unittest.mock import patch
from product_tool.apple_identity import identify,parse_model,classification
from product_tool.adapters.apple import AppleAdapter,parse_page
from product_tool.resolution import resolve_attributes
R=Path(__file__).resolve().parents[1]/'reports/apple_stage64'

class AppleModelScopes(unittest.TestCase):
 def test_public_catalog_row_is_identity_not_specs(self):
  from product_tool.apple_catalog import catalog_descriptions
  row='None Specified   MXP93LL/A   AirPods 4 with Active Noise Cancellation    179.00   179.00'
  self.assertEqual(catalog_descriptions(row,'MXP93LL/A'),{'AirPods 4 with Active Noise Cancellation'});self.assertFalse(catalog_descriptions(row,'MXP63LL/A'));self.assertFalse(catalog_descriptions(row,'A3056'))
 def model(self,file,key,config=None):return parse_model((R/file).read_text(encoding='utf-8'),'https://support.apple.com/en-us/121029',key,config or {})
 def test_model_not_commercial_configuration(self):
  doc,a=self.model('more_3.html','iphone16');self.assertEqual(doc.match_level,'model_confirmed');self.assertGreater(len(doc.attributes),15)
  self.assertFalse(any(x.name in {'Объём накопителя','Цвет'} for x in doc.attributes));self.assertTrue(a['candidates'])
 def test_mac_order_and_model_wireless_runtime_share_canonical_label(self):
  d,e=parse_page((R/'official_0.html').read_text(encoding='utf-8'),'https://www.apple.com/shop/product/fw123ll/a','FW123LL/A')
  class Response:
   ok=True;url='https://support.apple.com/en-us/122209';text=(R/'official_1.html').read_text(encoding='utf-8')
  with tempfile.TemporaryDirectory() as td,patch('product_tool.photo_metadata.inspect_saved_photo',return_value={'width':1000,'height':1000,'size_bytes':1234,'format':'JPG'}):
   adapter=AppleAdapter(fetch_log_path=Path(td)/'log.json');adapter.session=type('S',(),{'get':lambda *a,**kw:Response()})();adapter.check_manual=lambda *a:None;adapter.enrich_model(d,e,'FW123LL/A',1e12)
  name='Время беспроводной работы в интернете';self.assertIn(name,{a.name for a in d.attributes});self.assertIn(name,{a.name for a in adapter.model_document.attributes});self.assertNotIn('Время работы в интернете',{a.name for a in d.attributes})
 def test_other_generation_or_pro_rejected(self):
  d,a=self.model('official_2.html','iphone16');self.assertTrue(d.error);self.assertFalse(d.attributes)
 def test_mac_gpu_family_options_are_not_facts(self):
  d,a=self.model('official_1.html','air13m4');self.assertFalse(any(x.name=='Ядра GPU' for x in d.attributes));self.assertTrue(any('GPU' in x['value'] for x in a['candidates']))
 def test_m4_pro_is_not_base_m4(self):self.assertEqual(identify('MacBook Pro 14-inch M4 Pro'),'')
 def test_watch_46_scoped_dimensions(self):
  d,a=self.model('more_1.html','watch10',{'case_size':'46','material':'aluminum','connectivity':'GPS'})
  v={x.name:x.value for x in d.attributes};self.assertEqual(float(v['Высота товара, мм']),46);self.assertEqual(float(v['Ширина товара, мм']),39);self.assertAlmostEqual(float(v['Вес товара, кг']),.0364)
 def test_watch_42_does_not_inherit_46(self):
  d,a=self.model('more_1.html','watch10',{'case_size':'42','material':'aluminum','connectivity':'GPS'})
  v={x.name:x.value for x in d.attributes};self.assertEqual(float(v['Высота товара, мм']),42);self.assertEqual(float(v['Ширина товара, мм']),36);self.assertAlmostEqual(float(v['Вес товара, кг']),.030)
 def test_watch_unresolved_size_has_no_dimensions(self):
  d,a=self.model('more_1.html','watch10');self.assertFalse(any('товара' in x.name for x in d.attributes))
 def test_ipad_wifi_not_cellular_weight_or_runtime(self):
  d,a=self.model('more_0.html','ipada16',{'connectivity':'Wi-Fi'});v={x.name:x.value for x in d.attributes};self.assertAlmostEqual(float(v['Вес товара, кг']),.477);self.assertIn('Время работы в интернете по Wi-Fi',v);self.assertIn('Время воспроизведения видео',v);self.assertNotIn('Время работы через мобильную сеть',v)
 def test_ipad_cellular_separate_runtime(self):
  d,a=self.model('more_0.html','ipada16',{'connectivity':'Wi-Fi + Cellular'});v={x.name:x.value for x in d.attributes};self.assertAlmostEqual(float(v['Вес товара, кг']),.481);self.assertIn('Время работы через мобильную сеть',v)
 def test_iphone_three_runtime_modes(self):
  d,a=self.model('more_3.html','iphone16');v={x.name:x.value for x in d.attributes};self.assertIn('22 hours',v['Время воспроизведения видео']);self.assertIn('18 hours',v['Время потокового воспроизведения видео']);self.assertIn('80 hours',v['Время воспроизведения аудио'])
 def test_watch_battery_modes_and_fast_charge(self):
  d,a=self.model('more_1.html','watch10');v={x.name:x.value for x in d.attributes};self.assertIn('18 hours',v['Время работы в обычном режиме']);self.assertIn('36 hours',v['Время работы в режиме энергосбережения']);self.assertIn('Быстрая зарядка до 80%',v)
 def test_airpods_anc_exact_model_and_case_runtime(self):
  d,a=self.model('more_4.html','airpods4anc');self.assertFalse(d.error);self.assertTrue(any('с кейсом' in x.name for x in d.attributes));self.assertEqual(identify('AirPods 5 with Active Noise Cancellation'),'')
 def test_model_resolution_does_not_claim_full_sku(self):
  r=resolve_attributes([dict(source_key='apple_model',normalized_name='chipset',normalized_value='A18',unit='')],[dict(source_key='apple_model',match_level='model_confirmed',error='')])[0];self.assertEqual(r.status,'model_confirmed_official');self.assertFalse(r.full_sku_confirmed)
 def test_category_aware_gpu_and_weight(self):
  self.assertEqual(classification('gpu','mac'),'configuration-sensitive');self.assertEqual(classification('gpu','iphone'),'model-stable');self.assertEqual(classification('weight','ipad'),'configuration-sensitive')
 def test_same_a_number_cannot_confirm_commercial_configuration(self):
  d,e=parse_page((R/'official_3.html').read_text(encoding='utf-8'),'https://support.apple.com/en-us/108044','A3081');self.assertFalse(d.attributes);self.assertFalse(d.photo_candidates)
 def test_selection_storage_from_exact_record_not_family_capacity(self):
  b=(R/'official_5.html').read_text(encoding='utf-8');d,e=parse_page(b,'https://www.apple.com/shop/buy-iphone/iphone-16','MYAP3LL/A')
  class Response:ok=False;marker='offline';status_code=0
  with tempfile.TemporaryDirectory() as td:
   adapter=AppleAdapter(fetch_log_path=Path(td)/'log.json');adapter.session=type('S',(),{'get':lambda *a,**kw:Response()})();adapter.check_manual=lambda *a:None
   adapter.enrich_model(d,e,'MYAP3LL/A',1e12)
  self.assertEqual(e['configuration_fields']['storage'],'128GB');self.assertEqual(e['configuration_fields']['color'],'black');self.assertEqual(d.match_level,'full_sku')
 def test_different_storage_records_are_not_confirmed(self):
  b=(R/'official_5.html').read_text(encoding='utf-8').replace('128gb-black-unlocked','256gb-black-unlocked');d,e=parse_page(b,'https://www.apple.com/shop/buy-iphone/iphone-16','MYAP3LL/A')
  class Response:ok=False;marker='offline';status_code=0
  with tempfile.TemporaryDirectory() as td:
   adapter=AppleAdapter(fetch_log_path=Path(td)/'log.json');adapter.session=type('S',(),{'get':lambda *a,**kw:Response()})();adapter.check_manual=lambda *a:None;adapter.enrich_model(d,e,'MYAP3LL/A',1e12)
  self.assertNotIn('storage',e['configuration_fields']);self.assertFalse(e['configuration_complete'])

class AppleSharedPipeline(unittest.TestCase):
 def setUp(self):
  from product_tool import jobs,storage
  self.temp=tempfile.TemporaryDirectory();self.db=Path(self.temp.name)/'test.sqlite3';jobs.initialize(self.db)
  with storage._connection(self.db) as c:
   c.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('b','Apple','Products','{}',storage._now()))
   c.execute('INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (?,?,?,?,?,?,?,?,?,?)',('b',1,'iPhone 16','Apple','MYAP3LL/A','','phones',0,'[]','{}'))
  self.pid=storage.get_batch(self.db,'b')['products'][0]['id']
 def tearDown(self):self.temp.cleanup()
 def ready(self,candidates=None):
  from product_tool import jobs,card_evidence
  from product_tool.adapters.common import SourceDocument,RawAttribute,PhotoCandidate
  d,a=parse_model((R/'more_3.html').read_text(encoding='utf-8'),'https://support.apple.com/en-us/121029','iphone16',{})
  jobs.save_source_document(self.db,self.pid,d)
  jobs.save_source_document(self.db,self.pid,SourceDocument('apple','Apple','https://www.apple.com/shop/buy-iphone/iphone-16',match_level='full_sku',attributes=[RawAttribute('Цвет','Black'),RawAttribute('Объём накопителя','128GB')],photo_candidates=[PhotoCandidate('https://store.storeimages.cdn-apple.com/image.jpg','black')]))
  jobs.resolve_product(self.db,self.pid);jobs.set_photo_selection(self.db,self.pid,['black'],mode='exact')
  card_evidence.save(self.db,self.pid,'apple',dict(identity={'model':'model_confirmed','configuration':'exact_order_record','variant':'exact_model_color'},model_key='iphone16',configuration_complete=True,configuration_fields={'color':'Black','storage':'128GB'},exact_photo_assets=['black'],manual_status='Не проверена',configuration_candidates=[],photo_candidates=candidates or []))
 def test_manual_and_sku_absent_from_tech_specs_do_not_block(self):
  from product_tool import apple_pipeline,readiness
  self.ready();r=apple_pipeline.card_readiness(self.db,self.pid);self.assertEqual(r['verdict'],'export_ready');self.assertIn('manual_unverified',r['advisory_gaps']);self.assertGreater(r['model_specs'],10);self.assertEqual(readiness.card_readiness(self.db,self.pid)['verdict'],'export_ready')
 def test_photo_candidate_sheet_only_when_nonempty(self):
  from product_tool import exporter
  from openpyxl import load_workbook
  from io import BytesIO
  self.ready();b=load_workbook(BytesIO(exporter.export_batch(self.db,'b')));self.assertNotIn('Фото-кандидаты',b.sheetnames)
  self.ready([dict(url='https://store.storeimages.cdn-apple.com/family.jpg',reason='Other color',model_relation='exact_model',color_relation='different_color',width=1000,height=1000,file_size=1234,format='JPG')]);b=load_workbook(BytesIO(exporter.export_batch(self.db,'b')));s=b['Фото-кандидаты'];self.assertEqual(s.max_row,2);self.assertEqual(s['A2'].value,'MYAP3LL/A');self.assertEqual(s['I2'].value,1234);self.assertNotIn('family.jpg',' '.join(str(c.value) for row in b['Фотографии'] for c in row))
 def test_refurb_condition_not_model_identity(self):
  self.assertEqual(identify('Refurbished iPhone 15 Pro 128GB - Black Titanium'),'iphone15pro');self.assertEqual(identify('iPhone 15 Pro'),'iphone15pro')
 def test_configuration_missing_still_blocks_with_good_model_specs(self):
  from product_tool import apple_pipeline,card_evidence
  self.ready();e=card_evidence.load(self.db,self.pid,'apple');e['configuration_fields'].pop('storage');card_evidence.save(self.db,self.pid,'apple',e);self.assertEqual(apple_pipeline.card_readiness(self.db,self.pid)['verdict'],'not_ready')
