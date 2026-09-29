"""Stage 31: exact cable selection and unresolved charger, all ordinary-path checks offline."""
from __future__ import annotations

import gzip
import io
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

import _samsung_replay as R
from product_tool import jobs, samsung_readiness
from product_tool.coverage import classify as C, executor, planner
from product_tool.offline_guard import offline_only

ROOT=Path(__file__).resolve().parents[1]
STAGE=ROOT/"reports/source_census_2026-09-26_stage31"
CHARGER="EP-T4511XBEGEU"
CABLE="EP-DA705BBRGRU"

def saved_pages():
    directory=STAGE/"route_check/responses"
    pages={}
    for line in (directory/"index.jsonl").read_text(encoding="utf-8").splitlines():
        item=json.loads(line)
        if item["saved_as"]:
            with gzip.open(directory/item["saved_as"],"rt",encoding="utf-8",newline="") as handle:
                pages[item["url"]]=(item["status"],handle.read())
    return pages

class AccessorySelectionAndOrdinaryPath(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=TemporaryDirectory()
        units=R.catalog_units({CHARGER,CABLE})
        cls.cards=[(units[sku]["category"],sku) for sku in (CHARGER,CABLE)]
        with offline_only():
            cls.replay=R.Replay(extra_pages=saved_pages())
            cls.batch=R.run_products(Path(cls.tmp.name),cards=cls.cards,replay=cls.replay)
        cls.outcomes={o["article"]:o for o in cls.batch["outcomes"]}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_bounded_official_route_evidence(self):
        data=json.loads((STAGE/"raw/route_check_result.json").read_text(encoding="utf-8"))
        self.assertEqual(data["budget"]["spent"],5)
        self.assertEqual(data["stopped_hosts"],[])
        self.assertEqual(data["currently_listed_exact_charger_urls"],[])
        self.assertEqual(len(data["currently_listed_exact_cable_urls"]),1)
        charger,cable=data["pages"][CHARGER],data["pages"][CABLE]
        self.assertEqual(charger["identity"]["jsonld_sku"],"EP-T4511XBEGRU")
        self.assertNotIn("ept4511xbegeu",charger["product_data_codes"])
        self.assertEqual((cable["identity"]["level"],cable["identity"]["strength"],cable["identity"]["jsonld_sku"]),("full_sku","strong",CABLE))
        self.assertEqual((len(cable["specs"]),cable["photos"]["full_size"],cable["document_links"]),(5,4,[]))

    def test_planner_selects_only_the_verified_cable_and_stops_charger_before_request(self):
        units=[u for u in planner.build_plan()["units"] if u["family"]=="samsung"]
        self.assertEqual(len(units),1182)
        ready=[u for u in units if u["status"]==C.READY]
        self.assertEqual((len(ready),len({u["category"] for u in ready})),(20,20))
        cable=next(u for u in units if u["seller_sku"]==CABLE)
        charger=next(u for u in units if u["seller_sku"]==CHARGER)
        self.assertEqual(cable["status"],C.READY)
        self.assertEqual((charger["status"],charger["reason"]),(C.MANUAL,"not_selected_one_card_per_category"))
        with TemporaryDirectory() as tmp:
            replay=R.Replay(extra_pages=saved_pages())
            checkpoint=executor.run_units([charger],workdir=Path(tmp),checkpoint_path=Path(tmp)/"checkpoint.json",factories=executor.RunFactories(session=replay))
            entry=next(iter(checkpoint["units"].values()))
            self.assertEqual((entry["outcome"],entry["stopped_before_run_once"]),( "stopped",True))
            self.assertEqual(replay.calls,[])

    def test_job_export_and_page_scoped_instruction_gap(self):
        charger,cable=self.outcomes[CHARGER],self.outcomes[CABLE]
        self.assertEqual((charger["status"],cable["status"]),("needs_review","done"))
        db=self.batch["database"]
        charger_ready=samsung_readiness.card_readiness(db,charger["product_id"])
        cable_ready=samsung_readiness.card_readiness(db,cable["product_id"])
        self.assertEqual((charger_ready["verdict"],cable_ready["verdict"]),("not_ready","export_ready_with_gaps"))
        self.assertEqual((charger_ready["page_match_level"],cable_ready["page_match_level"]),("unknown","full_sku"))
        self.assertIn("instruction_no_exact_page",charger_ready["blocking_gaps"])
        self.assertNotIn("instruction_missing",charger_ready["blocking_gaps"])
        self.assertIn("не проверена",charger_ready["gap_reasons"]["instruction_no_exact_page"])
        self.assertEqual(len([f for f in jobs.get_facts(db,cable["product_id"]) if f["source_key"]=="samsung"]),5)
        self.assertEqual(len([p for p in jobs.get_photo_candidates(db,cable["product_id"]) if p["source_key"]=="samsung" and p["selected"]]),4)
        self.assertEqual(jobs.get_documents(db,cable["product_id"]),[])
        self.assertIn("instruction_missing",cable_ready["blocking_gaps"])
        self.assertEqual(self.replay.refused,[])
        self.assertEqual(self.batch["dealer_session"].calls,[])
        book=load_workbook(io.BytesIO(self.batch["export"]),read_only=True)
        self.assertIn("Готовность Samsung",book.sheetnames)
        self.assertIn(self.cards[1][0],book.sheetnames)
        self.assertIn(self.cards[0][0],book.sheetnames)

if __name__=="__main__":
    unittest.main()
