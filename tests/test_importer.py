"""Acceptance checks for the Excel preview, using only synthetic workbooks."""

from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook

from product_tool.importer import ColumnMapping, choose_sheet, preview_row, preview_xlsx


class ImporterAcceptanceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._temporary_directory = TemporaryDirectory()
        cls.addClassCleanup(cls._temporary_directory.cleanup)
        cls.sample = Path(cls._temporary_directory.name) / "synthetic_products.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Товары"
        sheet.append(["Группа товаров"])
        sheet.append(["Предмет", "Бренд", "Наименование", "Артикул продавца", "Модель"])
        sheet.append([None, None, None, None, None, "Размеры упаковки"])
        sheet.append(["Стиральная машина", "LG", "Стиральная машина LG F2V5HG1W", "F2V5HG1W", "F2V5HG1W"])
        sheet.append(["Стирально-сушильная машина", "LG", "Стиральная машина LG F4V5VG0W", "F4V5VG0W", "F4V5VG0W"])
        sheet.append(["Паровой шкаф", "LG", "Паровой шкаф LG S3WER", "S3WER.ALWPCOM", "S3WER"])
        sheet.append([None, None, "Смартфон Samsung SM-A566ELIHINS A56 12+256GB  (пример)"])
        sheet.append([None, None, "Ноутбук Lenovo 83M0003YRK"])
        sheet.append([None, None, "Ноутбук Apple MGN63RU"])
        sheet.append([None, None, "Телевизор Samsung UE55U8000FUXCE"])
        sheet.append([None, None, "Наушники Apple MGYM3RU/A"])
        workbook.save(cls.sample)
        workbook.close()

    def test_mixed_workbook_preserves_search_codes_and_flags_inferences(self) -> None:
        preview = preview_xlsx(self.sample)

        self.assertEqual(len(preview.products), 8)
        self.assertEqual(preview.mapping, ColumnMapping(
            brand=2, search_code=4, fallback_code=5,
            name=3, category=1, header_row=2,
        ))
        self.assertEqual(sum(preview.category_counts.values()), 8)
        self.assertEqual(len(preview.category_counts), 7)
        self.assertEqual([item.row_number for item in preview.products], list(range(4, 12)))
        self.assertEqual([item.needs_confirmation for item in preview.products],
                         [False] * 3 + [True] * 5)

        steam_cabinet = preview.products[2]
        self.assertEqual(steam_cabinet.search_code, "S3WER.ALWPCOM")
        self.assertEqual(steam_cabinet.alternate_code, "S3WER")
        self.assertEqual(preview.products[-1].search_code, "MGYM3RU/A")
        self.assertEqual(preview.products[3].search_code, "SM-A566ELIHINS")
        self.assertIn("12+256GB  (", preview.products[3].original_values[2])

    def test_manual_column_numbers_on_synthetic_workbook(self) -> None:
        preview = preview_xlsx(self.sample, column_overrides={
            "category": 1, "brand": 2, "name": 3,
            "search_code": 4, "fallback_code": 5,
        }, header_row=2)
        self.assertEqual(len(preview.products), 8)
        self.assertEqual(preview.products[0].brand_source, "column")
        self.assertEqual(preview.products[3].brand_source, "name")

    def test_name_only_row_keeps_exact_variant_code(self) -> None:
        product = preview_row(
            1, ("Телевизор Samsung UE55U8000FUXCE",),
            ColumnMapping(name=1),
        )
        self.assertIsNotNone(product)
        self.assertEqual(product.search_code, "UE55U8000FUXCE")
        self.assertEqual(product.category, "Телевизор")
        self.assertTrue(product.needs_confirmation)

    def test_headerless_sheet_can_be_selected_for_manual_mapping(self) -> None:
        class Sheet:
            title = "Данные"

            def iter_rows(self, **_kwargs):
                return iter((("LG", "TW4V7EB1W"),))

        class Workbook:
            worksheets = [Sheet()]
            sheetnames = ["Данные"]

            def __getitem__(self, _name):
                return self.worksheets[0]

        sheet, detected = choose_sheet(Workbook(), None, allow_headerless=True)
        self.assertEqual(sheet.title, "Данные")
        self.assertIsNone(detected)


if __name__ == "__main__":
    unittest.main()
