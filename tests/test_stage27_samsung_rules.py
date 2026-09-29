"""Stage 27: the owner's decisions on Samsung instructions and photos, the recorded Stage 8.6 evidence, the planner's one-product-per-category Samsung route, and the ordinary
create-process-export path, all offline on the saved responses (tests/_samsung_replay.py). No socket is opened."""
from __future__ import annotations

import io
import json
import unittest
from collections import Counter
from pathlib import Path
from tempfile import TemporaryDirectory

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import _samsung_replay as R
from product_tool import card_evidence, jobs, samsung_readiness
from product_tool.adapters.common import PhotoCandidate
from product_tool.adapters.samsung import instruction_acceptance, page_model_data, photo_variant_binding
from product_tool.adapters.samsung_source import RECORDED_DOCUMENTS, default_samsung_adapter, load_recorded_documents
from product_tool.coverage import classify as C
from product_tool.coverage import executor, planner
from product_tool.web import create_app

ROOT = Path(__file__).resolve().parents[1]
SELECTED = {
    "Телевизоры": "QE48S85HAEXCE", "Смартфоны": "SM-A376EZAGINS", "Пылесосы": "VC18M21D0VG/EV", "Стиральные машины": "WD10T754CBX/LD", "Микроволновые печи": "MS23K3614AK/BW", "Холодильники": "RB31FERNDSA",
    "Духовые шкафы": "NV7B4120ZAS/WT", "Сплит-системы": "AR80F09CABWNER", "Варочные панели": "NZ64T3506AK/WT", "Мониторы": "LS24D300GAIXCI", "Саундбары": "HW-Q800D",
    "Машины посудомоечные": "DW60M5050BB/WT", "Сушильные машины": "DV16DG8600BVLD", "Колонки": "MX-ST50B", "Роботы-пылесосы": "VR50T95735W/EV",
}
SELECTED.update({
    "Внутренние SSD-накопители": "MZ-77E500BW",
    "Flash-накопители": "MUF-64BE3/APC",
    "Внешние SSD-накопители": "MU-PC1T0H/WW",
    "Карты памяти": "MB-MC64HARU",
    "Кабели": "EP-DA705BBRGRU",
})

_PLAN = None


def plan():
    global _PLAN
    if _PLAN is None:
        _PLAN = planner.build_plan()
    return _PLAN


def facts_of(**over):
    base = {"content_status": "readable", "russian_by_text": True, "regulatory": None, "instruction_by_content": True, "conflicting_models": [], "tied_by_official_page": "exact_page",
            "names_catalog_model": {"exact": False, "family_mask": [], "tokens": ["X"]}, "names_link_model": {"model_name": "L1", "exact": False, "family_mask": []}}
    base.update(over)
    return base


class InstructionAcceptanceRule(unittest.TestCase):
    """Owner decision 1 and 2 (Stage 27)."""

    PAGE = {"model_code": "QE1", "model_name": "L1", "source": "test"}

    def test_a_russian_instruction_that_names_no_model_is_accepted_with_the_mark_only_on_an_exact_page(self):
        verdict = instruction_acceptance(facts_of(), self.PAGE, "QE1", "L1")
        self.assertEqual((verdict["basis"], verdict["accepted"]), ("exact_page_link_only", True))
        self.assertEqual(verdict["note"], "связь с моделью подтверждена точной официальной страницей; точный код в PDF не назван")
        for weaker in ("page_with_code_in_text_only", "base_model_page", "none"):
            verdict = instruction_acceptance(facts_of(tied_by_official_page=weaker), self.PAGE, "QE1", "L1")
            self.assertEqual((verdict["basis"], verdict["accepted"]), ("page_tie_not_exact", False), weaker)

    def test_a_conflicting_model_a_declaration_or_a_non_russian_text_is_never_accepted(self):
        self.assertFalse(instruction_acceptance(facts_of(conflicting_models=["ZZ9"]), self.PAGE, "QE1", "L1")["accepted"])
        self.assertFalse(instruction_acceptance(facts_of(regulatory="declaration"), self.PAGE, "QE1", "L1")["accepted"])
        self.assertFalse(instruction_acceptance(facts_of(russian_by_text=False), self.PAGE, "QE1", "L1")["accepted"])

    def test_a_file_without_extractable_text_is_a_manual_check(self):
        verdict = instruction_acceptance({"content_status": "text_not_extractable", "russian_by_text": None}, self.PAGE, "QE1", "L1")
        self.assertEqual((verdict["basis"], verdict["accepted"]), ("manual_check_text_not_extractable", False))

    def test_a_family_mask_is_accepted_and_the_mark_says_the_exact_code_is_not_named(self):
        verdict = instruction_acceptance(facts_of(names_catalog_model={"exact": False, "family_mask": ["AB12*"], "tokens": ["X"]}), {}, "AB12CD", "L1")
        self.assertEqual((verdict["basis"], verdict["accepted"]), ("family_mask_in_pdf", True))
        self.assertIn("точный код в PDF не назван", verdict["note"])

    def test_the_exact_code_needs_no_mark(self):
        verdict = instruction_acceptance(facts_of(names_catalog_model={"exact": True, "family_mask": [], "tokens": ["X"]}), {}, "AB12CD", "L1")
        self.assertEqual((verdict["basis"], verdict["accepted"], verdict["note"]), ("exact_code_in_pdf", True, ""))

    def test_another_code_in_the_pdf_needs_the_pages_own_data_and_never_a_general_rule(self):
        named = facts_of(names_catalog_model={"exact": False, "family_mask": [], "tokens": ["VC1"]}, names_link_model={"model_name": "SC1", "exact": True, "family_mask": []})
        page = {"model_code": "VC1/EV", "model_name": "SC1", "source": "test"}
        ok = instruction_acceptance(named, page, "VC1/EV", "SC1")
        self.assertEqual((ok["basis"], ok["accepted"], ok["code_relation"]["declared_by_page_data"]), ("code_relation_declared_by_page_data", True, True))
        for label, data in (("no data", {}), ("another model name", {**page, "model_name": "SC2"}), ("another model code", {**page, "model_code": "VC9/EV"})):
            verdict = instruction_acceptance(named, data, "VC1/EV", "SC1")
            self.assertEqual((verdict["basis"], verdict["accepted"]), ("link_code_relation_unverified", False), label)
        # the same two letters swapped with no page data behind them prove nothing either
        self.assertFalse(instruction_acceptance(named, {}, "SC1", "VC1")["accepted"])

    def test_the_pages_model_data_is_read_from_the_page_itself(self):
        for fragment, expected in (("vc18m21d0vg-ev", ("VC18M21D0VG/EV", "SC18M21D0VG")), ("ls24d300gaixci", ("LS24D300GAIXCI", "S24D300GAI")), ("qe48s85haexce", ("QE48S85HAEXCE", "QE48S85HAE"))):
            data = page_model_data(R.PAGES[R.page_url(fragment)][1])
            self.assertEqual((data["model_code"], data["model_name"]), expected)
        self.assertEqual(page_model_data("<html>digitalData.product.model_code = \"A\"</html>"), {})


class PhotosAreBoundToTheVariantOrConfirmedByAPerson(unittest.TestCase):
    def photo(self, code):
        return PhotoCandidate(f"https://images.samsung.com/is/image/samsung/p6pim/kz_ru/{code}/gallery/x-1?$Q90_2052_1641_JPG$", f"k-{code}")

    def test_binding_needs_the_assets_own_path_to_name_exactly_the_catalog_code(self):
        photos = [self.photo("nz64t3506ak-wt"), self.photo("nz64t3506ak-bb"), self.photo("sm-a376edggskz")]
        self.assertEqual([p.asset_key for p in photo_variant_binding(photos, "NZ64T3506AK/WT")], ["k-nz64t3506ak-wt"])
        self.assertEqual(photo_variant_binding([self.photo("rb31ferndsa-wt")], "RB31FERNDSA"), [])       # the model without its suffix is not the variant
        self.assertEqual(photo_variant_binding([self.photo("sm-a376edggskz")], "SM-A376EZAGINS"), [])
        self.assertEqual(photo_variant_binding([self.photo("x")], ""), [])

    def test_a_person_can_confirm_photos_the_pipeline_left_unselected_and_the_choice_survives_a_rerun(self):
        with TemporaryDirectory() as tmp:
            cards = [("Холодильники", "RB31FERNDSA")]
            first = R.run_products(Path(tmp), cards=cards)
            db, pid = first["database"], first["outcomes"][0]["product_id"]
            self.assertEqual(samsung_readiness.card_readiness(db, pid)["official_photos_selected"], 0)
            self.assertIn("photos_not_selected_variant_open", samsung_readiness.card_readiness(db, pid)["blocking_gaps"])
            keys = [p["asset_key"] for p in jobs.get_photo_candidates(db, pid, include_excluded=False) if p["source_key"] == "samsung"][:3]
            with TestClient(create_app(tmp)) as client:
                answer = client.post(f"/products/{pid}/photos", data={"action": "exact", "asset_keys": keys}, follow_redirects=False)
                self.assertEqual(answer.status_code, 303)
            readiness = samsung_readiness.card_readiness(db, pid)
            self.assertEqual(readiness["official_photos_selected"], 3)
            self.assertNotIn("photos_not_selected_variant_open", readiness["blocking_gaps"])
            self.assertIn("variant_code_only_in_page_text", readiness["blocking_gaps"])         # the variant is still not confirmed by the page: a photo choice does not change that
            self.assertEqual(readiness["verdict"], "needs_verification")
            # the same job again (a rerun of the ordinary path over the same database): the person's choice stays
            from product_tool import worker
            jobs.enqueue(db, pid, [1, 2, 3, 4, 6])
            worker.run_once(db, clock=lambda: 0.0, samsung_adapter_factory=lambda: default_samsung_adapter(Path(tmp), clock=lambda: 0.0, underlying=first["replay"], min_interval_seconds=0.0, pages_reader=first["replay"].pages_reader),
                            dns_adapter_factory=lambda: __import__("product_tool.adapters.dns", fromlist=["DnsAdapter"]).DnsAdapter(first["dealer_session"], clock=lambda: 0.0))
            self.assertEqual(sorted(p["asset_key"] for p in jobs.get_photo_candidates(db, pid, include_excluded=False) if p["source_key"] == "samsung" and p["selected"]), sorted(keys))


class RecordedStage86Evidence(unittest.TestCase):
    def test_the_recorded_entry_matches_the_exact_link_the_pages_own_markup_prints(self):
        recorded = load_recorded_documents()
        self.assertEqual(list(recorded), [item["href"] for item in json.loads(RECORDED_DOCUMENTS.read_text(encoding="utf-8"))["documents"]])
        entry = next(iter(recorded.values()))
        from product_tool.adapters.samsung import extract_document_links
        page = R.page_url("ms23k3614akbw")
        links = {link.href for link in extract_document_links(R.PAGES[page][1], page)}
        self.assertIn(entry["href"], links)
        for path in entry["evidence"]:
            self.assertTrue((ROOT / path).is_file(), path)
        self.assertEqual((entry["article"], entry["bytes"], entry["pages"]), ("MS23K3614AK/BW", 10572005, 80))
        stage86 = json.loads((ROOT / "reports/source_census_2026-09-22_stage8_6/assembly_result.json").read_text(encoding="utf-8"))
        self.assertEqual((entry["href"], entry["bytes"]), (stage86["url"], stage86["bytes_assembled"]))

    def test_an_entry_is_used_only_for_its_own_article_and_its_own_link(self):
        with TemporaryDirectory() as tmp:
            replay = R.Replay()
            adapter = default_samsung_adapter(Path(tmp), clock=lambda: 0.0, underlying=replay, min_interval_seconds=0.0, pages_reader=replay.pages_reader)
            page = R.page_url("ms23k3614akbw")
            document = adapter.find_source("MS23K3614AK/BW", deadline=1e9, category="Микроволновые печи")
            self.assertEqual(document.url, page)
            before = len(replay.calls)
            docs, reason = adapter.find_documents(document, "MS23K3614AK/BW", deadline=1e9)
            self.assertEqual((len(docs), reason, len(replay.calls) - before), (1, "", 0))
            self.assertIn("recorded", json.dumps(adapter.reports["MS23K3614AK/BW"]["documents"], default=str))
            # another article on the same page: the entry is not used, the file would have to be fetched (the replay refuses it)
            other = default_samsung_adapter(Path(tmp) / "other", clock=lambda: 0.0, underlying=replay, min_interval_seconds=0.0, pages_reader=replay.pages_reader)
            other_doc = other.find_source("MS23K3614AK/BW", deadline=1e9, category="Микроволновые печи")
            docs, _ = other.find_documents(other_doc, "MS23K3614AK/XX", deadline=1e9)
            self.assertEqual(docs, [])
            self.assertTrue(replay.refused)


class PlannerOffersOneSelectedProductPerCategory(unittest.TestCase):
    def test_exactly_the_selected_products_are_ready_and_every_other_samsung_row_waits(self):
        rows = [u for u in plan()["units"] if u["family"] == "samsung"]
        self.assertEqual(len(rows), 1182)
        ready = [u for u in rows if u["status"] == C.READY]
        self.assertEqual({(u["category"], u["seller_sku"]) for u in ready}, set(SELECTED.items()))
        self.assertEqual(len({u["category"] for u in ready}), len(ready))                      # one per category
        rest = Counter((u["status"], u["reason"]) for u in rows if u["status"] != C.READY)
        self.assertEqual(rest, {(C.MANUAL, "not_selected_one_card_per_category"): 1136, (C.MANUAL, "seller_sku_not_a_manufacturer_code"): 26})
        self.assertTrue(all(u["next_action"] == "manual_decision" for u in rows if u["status"] == C.MANUAL))

    def test_a_selected_row_is_ready_only_in_its_own_category(self):
        from product_tool.coverage.facts import load_config
        self.assertEqual({(i["category"], i["seller_sku"]) for i in load_config()["selected_products"]["samsung"]["products"]}, set(SELECTED.items()))
        units = [u for u in plan()["units"] if u["family"] == "samsung" and u["seller_sku"] in SELECTED.values()]
        for unit in units:
            self.assertEqual(unit["status"] == C.READY, SELECTED.get(unit["category"]) == unit["seller_sku"], unit["seller_sku"])

    def test_the_planner_touches_no_network_and_the_executor_runs_only_ready_units(self):
        ready = [u for u in plan()["units"] if u["family"] == "samsung" and u["status"] == C.READY and u["seller_sku"] in ("QE48S85HAEXCE", "HW-Q800D")]
        waiting = next(u for u in plan()["units"] if u["family"] == "samsung" and u["status"] == C.MANUAL and u["reason"] == "not_selected_one_card_per_category")
        with TemporaryDirectory() as tmp:
            replay = R.Replay()
            factories = executor.RunFactories(session=replay, samsung_adapter_factory=lambda: default_samsung_adapter(Path(tmp), clock=lambda: 0.0, underlying=replay, min_interval_seconds=0.0, pages_reader=replay.pages_reader))
            checkpoint = executor.run_units(ready + [waiting], workdir=Path(tmp), checkpoint_path=Path(tmp) / "checkpoint.json", factories=factories)
            outcomes = {v["seller_sku"]: v for v in checkpoint["units"].values()}
            self.assertEqual({k: v["outcome"] for k, v in outcomes.items() if k in ("QE48S85HAEXCE", "HW-Q800D")}, {"QE48S85HAEXCE": "card", "HW-Q800D": "card"})
            stopped = outcomes[waiting["seller_sku"]]
            self.assertEqual((stopped["outcome"], stopped["stop_code"], stopped["stopped_before_run_once"]), ("stopped", "not_selected_one_card_per_category", True))
            self.assertEqual(replay.refused, [])
            self.assertFalse(any(waiting["seller_sku"].lower() in call.lower() for call in replay.calls))

    def test_the_default_executor_transport_refuses_every_samsung_request(self):
        unit = next(u for u in plan()["units"] if u["family"] == "samsung" and u["seller_sku"] == "QE48S85HAEXCE")
        with TemporaryDirectory() as tmp:
            factories = executor.RunFactories()
            checkpoint = executor.run_units([unit], workdir=Path(tmp), checkpoint_path=Path(tmp) / "checkpoint.json", factories=factories)
            entry = next(iter(checkpoint["units"].values()))
            self.assertNotEqual(entry["outcome"], "card")
            self.assertTrue(factories.session.refused)
            self.assertTrue(all(url.startswith("https://www.samsung.com/kz_ru/") for url in factories.session.refused))


class OrdinaryPathCreateProcessExport(unittest.TestCase):
    """Upload -> confirm -> a search job per product -> run_once -> export, through the web app (not a helper that writes rows)."""

    @classmethod
    def setUpClass(cls):
        cls._tmp = TemporaryDirectory()
        cls.result = R.run_products(Path(cls._tmp.name), cards=[("Телевизоры", "QE48S85HAEXCE"), ("Холодильники", "RB31FERNDSA"), ("Смартфоны", "SM-A376EZAGINS")])

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_jobs_were_created_by_the_products_own_route_and_finished_by_the_worker(self):
        for outcome in self.result["outcomes"]:
            jobs_for = jobs.list_jobs(self.result["database"], outcome["product_id"])
            self.assertEqual(len(jobs_for), 1)
            self.assertEqual(jobs_for[0]["stages"], [1, 2, 3, 4, 6])
            self.assertIn(jobs_for[0]["status"], ("done", "needs_review"))
        self.assertEqual({o["article"]: o["status"] for o in self.result["outcomes"]}, {"QE48S85HAEXCE": "done", "RB31FERNDSA": "needs_review", "SM-A376EZAGINS": "needs_review"})

    def test_the_job_status_and_the_card_readiness_are_shown_apart_in_the_export(self):
        book = load_workbook(io.BytesIO(self.result["export"]))
        rows = {r[1]: r for r in book["Готовность Samsung"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual({a: r[2] for a, r in rows.items()}, {"QE48S85HAEXCE": "export_ready", "RB31FERNDSA": "needs_verification", "SM-A376EZAGINS": "needs_verification"})
        sources = {(r[1], r[2]): r[4] for r in book["Источники"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(sources[("RB31FERNDSA", "Samsung Казахстан")], "Артикул только в тексте страницы")
        self.assertEqual(sources[("SM-A376EZAGINS", "Samsung Казахстан")], "Базовая модель")
        photos = Counter(r[0] for r in book["Фотографии"].iter_rows(min_row=2, values_only=True))
        names = {o["article"]: o for o in self.result["outcomes"]}
        self.assertEqual(sum(photos.values()), 8)                                                  # the TV's gallery only; none of the fridge's or the phone's (Stage 30: the hob left this case)

    def test_the_web_card_page_and_the_export_agree_on_the_levels(self):
        with TestClient(create_app(Path(self._tmp.name))) as client:
            for outcome in self.result["outcomes"]:
                page = client.get(f"/products/{outcome['product_id']}")
                self.assertEqual(page.status_code, 200)
                self.assertIn("Samsung Казахстан", page.text)



BATCH4 = json.loads((ROOT / "reports/source_census_2026-09-25_stage27/raw/batch4_result.json").read_text(encoding="utf-8"))
DECLARATION = json.loads((ROOT / "reports/source_census_2026-09-25_stage27/raw/batch4_declaration.json").read_text(encoding="utf-8"))


class Batch4WasDeclaredFirstAndStayedInsideItsBudget(unittest.TestCase):
    """The four categories checked with REAL requests in Stage 27 (recorded in reports/source_census_2026-09-25_stage27)."""

    def test_the_budget_was_declared_before_the_first_request_and_never_exceeded(self):
        self.assertLess(DECLARATION["declared_at"], BATCH4["started_at"])
        budget = BATCH4["budget"]
        self.assertEqual((budget["max_total"], budget["max_per_row"]), (14, 7))
        self.assertLessEqual(budget["spent_total"], 14)
        self.assertEqual(budget["spent_total"], len(budget["log"]))
        for row in {e["row"] for e in budget["log"]}:
            self.assertLessEqual(sum(1 for e in budget["log"] if e["row"] == row), 7)
        self.assertEqual(BATCH4["stopped_hosts_after_run"], [])
        self.assertEqual(BATCH4["not_started"], [])
        self.assertEqual([e["status_code"] for e in BATCH4["fetch_log_entries"]], [200] * len(BATCH4["fetch_log_entries"]))

    def test_exactly_one_product_per_category_and_only_official_addresses(self):
        self.assertEqual({c["category"]: c["article"] for c in BATCH4["cards"]}, {"Машины посудомоечные": "DW60M5050BB/WT", "Сушильные машины": "DV16DG8600BVLD", "Колонки": "MX-ST50B", "Роботы-пылесосы": "VR50T95735W/EV"})
        for entry in BATCH4["budget"]["log"]:
            self.assertRegex(entry["url"], r"^https://(?:www\.samsung\.com/kz_ru/|org\.downloadcenter\.samsung\.com/downloadfile/ContentsFile\.aspx\?)")
        allowed = {DECLARATION["products"][c["category"]]["page_url_listed_in_saved_sitemaps"] for c in BATCH4["cards"]}
        pages = {e["url"] for e in BATCH4["budget"]["log"] if "/kz_ru/" in e["url"] and not e["url"].endswith(".xml")}
        self.assertEqual(pages, allowed)                                   # each page is the very address the official sitemap lists

    def test_the_declared_document_files_per_product_were_respected(self):
        for index, card in enumerate(BATCH4["cards"], start=1):
            files = [e for e in BATCH4["budget"]["log"] if e["row"] == str(index) and "ContentsFile" in e["url"]]
            self.assertLessEqual(len(files), 2, card["category"])


class Batch4OnTheOrdinaryPathOffline(unittest.TestCase):
    """The same four cards replayed offline through upload -> confirm -> jobs -> run_once -> export from the responses the real run recorded."""

    @classmethod
    def setUpClass(cls):
        cls._tmp = TemporaryDirectory()
        cls.batch = R.run_products(Path(cls._tmp.name), cards=R.CARDS4, recorded_documents={})      # the real Stage 27 run: no recorded evidence yet
        cls.db = cls.batch["database"]
        cls.by = {o["category"]: o for o in cls.batch["outcomes"]}
        cls.readiness = {c: samsung_readiness.card_readiness(cls.db, o["product_id"]) for c, o in cls.by.items()}
        cls.live = {c["category"]: c for c in BATCH4["cards"]}

    @classmethod
    def tearDownClass(cls):
        cls._tmp.cleanup()

    def test_no_request_beyond_the_saved_ones_and_none_to_the_dealer(self):
        self.assertEqual(self.batch["replay"].refused, [])
        self.assertEqual(self.batch["dealer_session"].calls, [])
        self.assertLessEqual(len(self.batch["replay"].calls), 12)

    def test_the_offline_replay_reproduces_what_the_real_run_read(self):
        for category, live in self.live.items():
            with self.subTest(category):
                r = self.readiness[category]
                self.assertEqual(r["page_match_level"], live["sources"]["samsung"]["match_level"])
                self.assertEqual(sum(1 for f in jobs.get_facts(self.db, self.by[category]["product_id"]) if f["source_key"] == "samsung"), live["facts"]["samsung"])
                self.assertEqual(r["official_photos"], live["photos"]["gallery_found"])
                live_states = [e["state"] for e in live["document_files_assessed"]]
                replay_states = [e["state"] for e in card_evidence.load(self.db, self.by[category]["product_id"], "samsung_documents")["documents"]]
                self.assertEqual(replay_states, live_states)

    def test_readiness_and_job_status_are_shown_apart(self):
        shown = {c: (self.by[c]["status"], r["verdict"]) for c, r in self.readiness.items()}
        self.assertEqual(shown, {"Машины посудомоечные": ("done", "export_ready_with_gaps"), "Сушильные машины": ("done", "export_ready"), "Колонки": ("done", "export_ready"), "Роботы-пылесосы": ("needs_review", "export_ready_with_gaps")})

    def test_the_dishwashers_russian_file_was_over_the_cap_and_is_reported_as_not_obtained(self):
        r = self.readiness["Машины посудомоечные"]
        self.assertEqual(r["blocking_gaps"], ["instruction_russian_candidate_not_verified"])            # Stage 28: a candidate not verified, not "no Russian instruction"
        self.assertIn("instruction_only_non_russian_file_saved", r["advisory_gaps"])
        states = [e["state"] for e in card_evidence.load(self.db, self.by["Машины посудомоечные"]["product_id"], "samsung_documents")["documents"]]
        self.assertEqual(states, ["reachable_but_not_a_complete_pdf", "instruction_saved"])
        doc = jobs.get_documents(self.db, self.by["Машины посудомоечные"]["product_id"])[0]
        self.assertEqual(doc["language"], "Английский")                                    # saved with the language its TEXT shows; never counted as Russian
        self.assertIn("samsung_instruction_check", r["open_reviews"])

    def test_the_dishwashers_colour_option_no_longer_collides_with_the_backlight_colour(self):
        self.assertEqual(jobs.result_counts(self.db, self.by["Машины посудомоечные"]["product_id"])["conflicts"], 0)
        self.assertEqual(self.live["Машины посудомоечные"]["resolved"]["conflicts"], 1)      # what the real run found before the vocabulary grew

    def test_the_robots_unassigned_weight_stays_a_conflict_for_a_person(self):
        # Stage 28: sizes and weights that an explicit manual table ties to the robot or to the station are split; what no evidence ties stays, with both source values
        counts = jobs.result_counts(self.db, self.by["Роботы-пылесосы"]["product_id"])
        self.assertEqual(counts["conflicts"], 1)
        names = sorted(r["normalized_name"] for r in jobs.get_resolved(self.db, self.by["Роботы-пылесосы"]["product_id"]) if r["conflict"])
        self.assertEqual(names, ["product_weight"])
        self.assertIn("unresolved_conflicts", self.readiness["Роботы-пылесосы"]["blocking_gaps"])

    def test_the_owner_rules_on_the_four_new_cards(self):
        dryer = self.readiness["Сушильные машины"]["instruction"]
        robot = self.readiness["Роботы-пылесосы"]["instruction"]
        speaker = self.readiness["Колонки"]["instruction"]
        self.assertEqual((dryer["acceptance_basis"], robot["acceptance_basis"], speaker["acceptance_basis"]), ("family_mask_in_pdf", "family_mask_in_pdf", "exact_code_in_pdf"))
        self.assertIn("точный код в PDF не назван", dryer["mark"])
        self.assertEqual(self.readiness["Колонки"]["advisory_gaps"], ["dealer_cross_check_missing"])
        for category in ("Сушильные машины", "Колонки", "Роботы-пылесосы", "Машины посудомоечные"):
            self.assertEqual(self.readiness[category]["official_photos_selected"], self.readiness[category]["official_photos"], category)   # full_sku pages: the gallery is selected

    def test_the_export_lists_the_four_cards_with_status_apart_from_readiness(self):
        book = load_workbook(io.BytesIO(self.batch["export"]))
        rows = {r[1]: r for r in book["Готовность Samsung"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual({a: r[2] for a, r in rows.items()}, {"DW60M5050BB/WT": "export_ready_with_gaps", "DV16DG8600BVLD": "export_ready", "MX-ST50B": "export_ready", "VR50T95735W/EV": "export_ready_with_gaps"})
        self.assertEqual(rows["DW60M5050BB/WT"][13], "not_a_russian_instruction_by_text: не принята")


if __name__ == "__main__":
    unittest.main()
