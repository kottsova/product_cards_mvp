"""Stage 55 three-state Russian instruction status and provenance guard."""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from product_tool import exporter, jobs, manual_status, storage
from product_tool.adapters.common import ProductDocument, SourceDocument


class ManualStatusTest(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.database = Path(self.temp.name) / "batch.sqlite3"
        jobs.initialize(self.database)
        with storage._connection(self.database) as connection:
            connection.execute(
                "INSERT INTO batches VALUES ('b','batch.xlsx','Items','{}','2026-01-01')"
            )
            self.product_id = connection.execute(
                "INSERT INTO products (batch_id,row_number,name,brand,search_code,"
                "alternate_code,category,needs_confirmation,issues_json,original_values_json) "
                "VALUES ('b',2,'Microwave','LG','MS2082F','','Microwave',0,'[]','{}')"
            ).lastrowid
        self.evidence = [{
            "kind": "official_manual_list", "official": True, "checked": True,
            "exact_model": True, "russian_found": False,
            "url": "https://www.lg.com/ncms/eu/api/v1/support/proxy/retrieveManualSoftwareList?locale=UK",
            "result": "Two English manuals"
        }]

    def test_only_completed_exact_official_inventory_allows_not_found(self):
        self.assertEqual(
            manual_status.russian_status(self.database, self.product_id, "MS2082F"),
            manual_status.UNCHECKED,
        )
        with self.assertRaises(ValueError):
            manual_status.record_completed_absence(
                self.database, self.product_id, "MS2082F",
                [{"kind": "regional_support", "official": True, "checked": True}]
            )
        manual_status.record_completed_absence(
            self.database, self.product_id, "MS2082F", self.evidence
        )
        self.assertEqual(
            manual_status.russian_status(self.database, self.product_id, "MS2082F"),
            manual_status.NOT_FOUND,
        )
        self.assertEqual(
            manual_status.russian_status(self.database, self.product_id, "MS2082F.NEW"),
            manual_status.UNCHECKED,
        )
        book = load_workbook(BytesIO(exporter.export_batch(self.database, "b")), read_only=True)
        try:
            row = list(book["Инструкции"].values)[1]
            self.assertEqual(row[-1], manual_status.NOT_FOUND)
        finally:
            book.close()

    def test_attended_dealer_rejection_survives_later_host_stop(self):
        url = "https://www.sulpak.kz/g/example_kit"
        reviewed = SourceDocument(
            "sulpak", "Sulpak", url, match_level="unknown",
            evidence="Attended browser: full kit not printed",
            html="<html><body>LG base model only</body></html>",
        )
        jobs.save_source_document(self.database, self.product_id, reviewed)
        jobs.save_source_document(
            self.database, self.product_id,
            SourceDocument("sulpak", "Sulpak", url, match_level="unknown",
                           error="policy_host_stopped")
        )
        saved = next(x for x in jobs.get_source_pages(self.database, self.product_id)
                     if x["source_key"] == "sulpak")
        self.assertEqual(saved["error"], "")
        self.assertEqual(saved["evidence"], reviewed.evidence)
        with storage._connection(self.database) as connection:
            statuses = [row["status"] for row in connection.execute(
                "SELECT status FROM fetch_attempts WHERE product_id=? AND source_id='sulpak' ORDER BY id",
                (self.product_id,)
            )]
        self.assertEqual(statuses, ["success", "error"])

    def test_verified_russian_document_overrides_prior_absence(self):
        manual_status.record_completed_absence(
            self.database, self.product_id, "MS2082F", self.evidence
        )
        url = "https://www.lg.com/kz/support/product-support/cs-MS2082F/"
        jobs.save_source_document(
            self.database, self.product_id,
            SourceDocument("lg_kz_support", "LG KZ support", url, match_level="full_sku")
        )
        jobs.save_documents(
            self.database, self.product_id, "lg_kz_support",
            [ProductDocument("Owner manual", "Русский", "", "", "https://gscs-b2c.lge.com/manual.pdf",
                             url, "MS2082F", "MS2082F", url, True)]
        )
        self.assertEqual(
            manual_status.russian_status(self.database, self.product_id, "MS2082F"),
            manual_status.VERIFIED,
        )


if __name__ == "__main__":
    unittest.main()
