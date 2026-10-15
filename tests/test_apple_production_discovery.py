"""Production-path counterexamples; synthetic transport here is never live evidence."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from product_tool import jobs, storage, worker, card_evidence, discovery_trace
from product_tool.apple_model_discovery import parse_native_specs, overlay_store, apple_candidate_type, iphone_model
from product_tool.adapters.apple import AppleAdapter
from product_tool.adapters.common import SourceDocument

SPECS = '''<h1>iPhone 18 Pro Technical Specifications</h1>
<div class="techspecs-section"><div class="techspecs-row"><div class="techspecs-rowheader">Chip</div>
<div class="techspecs-column" aria-colspan="2"><ul><li>A20 Pro chip</li><li>6-core CPU</li><li>7-core GPU</li><li>Dual 16-core Neural Engine</li></ul></div></div></div>
<div class="techspecs-section"><div class="techspecs-row"><div class="techspecs-rowheader">Display</div>
<div class="techspecs-column"><strong class="techspecs-small-heading">iPhone 18 Pro</strong><ul><li>6.3-inch OLED display</li><li>2622-by-1206-pixel resolution at 460 ppi</li></ul></div>
<div class="techspecs-column"><strong class="techspecs-small-heading">iPhone 18 Pro Max</strong><ul><li>6.9-inch OLED display</li><li>2868-by-1320-pixel resolution</li></ul></div></div></div>
<div class="techspecs-section"><div class="techspecs-row"><div class="techspecs-rowheader">Size and Weight</div>
<div class="techspecs-column"><strong class="techspecs-small-heading">iPhone 18 Pro</strong><figcaption class="diagram-text">Height: 150 mm</figcaption><figcaption class="diagram-text">Width: 71.9 mm</figcaption><div class="weight-wrapper"><p class="weight-copy">211 grams</p></div></div>
<div class="techspecs-column"><strong class="techspecs-small-heading">iPhone 18 Pro Max</strong><figcaption class="diagram-text">Height: 163.4 mm</figcaption></div></div></div>'''
PDP = '<h1>iPhone 18 Pro</h1><a href="/iphone-18-pro/specs/">Tech Specs</a><a href="/us/shop/goto/buy_iphone">Shop all</a><a href="/us/shop/goto/buy_iphone/iphone_18_pro">Buy</a>'
IMAGE = 'https://store.storeimages.cdn-apple.com/is/iphone-18-pro-finish-select-silver-202609.jpg'
PAYLOAD = dict(products=[dict(partNumber='TEST1LL/A',familyType='iphone18pro',dimensionCapacity='256gb',dimensionColor='silver',energyComplianceKey='a3714_energy_compliance',imageKey='iphone-18-pro-finish-select-silver-202609'),dict(partNumber='TEST2LL/A',familyType='iphone18promax',dimensionCapacity='256gb',dimensionColor='silver',energyComplianceKey='a3717_energy_compliance',imageKey='iphone-18-pro-max-finish-select-silver-202609')],imageDictionary={'iphone-18-pro-finish-select-silver-202609':dict(sources=[dict(srcSet=IMAGE)])})


class AppleProductionDiscovery(unittest.TestCase):
    def test_model_normalization_preserves_generation_and_pro_max(self):
        self.assertEqual(iphone_model('iPhone 18 Pro 256GB'), 'iPhone 18 Pro')
        self.assertEqual(iphone_model('iPhone 18 Pro Max 512GB'), 'iPhone 18 Pro Max')
        self.assertEqual(iphone_model('A3714'), '')

    def test_native_model_scope_excludes_max_and_unowned_cells(self):
        doc,audit=parse_native_specs(SPECS,'https://www.apple.com/iphone-18-pro/specs/','iPhone 18 Pro')
        self.assertFalse(doc.error)
        self.assertEqual(doc.match_level,'model_confirmed')
        values=' '.join(a.value for a in doc.attributes)
        self.assertIn('A20',values)
        self.assertNotIn('6.9',values)
        self.assertNotIn('163.4',values)
        self.assertTrue(audit['rejected_cells'])
        unowned=SPECS.replace('<strong class="techspecs-small-heading">iPhone 18 Pro</strong>','')
        other,_=parse_native_specs(unowned,'https://www.apple.com/iphone-18-pro/specs/','iPhone 18 Pro')
        self.assertFalse(any(a.name=='Высота товара, мм' for a in other.attributes))

    def test_wrong_generation_rejected_without_facts(self):
        doc,_=parse_native_specs(SPECS,'https://www.apple.com/iphone-18-pro/specs/','iPhone 17 Pro')
        self.assertTrue(doc.error)
        self.assertFalse(doc.attributes)

    def test_shared_display_continuation_and_controls_are_not_conflicts(self):
        body=SPECS+'''<div class="techspecs-section"><div class="techspecs-row"><div class="techspecs-rowheader">Display</div>
        <div class="techspecs-column" aria-colspan="2"><p>HDR display</p><p>Always-On display</p></div></div></div>
        <div class="techspecs-section"><div class="techspecs-row"><div class="techspecs-rowheader">External Buttons and Connectors</div>
        <div class="techspecs-column" aria-colspan="2"><strong class="techspecs-small-heading">Action button features</strong><p>Silent Mode, Camera</p></div></div>
        <div class="techspecs-row"><div class="techspecs-column" aria-colspan="2"><strong class="techspecs-small-heading">Camera controls</strong><p>Exposure, Zoom</p></div></div></div>'''
        d,_=parse_native_specs(body,'https://www.apple.com/iphone-18-pro/specs/','iPhone 18 Pro')
        tech=[a for a in d.attributes if a.name=='Технология экрана']
        self.assertEqual(len(tech),1)
        self.assertIn('HDR',tech[0].value)
        self.assertEqual([a.value for a in d.attributes if a.name=='Диагональ экрана, дюймы'],['6.3'])
        self.assertIn('Функции кнопки Action',{a.name for a in d.attributes})
        self.assertIn('Функции Camera Control',{a.name for a in d.attributes})

    def test_apple_search_classifier_admits_apple_not_lenovo(self):
        self.assertEqual(apple_candidate_type('https://support.apple.com/en-us/148590'),'support')
        self.assertEqual(apple_candidate_type('https://www.apple.com/iphone-18-pro/specs/'),'product')
        self.assertEqual(apple_candidate_type('https://apple.com.evil.test/iphone-18-pro/'),'unknown')
        self.assertEqual(apple_candidate_type('https://www.apple.com/search/'),'unknown')

    def test_wrong_hardware_or_color_does_not_confirm_configuration_or_photos(self):
        for article,color in [('A3717','silver'),('A3714','black')]:
            with self.subTest(article=article,color=color),tempfile.TemporaryDirectory() as td:
                a=AppleAdapter(fetch_log_path=Path(td)/'fetch.json');a.input_configuration={'color':color}
                ev={'identity':{},'exact_photo_assets':[]};doc=SourceDocument('apple','Apple','https://www.apple.com/iphone-18-pro/',match_level='model_confirmed')
                overlay_store(a,doc,ev,PAYLOAD,'https://www.apple.com/shop/buy-iphone/iphone-18-pro',article,'iPhone 18 Pro 256GB','iPhone 18 Pro')
                self.assertFalse(doc.attributes)
                self.assertFalse(doc.photos)
                self.assertFalse(ev.get('configuration_complete'))

    def test_single_ui_queue_runs_discovery_before_readiness_with_hardware_code(self):
        from fastapi.testclient import TestClient
        from product_tool.web import create_app
        responses={
          'https://www.apple.com/iphone/':'<a href="/iphone-18-pro/">iPhone 18 Pro</a>',
          'https://www.apple.com/iphone-18-pro/':PDP,
          'https://www.apple.com/iphone-18-pro/specs/':SPECS,
          'https://www.apple.com/us/shop/goto/buy_iphone/iphone_18_pro':'<script>window.PRODUCT_SELECTION_BOOTSTRAP = {productSelectionData: '+json.dumps(PAYLOAD)+'};</script>'}
        calls=[]
        class Session:
            def get(self,url,**kwargs):
                from product_tool.adapters.policy_session import PolicyResponse
                calls.append(url)
                return PolicyResponse(url,200,responses[url])
        with tempfile.TemporaryDirectory() as td:
            db=Path(td)/'batches.sqlite3';jobs.initialize(db)
            with storage._connection(db) as c:
                c.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('b','input.xlsx','Sheet','{}',storage._now()))
                c.execute('INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json,identity_json) VALUES (?,?,?,?,?,?,?,?,?,?,?)',('b',2,'iPhone 18 Pro 256GB','Apple','A3714','','смартфоны',0,'[]','[]',json.dumps({'variant_attributes':{'color':'Серебристый'}})))
            with TestClient(create_app(td,start_worker=False)) as client:
                response=client.post('/products/1/search',data={'stages':['1','3','4']})
                self.assertEqual(response.status_code,200)
                with patch('product_tool.adapters.apple.PolicyAwareSession',return_value=Session()),patch('product_tool.photo_metadata.inspect_saved_photo',return_value={'width':1000,'height':1000,'size_bytes':1234,'format':'JPG','sha256':'f'*64}):
                    self.assertTrue(worker.run_once(db))
                html=client.get('/products/1').text
                self.assertIn('discovery_started',html)
                self.assertIn('readiness_passed',html)
            self.assertEqual(calls,list(responses))
            ready=__import__('product_tool.apple_pipeline',fromlist=['card_readiness']).card_readiness(db,1)
            self.assertEqual(ready['verdict'],'export_ready')
            events=discovery_trace.for_job(db,jobs.list_jobs(db,1)[0]['id'])
            names=[e['event'] for e in events]
            self.assertLess(names.index('discovery_started'),names.index('apple_model_specs'))
            self.assertLess(names.index('apple_extraction'),names.index('readiness_passed'))
            self.assertTrue(all(e['timestamp'] for e in events))
            ev=card_evidence.load(db,1,'apple')
            self.assertEqual(ev['identity']['part_number'],'')
            self.assertEqual(ev['retail_candidates'],['TEST1LL/A'])
            self.assertTrue(all(not x['full_sku_confirmed'] for x in jobs.get_resolved(db,1)))
            from product_tool.photo_inventory import enriched_photos
            gallery=enriched_photos(db,1)
            self.assertTrue(gallery[0]['byte_verified'])
            self.assertEqual(gallery[0]['official_color'],'silver')


class ImportMappingRecovery(unittest.TestCase):
    def test_actual_header_and_brand_mapping_collision(self):
        from openpyxl import Workbook
        from product_tool.importer import detect_mapping,preview_row,ColumnMapping
        w=Workbook();w.active.append(['категория','бренд',None,'Наименование на сайте ','Цвет'])
        mapping=detect_mapping(w.active)
        self.assertEqual(mapping.name,4)
        row=('смартфоны','Apple','A3714','iPhone 18 Pro 256GB','Серебристый')
        bad=preview_row(2,row,ColumnMapping(brand=2,search_code=3,fallback_code=2,name=3,category=1,header_row=1))
        self.assertTrue(bad.needs_confirmation)
        self.assertTrue(any('Колонка бренда' in issue for issue in bad.issues))

    def test_supplier_color_survives_upload_confirmation_as_input_only(self):
        from fastapi.testclient import TestClient
        from product_tool.web import create_app
        from openpyxl import Workbook
        from io import BytesIO
        w=Workbook();w.active.append(['категория','бренд',None,'Наименование на сайте','Цвет'])
        w.active.append(['смартфоны','Apple','A3714','iPhone 18 Pro 256GB','Серебристый'])
        body=BytesIO();w.save(body)
        with tempfile.TemporaryDirectory() as td,TestClient(create_app(td,start_worker=False)) as client:
            response=client.post('/upload',files={'file':('supplier.xlsx',body.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
            draft=response.url.path.split('/')[-1]
            # The source hardware column has no heading: select it explicitly.
            response=client.post(f'/drafts/{draft}/mapping',data={'sheet_name':'Sheet','header_row':'1','brand':'2','search_code':'3','name':'4','category':'1','fallback_code':''})
            self.assertEqual(response.status_code,200)
            response=client.post(f'/drafts/{draft}/confirm',data={'brand_2':'Apple','search_code_2':'A3714','alternate_code_2':'','category_2':'смартфоны'})
            self.assertEqual(response.status_code,200)
            product=storage.get_batch(Path(td)/'batches.sqlite3',draft)['products'][0]
            self.assertEqual(product['name'],'iPhone 18 Pro 256GB')
            self.assertEqual(json.loads(product['identity_json'])['variant_attributes']['color'],'Серебристый')
            self.assertFalse(jobs.get_facts(Path(td)/'batches.sqlite3',product['id']))
