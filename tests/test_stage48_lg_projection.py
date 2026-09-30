"""Output-only presence/refinement projection for LG cards and Excel."""
from __future__ import annotations

from contextlib import closing
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from sqlite3 import connect
from tempfile import TemporaryDirectory
import unittest

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from product_tool import attribute_projection, exporter, jobs
from product_tool.web import create_app


BATCH_ID = "ae3d2cb381744ba8a811e54231233254"
DATABASE = Path(__file__).resolve().parents[1] / "data" / "batches.sqlite3"


def _row(name: str, value: str, unit: str = "", *, full: bool = True,
         conflict: bool = False, raw: str | None = None) -> dict:
    fact = {
        "raw_name": name, "raw_value": raw if raw is not None else value,
        "normalized_value": value, "unit": unit,
        "display_value": "Да" if (value, unit) == ("true", "bool") else value,
        "source_key": "lg_kz", "site_name": "LG Kazakhstan",
    }
    return {
        "normalized_name": name, "display_name": name, "raw_names": [name],
        "sources": {"lg_kz": fact},
        "resolved": {
            "id": len(name), "normalized_name": name, "selected_value": value,
            "selected_unit": unit, "display_value": fact["display_value"],
            "selected_source": "lg_kz", "display_source": "LG Kazakhstan",
            "status": "full_sku_lg" if full else "official_base_only",
            "full_sku_confirmed": int(full), "conflict": int(conflict),
            "reason": "From official page.",
        },
    }


class FinalAttributeProjectionTests(unittest.TestCase):
    def test_substantive_value_replaces_presence_in_final_rows_only(self):
        rows = [
            _row("bluetooth", "true", "bool", raw="●"),
            _row("bluetooth_version", "5.0"),
            _row("wi_fi", "true", "bool", raw="●"),
            _row("wi_fi_version", "wi-fi 6"),
            _row("usb", "true", "bool", raw="●"),
            _row("usb_ports", "3"),
        ]
        original = deepcopy(rows)
        final = {r["normalized_name"]: r for r in attribute_projection.project_final_rows(rows)}
        self.assertEqual(set(final), {"bluetooth", "wi_fi", "usb"})
        self.assertEqual(final["bluetooth"]["resolved"]["display_value"], "5.0")
        self.assertEqual(final["wi_fi"]["resolved"]["display_value"], "Wi-Fi 6")
        self.assertEqual(final["usb"]["resolved"]["display_value"], "3 порта")
        self.assertEqual(final["bluetooth"]["resolved"]["id"], len("bluetooth_version"))
        self.assertEqual(final["bluetooth"]["sources"]["lg_kz"]["raw_value"], "●; 5.0")
        self.assertEqual(rows, original)

    def test_presence_without_refinement_and_conflicts_are_not_collapsed(self):
        cases = (
            [_row("bluetooth", "true", "bool")],
            [_row("bluetooth", "false", "bool"), _row("bluetooth_version", "5.0")],
            [_row("bluetooth", "true", "bool"), _row("bluetooth_version", "5.0", conflict=True)],
            [_row("bluetooth", "true", "bool"), _row("bluetooth_version", "5.0", full=False)],
        )
        self.assertEqual(attribute_projection.project_final_rows(cases[0])[0]["resolved"]["display_value"], "Есть")
        for rows in cases:
            with self.subTest(rows=rows):
                projected = attribute_projection.project_final_rows(rows)
                self.assertEqual(len(projected), len(rows))
                self.assertEqual(projected[0]["resolved"]["selected_value"], rows[0]["resolved"]["selected_value"])

    def test_real_lg_card_and_export_retain_raw_audit(self):
        if not DATABASE.exists():
            self.skipTest("Saved working batch unavailable")
        final = attribute_projection.final_attribute_rows(DATABASE, 14)
        selected = {r["normalized_name"]: r for r in final}
        self.assertIn("bluetooth", selected)
        self.assertNotIn("bluetooth_version", selected)
        self.assertEqual(selected["bluetooth"]["resolved"]["display_value"], "4.0")
        raw = {r["normalized_name"]: r for r in jobs.comparison_rows(DATABASE, 14)}
        self.assertIn("bluetooth_version", raw)
        self.assertEqual(raw["bluetooth"]["sources"]["lg_kz"]["raw_value"], "●")
        book = load_workbook(BytesIO(exporter.export_batch(DATABASE, BATCH_ID)), read_only=True)
        try:
            product_sheet = next(sheet for sheet in book if
                                 any(row[3] == "ON66" for row in sheet.iter_rows(min_row=2, values_only=True)
                                     if len(row) > 3))
            headers = [cell.value for cell in product_sheet[1]]
            self.assertEqual(headers.count("Bluetooth"), 1)
            audit = next(sheet for sheet in book if
                         sheet.title == "Проверка источников")
            observed = [row for row in audit.iter_rows(min_row=2, values_only=True)
                        if row[1] == "ON66" and row[2] in {"Bluetooth", "Версия Bluetooth"}]
            self.assertEqual(len(observed), 2)
            self.assertTrue(any(row[3] == "●" for row in observed))
            self.assertTrue(any(row[3] == "ver 4.0" for row in observed))
        finally:
            book.close()


    def test_ordinary_product_page_uses_projection(self):
        if not DATABASE.exists():
            self.skipTest("Saved working batch unavailable")
        with TemporaryDirectory() as temporary:
            root = Path(temporary)
            with closing(connect(f"file:{DATABASE.as_posix()}?mode=ro", uri=True)) as source:
                with closing(connect(root / "batches.sqlite3")) as target:
                    source.backup(target)
            with TestClient(create_app(root, start_worker=False)) as client:
                response = client.get("/products/14")
                self.assertEqual(response.status_code, 200)
                self.assertIn("Bluetooth", response.text)
                self.assertIn("4.0", response.text)
                # The user-facing value is the meaningful refinement; the
                # original presence marker and version remain in evidence.
                from bs4 import BeautifulSoup
                soup = BeautifulSoup(response.text, "html.parser")
                row = next(tr for tr in soup.select("#comparison tbody tr")
                           if tr.find("th") and tr.find("th").text.strip() == "Bluetooth")
                self.assertIn("4.0", row.find_all("td")[0].text)
                self.assertNotIn("●", row.find_all("td")[0].text)
                evidence = row.select_one(".evidence-detail").text
                self.assertIn("●", evidence)
                self.assertIn("ver 4.0", evidence)


if __name__ == "__main__":
    unittest.main()



