"""Stage 56: delayed support discovery and dealer columns are evidence-driven."""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from product_tool import exporter,jobs,lg_discovery_pipeline,storage,worker
from product_tool.adapters.common import RawAttribute,SourceDocument
from product_tool.adapters.lg import LGAdapter,LGRUAdapter

class Stage56Tests(unittest.TestCase):
    def setUp(self):
        self.temp=TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.db=Path(self.temp.name)/"jobs.sqlite3"
        jobs.initialize(self.db)
        with storage._connection(self.db) as connection:
            connection.execute("INSERT INTO batches VALUES ('b','sample.xlsx','Items','{}','2026-01-01')")
            self.pid=connection.execute(
                "INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) "
                "VALUES ('b',2,'Sample','LG','ZX123','','TV',0,'[]','{}')").lastrowid

    def test_pdp_fallback_precedes_regional_support_search(self):
        order=[]
        calls=[]
        class Browser:
            trace_callback=None
            def _trace(self,**event):pass
            def search_google(self,*args):return None
            def close(self):pass
        class Dealer:
            source_key="sulpak";site_name="Sulpak"
            def find_source(self,*args,**kwargs):
                return SourceDocument("sulpak","Sulpak","",match_level="unknown")
        class DNS:
            def find_source(self,*args,**kwargs):
                return SourceDocument("dns","DNS","",match_level="dealer_url_needed")
        browser=Browser()
        kz=LGAdapter(http=object(),clock=lambda:0,browser_search=browser)
        ru=LGRUAdapter(http=object(),clock=lambda:0,browser_search=browser)
        def regional(key):
            def find(*args,**kwargs):
                calls.append(kwargs)
                order.append(key)
                return SourceDocument("lg_"+key,key,"",match_level="mismatch")
            return find
        def sitemap(*args,**kwargs):
            order.append("sitemap")
            return lg_discovery_pipeline.SitemapFallbackOutcome()
        def web(*args,**kwargs):
            order.append("web")
            return lg_discovery_pipeline.GoogleFallbackOutcome()
        def support(doc,*args,**kwargs):
            order.append("support")
            return doc
        jobs.enqueue(self.db,self.pid,[1])
        with patch.object(kz,"find_source",side_effect=regional("kz")), \
             patch.object(ru,"find_source",side_effect=regional("ru")), \
             patch.object(lg_discovery_pipeline,"run_sitemap_fallback",side_effect=sitemap), \
             patch.object(lg_discovery_pipeline,"run_web_fallback",side_effect=web), \
             patch.object(worker,"_augment_with_browser_search",side_effect=support):
            self.assertTrue(worker.run_once(self.db,
                adapter_factory=lambda:(kz,ru,Dealer(),browser),
                lg_sitemap_discovery_factory=lambda:object(),
                dns_adapter_factory=DNS,clock=lambda:0))
        self.assertEqual(order,["kz","ru","sitemap","web","support","support"])
        self.assertEqual([c["support_search"] for c in calls],[False,False])

    def test_sulpak_excel_column_requires_a_fact(self):
        def headers():
            book=load_workbook(BytesIO(exporter.export_batch(self.db,"b")),read_only=True)
            try:return [cell.value for cell in book["Проверка источников"][1]]
            finally:book.close()
        self.assertNotIn("Sulpak",headers())
        jobs.save_source_document(self.db,self.pid,SourceDocument(
            "sulpak","Sulpak","https://example.org/item",match_level="full_sku",
            attributes=[RawAttribute("Power","100")]))
        self.assertIn("Sulpak",headers())

if __name__=="__main__":unittest.main()
