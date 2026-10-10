"""Frozen real-page contracts plus adversarial configuration/component boundaries."""
import hashlib,json,sqlite3,tempfile,unittest
from pathlib import Path
from collections import defaultdict
from io import BytesIO
from unittest.mock import patch
from product_tool.adapters.published_page import embedded_objects,published_links,two_column_specs
from product_tool.adapters.xiaomi import XiaomiAdapter,manual_assessment
from product_tool.adapters.common import SourceDocument,RawAttribute,PhotoCandidate
from product_tool.xiaomi_page import parse_page,charging_fields
from product_tool.xiaomi_identity import model_matches,region_matches,fact_scope
from product_tool.normalization import normalize_facts
from product_tool import jobs,card_evidence,worker,readiness
from product_tool.xiaomi_pipeline import card_readiness
R=Path(__file__).resolve().parents[1]/'reports/xiaomi_stage74'

def row(label,value):
    return '<div platform="pc"><span class="xm-text f-regular" data-key="spec_1">'+label+'</span><span class="xm-text f-regular" data-key="spec_2">'+value+'</span></div>'
def page(body,title='Xiaomi 15 Specs - Xiaomi Global'):
    return '<title>'+title+'</title>'+body
def parsed(body,config='region=Global',url='https://www.mi.com/global/product/xiaomi-15/specs/'):
    return parse_page(page(body),url,'Xiaomi 15','Xiaomi 15 ['+config+']')

class Stage74Tests(unittest.TestCase):
    def test_frozen_inputs_and_category_mix(self):
        raw=(R/'dataset.json').read_bytes();self.assertEqual(hashlib.sha256(raw).hexdigest(),(R/'dataset.sha256').read_text().strip())
        rows=json.loads(raw)['rows'];self.assertEqual(len(rows),10);self.assertNotIn('Xiaomi 14',[x['model'] for x in rows]);self.assertEqual(len({x['model'] for x in rows}),10);self.assertFalse(any('url' in x for x in rows))
    def test_ordinary_baseline_was_unseeded_and_empty(self):
        b=json.loads((R/'baseline.json').read_text(encoding='utf8'));self.assertEqual(b['facts'],[]);self.assertEqual(b['trace'],[]);self.assertEqual(b['job']['status'],'error')
    def test_exact_line_generation_and_connectivity(self):
        for wrong in ['Xiaomi 15 Ultra Specs','Redmi 15 Specs','POCO 15 Specs','Xiaomi 15 Pro Specs','Xiaomi 15 5G Specs']:
            self.assertFalse(model_matches('Xiaomi 15',wrong))
        self.assertTrue(model_matches('Xiaomi Router AX3000T','All Specs, Features of Xiaomi Router AX3000T | Xiaomi'))
    def test_ram_storage_options_never_exact_for_two_queries(self):
        for variant in ['RAM=12GB;storage=256GB;region=Global','RAM=12GB;storage=512GB;region=Global']:
            d,e=parsed(row('Storage & RAM','12GB + 256GB | 12GB + 512GB'),variant);self.assertFalse(d.attributes);self.assertEqual(e['identity']['configuration_relation'],'partial');self.assertIsNone(e['identity']['retail_sku'])
    def test_color_and_color_dependent_dimensions_require_relation(self):
        for color in ['Black','White']:
            d,e=parsed(row('Color','Black White')+row('Dimensions','Black 191g Liquid Silver 192g'),'color='+color+';region=Global');self.assertFalse(d.attributes);self.assertEqual(len(e['rejected_specs']),2)
    def test_global_cn_and_eea_never_transfer_nfc_bands_charger(self):
        body=row('NFC','Yes')+row('Network & Connectivity','LTE B20')+row('Package Contents','Adapter included')+row('Processor','Snapdragon')
        for market in ['CN','EEA','de']:
            d,e=parsed(body,'region='+market);self.assertEqual([x.name for x in d.attributes],['Processor']);self.assertEqual(len(e['rejected_specs']),3)
        self.assertFalse(region_matches('EEA','global'));self.assertFalse(region_matches('EEA','de'))
    def test_explicit_vary_by_markets_is_candidate(self):
        d,e=parsed(row('Broadcasting system','DVB-T2 *Vary by markets'));self.assertFalse(d.attributes)
    def test_display_color_count_is_not_retail_color(self):
        d,e=parsed(row('Colours','16.7 million')+row('Color gamut','DCI-P3 94%'));self.assertEqual(len(d.attributes),2)
    def test_opaque_number_not_promoted_to_sku(self):
        d,e=parse_page(page(row('Model','24129PN74G')),'https://www.mi.com/global/product/xiaomi-15/specs/','24129PN74G','Xiaomi 15 [region=Global]');self.assertIsNone(e['identity']['retail_sku']);self.assertEqual(e['identity']['identifier_kind'],'opaque_code');self.assertEqual(e['model_numbers'][0]['kind'],'hardware_model_number')
    def test_bundle_station_not_transferred(self):
        body='<div><div><span class="xm-text f-bold" data-key="spec_0">Omni Station</span></div><div>'+row('Station Size','340×420×563mm')+'</div></div>'
        d,e=parsed(body,'bundle=robot-only;region=Global');self.assertFalse(d.attributes);self.assertEqual(e['rejected_specs'][0]['scope'],'model')
    def test_mode_specific_maximums_remain_separate(self):
        d,e=parsed(row('Runtime (Eco)','Up to 60 minutes')+row('Runtime (Turbo)','Up to 10 minutes'));self.assertEqual([x.value for x in d.attributes],['Up to 60 minutes','Up to 10 minutes']);self.assertNotEqual(d.attributes[0].name,d.attributes[1].name)
    def test_wired_wireless_reverse_capacity_are_separate(self):
        fields=charging_fields(RawAttribute('Battery & Charging','5240mAh (typ)\n90W HyperCharge\n50W wireless HyperCharge\n10W reverse charging\n*chargers sold separately'))
        self.assertEqual([x.name for x in fields],['Battery capacity','Wired charging','Wireless charging','Reverse charging']);self.assertIn('(typ)',fields[0].value)
    def test_continuation_value_does_not_become_fake_row(self):
        body='<div><div><span class="xm-text f-bold" data-key="spec_1">Security</span></div><div><span class="xm-text f-bold" data-key="spec_2">Fingerprint</span><span class="xm-text f-regular" data-key="spec_3">Face unlock</span></div></div>'
        d,e=parsed(body);self.assertEqual(len(d.attributes),1);self.assertEqual(d.attributes[0].value,'Fingerprint\nFace unlock')
    def test_real_specs_all_ten_offline_no_false_duplicate_components(self):
        data=json.loads((R/'corrected_live/results.json').read_text(encoding='utf8'))
        for r in data['rows']:
            source=next(s for s in r['sources'] if s['source_key']=='xiaomi_model');capture=next(c for c in r['evidence']['captures'] if c['url']==source['url'])
            d,e=parse_page((R/'corrected_live/captures'/capture['file']).read_text(encoding='utf8'),source['url'],r['input']['article'],r['input']['name'],r['input']['category']);self.assertEqual(len(e['raw_specs']),len(e['accepted_specs'])+len(e['rejected_specs']))
            names=defaultdict(set)
            for f in normalize_facts(d.attributes):names[f.normalized_name].add(f.normalized_value)
            self.assertFalse(any(len(v)>1 for v in names.values()),r['id']);self.assertFalse(any(fact_scope(x.name)=='configuration' for x in d.attributes))
    def test_safety_qr_reference_is_not_a_guide(self):
        check,heads=manual_assessment(['Xiaomi 15\nSafety Information\nTo view the user guide, scan the QR code.\nPress power. Connect charging.'],['Xiaomi 15']);self.assertFalse(check['accepted']);self.assertFalse(heads)
    def test_russian_quick_guide_and_combined_english_guide(self):
        check,_=manual_assessment(['Xiaomi 15\nРуководство пользователя\nНажмите кнопку питания. Подключите зарядку. Следуйте инструкции.'],['Xiaomi 15']);self.assertTrue(check['accepted'])
        check,_=manual_assessment(['Xiaomi 15\nSafety Information','Xiaomi Quick Start Guide and Warranty Card\nPress power. Connect charging.'],['Xiaomi 15']);self.assertTrue(check['accepted'])
    def test_router_instruction_contract_accepts_operations(self):
        check,_=manual_assessment(['Router AX3000T User Manual\nConnect cable. Reset the router. Configure network.'],['Router AX3000T']);self.assertTrue(check['accepted'])
    def test_safety_language_does_not_become_guide_language(self):
        russian='Xiaomi 15\nСведения о безопасности\n'+('Это сведения о безопасности устройства для пользователя. '*20)
        english='Xiaomi Quick Start Guide and Warranty Card\n'+('Press the power button to turn on the phone. Connect the charger to the port. '*20)
        check,_=manual_assessment([russian,english],['Xiaomi 15']);self.assertTrue(check['accepted']);self.assertNotIn('ru',check['guide_languages']['present']);self.assertIn('en',check['guide_languages']['present'])
    def test_embedded_json_is_parsed_without_execution_and_links_not_constructed(self):
        html='<script>window.__PRELOADED_STATE__ = '+json.dumps({'cards':[{'text':'Example','link':'https://www.mi.com/global/product/published-only/'}]})+';</script><script>window.__BAD__ = alert(1);</script>'
        self.assertEqual(len(list(embedded_objects(html))),1);self.assertEqual(published_links(html,'https://www.mi.com/global/'),[('https://www.mi.com/global/product/published-only/','Example')])
    def test_worker_default_dispatch_and_optional_manual(self):
        with tempfile.TemporaryDirectory() as temp:
            db=Path(temp)/'batches.sqlite3';jobs.initialize(db)
            with sqlite3.connect(db) as c:
                c.execute("INSERT INTO batches VALUES ('xm','test.xlsx','Products','{}','2026-10-10')")
                c.execute("INSERT INTO products (id,batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (1,'xm',2,'Xiaomi 15 [region=Global]','Xiaomi','Xiaomi 15','','Смартфоны',0,'[]','{}')")
            c.close()
            class Adapter:
                extra_sources=[]
                def __init__(self,**kwargs):pass
                def find_source(self,article,**kw):
                    self.reports={article:{'identity':{'model':'Xiaomi 15','model_relation':'model_confirmed','configuration_relation':'configuration_confirmed'},'exact_photo_assets':['render'],'manual_status':'Проверена, не найдена'}}
                    return SourceDocument('xiaomi_model','Xiaomi Official','https://www.mi.com/global/product/test/specs/',found_model='Xiaomi 15',match_level='model_confirmed',attributes=[RawAttribute('Processor','Chip'),RawAttribute('Display','AMOLED'),RawAttribute('Battery capacity','5000mAh')],photo_candidates=[PhotoCandidate('https://i02.appmifile.com/render.jpg','render',kind='product_gallery')])
                def find_documents(self,*args):return []
            # Test selection evidence is explicitly controlled; no real requests.
            with patch('product_tool.xiaomi_pipeline.XiaomiAdapter',Adapter),patch('product_tool.xiaomi_pipeline.photo_metadata.inspect_saved_photo',return_value={'width':800,'height':800,'size_bytes':2000,'format':'PNG'}),patch('product_tool.xiaomi_pipeline.DnsAdapter') as dealer:
                dealer.return_value.find_source.return_value=SourceDocument('dns','DNS','',error='Not found')
                jobs.enqueue(db,1,[1,3,4,6]);worker.run_once(db)
            self.assertEqual(jobs.list_jobs(db,1)[0]['status'],'done');self.assertNotIn('manual_unverified',card_readiness(db,1)['blocking_gaps'])
    def test_ui_excel_show_candidates_and_russian_groups(self):
        from fastapi.testclient import TestClient
        from product_tool.web import create_app
        from product_tool.exporter import export_batch
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as temp:
            path=Path(temp);db=path/'batches.sqlite3';jobs.initialize(db)
            captured=json.loads((R/'corrected_live/results.json').read_text(encoding='utf8'))['rows']
            with sqlite3.connect(db) as c:
                c.execute("INSERT INTO batches VALUES ('xm74','test.xlsx','Products','{}','2026-10-10')")
                for r in captured:
                    p=r['input'];c.execute("INSERT INTO products (id,batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (?,'xm74',?,?,?,?,'',?,0,'[]','{}')",(r['id'],r['id']+1,p['name'],p['brand'],p['article'],p['category']))
            c.close()
            for r in captured:
                for source in r['sources']:
                    attrs=[RawAttribute(f['raw_name'],f['raw_value'],f['section']) for f in r['facts'] if f['source_key']==source['source_key']]
                    photos=[PhotoCandidate(f['url'],f['asset_key'],kind=f['kind'],excluded_reason=f['excluded_reason']) for f in r['photos'] if f['source_key']==source['source_key']]
                    doc=SourceDocument(source['source_key'],source['site_name'],source['url'],found_model=source['found_model'],match_level=source['match_level'],evidence=source['evidence'],attributes=attrs,photo_candidates=photos)
                    jobs.save_source_document(db,r['id'],doc)
                jobs.set_photo_selection(db,r['id'],[],mode='none');card_evidence.save(db,r['id'],'xiaomi',r['evidence']);jobs.resolve_product(db,r['id'])
            with TestClient(create_app(path,start_worker=False)) as client:
                response=client.get('/products/1');self.assertEqual(response.status_code,200);self.assertIn('Модель и конфигурация Xiaomi',response.text);self.assertIn('Аккумулятор',response.text);self.assertIn('Коммерческий SKU',response.text)
            book=load_workbook(BytesIO(export_batch(path/'batches.sqlite3','xm74')));self.assertIn('Xiaomi конфигурация',book.sheetnames);self.assertGreater(book['Xiaomi фото-кандидаты'].max_row,1);self.assertEqual(book['Фотографии'].max_row,1);book.close()
