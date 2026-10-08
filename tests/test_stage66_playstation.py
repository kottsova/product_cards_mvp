"""Hardware/order separation and adversarial PlayStation scoping controls."""
import json,tempfile,unittest
from unittest.mock import patch
from pathlib import Path
from product_tool.adapters.playstation import parse_page,parse_hardware_specs
from product_tool.playstation_identity import codes,hardware_family,document_type,official,model_key
from product_tool.resolution import resolve_attributes
from product_tool import jobs,storage,worker,card_evidence,readiness,exporter
from product_tool.adapters.common import SourceDocument,RawAttribute
R=Path(__file__).resolve().parents[1]/'reports/playstation_stage66'

class PlayStationScopes(unittest.TestCase):
    def test_combined_codes_and_regional_suffix(self):
        self.assertEqual(codes('CFI-2016A/B Quick Start'),{'CFI-2016A','CFI-2016B'})
        self.assertNotIn('CFI-2015A',codes('CFI-2016A/B'))
        self.assertNotIn('CFI-ZWH2EC',codes('CFI-ZWH2'))
    def test_chassis_routing_is_not_collapsed(self):
        self.assertEqual(hardware_family('https://www.playstation.com/ps5-docs/1000a/foo.pdf'),'ps5_original')
        self.assertEqual(hardware_family('https://www.playstation.com/ps5-docs/2000ab/foo.pdf'),'ps5_slim')
        self.assertEqual(hardware_family('https://www.playstation.com/ps5-docs/7000/foo.pdf'),'pro')
    def test_manual_types_remain_separate(self):
        self.assertEqual(document_type('CFI-2016 Safety Guide'),'Safety Guide')
        self.assertEqual(document_type('CFI-2016 Quick Start Guide'),'Quick Start')
        self.assertEqual(document_type('Руководство по эксплуатации CFI-ZCT1W'),'User Guide')
    def test_only_official_https_host(self):
        for u in ('https://playstation.com.evil.test/','https://evilplaystation.com/','https://user@playstation.com/','http://playstation.com/'):
            self.assertFalse(official(u))
    def test_bundle_requires_own_order_code(self):
        body=(R/'direct_bundle.html').read_text(encoding='utf-8');u='https://direct.playstation.com/en-gb/buy-consoles/fortnite-bundle'
        d,e=parse_page(body,u,'1000049849-GB','ps5');self.assertEqual(d.match_level,'full_sku')
        self.assertEqual(e['configuration_fields']['storage'],'825 GB');self.assertIn('voucher',e['configuration_fields']['bundle_contents'])
        other,oe=parse_page(body,u,'CFI-2016A','ps5');self.assertNotEqual(other.match_level,'full_sku');self.assertFalse(oe['configuration_fields']);self.assertFalse(oe['exact_photo_assets'])
    def test_gameplay_is_candidate_but_exact_bundle_packaging_is_allowed(self):
        body=(R/'direct_bundle.html').read_text(encoding='utf-8')
        d,e=parse_page(body,'https://direct.playstation.com/en-gb/buy-consoles/fortnite-bundle','1000049849-GB','ps5')
        self.assertTrue(e['exact_photo_assets']);self.assertTrue(any('gameplay' in p['url'] for p in e['photo_candidates']))
        self.assertFalse(any('gameplay' in u for u in d.photos))
        other,oe=parse_page(body,'https://direct.playstation.com/en-gb/buy-consoles/fortnite-bundle','CFI-2016A','ps5')
        self.assertFalse(oe['exact_photo_assets'])
    def test_black_order_not_white_order(self):
        body=(R/'direct_black.html').read_text(encoding='utf-8');u='https://direct.playstation.com/en-gb/buy-accessories/dualsense-black'
        d,e=parse_page(body,u,'1000050213-GB','dualsense');self.assertEqual(e['configuration_fields']['color'],'Чёрный (Midnight Black)')
        self.assertEqual(e['configuration_fields']['bundle_contents'],'DualSense® Wireless Controller; User manual')
        d,e=parse_page(body,u,'CFI-ZCT1W','dualsense');self.assertNotEqual(d.match_level,'full_sku');self.assertFalse(e['configuration_fields'])
    def test_gallery_dedup_and_named_features_not_false_conflicts(self):
        body=(R/'direct_black.html').read_text(encoding='utf-8')
        d,e=parse_page(body,'https://direct.playstation.com/en-gb/buy-accessories/dualsense-black','1000050213-GB','dualsense')
        self.assertEqual(len(d.photos),len(set(d.photos)))
        self.assertGreaterEqual(len(d.photos),2)
        self.assertFalse(any('bundle' in u.lower() or 'cable-and' in u.lower() for u in d.photos))
        self.assertTrue(any('USB-Bundle' in p['url'] for p in e['photo_candidates']))
        for label in ('Haptic Feedback','Adaptive Triggers'):
            values={a.value for a in d.attributes if a.name==label}
            self.assertLessEqual(len(values),1)
        self.assertTrue(e['raw_specs'])
    def test_negative_feature_is_not_confirmed_present(self):
        d,e=parse_page('<h1>DualSense Wireless Controller</h1><h3>Haptic Feedback</h3><p>Not supported with this device</p>','https://www.playstation.com/en-gb/accessories/dualsense/','CFI-ZCT1W','dualsense')
        self.assertFalse(d.attributes);self.assertTrue(e['configuration_candidates'])
    def test_other_family_and_support_are_rejected(self):
        body=(R/'pro.html').read_text(encoding='utf-8')
        for u,key in [('https://www.playstation.com/en-gb/ps5/pro/','dualsense'),('https://www.playstation.com/en-gb/support/ps5/pro/','pro')]:
            d,e=parse_page(body,u,'CFI-7021',key);self.assertTrue(d.error);self.assertFalse(d.attributes)
    def test_revision_dimensions_and_controller_not_mixed(self):
        vals=[]
        for code in ('CFI-1216A','CFI-1216B'):
            d,raw,c=parse_hardware_specs((R/(code+'_text.txt')).read_text(encoding='utf-8'),'https://www.playstation.com/ps5-docs/1200'+code[-1].lower()+'/safety.pdf',code)
            v={a.name:a.value for a in d.attributes};vals.append(v)
            self.assertNotIn('Ёмкость аккумулятора',v);self.assertNotIn('280',v['Вес товара']);self.assertTrue(raw)
        self.assertEqual(vals[0]['Высота товара, мм'],'104 мм');self.assertEqual(vals[1]['Высота товара, мм'],'92 мм')
        self.assertIn('3.9',vals[0]['Вес товара']);self.assertIn('3.4',vals[1]['Вес товара'])
    def test_combined_slim_table_withheld(self):
        d,raw,c=parse_hardware_specs((R/'CFI-2016A_B_text.txt').read_text(encoding='utf-8'),'https://www.playstation.com/ps5-docs/2000ab/safety.pdf','CFI-2016B')
        self.assertFalse(d.attributes);self.assertTrue(c)
    def test_package_and_stand_dimensions_are_not_console_dimensions(self):
        for label in ('External dimensions (packaging)','External dimensions with stand'):
            d,raw,c=parse_hardware_specs('Specifications\n'+label+'\n400 × 200 × 300 mm (width × height × depth)\n','https://www.playstation.com/ps5-docs/1200a/safety.pdf','CFI-1216A')
            self.assertFalse(d.attributes)
    def test_hardware_code_never_becomes_retail_sku(self):
        f=dict(source_key='playstation_hardware',normalized_name='storage',normalized_value='825 GB',unit='')
        r=resolve_attributes([f],[dict(source_key='playstation_hardware',match_level='hardware_confirmed',error='')])[0]
        self.assertEqual(r.status,'hardware_confirmed_official');self.assertFalse(r.full_sku_confirmed)
    def test_color_and_box_contents_do_not_fill_missing_technical_specs(self):
        from product_tool import playstation_pipeline as ps
        ev=dict(identity={'model':'model_confirmed'},configuration_complete=True,exact_photo_assets=['photo'],manual_status='Не проверена')
        values=[dict(normalized_name=n,status='full_sku_official',conflict=False) for n in ('color','комплектация','haptic_feedback','adaptive_triggers')]
        photos=[dict(selected=True,asset_key='photo',source_key='playstation',kind='product_gallery')]
        with patch.object(card_evidence,'load',return_value=ev),patch.object(jobs,'get_resolved',return_value=values),patch.object(jobs,'get_photo_candidates',return_value=photos),patch.object(jobs,'result_counts',return_value={'conflicts':0}):
            r=ps.card_readiness(None,1)
        self.assertEqual(r['confirmed_specs'],2);self.assertEqual(r['verdict'],'not_ready');self.assertEqual(r['blocking_gaps'],['specifications_missing'])
    def test_model_cannot_confirm_configuration(self):
        for name in ('color','storage','комплектация','product_dimensions','product_weight','объем_накопителя','product_dimensions__width'):
            f=dict(source_key='playstation_model',normalized_name=name,normalized_value='White',unit='')
            r=resolve_attributes([f],[dict(source_key='playstation_model',match_level='model_confirmed',error='')])[0]
            self.assertEqual(r.status,'needs_review');self.assertFalse(r.selected_value);self.assertFalse(r.full_sku_confirmed)
    def test_scope_has_same_common_pipeline_and_manual_advisory(self):
        class Adapter:
            extra_documents=[]
            reports={'CFI-ZCT1W':dict(identity={'model':'model_confirmed','configuration':'unproven','hardware':'official_hardware_model_code'},configuration_complete=False,manual_status='Не проверена',exact_photo_assets=[])}
            def find_source(self,*a,**k):return SourceDocument('playstation_model','PlayStation','https://www.playstation.com/en-gb/accessories/dualsense-wireless-controller/',match_level='model_confirmed',attributes=[RawAttribute('Haptic Feedback','Да')])
        class Dealer:
            def find_source(self,*a,**k):return SourceDocument('dns','DNS','',match_level='dealer_url_needed',evidence='Exact dealer URL needed')
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/'batches.sqlite3';jobs.initialize(p)
            with storage._connection(p) as db:
                db.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('ps66','Test','Products','{}',storage._now()))
                db.execute('INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (?,?,?,?,?,?,?,?,?,?)',('ps66',1,'DualSense','Sony Interactive Entertainment','CFI-ZCT1W','','controllers',0,'[]','{}'))
            pid=storage.get_batch(p,'ps66')['products'][0]['id'];jobs.enqueue(p,pid,[1,2,3,4,6])
            worker.run_once(p,playstation_adapter_factory=Adapter,dns_adapter_factory=Dealer)
            r=readiness.card_readiness(p,pid);self.assertEqual(r['verdict'],'not_ready');self.assertIn('configuration_unresolved',r['blocking_gaps']);self.assertNotIn('manual_unverified',r['blocking_gaps']);self.assertIn('manual_unverified',r['advisory_gaps'])
            self.assertEqual(jobs.list_jobs(p,pid)[0]['status'],'needs_review');self.assertTrue(jobs.get_facts(p,pid))
            from io import BytesIO
            from openpyxl import load_workbook
            b=load_workbook(BytesIO(exporter.export_batch(p,'ps66')));self.assertIn('Готовность PlayStation',b.sheetnames)

if __name__=='__main__':unittest.main()
