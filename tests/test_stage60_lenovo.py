"""Stage60 identity traps and configuration evidence safety; no real network."""
import json,unittest,tempfile
from pathlib import Path
from product_tool.adapters.lenovo import support_identity,split_specs,catalog_paths,LenovoAdapter,typed_manual
from product_tool.adapters.common import SourceDocument,RawAttribute,PhotoCandidate
from product_tool import jobs,storage,worker,card_evidence,lenovo_pipeline,attribute_projection,exporter
from product_tool.resolution import resolve_attributes
from openpyxl import load_workbook
from io import BytesIO

class LenovoIdentityTests(unittest.TestCase):
    def test_exact_mtm_requires_returned_configuration_body(self):
        info={'Model':'21ML005BUS','ProductName':'ThinkPad T14 Gen5','Description':'Processor: Core Ultra 5 135U: : Memory: 16 GB'}
        self.assertEqual(support_identity(info,'21ML005BUS')['relation'],'exact_mtm')
        self.assertFalse(support_identity(info,'21ML001UCK')['configuration_resolved'])
        self.assertFalse(support_identity({**info,'Model':'21MLZZZZUS','Description':'21ML'},'21MLZZZZUS')['configuration_resolved'])
        self.assertFalse(support_identity({**info,'Model':'21ML','Description':'21ML'},'21ML005BUS')['configuration_resolved'])
    def test_cto_does_not_identify_ordered_configuration(self):
        info={'Model':'21MLCTO1WW','Description':'Processor: Core Ultra 5: : Memory: 16 GB'}
        self.assertFalse(support_identity(info,'21MLCTO1WW')['configuration_resolved'])
    def test_family_multi_configuration_options_never_become_product_facts(self):
        info={'Description':'Processor: Intel Core Ultra 5 or Ultra 7: : Memory: up to 64 GB: : Display: optional OLED: : Weight: starting at 1.4 kg: : Hard Drive: 512 GB'}
        facts,candidates=split_specs(info,'family_model')
        self.assertFalse(facts);self.assertEqual(len(candidates),5)
        facts,candidates=split_specs(info,'exact_mtm')
        self.assertEqual([f.name for f in facts],['Hard Drive']);self.assertEqual(len(candidates),4)
    def test_configuration_fields_not_support_service_content(self):
        facts,_=split_specs({'Description':'Processor: Core i7: : Memory: 16 GB: : Included Warranty: 3YR Onsite: : Wired Network: 1x: : BIOS: update'},'exact_mtm')
        self.assertEqual([f.name for f in facts],['Processor','Memory'])
    def test_optional_ports_configured_is_installed_not_family_choice(self):
        facts,_=split_specs({'Description':'Optional Ports (configured): 1x Nano-SIM'},'exact_mtm')
        self.assertEqual(len(facts),1)
    def test_catalog_slug_uses_official_hierarchy_not_model_manifest(self):
        tree=[{'n':'Laptops and netbooks','k':'a','o':[{'n':'ThinkPad T Series laptops','k':'b','o':[{'n':'ThinkPad T14 Gen 5 Type 21ML 21MM','k':'c','o':[{'n':'21ML','k':'d'}]}]}]}]
        paths=catalog_paths('var e ='+json.dumps(tree)+';', '21ML005BUS')
        self.assertEqual(paths[0]['path'],'laptops-and-netbooks/thinkpad-t-series-laptops/thinkpad-t14-gen-5-type-21ml-21mm/21ml')
        self.assertFalse(catalog_paths('var e ='+json.dumps(tree),'21MM005BUS'))
    def test_family_store_does_not_promote_request_or_jsonld_family_mpn(self):
        class Response:
            text='<script type="application/ld+json">{"@type":"Product","mpn":"LEN101T0089"}</script>'
            url='https://www.lenovo.com/us/en/p/len101t0089'
        adapter=LenovoAdapter(session=object())
        r=adapter.store_evidence(Response(),'21ML005BUS')
        self.assertEqual(r['relation'],'family_model');self.assertFalse(r['configuration_confirmed'])
    def test_hmm_is_not_user_guide(self):
        self.assertEqual(typed_manual('Hardware Maintenance Manual'),'Hardware Maintenance Manual')
        self.assertEqual(typed_manual('(Russian) User Guide'),'User Guide')
    def test_resolution_rejects_family_and_error_pages(self):
        fact={'source_key':'lenovo_support','normalized_name':'memory','normalized_value':'16 GB','unit':'text'}
        for page in ({'match_level':'base_model','error':''},{'match_level':'full_sku','error':'blocked'}):
            result=resolve_attributes([fact],[{'source_key':'lenovo_support',**page}])[0]
            self.assertFalse(result.full_sku_confirmed)
        self.assertTrue(resolve_attributes([fact],[{'source_key':'lenovo_support','match_level':'full_sku','error':''}])[0].full_sku_confirmed)
    def test_shared_old_parser_accepts_only_requested_official_site(self):
        from types import SimpleNamespace
        from urllib.parse import quote
        from product_tool.census.old_parser_google import OldParserGoogleBrowser
        from product_tool.census.browser_runtime import BrowserFailure
        driver=OldParserGoogleBrowser(Path('unused'))
        driver.page=SimpleNamespace(url='https://www.google.com/search')
        driver.search_function=lambda page,query:[('https://psref.lenovo.com/Detail/test?M=21ML005BUS','Lenovo 21ML005BUS')]
        driver.candidate_function=lambda link,title,domain,brand,name,article:domain=='psref.lenovo.com' and article=='21ML005BUS'
        result=driver.call('goto',url='https://www.google.com/search?q='+quote('site:psref.lenovo.com "21ML005BUS"'),search_result_hosts=['psref.lenovo.com'])
        self.assertTrue(result['projection']['fragments'][0]['old_parser_candidate'])
        with self.assertRaises(BrowserFailure):
            driver.call('goto',url='https://www.google.com/search?q='+quote('site:evil.test "21ML005BUS"'),search_result_hosts=['psref.lenovo.com'])
    def test_ui_identification_uses_exact_lenovo_not_lg_label(self):
        self.assertIn('Точный MTM Lenovo',jobs.identification_status([{'source_key':'lenovo_support','url':'https://pcsupport.lenovo.com/model','match_level':'full_sku','error':''}]))
    def test_family_photos_and_unverified_manuals_are_never_verified(self):
        self.assertFalse(lenovo_pipeline.photo_verified({'asset_key':'photo','kind':'product_gallery'},{'photos_relation':'family_model'}))
        self.assertFalse(lenovo_pipeline.document_verified({'direct_url':'manual'},{'manuals':[{'url':'manual','type':'Hardware Maintenance Manual'}]}))

class LenovoWorkerTests(unittest.TestCase):
    def test_shared_worker_storage_russian_projection_and_candidate_export(self):
        with tempfile.TemporaryDirectory() as folder:
            db=Path(folder)/'audit.sqlite3';jobs.initialize(db)
            with storage._connection(db) as conn:
                conn.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('lenovo-test','test','Products','{}',storage._now()))
                pid=conn.execute('INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (?,?,?,?,?,?,?,?,?,?)',('lenovo-test',1,'ThinkPad','Lenovo','21ML005BUS','','Laptops',0,'[]','{}')).lastrowid
            class Adapter:
                source_key='lenovo_support'
                reports={'21ML005BUS':{'configuration_candidates':[{'section':'Configuration','raw_label':'Weight','value':'starting at 1.4 kg','scope':'family_option','reason':'not exact'}],'manuals':[],'manual_status':'Не проверена'}}
                def find_source(self,article,**kwargs):return SourceDocument('lenovo_support','Lenovo Support','https://pcsupport.lenovo.com/model',found_model=article,match_level='full_sku',attributes=[RawAttribute('Processor','Core Ultra 5','Configuration'),RawAttribute('Memory','16 GB','Configuration')],photo_candidates=[PhotoCandidate('https://download.lenovo.com/family.png','family','product_gallery')])
                def find_documents(self,*args):return []
            class Dealer:
                def find_source(self,*args,**kwargs):return SourceDocument('dns','DNS','',match_level='dealer_url_needed')
            jobs.enqueue(db,pid,[1,2,3,4,6])
            self.assertTrue(worker.run_once(db,lenovo_adapter_factory=Adapter,dns_adapter_factory=Dealer))
            self.assertEqual(jobs.list_jobs(db,pid)[0]['status'],'done')
            rows=attribute_projection.final_attribute_rows(db,pid)
            self.assertEqual({r['display_name'] for r in rows},{'Процессор','Оперативная память'})
            self.assertEqual(lenovo_pipeline.card_readiness(db,pid)['verdict'],'not_ready')
            book=load_workbook(BytesIO(exporter.export_batch(db,'lenovo-test')))
            self.assertIn('Конфигурации-кандидаты',book.sheetnames)
            self.assertEqual(book['Фотографии'].max_row,1)
            self.assertGreater(book['Фото-кандидаты'].max_row,1)
            self.assertNotIn('Вес товара, кг',[c.value for c in book['Laptops'][1]])
            self.assertEqual(len(jobs.get_facts(db,pid)),2)

if __name__=='__main__':unittest.main()
