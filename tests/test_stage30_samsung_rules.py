"""Stage 30: the owner's decisions after Stage 29, offline on the saved responses (tests/_samsung_replay.py). No socket is opened.

  1. a multilingual leaflet with a short Russian section is kept as a BRIEF GUIDE (pages, "model not named in the text"), never a full Russian instruction, and it does not close the full-manual gap;
  2. the page's own single `modelCode`/`shopSKU` product data, exactly equal to the catalog code with no contradicting code, is strong evidence of the variant;
  3. Fit 3 is not chosen (a page in the sitemap does not prove a novelty);
  and the dealer link request for a Samsung row stays open until a FULL Russian instruction is confirmed.
"""
from __future__ import annotations

import gzip
import io
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

import _samsung_replay as R
from product_tool import card_evidence, jobs, samsung_pipeline, samsung_readiness
from product_tool.adapters import samsung
from product_tool.adapters.samsung_source import levels_note, match_level_of

ROOT = Path(__file__).resolve().parents[1]
S29 = ROOT / "reports/source_census_2026-09-26_stage29"
S30 = ROOT / "reports/source_census_2026-09-26_stage30"


def page(fragment: str) -> str:
    return R.PAGES[R.page_url(fragment)][1]


def leaflet_pages() -> list[str]:
    record = next(r for r in json.loads((S29 / "docs_extract/index.json").read_text(encoding="utf-8")) if "Installation_Guide" in r["url"])
    with gzip.open(S29 / "docs_extract" / record["text_file"], "rt", encoding="utf-8") as handle:
        return handle.read().split("\n\f\n")


class ProductDataCodeIsStrongVariantEvidence(unittest.TestCase):
    def identity(self, fragment: str, article: str):
        return samsung.extract_identity(page(fragment), R.page_url(fragment), article)

    def test_the_external_ssd_and_the_hob_are_confirmed_by_their_own_product_data(self):
        for fragment, article in (("portable-ssd-t7", "MU-PC1T0H/WW"), ("hob-nz64t3506ak-wt", "NZ64T3506AK/WT")):
            with self.subTest(article):
                identity = self.identity(fragment, article)
                self.assertEqual((identity.level, identity.evidence_strength, match_level_of(identity)), ("full_sku", "strong", "full_sku"))
                self.assertIn("product_data_code", identity.evidence)
                self.assertEqual(samsung.product_data_codes(page(fragment)), {samsung.norm(article)})

    def test_a_page_whose_data_differs_from_the_catalog_code_stays_open(self):
        fridge = self.identity("rbf310g-3050-310l-silver-rb31fernds", "RB31FERNDSA")
        self.assertEqual((match_level_of(fridge), "product_data_code" in fridge.evidence), ("code_in_page_text", False))     # the page's data says RB31FERNDSA/WT
        self.assertEqual(samsung.product_data_codes(page("rbf310g-3050-310l-silver-rb31fernds")), {"rb31ferndsawt"})
        phone = self.identity("galaxy-a37-5g-awesome-graygreen-256gb-sm-a376edgg", "SM-A376EZAGINS")
        self.assertEqual((phone.level, "product_data_code" in phone.evidence), ("base_model", False))

    def test_a_second_code_or_a_conflicting_sku_takes_the_confirmation_away(self):
        html = page("hob-nz64t3506ak-wt")
        extra = html.replace("</body>", '<input type="hidden" name="apiChangeModelCode" value="NZ64T3506AK/OTHER"/></body>')
        self.assertEqual(match_level_of(samsung.extract_identity(extra, R.page_url("hob-nz64t3506ak-wt"), "NZ64T3506AK/WT")), "code_in_page_text")
        self.assertEqual(len(samsung.product_data_codes(extra)), 2)
        # a template placeholder is not data
        self.assertEqual(samsung.product_data_codes('<input name="modelCode" value="{{item.modelCode}}"/><a data-model-code="{{x}}"></a>'), set())
        # no data at all: the text alone stays text-only
        bare = html
        for marker in ('name="modelCode"', 'name="apiChangeModelCode"', 'id="shopSKU"', 'name="apiChangeShopSKU"', 'name="originShopSku"', "data-model-code=", "data-shop-sku="):
            bare = bare.replace(marker, marker.replace("=", "_x=", 1) if "=" in marker else marker + "_x")
        self.assertEqual(samsung.product_data_codes(bare), set())
        self.assertEqual(match_level_of(samsung.extract_identity(bare, R.page_url("hob-nz64t3506ak-wt"), "NZ64T3506AK/WT")), "code_in_page_text")


class BriefGuideIsKeptAsSuch(unittest.TestCase):
    def test_the_leaflet_is_a_brief_guide_with_its_pages_and_is_not_a_russian_instruction(self):
        facts = samsung.assess_samsung_document(leaflet_pages(), "MZ-77E500BW", "MZ-77E500", "full_sku")
        self.assertEqual(facts["russian_section"]["pages"], [8])
        self.assertTrue(samsung.is_brief_guide(facts))
        self.assertFalse(facts["russian_by_text"])
        acceptance = samsung.instruction_acceptance(facts, {}, "MZ-77E500BW", "MZ-77E500")
        self.assertEqual((acceptance["basis"], acceptance["accepted"]), ("brief_guide_russian_section", False))
        for text in ("краткая памятка", "стр. 8", "модель в тексте не названа", "полной русской инструкцией не считается"):
            self.assertIn(text, acceptance["note"])
        self.assertIn("русский раздел есть", levels_note({**facts, "acceptance": acceptance}))
        self.assertNotIn("не русский по тексту", levels_note({**facts, "acceptance": acceptance}))

    def test_a_short_russian_legal_notice_in_an_english_manual_is_not_a_brief_guide(self):
        english = ["Installation guide. Please read this manual and the safety instructions before installing. " * 30]
        notice = ["Уведомление о соответствии: изделие произведено в соответствии с требованиями. " * 8]
        facts = samsung.assess_samsung_document(english + notice, "MZ-77E500BW", "MZ-77E500", "full_sku")
        self.assertFalse(samsung.is_brief_guide(facts))
        self.assertEqual(samsung.instruction_acceptance(facts, {}, "MZ-77E500BW", "MZ-77E500")["basis"], "not_a_russian_instruction_by_text")

    def test_a_full_russian_instruction_is_still_russian_and_not_a_brief_guide(self):
        russian = ["Руководство пользователя. Перед использованием внимательно прочитайте меры предосторожности. Подключите изделие и нажмите кнопку. Не допускайте попадания воды. " * 60]
        facts = samsung.assess_samsung_document(russian, "MZ-77E500BW", "MZ-77E500", "full_sku")
        self.assertTrue(facts["russian_by_text"])
        self.assertFalse(samsung.is_brief_guide(facts))

    def test_the_gap_texts_do_not_say_all_files_are_not_russian_or_that_the_manual_was_found(self):
        gap = samsung_readiness.GAP_TEXT
        self.assertNotEqual(gap["instruction_full_russian_manual_missing"], gap["instruction_language_not_russian"])
        self.assertIn("краткая памятка", gap["instruction_full_russian_manual_missing"])
        self.assertIn("русский раздел", gap["instruction_full_russian_manual_missing"])
        self.assertIn("не закрывает пробел по полной русской инструкции", gap["instruction_brief_guide_russian_section"])
        self.assertNotIn("найден", gap["instruction_full_russian_manual_missing"])
        self.assertIn("русского раздела в них нет", gap["instruction_language_not_russian"])          # said only when there is none


class OrdinaryPathAfterTheDecisions(unittest.TestCase):
    CARDS = (("Внутренние SSD-накопители", "MZ-77E500BW"), ("Flash-накопители", "MUF-64BE3/APC"), ("Внешние SSD-накопители", "MU-PC1T0H/WW"), ("Карты памяти", "MB-MC64HARU"),
             ("Варочные панели", "NZ64T3506AK/WT"), ("Холодильники", "RB31FERNDSA"), ("Телевизоры", "QE48S85HAEXCE"))

    @classmethod
    def setUpClass(cls):
        cls.tmp = TemporaryDirectory()
        cls.batch = R.run_products(Path(cls.tmp.name), cards=cls.CARDS)
        cls.by = {o["article"]: o for o in cls.batch["outcomes"]}
        cls.readiness = {a: samsung_readiness.card_readiness(cls.batch["database"], o["product_id"]) for a, o in cls.by.items()}

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def dealer(self, article):
        return next(s for s in jobs.get_source_pages(self.batch["database"], self.by[article]["product_id"]) if s["source_key"] == "dns")

    def test_the_external_ssd_and_the_hob_are_now_confirmed_and_their_photos_selected(self):
        ssd, hob = self.readiness["MU-PC1T0H/WW"], self.readiness["NZ64T3506AK/WT"]
        self.assertEqual((ssd["page_match_level"], ssd["verdict"], ssd["official_photos_selected"], ssd["official_photos"]), ("full_sku", "export_ready_with_gaps", 12, 12))
        self.assertEqual(ssd["blocking_gaps"], ["instruction_missing"])
        self.assertEqual((hob["page_match_level"], hob["verdict"]), ("full_sku", "export_ready_with_gaps"))
        self.assertEqual(hob["blocking_gaps"], ["instruction_text_not_extractable_manual_check"])       # the scanned instruction still needs a person
        self.assertEqual((self.by["MU-PC1T0H/WW"]["status"], self.by["NZ64T3506AK/WT"]["status"]), ("done", "done"))
        self.assertEqual(self.readiness["RB31FERNDSA"]["verdict"], "needs_verification")               # the /WT suffix keeps the fridge open
        self.assertEqual(self.by["RB31FERNDSA"]["status"], "needs_review")

    def test_the_brief_guide_is_saved_with_pages_and_keeps_the_full_manual_gap_open(self):
        r = self.readiness["MZ-77E500BW"]
        self.assertEqual((r["instruction"]["saved"], r["instruction"]["russian_saved"], r["instruction"]["accepted"]), (1, 0, False))
        self.assertEqual([b["pages"] for b in r["instruction"]["brief_guides"]], [[8]])
        self.assertEqual(r["blocking_gaps"], ["instruction_full_russian_manual_missing"])
        self.assertIn("instruction_brief_guide_russian_section", r["advisory_gaps"])
        self.assertNotIn("instruction_language_not_russian", r["gaps"])
        self.assertIn("стр. 8", r["gap_reasons"]["instruction_brief_guide_russian_section"])
        self.assertIn("модель в тексте не названа", r["gap_reasons"]["instruction_brief_guide_russian_section"])
        documents = jobs.get_documents(self.batch["database"], self.by["MZ-77E500BW"]["product_id"])
        self.assertEqual(len(documents), 1)
        self.assertNotEqual(documents[0]["language"], "Русский")
        self.assertTrue(documents[0]["title"].startswith("Краткая памятка Samsung"))
        book = load_workbook(io.BytesIO(self.batch["export"]))
        self.assertTrue(any("Краткая памятка" in str(row[1]) and "стр. 8" in str(row[1]) for row in book["Инструкции"].iter_rows(min_row=2, values_only=True)))

    def test_the_dealer_link_request_stays_open_until_a_full_russian_instruction_is_confirmed(self):
        for article in ("MZ-77E500BW", "MUF-64BE3/APC", "MU-PC1T0H/WW", "MB-MC64HARU", "NZ64T3506AK/WT", "RB31FERNDSA"):
            with self.subTest(article):
                source = self.dealer(article)
                self.assertEqual(source["match_level"], "dealer_url_needed")
                self.assertIn("полная русская инструкция", source["evidence"])
                self.assertIn("не будет угадана", source["evidence"])
        self.assertEqual(self.dealer("QE48S85HAEXCE")["match_level"], "not_needed")                      # an accepted Russian instruction closes it
        self.assertEqual(self.batch["dealer_session"].calls, [])                                        # a request text is not a request
        self.assertEqual(self.batch["replay"].refused, [])

    def test_the_brief_guide_and_an_english_file_do_not_count_as_a_full_instruction(self):
        for article, expected in (("MZ-77E500BW", False), ("QE48S85HAEXCE", True), ("RB31FERNDSA", False)):
            self.assertEqual(samsung_readiness.full_russian_instruction_confirmed(self.batch["database"], self.by[article]["product_id"]), expected, article)
        self.assertEqual(samsung_pipeline.dealer_missing_fields(self.batch["database"], self.by["QE48S85HAEXCE"]["product_id"], [1, 2, 3, 4, 6]), [])
        self.assertEqual(samsung_pipeline.dealer_missing_fields(self.batch["database"], self.by["MZ-77E500BW"]["product_id"], [1, 2, 3, 4, 6]), ["полная русская инструкция"])
        self.assertEqual(samsung_pipeline.dealer_missing_fields(self.batch["database"], self.by["MZ-77E500BW"]["product_id"], [1, 2, 3, 4]), [])      # the stage was not requested

    def test_only_the_saved_official_addresses_were_used(self):
        for url in self.batch["replay"].calls:
            self.assertRegex(url, r"^https://(?:www\.samsung\.com/kz_ru/|org\.downloadcenter\.samsung\.com/downloadfile/ContentsFile\.aspx\?)")


class BeforeAfterRecord(unittest.TestCase):
    RECORD = json.loads((S30 / "raw/before_after.json").read_text(encoding="utf-8"))["cards"]

    def test_nineteen_cards_and_only_these_changed(self):
        self.assertEqual(len(self.RECORD), 19)
        changed = {r["article"]: r["changed"] for r in self.RECORD if r["changed"]}
        self.assertEqual(sorted(changed), sorted(["MU-PC1T0H/WW", "NZ64T3506AK/WT", "MZ-77E500BW", "RB31FERNDSA"]))
        self.assertEqual(changed["RB31FERNDSA"], ["dealer_link_request"])
        self.assertIn("variant", changed["MU-PC1T0H/WW"])
        self.assertIn("document", changed["MZ-77E500BW"])

    def test_the_offline_run_made_no_dealer_request_and_no_unsaved_request(self):
        calls = json.loads((S30 / "raw/replay_calls.json").read_text(encoding="utf-8"))
        self.assertEqual((calls["dealer_session_calls"], calls["refused"]), ([], []))


class WearablesNotChosen(unittest.TestCase):
    DECISION = json.loads((S30 / "raw/wearables_decision.json").read_text(encoding="utf-8"))

    def test_fit3_and_watch7_are_not_chosen_and_nothing_was_requested(self):
        self.assertEqual((self.DECISION["requests_made"], self.DECISION["watch"]["chosen"], self.DECISION["band"]["chosen"]), (0, False, False))
        self.assertIn("не доказывает новизну", self.DECISION["band"]["reason"])
        self.assertEqual(self.DECISION["group_5b_run"], False)


if __name__ == "__main__":
    unittest.main()
