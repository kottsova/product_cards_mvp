"""Real public Xbox captures and negative controls for scoped identity transfer."""
import hashlib,json,tempfile,time,unittest
from pathlib import Path
from product_tool.xbox_identity import identifier,family,official,configuration_sensitive
from product_tool.xbox_page import parse_page,assignment
from product_tool.xbox_documents import document_role
from product_tool.adapters.xbox import sitemap_urls,XboxTransport
from product_tool.adapters.policy_session import PolicyResponse
from product_tool.adapters.common import RawAttribute,SourceDocument,PhotoCandidate
from product_tool import jobs,storage,card_evidence,worker,xbox_pipeline,attribute_projection,exporter
from product_tool.normalization import normalize_fact
from product_tool.resolution import resolve_attributes

R=Path('reports/xbox_stage68/observed')
WHITE='https://www.microsoft.com/en-us/d/xbox-wireless-controller/8xn59crbsqgz'

def page(file,url,article,expected,region='en-US'):
    return parse_page((R/file).read_text(encoding='utf8'),url,article,expected,region)

class XboxIdentityTests(unittest.TestCase):
    def test_typed_store_ids_not_manufacturer_parts(self):
        i=identifier('8XN59CRBSQGZ/KPRJ');self.assertEqual(i['sku_id'],'KPRJ');self.assertEqual(i['product_id'],'8XN59CRBSQGZ');self.assertFalse(i['hardware_model']);self.assertFalse(i['part_number'])
        self.assertEqual(identifier('1883')['type'],'hardware_model_number');self.assertEqual(identifier('Xbox Series X')['type'],'commercial_configuration')
    def test_malicious_and_replacement_urls_not_identity(self):
        self.assertFalse(official('https://microsoft.com.evil.example/d/item'));self.assertFalse(official('https://user@www.microsoft.com/d/item'));self.assertFalse(family('Replacement input PCBA for Xbox Wireless Controller'))
    def test_series_x_cannot_confirm_series_s(self):
        d,e=page('22045379f986be6093bc.html','https://www.xbox.com/en-US/consoles/xbox-series-x','Xbox Series S','series_s')
        self.assertTrue(d.error);self.assertEqual(e['identity']['model'],'unproven');self.assertFalse(d.attributes)
    def test_commercial_name_not_exact(self):
        d,e=page('d97e89f0ba470112cce0.html','https://www.xbox.com/en-US/consoles/xbox-series-s','Xbox Series S 512GB','series_s')
        self.assertEqual(e['identity']['model'],'model_confirmed');self.assertFalse(e['configuration_complete']);self.assertFalse(e['configuration_fields']);self.assertTrue(e['configuration_candidates'])
    def test_other_region_does_not_confirm_regional_sku(self):
        d,e=page('21a3e8202858c6ba83fb.html',WHITE,'8XN59CRBSQGZ/KPRJ','controller','en-GB')
        self.assertEqual(d.source_key,'xbox_model');self.assertFalse(e['exact_photo_assets']);self.assertFalse(e['configuration_fields'])
    def test_product_id_does_not_select_default_color(self):
        d,e=page('21a3e8202858c6ba83fb.html',WHITE,'8XN59CRBSQGZ','controller')
        self.assertFalse(e['configuration_complete']);self.assertFalse(d.photos);self.assertEqual(d.source_key,'xbox_model')
        self.assertEqual(d.found_model,'Xbox Wireless Controller');self.assertNotIn('Carbon Black',d.description)
    def test_current_sku_differs_from_old_catalog_color_id(self):
        d,e=page('21a3e8202858c6ba83fb.html',WHITE,'8XN59CRBSQGZ/LK4W','controller')
        self.assertEqual(e['identity']['configuration'],'unproven');self.assertFalse(d.photos)
    def test_same_controller_different_color(self):
        white,we=page('21a3e8202858c6ba83fb.html',WHITE,'8XN59CRBSQGZ/KPRJ','controller')
        black,be=page('21a3e8202858c6ba83fb.html',WHITE,'8XN59CRBSQGZ/QP6M','controller')
        self.assertEqual(we['configuration_fields']['color'],'Robot White');self.assertEqual(be['configuration_fields']['color'],'Carbon Black')
        self.assertIn('Robot White',white.found_model);self.assertNotIn('Carbon Black',white.description)
        self.assertIn('Carbon Black',black.found_model);self.assertNotIn('Robot White',black.description)
        self.assertIn('Carbon Black',we['page_default_title'])
        self.assertTrue(white.photos);self.assertTrue(black.photos);self.assertFalse(set(white.photos)&set(black.photos))
        self.assertTrue(any('Black-b00' in p['url'] for p in we['photo_candidates']))
        self.assertTrue(all('White' in u for u in white.photos));self.assertTrue(all('Black' in u for u in black.photos))
    def test_returned_other_product_id_not_accepted(self):
        d,e=page('21a3e8202858c6ba83fb.html',WHITE,'8QRF79K7JSR6/ZHP2','controller')
        self.assertFalse(e['configuration_complete']);self.assertFalse(e['identifiers'])
    def test_special_color_sku_is_own(self):
        d,e=page('177ce9a0d7f03a67074f.html','https://www.microsoft.com/en-us/d/xbox-wireless-controller-arctic-camo-special-edition/8qrf79k7jsr6','8QRF79K7JSR6/ZHP2','controller')
        self.assertEqual(e['configuration_fields']['color'],'Arctic Camo');self.assertEqual(e['identity']['configuration'],'exact_store_sku');self.assertTrue(d.photos)
    def test_storage_512_and_1tb_are_scoped(self):
        d,s=page('cd2f186ba023a8bdea79.html','https://www.microsoft.com/en-us/d/xbox-series-s-starter-bundle/8zb8z8mm10v0','8ZB8Z8MM10V0','series_s')
        d,t=page('2032e5708c7973ee9949.html','https://www.microsoft.com/en-us/d/xbox-series-s-1tb-white/8r02lkf9q26r','8R02LKF9Q26R','series_s')
        self.assertIn('512GB',s['configuration_fields']['storage']);self.assertIn('1TB',t['configuration_fields']['storage'])
        self.assertNotIn('game_pass',t['configuration_fields']);self.assertNotIn('bundle',t['configuration_fields'])
    def test_expansion_is_not_usable_storage(self):
        d,e=page('d97e89f0ba470112cce0.html','https://www.xbox.com/en-US/consoles/xbox-series-s','1883','series_s')
        names=[x.name for x in d.attributes]
        self.assertIn('Поддержка карты расширения',names);self.assertIn('Поддержка внешних USB-накопителей',names)
        self.assertNotIn('Встроенный накопитель',names);self.assertNotIn('Доступное пользователю место',names)
    def test_bundle_kit_not_borrowed_from_base(self):
        d,e=page('e39314409941dfa9650b.html','https://www.microsoft.com/en-us/d/xbox-series-x-diablo-iv-bundle/8n1bb8dsbknt','8N1BB8DSBKNT','series_x')
        self.assertIn('Diablo',e['configuration_fields']['included_game']);self.assertNotIn('game_pass',e['configuration_fields'])
        self.assertNotIn('second_controller',e['configuration_fields']);self.assertNotIn('packaging',e['configuration_fields'])
        self.assertTrue(d.photos);self.assertTrue(all(e['photo_scopes'][p]['kind']=='bundle_gallery' for p in e['exact_photo_assets']))
        base,be=page('22045379f986be6093bc.html','https://www.xbox.com/en-US/consoles/xbox-series-x','Xbox Series X','series_x')
        self.assertNotIn('bundle_contents',be['configuration_fields'])
    def test_revision_physical_values_withheld(self):
        d,e=page('d97e89f0ba470112cce0.html','https://www.xbox.com/en-US/consoles/xbox-series-s','1883','series_s')
        self.assertFalse(any(x.name in {'Вес товара','Ширина товара','Высота товара','Глубина товара'} for x in d.attributes));self.assertTrue(any(x['raw_label'].lower()=='weight' for x in e['configuration_candidates']))
        self.assertFalse(any('GHz' in x.value for x in d.attributes));self.assertTrue(any(x.name=='Архитектура процессора' and x.value=='Zen 2' for x in d.attributes))
    def test_no_support_procedure_description(self):
        d,e=page('ff568285b2dd532fb4e9.html','https://www.xbox.com/en-US/accessories/controllers/elite-wireless-controller-series-2','1797','elite')
        self.assertNotIn('Drivers available',d.description);self.assertNotIn('manually',d.description);self.assertFalse(any(x.name=='Вес товара' for x in d.attributes))
        d,e=parse_page('<h1>Xbox Series S</h1><p>Reset and repair</p>','https://support.xbox.com/en-US/help/hardware-network/console/reset-console','1883','series_s','en-US');self.assertTrue(d.error);self.assertFalse(d.description)
    def test_safety_and_setup_not_user_guide(self):
        self.assertEqual(document_role('Xbox Series X PRODUCT GUIDE Important Safety and Regulatory Information'),'Safety/Regulatory')
        self.assertEqual(document_role('Quick Start Guide'),'Setup Guide');self.assertEqual(document_role('User Guide'),'User Guide')
    def test_store_sitemap_keeps_unicode_alternate_regions(self):
        z=PolicyResponse('https://www.microsoft.com/store/xbox-accessory-01.xml',200,(R/'382aafdac993fd9958c2.html').read_text(encoding='utf8'))
        kind,urls=sitemap_urls(z);self.assertEqual(kind,'urlset');self.assertTrue(any('/en-us/d/xbox-wireless-controller/8xn59crbsqgz' in u for u in urls));self.assertTrue(any('tradløs' in u for u in urls))
    def test_plain_xml_gz_url_and_real_gzip_both(self):
        import gzip
        xml='<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9"><url><loc>https://www.xbox.com/en-US/consoles/xbox-series-s</loc></url></urlset>'
        a=sitemap_urls(PolicyResponse('https://www.xbox.com/sitemap/a.xml.gz',200,xml))
        b=sitemap_urls(PolicyResponse('https://www.xbox.com/sitemap/a.xml.gz',200,gzip.compress(xml.encode()).decode('latin1')))
        self.assertEqual(a,b)
    def test_frozen_inputs_unchanged_and_ten(self):
        p=Path('reports/xbox_stage68/dataset.json');self.assertEqual(hashlib.sha256(p.read_bytes()).hexdigest(),'556dfe669181b1959e1d21dd52ed491f78719744bdbdab3442bef981de76191c');self.assertEqual(len(json.loads(p.read_text())['rows']),10)
    def test_assignment_is_data_not_js_execution(self):
        self.assertIsNone(assignment('window.__BuyBox__=evil()','window.__BuyBox__'));self.assertEqual(assignment('window.__BuyBox__={"product":{}}; throw new Error()', 'window.__BuyBox__'),{'product':{}})

class XboxPipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.db=Path(self.tmp.name)/'batches.sqlite3';jobs.initialize(self.db)
        with storage._connection(self.db) as c:
            c.execute('INSERT INTO batches VALUES (?,?,?,?,?)',('x','Xbox','Products','{}',storage._now()))
            c.execute('INSERT INTO products (batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (?,?,?,?,?,?,?,?,?,?)',('x',1,'Xbox Wireless Controller Robot White','Xbox','8XN59CRBSQGZ/KPRJ','','accessories',0,'[]','{}'))
        self.pid=storage.get_batch(self.db,'x')['products'][0]['id']
    def tearDown(self):self.tmp.cleanup()
    def install(self):
        d,e=page('21a3e8202858c6ba83fb.html',WHITE,'8XN59CRBSQGZ/KPRJ','controller');e['manual_status']='Не проверена';e['configuration_scope']='retail_variant'
        jobs.save_source_document(self.db,self.pid,d);card_evidence.save(self.db,self.pid,'xbox',e);jobs.set_photo_selection(self.db,self.pid,e['exact_photo_assets'],mode='exact');jobs.resolve_product(self.db,self.pid);return d,e
    def test_optional_manual_does_not_block_readiness(self):
        self.install();r=xbox_pipeline.card_readiness(self.db,self.pid);self.assertEqual(r['verdict'],'export_ready');self.assertIn('manual_unverified',r['advisory_gaps'])
    def test_resolver_rejects_family_configuration_even_if_injected(self):
        self.install();d=SourceDocument('xbox_model','Xbox', 'https://www.xbox.com/en-US/consoles/xbox-series-s',match_level='model_confirmed',attributes=[RawAttribute('Цвет','Other'),RawAttribute('Вес товара','999 kg'),RawAttribute('Встроенный накопитель','2 TB')]);jobs.save_source_document(self.db,self.pid,d);jobs.resolve_product(self.db,self.pid)
        rows={r['normalized_name']:r for r in jobs.get_resolved(self.db,self.pid)};self.assertNotEqual(rows['product_weight']['status'],'model_confirmed_official');self.assertFalse(rows['product_weight']['selected_value']);self.assertEqual(rows['color']['selected_value'],'robot white')
    def test_metric_axes_and_weight_use_common_keys(self):
        d=SourceDocument('xbox_configuration','Xbox',WHITE,match_level='full_sku',attributes=[RawAttribute('Ширина товара','15.1 cm'),RawAttribute('Высота товара','30.1 cm'),RawAttribute('Глубина товара','15.1 cm'),RawAttribute('Вес товара','320 g')]);jobs.save_source_document(self.db,self.pid,d);jobs.resolve_product(self.db,self.pid)
        r={x['normalized_name']:x for x in attribute_projection.final_attribute_rows(self.db,self.pid)}
        self.assertEqual(r['product_dimensions__height']['resolved']['selected_unit'],'mm');self.assertEqual(r['product_dimensions__height']['resolved']['selected_value'],'301');self.assertEqual(r['product_weight']['resolved']['selected_unit'],'kg');self.assertEqual(r['product_weight']['resolved']['selected_value'],'0.32')
    def test_native_export_contains_scope_candidates(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.install();b=load_workbook(BytesIO(exporter.export_batch(self.db,'x')))
        self.assertIn('Идентификаторы Xbox',b.sheetnames);self.assertIn('Фото-кандидаты',b.sheetnames);self.assertIn('Документы Xbox',b.sheetnames);self.assertEqual(b['Готовность Xbox'].cell(2,3).value,'Готова')
        self.assertNotIn('Black-b00',' '.join(str(c.value) for row in b['Фотографии'] for c in row));b.close()
    def test_wrong_photo_source_not_verified(self):
        d,e=self.install();p=jobs.get_photo_candidates(self.db,self.pid)[0];self.assertTrue(xbox_pipeline.photo_verified(p,e));self.assertFalse(xbox_pipeline.photo_verified({**p,'source_key':'dns'},e))
    def test_branded_terms_in_final_values_preserve_case(self):
        d,e=page('0cb77f891ef9dafe07c1.html','https://www.microsoft.com/en-us/d/xbox-wireless-headset/9203q8w23lhn','9203Q8W23LHN/HFXZ','headset');jobs.save_source_document(self.db,self.pid,d);jobs.resolve_product(self.db,self.pid)
        original=jobs.get_resolved(self.db,self.pid);rows=attribute_projection.final_attribute_rows(self.db,self.pid);shown=next(x['resolved']['display_value'] for x in rows if x['display_name']=='Spatial Sound')
        self.assertIn('Dolby Atmos',shown);self.assertIn('Windows Sonic',shown);self.assertIn('Xbox One',shown);self.assertEqual(original,jobs.get_resolved(self.db,self.pid))

if __name__=='__main__':unittest.main()
