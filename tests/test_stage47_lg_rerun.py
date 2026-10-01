from __future__ import annotations

import sqlite3
import tempfile
import unittest
from pathlib import Path

from product_tool import jobs, worker
from product_tool.adapters.common import PhotoCandidate, ProductDocument, RawAttribute, SourceDocument


class EmptyOfficial:
    def __init__(self, key):
        self.source_key = self.site_name = key

    def find_source(self, *args, **kwargs):
        return SourceDocument(self.source_key, self.site_name, "", match_level="unknown")


class EmptyDealer:
    source_key = site_name = "sulpak"

    def find_source(self, *args, **kwargs):
        return SourceDocument(self.source_key, self.site_name, "", match_level="unknown")


class EmptyDns:
    source_key = site_name = "dns"
    document_urls = {}

    def find_source(self, *args, **kwargs):
        return SourceDocument(self.source_key, self.site_name, "", match_level="dealer_url_needed")

    def find_documents(self, *args, **kwargs):
        return [], "no observed URL"


class UnreachableSupport:
    candidate_urls = ()
    candidate_pages = {}
    candidate_payloads = {}

    def find_source(self, *args, **kwargs):
        return SourceDocument("lg_kz_support", "LG KZ support", self.candidate_urls[0],
                              match_level="full_sku", html='<div data-product-id="TEST.USAR"></div>')

    def find_documents(self, *args, **kwargs):
        return [], {"outcome": "no_verified_russian_instruction",
                    "files": [{"state": "unreachable", "url": "https://gscs-b2c.lge.com/manual"}]}


class NoCandidatesSupport(UnreachableSupport):
    def find_documents(self, *args, **kwargs):
        return [], {"outcome": "no_manual_candidates", "files": []}


class RetainedEvidenceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.db = Path(self.temp.name) / "batches.sqlite3"
        jobs.initialize(self.db)
        with sqlite3.connect(self.db) as con:
            con.execute("INSERT INTO batches(id,filename,sheet_name,mapping_json,confirmed_at) VALUES ('b','x.xlsx','Sheet1','{}','now')")
            con.execute("INSERT INTO products(id,batch_id,row_number,name,brand,search_code,alternate_code,category,needs_confirmation,issues_json,original_values_json) VALUES (4,'b',2,'test','LG','TEST.USAR','','other',0,'[]','{}')")
        con.close()

    def test_empty_refresh_preserves_verified_product_evidence(self):
        page = "https://www.lg.com/kz/product/test"
        ru_page = "https://www.lg.com/ru/product/test"
        jobs.save_source_document(self.db, 4, SourceDocument("lg_kz", "LG KZ", page,
            match_level="full_sku", attributes=[RawAttribute("Power", "100 W")],
            photo_candidates=[PhotoCandidate("https://www.lg.com/photo.jpg", "photo", "product_gallery")],
            html="<h1>TEST.USAR</h1>"))
        jobs.save_source_document(self.db, 4, SourceDocument("lg_ru", "LG RU", ru_page,
            match_level="full_sku", html="<h1>TEST.USAR</h1>"))
        jobs.save_documents(self.db, 4, "lg_ru", [ProductDocument("Guide", "\u0420\u0443\u0441\u0441\u043a\u0438\u0439", "", "",
            "https://gscs-b2c.lge.com/manual", "https://www.lg.com/ru/support/test",
            "TEST", "TEST.USAR", ru_page)])
        job = jobs.enqueue(self.db, 4, [1, 2, 3, 4, 6])
        self.assertTrue(worker.run_once(self.db,
            adapter_factory=lambda: (EmptyOfficial("lg_kz"), EmptyOfficial("lg_ru"), EmptyDealer()),
            dns_adapter_factory=EmptyDns))
        self.assertEqual(jobs.list_jobs(self.db, 4)[0]["id"], job)
        self.assertEqual(jobs.list_jobs(self.db, 4)[0]["status"], "done")
        self.assertEqual(len(jobs.get_facts(self.db, 4)), 1)
        self.assertEqual(len(jobs.get_photo_candidates(self.db, 4)), 1)
        self.assertEqual(len(jobs.get_documents(self.db, 4)), 1)
        self.assertEqual(next(x for x in jobs.get_source_pages(self.db, 4) if x["source_key"] == "lg_kz")["match_level"], "full_sku")

    def test_unreachable_support_pdf_keeps_existing_verified_document(self):
        url = "https://www.lg.com/kz/support/product-support/cs-TEST.USAR/"
        jobs.save_source_document(self.db, 4, SourceDocument("lg_kz_support", "LG KZ support", url,
            match_level="full_sku", html='<div data-product-id="TEST.USAR"></div>'))
        jobs.save_documents(self.db, 4, "lg_kz_support", [ProductDocument("Guide", "\u0420\u0443\u0441\u0441\u043a\u0438\u0439", "", "",
            "https://gscs-b2c.lge.com/manual", url, "TEST", "TEST.USAR", url)])
        jobs.enqueue(self.db, 4, [1, 2, 3, 4, 6])
        self.assertTrue(worker.run_once(self.db,
            adapter_factory=lambda: (EmptyOfficial("lg_kz"), EmptyOfficial("lg_ru"), EmptyDealer()),
            dns_adapter_factory=EmptyDns, lg_support_adapter_factory=UnreachableSupport))
        self.assertEqual(len(jobs.get_documents(self.db, 4)), 1)
        self.assertEqual(jobs.get_documents(self.db, 4)[0]["source_key"], "lg_kz_support")
        events = jobs.list_events(self.db, jobs.list_jobs(self.db, 4)[0]["id"])
        self.assertTrue(any("Ранее проверенная русская инструкция сохранена" in e["message"] for e in events))
        message = jobs.list_jobs(self.db, 4)[0]["message"]
        self.assertIn("подтверждает модель/комплект", message)
        self.assertIn("характеристиками и фото", message)

    def test_render_without_manual_list_does_not_erase_previously_verified_pdf(self):
        url = "https://www.lg.com/kz/support/product-support/cs-TEST.USAR/"
        jobs.save_source_document(self.db, 4, SourceDocument(
            "lg_kz_support", "LG KZ support", url, match_level="full_sku",
            html='<div data-product-id="TEST.USAR"></div>'))
        jobs.save_documents(self.db, 4, "lg_kz_support", [ProductDocument(
            "Guide", "Русский", "", "",
            "https://gscs-b2c.lge.com/manual", url, "TEST", "TEST.USAR", url)])
        jobs.enqueue(self.db, 4, [1, 6])
        self.assertTrue(worker.run_once(self.db,
            adapter_factory=lambda: (EmptyOfficial("lg_kz"), EmptyOfficial("lg_ru"), EmptyDealer()),
            dns_adapter_factory=EmptyDns, lg_support_adapter_factory=NoCandidatesSupport))
        self.assertEqual(len(jobs.get_documents(self.db, 4)), 1)
        events = jobs.list_events(self.db, jobs.list_jobs(self.db, 4)[0]["id"])
        self.assertTrue(any("Ранее проверенная" in e["message"] for e in events))

    def test_printed_product_support_link_reaches_existing_adapter(self):
        product_url = "https://www.lg.com/kz/product/test"
        support_url = "https://www.lg.com/kz/support/product-support/cs-TEST.USAR"
        html = ('<a href="/kz/support/product-support/">generic support</a>'
                '<a href="/kz/support/product-support/cs-TEST.USAR#manual-tab">manuals</a>'
                '<a href="https://example.org/kz/support/product-support/cs-TEST.USAR">off host</a>')
        jobs.save_source_document(self.db, 4, SourceDocument("lg_kz", "LG KZ", product_url,
            match_level="full_sku", html=html))
        support = UnreachableSupport()
        jobs.enqueue(self.db, 4, [1])
        self.assertTrue(worker.run_once(self.db,
            adapter_factory=lambda: (EmptyOfficial("lg_kz"), EmptyOfficial("lg_ru"), EmptyDealer()),
            dns_adapter_factory=EmptyDns, lg_support_adapter_factory=lambda: support))
        self.assertEqual(support.candidate_urls, (support_url,))
        page = next(x for x in jobs.get_source_pages(self.db, 4) if x["source_key"] == "lg_kz_support")
        self.assertEqual(page["url"], support_url)

    def test_excel_uses_article_when_imported_name_is_empty(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from product_tool.exporter import export_batch

        with sqlite3.connect(self.db) as con:
            con.execute("UPDATE products SET name='' WHERE id=4")
        con.close()
        url = "https://www.lg.com/kz/support/product-support/cs-TEST.USAR/"
        jobs.save_source_document(self.db, 4, SourceDocument("lg_kz_support", "LG KZ support", url,
            match_level="full_sku", html='<div data-product-id="TEST.USAR"></div>'))
        jobs.save_documents(self.db, 4, "lg_kz_support", [ProductDocument("Guide", "Русский", "", "",
            "https://gscs-b2c.lge.com/manual", url, "TEST", "TEST.USAR", url)])
        book = load_workbook(BytesIO(export_batch(self.db, "b")), read_only=True)
        try:
            self.assertEqual(list(book["Инструкции"].values)[1][0], "TEST.USAR")
            self.assertEqual(list(book["Готовность LG"].values)[1][0], "TEST.USAR")
        finally:
            book.close()

    def test_actual_mismatch_page_cannot_replace_confirmed_exact_page(self):
        old = SourceDocument("lg_kz", "LG KZ", "https://www.lg.com/kz/product/old",
                             match_level="full_sku", attributes=[RawAttribute("Power", "100 W")])
        new = SourceDocument("lg_kz", "LG KZ", "https://www.lg.com/kz/product/new",
                             match_level="mismatch", html="<h1>OTHER</h1>")
        jobs.save_source_document(self.db, 4, old)
        jobs.save_source_document(self.db, 4, new)
        page = next(x for x in jobs.get_source_pages(self.db, 4) if x["source_key"] == "lg_kz")
        self.assertEqual((page["url"], page["match_level"], len(jobs.get_facts(self.db, 4))),
                         (old.url, "full_sku", 1))


if __name__ == "__main__":
    unittest.main()
