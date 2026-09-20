"""Read a supplier Excel file and preview product identities before parsing sites.

This module does not edit the uploaded workbook or call product websites.
Column numbers shown to users are one-based, just like Excel.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5_000
MAX_COLUMNS = 100

HEADER_ALIASES = {
    "brand": {"бренд", "марка", "производитель", "brand", "manufacturer"},
    "search_code": {"артикул продавца", "артикул", "код товара", "sku", "part number"},
    "fallback_code": {"модель", "код модели", "model", "model number"},
    "name": {"наименование", "название", "название товара", "товар", "product name"},
    "category": {"предмет", "категория", "тип товара", "category"},
}

KNOWN_BRANDS = ("Samsung", "Lenovo", "Apple", "LG")
CODE_PATTERN = re.compile(r"(?<![\w])([A-Za-z0-9][A-Za-z0-9./_-]{5,})(?![\w])")


@dataclass(frozen=True)
class ColumnMapping:
    brand: int | None = None
    search_code: int | None = None
    fallback_code: int | None = None
    name: int | None = None
    category: int | None = None
    header_row: int = 0

    def __post_init__(self) -> None:
        for field in ("brand", "search_code", "fallback_code", "name", "category"):
            number = getattr(self, field)
            if number is not None and not 1 <= number <= MAX_COLUMNS:
                raise ValueError(f"Номер колонки {field} должен быть от 1 до {MAX_COLUMNS}")
        if self.header_row < 0:
            raise ValueError("Номер строки заголовков не может быть отрицательным")
        if self.name is None and self.brand is None and self.search_code is None:
            raise ValueError("Укажите хотя бы колонку названия, бренда или артикула")


@dataclass(frozen=True)
class ProductPreview:
    row_number: int
    brand: str
    search_code: str
    alternate_code: str
    category: str
    name: str
    brand_source: str
    code_source: str
    category_source: str
    needs_confirmation: bool
    issues: tuple[str, ...]
    original_values: tuple[Any, ...]


@dataclass(frozen=True)
class ImportPreview:
    file_name: str
    sheet_name: str
    mapping: ColumnMapping
    products: tuple[ProductPreview, ...]
    category_counts: dict[str, int]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return re.sub(r"\s+", " ", str(value)).strip()


def original_cell(value: Any) -> Any:
    """Keep source values intact, converting only values JSON cannot represent."""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def header_key(value: Any) -> str:
    return re.sub(r"[^\w]+", " ", clean(value).casefold()).strip()


def detect_mapping(worksheet: Any) -> ColumnMapping:
    best: tuple[int, int, dict[str, int]] | None = None
    for row_number, row in enumerate(worksheet.iter_rows(max_row=12, values_only=True), start=1):
        found: dict[str, int] = {}
        for index, value in enumerate(row[:MAX_COLUMNS], start=1):
            name = header_key(value)
            for field, aliases in HEADER_ALIASES.items():
                if name in aliases and field not in found:
                    found[field] = index
        score = len(found) + (2 if "name" in found else 0)
        if score and (best is None or score > best[0]):
            best = (score, row_number, found)
    if best is None:
        raise ValueError("Заголовки не найдены. Укажите номера колонок вручную.")
    _, row_number, found = best
    return ColumnMapping(**found, header_row=row_number)


def choose_sheet(
    workbook: Any, sheet_name: str | None, allow_headerless: bool = False
) -> tuple[Any, ColumnMapping | None]:
    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Лист «{sheet_name}» не найден")
        sheet = workbook[sheet_name]
        try:
            return sheet, detect_mapping(sheet)
        except ValueError:
            if allow_headerless:
                return sheet, None
            raise
    for sheet in workbook.worksheets:
        try:
            return sheet, detect_mapping(sheet)
        except ValueError:
            continue
    if allow_headerless and workbook.worksheets:
        return workbook.worksheets[0], None
    raise ValueError("Ни на одном листе не найдены заголовки. Укажите лист и номера колонок.")


def value_at(row: tuple[Any, ...], column: int | None) -> str:
    return clean(row[column - 1]) if column and column <= len(row) else ""


def brand_from_name(name: str) -> str:
    for brand in KNOWN_BRANDS:
        if re.search(rf"(?<!\w){re.escape(brand)}(?!\w)", name, re.IGNORECASE):
            return brand
    return ""


def code_candidates(name: str) -> list[str]:
    candidates: list[str] = []
    for match in CODE_PATTERN.finditer(name):
        code = match.group(1)
        if not re.search(r"[A-Za-z]", code) or not re.search(r"\d", code):
            continue
        if re.fullmatch(r"\d+(?:GB|TB|MB)", code, re.IGNORECASE):
            continue
        if code.casefold() in {brand.casefold() for brand in KNOWN_BRANDS}:
            continue
        candidates.append(code)
    return list(dict.fromkeys(candidates))


def category_from_name(name: str) -> str:
    lowered = name.casefold()
    rules = (
        (r"стиральн\w*\s+машин\w*\s+с\s+сушк", "Стирально-сушильная машина"),
        (r"стиральн\w*\s+машин", "Стиральная машина"),
        (r"паровой\s+шкаф", "Паровой шкаф"),
        (r"смартфон", "Смартфон"),
        (r"ноутбук", "Ноутбук"),
        (r"телевизор", "Телевизор"),
        (r"наушник", "Наушники"),
    )
    return next((category for pattern, category in rules if re.search(pattern, lowered)), "")


def preview_row(row_number: int, raw: tuple[Any, ...], mapping: ColumnMapping) -> ProductPreview | None:
    explicit_name = value_at(raw, mapping.name)
    explicit_brand = value_at(raw, mapping.brand)
    explicit_code = value_at(raw, mapping.search_code)
    fallback_code = value_at(raw, mapping.fallback_code)
    explicit_category = value_at(raw, mapping.category)
    if not any((explicit_name, explicit_brand, explicit_code, fallback_code, explicit_category)):
        return None

    issues: list[str] = []
    brand = explicit_brand or brand_from_name(explicit_name)
    brand_source = "column" if explicit_brand else "name" if brand else "missing"
    candidates = code_candidates(explicit_name) if not (explicit_code or fallback_code) else []
    search_code = explicit_code or fallback_code or (candidates[0] if candidates else "")
    code_source = (
        "column" if explicit_code else "fallback_column" if fallback_code else
        "name" if search_code else "missing"
    )
    alternate_code = fallback_code if explicit_code and fallback_code != explicit_code else ""
    category = explicit_category or category_from_name(explicit_name)
    category_source = "column" if explicit_category else "name" if category else "missing"

    if not brand:
        issues.append("Не определён бренд")
    if not search_code:
        issues.append("Не определён артикул или модель")
    if not category:
        issues.append("Не определена категория")
    if len(candidates) > 1:
        issues.append("В названии несколько возможных кодов")
    needs_confirmation = bool(issues) or any(
        source == "name" for source in (brand_source, code_source, category_source)
    )
    return ProductPreview(
        row_number=row_number,
        brand=brand,
        search_code=search_code,
        alternate_code=alternate_code,
        category=category,
        name=explicit_name,
        brand_source=brand_source,
        code_source=code_source,
        category_source=category_source,
        needs_confirmation=needs_confirmation,
        issues=tuple(issues),
        original_values=raw,
    )


def preview_xlsx(
    path: str | Path,
    *,
    sheet_name: str | None = None,
    column_overrides: dict[str, int | None] | None = None,
    header_row: int | None = None,
) -> ImportPreview:
    file_path = Path(path)
    if file_path.suffix.lower() != ".xlsx":
        raise ValueError("На первом этапе поддерживаются файлы .xlsx")
    if not file_path.is_file():
        raise ValueError(f"Файл не найден: {file_path}")
    if file_path.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("Файл больше 10 МБ")
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheet, detected = choose_sheet(workbook, sheet_name, bool(column_overrides))
        settings = asdict(detected) if detected else {
            "brand": None, "search_code": None, "fallback_code": None,
            "name": None, "category": None, "header_row": 0,
        }
        for field, number in (column_overrides or {}).items():
            if field not in settings or field == "header_row":
                raise ValueError(f"Неизвестное поле сопоставления: {field}")
            settings[field] = number
        if header_row is not None:
            settings["header_row"] = header_row
        mapping = ColumnMapping(**settings)
        products: list[ProductPreview] = []
        for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            if row_number <= mapping.header_row:
                continue
            if row_number - mapping.header_row > MAX_ROWS + 10:
                raise ValueError(f"Файл содержит больше {MAX_ROWS} строк для обработки")
            raw = tuple(original_cell(value) for value in cells[:MAX_COLUMNS])
            item = preview_row(row_number, raw, mapping)
            if item:
                products.append(item)
        if not products:
            raise ValueError("После заголовков товары не найдены")
        counts: dict[str, int] = {}
        for item in products:
            key = item.category or "Категория не определена"
            counts[key] = counts.get(key, 0) + 1
        return ImportPreview(file_path.name, sheet.title, mapping, tuple(products), counts)
    finally:
        workbook.close()


def main() -> None:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description="Предпросмотр бренда, кода и категории из Excel")
    parser.add_argument("file", help="Путь к файлу .xlsx")
    parser.add_argument("--sheet", help="Название листа")
    parser.add_argument("--brand-col", type=int, help="Номер колонки бренда, начиная с 1")
    parser.add_argument("--code-col", type=int, help="Номер колонки основного артикула")
    parser.add_argument("--fallback-col", type=int, help="Номер запасной колонки модели")
    parser.add_argument("--name-col", type=int, help="Номер колонки названия")
    parser.add_argument("--category-col", type=int, help="Номер колонки категории")
    parser.add_argument("--header-row", type=int, help="Номер строки заголовков; 0 если их нет")
    args = parser.parse_args()
    flag_to_field = {
        "brand_col": "brand", "code_col": "search_code", "fallback_col": "fallback_code",
        "name_col": "name", "category_col": "category",
    }
    overrides = {
        field: getattr(args, flag)
        for flag, field in flag_to_field.items()
        if getattr(args, flag) is not None
    }
    result = preview_xlsx(
        args.file, sheet_name=args.sheet, column_overrides=overrides,
        header_row=args.header_row,
    )
    concise = result.to_dict()
    for product in concise["products"]:
        product.pop("original_values")
    print(json.dumps(concise, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
