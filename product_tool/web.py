"""Local web UI for uploading, reviewing, and confirming Excel batches."""

from __future__ import annotations

from dataclasses import asdict
import logging
import os
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import FastAPI, File, Request, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import load_workbook

from . import jobs, storage
from .importer import MAX_COLUMNS, MAX_FILE_BYTES, MAX_ROWS, ImportPreview, preview_xlsx


LOGGER = logging.getLogger(__name__)
PACKAGE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = PACKAGE_DIR.parent
COLUMN_FIELDS = ("category", "brand", "name", "search_code", "fallback_code")
COLUMN_LABELS = {
    "category": "Категория",
    "brand": "Бренд",
    "name": "Название",
    "search_code": "Основной артикул",
    "fallback_code": "Запасная модель",
}


def _sheet_names(path: Path) -> list[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return workbook.sheetnames
    finally:
        workbook.close()


def _parse_mapping(form: Any, sheets: list[str]) -> tuple[str, dict[str, int | None]]:
    sheet_name = str(form.get("sheet_name", "")).strip()
    if sheet_name not in sheets:
        raise ValueError("Выберите лист из загруженного Excel.")
    mapping: dict[str, int | None] = {}
    for field in COLUMN_FIELDS:
        raw = str(form.get(field, "")).strip()
        if not raw:
            mapping[field] = None
            continue
        try:
            number = int(raw)
        except ValueError as exc:
            raise ValueError(f"Поле «{COLUMN_LABELS[field]}»: укажите номер колонки.") from exc
        if not 1 <= number <= MAX_COLUMNS:
            raise ValueError(
                f"Поле «{COLUMN_LABELS[field]}»: номер колонки должен быть от 1 до {MAX_COLUMNS}."
            )
        mapping[field] = number
    raw_header = str(form.get("header_row", "")).strip()
    try:
        header_row = int(raw_header)
    except ValueError as exc:
        raise ValueError("Укажите номер строки заголовков; 0 означает отсутствие заголовков.") from exc
    if not 0 <= header_row <= MAX_ROWS:
        raise ValueError(f"Строка заголовков должна быть от 0 до {MAX_ROWS}.")
    mapping["header_row"] = header_row
    if not any(mapping[field] for field in ("name", "brand", "search_code")):
        raise ValueError("Укажите хотя бы колонку названия, бренда или основного артикула.")
    return sheet_name, mapping


def _preview(path: Path, draft: dict[str, Any]) -> ImportPreview:
    mapping = draft["mapping"]
    if mapping is None:
        return preview_xlsx(path, sheet_name=draft["sheet_name"])
    return preview_xlsx(
        path,
        sheet_name=draft["sheet_name"],
        column_overrides={field: mapping[field] for field in COLUMN_FIELDS},
        header_row=mapping["header_row"],
    )


def create_app(data_dir: str | Path | None = None) -> FastAPI:
    """Create an app with its own data directory, including for isolated tests."""
    root = Path(data_dir or os.environ.get("PRODUCT_CARDS_DATA_DIR") or PROJECT_DIR / "data")
    root = root.resolve()
    uploads = root / "uploads"
    uploads.mkdir(parents=True, exist_ok=True)
    database = root / "batches.sqlite3"
    jobs.initialize(database)
    templates = Jinja2Templates(directory=str(PACKAGE_DIR / "templates"))

    app = FastAPI(title="Product Cards MVP", docs_url=None, redoc_url=None)
    app.mount("/static", StaticFiles(directory=str(PACKAGE_DIR / "static")), name="static")

    def render(request: Request, name: str, status_code: int = 200, **context: Any) -> HTMLResponse:
        return templates.TemplateResponse(
            request=request, name=name, context=context, status_code=status_code
        )

    def error_page(request: Request, message: str, status_code: int = 400) -> HTMLResponse:
        return render(request, "error.html", status_code, message=message)

    def preview_page(
        request: Request,
        draft_id: str,
        *,
        message: str = "",
        status_code: int = 200,
        mapping_values: dict[str, Any] | None = None,
        edits: dict[int, dict[str, str]] | None = None,
    ) -> HTMLResponse:
        draft = storage.get_draft(database, draft_id)
        if draft is None:
            return error_page(request, "Черновик не найден или партия уже подтверждена.", 404)
        path = uploads / f"{draft_id}.xlsx"
        try:
            sheets = _sheet_names(path)
        except Exception:
            LOGGER.exception("Could not reopen uploaded workbook")
            return error_page(request, "Не удалось открыть сохранённый Excel. Загрузите файл снова.", 400)

        preview: ImportPreview | None = None
        try:
            preview = _preview(path, draft)
        except ValueError as exc:
            if not message:
                message = str(exc)
        except Exception:
            LOGGER.exception("Could not preview workbook")
            if not message:
                message = "Не удалось прочитать товары. Проверьте файл и номера колонок."

        mapping = mapping_values or draft["mapping"] or (
            asdict(preview.mapping) if preview else {**dict.fromkeys(COLUMN_FIELDS), "header_row": 0}
        )
        selected_sheet = (
            str(mapping_values.get("sheet_name")) if mapping_values and mapping_values.get("sheet_name")
            else draft["sheet_name"] or (preview.sheet_name if preview else sheets[0])
        )
        return render(
            request,
            "preview.html",
            status_code,
            draft=draft,
            sheets=sheets,
            selected_sheet=selected_sheet,
            mapping=mapping,
            column_fields=COLUMN_FIELDS,
            column_labels=COLUMN_LABELS,
            preview=preview,
            edits=edits or {},
            message=message,
        )

    @app.get("/", response_class=HTMLResponse)
    def home(request: Request) -> HTMLResponse:
        return render(request, "index.html", batches=storage.list_batches(database), message="")

    @app.post("/upload", response_class=HTMLResponse)
    async def upload(request: Request, file: UploadFile | None = File(default=None)) -> HTMLResponse:
        batches = storage.list_batches(database)
        filename = (file.filename if file else "") or ""
        filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
        if not filename:
            return render(
                request, "index.html", 400, batches=batches, message="Выберите файл .xlsx."
            )
        if Path(filename).suffix.lower() != ".xlsx":
            return render(
                request, "index.html", 400, batches=batches,
                message="Поддерживаются только файлы .xlsx.",
            )
        assert file is not None
        try:
            contents = await file.read(MAX_FILE_BYTES + 1)
        finally:
            await file.close()
        if not contents:
            return render(request, "index.html", 400, batches=batches, message="Файл пустой.")
        if len(contents) > MAX_FILE_BYTES:
            return render(
                request, "index.html", 400, batches=batches,
                message="Файл больше 10 МБ. Выберите книгу меньшего размера.",
            )

        draft_id = uuid4().hex
        path = uploads / f"{draft_id}.xlsx"
        path.write_bytes(contents)
        try:
            if not _sheet_names(path):
                raise ValueError("В книге нет листов")
        except Exception:
            path.unlink(missing_ok=True)
            LOGGER.warning("Rejected unreadable xlsx upload")
            return render(
                request, "index.html", 400, batches=batches,
                message="Не удалось открыть Excel. Проверьте, что это исправный файл .xlsx.",
            )
        try:
            storage.create_draft(database, draft_id, filename)
        except Exception:
            path.unlink(missing_ok=True)
            raise
        return RedirectResponse(f"/drafts/{draft_id}", status_code=303)

    @app.get("/drafts/{draft_id}", response_class=HTMLResponse)
    def show_preview(request: Request, draft_id: str) -> HTMLResponse:
        return preview_page(request, draft_id)

    @app.post("/drafts/{draft_id}/mapping", response_class=HTMLResponse)
    async def change_mapping(request: Request, draft_id: str) -> HTMLResponse:
        draft = storage.get_draft(database, draft_id)
        if draft is None:
            return error_page(request, "Черновик не найден или партия уже подтверждена.", 404)
        form = await request.form()
        values = {field: str(form.get(field, "")) for field in (*COLUMN_FIELDS, "header_row", "sheet_name")}
        try:
            sheet_name, mapping = _parse_mapping(form, _sheet_names(uploads / f"{draft_id}.xlsx"))
        except ValueError as exc:
            return preview_page(request, draft_id, message=str(exc), status_code=400, mapping_values=values)
        storage.update_draft_mapping(database, draft_id, sheet_name, mapping)
        return RedirectResponse(f"/drafts/{draft_id}", status_code=303)

    @app.post("/drafts/{draft_id}/confirm", response_class=HTMLResponse)
    async def confirm(request: Request, draft_id: str) -> HTMLResponse:
        draft = storage.get_draft(database, draft_id)
        if draft is None:
            return error_page(request, "Черновик не найден или партия уже подтверждена.", 404)
        try:
            preview = _preview(uploads / f"{draft_id}.xlsx", draft)
        except ValueError as exc:
            return preview_page(request, draft_id, message=str(exc), status_code=400)
        except Exception:
            LOGGER.exception("Could not confirm workbook")
            return preview_page(
                request, draft_id, message="Не удалось прочитать Excel. Проверьте файл и колонки.",
                status_code=400,
            )

        form = await request.form(max_fields=MAX_ROWS * 4 + 20)
        products: list[dict[str, Any]] = []
        edits: dict[int, dict[str, str]] = {}
        errors: list[str] = []
        for item in preview.products:
            row = item.row_number
            values = {
                field: str(form.get(f"{field}_{row}", "")).strip()
                for field in ("brand", "search_code", "alternate_code", "category")
            }
            edits[row] = values
            if not values["brand"] or not values["search_code"]:
                errors.append(f"Строка {row}: заполните бренд и основной артикул.")
            if any(len(value) > 500 for value in values.values()):
                errors.append(f"Строка {row}: значение длиннее 500 символов.")
            products.append({
                "row_number": row,
                "name": item.name,
                **values,
                "needs_confirmation": item.needs_confirmation,
                "issues": item.issues,
                "original_values": item.original_values,
            })
        if errors:
            message = " ".join(errors[:5])
            if len(errors) > 5:
                message += f" И ещё ошибок: {len(errors) - 5}."
            return preview_page(
                request, draft_id, message=message, status_code=400, edits=edits
            )
        storage.confirm_draft(database, draft_id, preview, products)
        return RedirectResponse(f"/batches/{draft_id}", status_code=303)

    @app.get("/batches/{batch_id}", response_class=HTMLResponse)
    def show_batch(request: Request, batch_id: str) -> HTMLResponse:
        batch = storage.get_batch(database, batch_id)
        if batch is None:
            return error_page(request, "Партия не найдена.", 404)
        groups: dict[str, list[dict[str, Any]]] = {}
        for product in batch["products"]:
            category = product["category"] or "Категория не определена"
            groups.setdefault(category, []).append(product)
        return render(request, "batch.html", batch=batch, groups=groups)


    def product_page(
        request: Request, product_id: int, *, message: str = "", status_code: int = 200
    ) -> HTMLResponse:
        product = jobs.get_product(database, product_id)
        if product is None:
            return error_page(request, "Товар не найден.", 404)
        history = jobs.list_jobs(database, product_id)
        latest = history[0] if history else None
        return render(
            request, "product.html", status_code,
            product=product,
            result=jobs.get_result(database, product_id),
            latest=latest,
            events=jobs.list_events(database, latest["id"]) if latest else [],
            stages=jobs.STAGE_NAMES,
            message=message,
            active=bool(latest and latest["status"] in jobs.ACTIVE),
        )

    @app.get("/products/{product_id}", response_class=HTMLResponse)
    def show_product(request: Request, product_id: int) -> HTMLResponse:
        return product_page(request, product_id)

    @app.post("/products/{product_id}/search", response_class=HTMLResponse)
    async def start_product_search(request: Request, product_id: int) -> HTMLResponse:
        if jobs.get_product(database, product_id) is None:
            return error_page(request, "Товар не найден.", 404)
        form = await request.form()
        try:
            selected = [int(value) for value in form.getlist("stages")]
        except (TypeError, ValueError):
            return product_page(
                request, product_id, message="Выберите этапы 1–4.", status_code=400
            )
        try:
            jobs.enqueue(database, product_id, selected)
        except ValueError as exc:
            return product_page(request, product_id, message=str(exc), status_code=400)
        return RedirectResponse(f"/products/{product_id}", status_code=303)
    return app