"""Read the cleaned assortment and build reproducible coverage records."""

from __future__ import annotations

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CATEGORY_GROUP_RULES: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("mobile_wearables", ("смартфон", "телефон", "планшет", "смарт-часы", "фитнес-браслет", "умные часы")),
    ("computers_components", ("ноутбук", "компьютер", "монитор", "видеокарт", "материнск", "оперативн", "накопител", "ssd", "процессор", "корпус", "кулер", "блок питания", "мыш", "клавиатур", "роутер", "сетев")),
    ("tv_audio", ("телевизор", "саундбар", "колонк", "акустик", "наушник", "микрофон", "проектор")),
    ("major_home_appliances", ("холодиль", "морозиль", "стираль", "сушиль", "посудомоеч", "духов", "варочн", "вытяжк", "плита", "микроволнов")),
    ("small_home_appliances", ("пылесос", "кофе", "чайник", "блендер", "миксер", "мультиварк", "тостер", "утюг", "фен", "бритв", "триммер")),
    ("tools_garden", ("дрель", "шуруповерт", "перфорат", "шлифов", "пила", "лобзик", "инструмент", "газон", "триммер сад", "мойка высокого")),
    ("climate", ("кондиционер", "обогрев", "увлажн", "очистител воздуха", "вентилятор")),
    ("gaming", ("игров", "геймпад", "джойстик", "кресло")),
    ("accessories", ("чехол", "кабель", "заряд", "переходник", "держатель", "сумка", "ремешок", "кронштейн")),
)


def category_group(category: str) -> str:
    value = str(category or "").casefold()
    return next((group for group, tokens in CATEGORY_GROUP_RULES if any(token in value for token in tokens)), "other")


@dataclass(frozen=True)
class CatalogCoverage:
    brand: str
    brand_canonical: str
    category: str
    category_group: str
    catalog_market: str
    unique_products: int
    source_rows: int
    sample_seller_skus: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        value = asdict(self)
        value["sample_seller_skus"] = list(self.sample_seller_skus)
        return value

    @property
    def market(self) -> str:
        """Backward-compatible alias for pre-v2 report consumers."""
        return self.catalog_market


@dataclass(frozen=True)
class CatalogCensus:
    source_file: str
    catalog_market: str
    coverages: tuple[CatalogCoverage, ...]
    expected_totals: dict[str, int]
    observed_totals: dict[str, int]

    @property
    def brands(self) -> tuple[str, ...]:
        return tuple(sorted({item.brand for item in self.coverages if item.brand}, key=str.casefold))

    @property
    def unresolved_brand_coverages(self) -> tuple[CatalogCoverage, ...]:
        return tuple(item for item in self.coverages if not item.brand)

    @property
    def market(self) -> str:
        return self.catalog_market


def _context_values(workbook) -> dict[str, Any]:
    rows = workbook["Контекст"].iter_rows(values_only=True)
    next(rows, None)
    return {str(key): value for key, value in rows if key}


def load_catalog_coverage(path: str | Path, *, market: str = "unknown") -> CatalogCensus:
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        context = _context_values(workbook)
        sheet = workbook["Бренд_Категория"]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        expected_header = ("Бренд", "Категория", "Уникальных товаров", "Строк в исходных выгрузках", "Примеры артикулов продавца")
        if tuple(header or ()) != expected_header:
            raise ValueError(f"Unexpected Бренд_Категория header: {header}")
        coverages: list[CatalogCoverage] = []
        for brand, category, unique_products, source_rows, samples in rows:
            display_brand = " ".join(str(brand or "").split())
            display_category = " ".join(str(category or "").split())
            if not display_category:
                continue
            sample_values = tuple(part.strip() for part in str(samples or "").split(";") if part.strip())
            coverages.append(CatalogCoverage(
                display_brand, display_brand.casefold(), display_category,
                category_group(display_category), market, int(unique_products or 0),
                int(source_rows or 0), sample_values,
            ))
        brands = {item.brand_canonical for item in coverages if item.brand}
        categories = {item.category.casefold() for item in coverages}
        observed = {
            "unique_products": sum(item.unique_products for item in coverages),
            "source_rows": sum(item.source_rows for item in coverages),
            "brands": len(brands),
            "categories": len(categories),
            "brand_category_pairs": len(coverages),
        }
        expected = {
            "unique_products": int(context["Уникальных товаров после фильтра"]),
            "source_rows": int(context["Рабочих строк после фильтра"]),
            "brands": int(context["Брендов после фильтра"]),
            "categories": int(context["Категорий после фильтра"]),
            "brand_category_pairs": int(context["Связок бренд + категория после фильтра"]),
        }
        if observed != expected:
            raise ValueError(f"Catalog totals do not reconcile: observed={observed}, expected={expected}")
        return CatalogCensus(str(path.resolve()), market, tuple(coverages), expected, observed)
    finally:
        workbook.close()
