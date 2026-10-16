"""Excel export with Russian resolved values and audit sheets."""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
import re,json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from . import jbl_pipeline, lenovo_pipeline, lenovo_presentation, attribute_projection, bosch_presentation, bosch_readiness, card_evidence, jobs, lg_batch, lg_presentation, manual_status, photo_metadata, samsung_readiness, storage
from .adapters.lg import lg_base_model
from .lg_identity import document_tied_to_article, photo_tied_to_article
from .display import display_access_error, display_name_ru, display_source, display_status, display_value

def _title(value,used):
    base=re.sub(r"[\[\]:*?/\\]"," ",value or "Категория не определена").strip()[:31] or "Категория"; candidate=base; n=2
    while candidate in used: suffix=f" {n}"; candidate=base[:31-len(suffix)]+suffix; n+=1
    used.add(candidate); return candidate

def _headers(sheet,values):
    sheet.append(values)
    for cell in sheet[1]: cell.font=Font(bold=True)
    sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions


def _photo_cells(item: dict) -> list[str | int]:
    size = item.get("verified_bytes")
    return [
        item.get("verified_width") or "не определено",
        item.get("verified_height") or "не определено",
        photo_metadata.format_file_size(size),
        item.get("verified_format") or "не определено",
    ]

def export_batch(database: Path,batch_id: str)->bytes:
    batch=storage.get_batch(database,batch_id)
    if not batch: raise ValueError("Партия не найдена.")
    from .xiaomi_identity import BRANDS as XIAOMI_BRANDS
    xiaomi_only=all(p['brand'].strip().casefold() in XIAOMI_BRANDS or p['brand'].strip().casefold()=='dreame' for p in batch['products'])
    book=Workbook(); book.remove(book.active); used=set(); groups={}
    for p in batch["products"]: groups.setdefault(p["category"] or "Категория не определена",[]).append(p)
    for category,products in groups.items():
        sheet=book.create_sheet(_title(category,used)); rows={p["id"]:attribute_projection.final_attribute_rows(database,p["id"]) for p in products}
        # The common presentation vocabulary supplies semantic labels/scopes.
        # Language aliases of the same visible concept share one product column;
        # conflicting values are exposed rather than choosing the last alias.
        keys=sorted({r['display_name'] for rr in rows.values() for r in rr},key=lambda k:(min(r.get('parent_group','') for rr in rows.values() for r in rr if r['display_name']==k),k.casefold()))
        if 'Цвет' not in keys:keys.append('Цвет')
        labels={k:k for k in keys}
        _headers(sheet,["Строка","Название","Бренд","Поисковый идентификатор" if xiaomi_only else "Полный артикул","Модель" if xiaomi_only else "Базовая модель",*[labels[k] for k in keys]])
        for p in products:
            is_lg = lg_batch.is_lg(p)
            visible={}
            for r in rows[p['id']]:
                value=r.get('resolved') or {};shown=value.get('display_value','')
                if value.get('conflict'):
                    shown='Расхождение: '+'; '.join(dict.fromkeys(f.get('display_value','') for f in r.get('sources',{}).values()))
                elif (is_lg or p['brand'].strip().upper() in {'SAMSUNG','LENOVO'}) and not (value.get('full_sku_confirmed') or value.get('status') in {'model_confirmed_official','hardware_confirmed_official'}):shown=''
                if shown:visible.setdefault(r['display_name'],[]).append(shown)
            resolved={k:(vals[0] if len(vals)==1 else 'Расхождение: '+'; '.join(vals)) for k,items in visible.items() for vals in [list(dict.fromkeys(items))]}
            from .canonical_attributes import color_result
            from .photo_inventory import evidence_kind
            resolved['Цвет']=color_result(p,rows[p['id']],jobs.get_source_pages(database,p['id']),card_evidence.load(database,p['id'],evidence_kind(p)) or {})['value']
            model=lg_base_model(p["search_code"]) if p["brand"].strip().upper()=="LG" else ""
            if p['brand'].strip().casefold()=='dreame':
                identity=(card_evidence.load(database,p['id'],'dreame') or {}).get('identity',{})
                model=identity.get('model','') if identity.get('model_relation')=='model_confirmed' else ''
            sheet.append([p["row_number"],(p["name"] or p["search_code"]),p["brand"],p["search_code"],model,*[resolved.get(k,"") for k in keys]])
        if any(p['brand'].strip().casefold()=='dreame' for p in products):
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width=min(60,max(14,max(len(str(c.value or '')) for c in column)+2))
            for row in sheet.iter_rows():
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
                sheet.row_dimensions[row[0].row].height=48
    has_lg=any(lg_batch.is_lg(p) for p in batch["products"])
    has_lg_global=any(s["source_key"]=="lg_global" for p in batch["products"] for s in jobs.get_source_pages(database,p["id"]))
    has_sulpak_facts=any(f["source_key"]=="sulpak" for p in batch["products"] for f in jobs.get_facts(database,p["id"]))
    has_bosch=any(s["source_key"]=="bosch_home" for p in batch["products"] for s in jobs.get_source_pages(database,p["id"]))
    has_samsung=any(s["source_key"]=="samsung" for p in batch["products"] for s in jobs.get_source_pages(database,p["id"]))  # the Samsung column and sheet exist only when the batch has a Samsung page
    check=book.create_sheet(_title("Проверка источников",used)); _headers(check,["Товар","Полный артикул","Характеристика",*(["LG Казахстан","LG Россия"] if has_lg else []),*(["Sulpak"] if has_sulpak_facts else []),*(["LG other region"] if has_lg_global else []),*(["Samsung"] if has_samsung else []),*(["Bosch Home"] if has_bosch else []),"Итог","Причина","Статус","Все спорные значения","\u0418\u0441\u0445\u043e\u0434\u043d\u044b\u0435 \u043d\u0430\u0437\u0432\u0430\u043d\u0438\u044f"])
    for p in batch["products"]:
        conflict_facts = {}
        for fact in jobs.get_facts(database,p["id"]):
            detail = f"{fact['site_name']}: {fact['raw_value']}"
            if fact.get("section"):
                detail += f" (раздел «{fact['section']}»)"
            values = conflict_facts.setdefault(fact["normalized_name"], [])
            if detail not in values:
                values.append(detail)
        for row in jobs.comparison_rows(database,p["id"]):
            s=row["sources"]; r=(lg_presentation.safe_composite_dimensions(row).get("resolved")
                                  if lg_batch.is_lg(p) else row.get("resolved")) or {}
            detail = "; ".join(conflict_facts.get(row["normalized_name"], [])) if r.get("conflict") else ""
            label = (bosch_presentation.label(row) if p["brand"].strip().upper() == "BOSCH" else row["display_name"])
            if r.get("conflict") and label == "Дополнительная характеристика" and row["raw_names"]:
                label = row["raw_names"][0]
            check.append([(p["name"] or p["search_code"]),p["search_code"],label,*([s.get("lg_kz",{}).get("raw_value",""),s.get("lg_ru",{}).get("raw_value","")] if has_lg else []),*([s.get("sulpak",{}).get("raw_value","")] if has_sulpak_facts else []),*([s.get("lg_global",{}).get("raw_value","")] if has_lg_global else []),*([s.get("samsung",{}).get("raw_value","")] if has_samsung else []),*([s.get("bosch_home",{}).get("raw_value","")] if has_bosch else []),r.get("display_value",""),r.get("reason",""),r.get("display_status",""),detail, "; ".join(dict.fromkeys(row["raw_names"]))])
    semantic_rows=[]
    for p in batch['products']:
        semantic=card_evidence.load(database,p['id'],'semantic_decomposition') or {}
        for source,source_evidence in semantic.get('sources',{}).items():
            for block in source_evidence.get('decomposed',[]):
                for child in block['children']:
                    semantic_rows.append([p['name'] or p['search_code'],source_evidence.get('url') or source,block['section'],block['raw_name'],block['raw_value'],child['key'],child['label'],child['value'],child['unit'],block['rule']])
    if semantic_rows:
        raw_sheet=book.create_sheet(_title('Семантика и исходные значения',used))
        _headers(raw_sheet,['Товар','Источник','Раздел','Исходное поле','Полное исходное значение','Canonical key','Характеристика','Нормализованное значение','Единица','Правило'])
        for row in semantic_rows:raw_sheet.append(row)
        raw_sheet.column_dimensions['E'].width=80
        for row in raw_sheet:
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    source_sheet=book.create_sheet(_title("Источники",used)); _headers(source_sheet,["Товар","Полный артикул","Сайт","Найденная модель","Уровень совпадения","Доказательство","Дата","Ошибка","URL"])
    attributes_sheet=book.create_sheet(_title('Атрибуты',used));_headers(attributes_sheet,['Товар','Категория','Группа','Основные / особенности','Характеристика','Значение','Статус','Источник'])
    schema_sheet=book.create_sheet(_title('Схема категорий',used));_headers(schema_sheet,['Категория','Группа','Характеристика','Основные / особенности','Бренды'])
    observed_schema={}
    from .card_presentation import present_card_rows
    for p in batch['products']:
        entries=attribute_projection.final_attribute_rows(database,p['id'])
        presented=present_card_rows(entries,jobs.get_facts(database,p['id']),jobs.get_source_pages(database,p['id']),p['category'])
        for role,sections in presented['groups'].items():
            for group,items in sections.items():
                for item in items:
                    value=item.get('resolved') or {};shown=value.get('display_value','')
                    if value.get('conflict'):shown='Расхождение: '+'; '.join(dict.fromkeys(f.get('display_value','') for f in item.get('sources',{}).values()))
                    attributes_sheet.append([p['name'] or p['search_code'],p['category'],group,'Основные' if role=='core' else 'Особенности',item['display_name'],shown,value.get('display_status','Не подтверждено'),'; '.join(item.get('sources',{}))])
                    key=(p['category'],group,item['display_name'],role);observed_schema.setdefault(key,set()).add(p['brand'])
    for (cat,group,label,role),brands in sorted(observed_schema.items()):schema_sheet.append([cat,group,label,'Основные' if role=='core' else 'Особенности',', '.join(sorted(brands))])
    for p in batch["products"]:
        for s in jobs.get_source_pages(database,p["id"]):
            if s["source_key"] in {"lg_kz_support", "lg_ru_support"}:
                level = {"full_sku":"Связь инструкции с артикулом","base_model":"Поддержка семейства/другого варианта","component_only":"Поддержка одного компонента"}.get(s["match_level"],"Кандидат поддержки")
            else:
                level = {"full_sku":"Полный артикул","code_in_page_text":"Артикул только в тексте страницы","base_model":"Базовая модель","mismatch":"Несоответствие","unknown":"Не проверено"}.get(s["match_level"],s["match_level"])
            if s["source_key"] == "sulpak" and s["match_level"] == "unknown" and s["evidence"] and not s["error"]:
                level = "Комплект не подтверждён"
            source_sheet.append([(p["name"] or p["search_code"]),p["search_code"],s["site_name"],s["found_model"],level,s["evidence"],s["fetched_at"],display_access_error(s["error"]),s["url"]])
    docs=book.create_sheet(_title("Инструкции",used)); _headers(docs,["Товар","Документ","Язык","Дата","Размер","Основной","Модель поддержки","Источник","Прямая ссылка","Связь с артикулом","Тип источника","Статус проверки"])
    photo_columns=["Ширина, px","Высота, px","Размер файла","Формат"]
    photos=book.create_sheet(_title("Фотографии",used)); _headers(photos,["Товар","Источник","Тип","URL","Подтверждение варианта",*photo_columns,"Роль","Выбрано для карточки","SHA256","Проверка байтов","Источник галереи"])
    photo_candidates = None
    from . import playstation_pipeline
    from .playstation_identity import BRANDS as PLAYSTATION_BRANDS
    from . import xbox_pipeline
    from .xbox_identity import BRANDS as XBOX_BRANDS
    from . import xiaomi_pipeline
    from .xiaomi_identity import BRANDS as XIAOMI_BRANDS
    from . import razer_pipeline
    from . import electrolux_pipeline
    for p in batch["products"]:
        el=p['brand'].strip().casefold()=='electrolux'
        el_evidence=card_evidence.load(database,p['id'],'electrolux') if el else None
        source_pages = jobs.get_source_pages(database, p["id"])
        ps=p['brand'].strip().casefold() in PLAYSTATION_BRANDS
        ps_evidence=card_evidence.load(database,p['id'],'playstation') if ps else None
        xb=p['brand'].strip().casefold() in XBOX_BRANDS
        xb_evidence=card_evidence.load(database,p['id'],'xbox') if xb else None
        lg = lg_batch.is_lg(p)
        samsung = p["brand"].strip().upper() == "SAMSUNG"
        xm = p['brand'].strip().casefold() in XIAOMI_BRANDS or p['brand'].strip().casefold()=='dreame'
        xm_evidence=card_evidence.load(database,p['id'],'dreame' if p['brand'].strip().casefold()=='dreame' else 'xiaomi') if xm else None
        rz = p['brand'].strip().casefold() == 'razer'
        rz_evidence=card_evidence.load(database,p['id'],'razer') if rz else None
        jbl = p["brand"].strip().upper() == "JBL"
        jbl_evidence = card_evidence.load(database,p["id"],"jbl") if jbl else None
        lenovo = p["brand"].strip().upper() == "LENOVO"
        lenovo_evidence = card_evidence.load(database,p["id"],"lenovo") if lenovo else None
        saved_documents = (manual_status.effective_documents(database,p["id"],p["search_code"])
                           if lg else jobs.get_documents(database,p["id"]))
        for d in saved_documents:
            tied = (electrolux_pipeline.document_verified(d,el_evidence) if el else xiaomi_pipeline.document_verified(d,xm_evidence) if xm else d.get("identity_confirmed", False) if lg else
                    bosch_presentation.document_verified(d, card_evidence.load(database,p["id"],"bosch_home"), source_pages)
                    if p["brand"].strip().upper() == "BOSCH" else
                    samsung_readiness.document_verified(d, card_evidence.load(database,p["id"],"samsung_documents"))
                    if samsung else lenovo_pipeline.document_verified(d, lenovo_evidence) if lenovo else jbl_pipeline.document_verified(d, jbl_evidence) if jbl else True)
            docs.append([(p["name"] or p["search_code"]),d["title"],d["language"],d["document_date"],d["size"],"Да" if d["is_primary"] else "Нет",d["support_model"],d["source_url"],d["direct_url"],
                         "Подтверждена" if tied else "PDF проверен, связь с вариантом не подтверждена",
                         "LG" if d["source_key"].startswith("lg_") else display_source(
                             d["source_key"], next((page["site_name"] for page in source_pages
                                                    if page["source_key"] == d["source_key"]), "")),
                         "Проверена" if tied else "Не проверена"])
        if lg or (samsung and manual_status.completed_search(database,p["id"])):
            status = manual_status.russian_status(database, p["id"], p["search_code"])
            if status != manual_status.VERIFIED:
                audit = manual_status.completed_search(database, p["id"])
                support_url = next((item.get("support_url") or item.get("url", "")
                                    for item in audit["evidence"] if item.get("official")), "") if audit else ""
                if samsung and not support_url:
                    support_url = (card_evidence.load(database,p["id"],"samsung_documents") or {}).get("support", {}).get("url", "")
                reason = manual_status.unchecked_reason(database, p["id"], p["search_code"])
                docs.append([(p["name"] or p["search_code"]), "", "", "", "", "", "",
                             support_url, "", reason, "LG" if lg and audit else "Samsung" if samsung else "", status])
        from . import photo_inventory
        priority=p['brand'].strip().casefold() in {'jbl','samsung','lg','apple','playstation','xbox','razer','hyperx','xiaomi','redmi','poco','dreame','electrolux'}
        gallery=photo_inventory.enriched_photos(database,p['id']) if priority else jobs.get_photo_candidates(database,p['id'],include_excluded=p['brand'].strip().upper() in ('BOSCH','SAMSUNG'))
        # Enriched entries already contain byte metadata. Re-projecting the
        # entire gallery for every image made large batch downloads quadratic.
        rich_by_id={x['id']:x for x in (gallery if priority else photo_inventory.enriched_photos(database,p['id'],dedupe=False))}
        for item in gallery:
            tied = electrolux_pipeline.photo_verified(item,el_evidence) if el else xiaomi_pipeline.photo_verified(item,xm_evidence) if xm else razer_pipeline.photo_verified(item,rz_evidence) if rz else xbox_pipeline.photo_verified(item,xb_evidence) if xb else playstation_pipeline.photo_verified(item,ps_evidence) if ps else photo_tied_to_article(item, source_pages) if lg else samsung_readiness.photo_verified(item, source_pages, p["search_code"]) if samsung else lenovo_pipeline.photo_verified(item, lenovo_evidence) if lenovo else jbl_pipeline.photo_verified(item, jbl_evidence) if jbl else True
            if priority:tied=item['identity_confirmed']
            kind = {"product_gallery":"Товарная галерея","feature":"Особенности и инфографика","marketing":"Маркетинговые материалы"}.get(item["kind"],item["kind"])
            if item['kind']=='excluded' and priority:continue
            if item["kind"] == "excluded" and p["brand"].strip().upper() in ("BOSCH", "SAMSUNG"):
                if photo_candidates is None:
                    photo_candidates = book.create_sheet(_title("Фото-кандидаты",used))
                    _headers(photo_candidates,["Товар","Источник","Тип","URL","Выбрано для просмотра","Статус",*photo_columns])
                reason = {"base_model_only": "Только базовая модель",
                          "not_model_bound": "Нет доказанной связи с точной моделью"}.get(
                              item["excluded_reason"], item["excluded_reason"] or "Исключено")
                photo_candidates.append([(p["name"] or p["search_code"]),item["site_name"],
                                         "Исключённый кандидат",item["url"],"Нет",reason,*_photo_cells(item)])
                continue
            if tied and item['kind']!='excluded':
                from . import photo_inventory
                rich=rich_by_id.get(item['id'],item)
                photos.append([(p["name"] or p["search_code"]),item["site_name"],kind,item["url"],"Фото модели; цвет поставщика не подтверждён" if xm else "Цвет подтверждён" if lenovo else "Подтверждён",*_photo_cells(item),rich.get('role_label','Не определена'),'Да' if item['selected'] else 'Нет',rich.get('sha256',''),'Проверены' if rich.get('byte_verified') else 'Не проверены',rich.get('source_url','')])
            elif item["selected"] and not tied:
                if photo_candidates is None:
                    photo_candidates = book.create_sheet(_title("Фото-кандидаты",used))
                    _headers(photo_candidates,["Товар","Источник","Тип","URL","Выбрано для просмотра","Статус",*photo_columns])
                photo_candidates.append([(p["name"] or p["search_code"]),item["site_name"],kind,item["url"],
                                         "Да" if item["selected"] else "Нет","Связь с артикулом не подтверждена",*_photo_cells(item)])
            elif (priority or samsung or lenovo or jbl or rz) and item["kind"] != "excluded":
                if photo_candidates is None:
                    photo_candidates = book.create_sheet(_title("Фото-кандидаты",used))
                    _headers(photo_candidates,["Товар","Источник","Тип","URL","Выбрано для просмотра","Статус",*photo_columns])
                photo_candidates.append([(p["name"] or p["search_code"]),item["site_name"],kind,item["url"],
                                         "Нет","Связь с артикулом не подтверждена" if not tied else "Не выбрано",*_photo_cells(item)])
    lenovo_products = [p for p in batch["products"] if p["brand"].strip().upper() == "LENOVO"]
    if lenovo_products:
        candidates=book.create_sheet(_title("Конфигурации-кандидаты",used))
        _headers(candidates,["Товар","Артикул","Раздел","Исходное поле","Значение","Связь","Причина"])
        manual_candidates=book.create_sheet(_title("Документы-кандидаты Lenovo",used))
        _headers(manual_candidates,["Товар","Артикул","Тип","Название","Язык","URL","Связь","Проверка"])
        ready=book.create_sheet(_title("Готовность Lenovo",used))
        _headers(ready,["Товар","Артикул","Готовность","Пробелы","Русская инструкция"])
        for p in lenovo_products:
            ev=card_evidence.load(database,p["id"],"lenovo") or {}
            for d in ev.get("manuals",[]):
                docs.append([p['name'],d.get('title',d['type']),d.get('language','Не проверена'),'','','Да' if d.get('verified') else 'Нет',ev.get('model_name',''),d['url'],d.get('chapter_url',d['url']),d.get('model_relation','unverified'),d['type'],'Проверена' if d.get('verified') else 'Не проверена'])
            for c in ev.get("configuration_candidates",[]):
                candidates.append([p["name"],p["search_code"],lenovo_presentation.section(c["section"]),c["raw_label"],c["value"],lenovo_presentation.SCOPES.get(c["scope"],c["scope"]),"Условное значение; требуется проверка для полного артикула"])
            for d in ev.get("manuals",[]):
                manual_candidates.append([p["name"],p["search_code"],d["type"],d["title"],d["language"],d["url"],lenovo_presentation.SCOPES.get(d["relation"],d["relation"]),"Проверена" if d["verified"] else "Не проверена"])
            r=lenovo_pipeline.card_readiness(database,p["id"])
            ready.append([p["name"],p["search_code"],lenovo_presentation.VERDICTS.get(r["verdict"],r["verdict"]),"; ".join(lenovo_presentation.gap_labels(r)),r["manual_status"]])
    if any(lg_batch.is_lg(p) for p in batch["products"]):
        ready=book.create_sheet(_title("Готовность LG",used))
        _headers(ready,["Товар","Полный артикул","Статус задания","Готовность карточки","Проверенные этапы","Причины пробелов","Справочно"])
        for p in batch["products"]:
            if not lg_batch.is_lg(p): continue
            history=jobs.list_jobs(database,p["id"])
            latest=history[0] if history else None
            summary=lg_batch.card_summary(database,p["id"],latest)
            ready.append([(p["name"] or p["search_code"]),p["search_code"],f"{lg_batch.JOB_LABELS.get(latest['status'], latest['status'])} ({latest['status']})" if latest else "Не запускалось",
                          f"{summary['label']} ({summary['verdict']})",
                          ", ".join(jobs.STAGE_NAMES[n] for n in latest["stages"]) if latest else "",
                          " ".join(summary["reasons"]),
                          " ".join(summary.get("advisories",[]))])
    if has_samsung:
        ready=book.create_sheet(_title("Готовность Samsung",used))
        include_dealer = any(samsung_readiness.card_readiness(database, p["id"])["dealer_disputes"] for p in batch["products"] if p["brand"].strip().upper() == "SAMSUNG")
        headers=["Товар","Полный артикул","Вердикт","Совпадение страницы","Открыто по варианту","Блокирующие пробелы","Справочно","Инструкция: русский по тексту","Инструкция: связь со страницей","Инструкция: код каталога в тексте","Фото выбрано / найдено","Расхождения с дилером","Открытые проверки","Инструкция: основание принятия","Причины пробелов"]
        if not include_dealer: headers.pop(11)
        _headers(ready, headers)
        for p in batch["products"]:
            if not any(s["source_key"]=="samsung" for s in jobs.get_source_pages(database,p["id"])): continue
            r=samsung_readiness.card_readiness(database,p["id"]); i=r["instruction"]
            yn=lambda v:"" if v is None else ("Да" if v else "Нет")
            named="" if not i["saved"] else "точный код" if i["names_catalog_code_exactly"] else "пометка: "+(i["mark"] or "точный код в PDF не назван")
            row=[(p["name"] or p["search_code"]),p["search_code"],r["verdict"],r["page_match_level"],"; ".join(r["open_variant_differences"]),", ".join(r["blocking_gaps"]),", ".join(r["advisory_gaps"]),yn(i["russian_by_text"]),i["tied_by_official_page"] or "",named,f"{r['official_photos_selected']} / {r['official_photos']}","; ".join(d["field"] for d in r["dealer_disputes"]),", ".join(r["open_reviews"]),(i["acceptance_basis"]+(": принята" if i["accepted"] else ": не принята")) if i["acceptance_basis"] else "","; ".join(f"{k}: {v}" for k,v in r["gap_reasons"].items())]
            if not include_dealer: row.pop(11)
            ready.append(row)
    if has_bosch:
        ready=book.create_sheet(_title("\u0413\u043e\u0442\u043e\u0432\u043d\u043e\u0441\u0442\u044c Bosch Home",used))
        _headers(ready,["\u0422\u043e\u0432\u0430\u0440","\u041a\u043e\u0434 \u043a\u0430\u0442\u0430\u043b\u043e\u0433\u0430","\u0421\u0442\u0430\u0442\u0443\u0441 \u0437\u0430\u0434\u0430\u043d\u0438\u044f","\u0413\u043e\u0442\u043e\u0432\u043d\u043e\u0441\u0442\u044c \u043a\u0430\u0440\u0442\u043e\u0447\u043a\u0438","GTIN","E-Nr","\u0420\u0435\u0432\u0438\u0437\u0438\u044f","\u0425\u0430\u0440\u0430\u043a\u0442\u0435\u0440\u0438\u0441\u0442\u0438\u043a\u0438","\u0424\u043e\u0442\u043e \u0432\u044b\u0431\u0440\u0430\u043d\u043e / \u043d\u0430\u0439\u0434\u0435\u043d\u043e","\u0421\u0435\u043c\u0435\u0439\u0441\u0442\u0432\u043e \u0440\u0443\u043a\u043e\u0432\u043e\u0434\u0441\u0442\u0432\u0430","\u0422\u043e\u0447\u043d\u044b\u0439 \u043a\u043e\u0434 \u0432 PDF","\u041f\u0440\u043e\u0447\u0438\u0435 PDF \u0431\u0435\u0437 \u043f\u0440\u043e\u0432\u0435\u0440\u043a\u0438","\u041f\u0440\u043e\u0431\u0435\u043b\u044b"])
        for p in batch["products"]:
            if not any(src["source_key"]=="bosch_home" for src in jobs.get_source_pages(database,p["id"])): continue
            r=bosch_readiness.card_readiness(database,p["id"])
            history=jobs.list_jobs(database,p["id"])
            status=history[0]["status"] if history else ""
            ready.append([(p["name"] or p["search_code"]),p["search_code"],status,r["verdict"],r["gtin"],r["page_enr"] or "",r["revision_status"],r["official_facts"],f"{r['official_photos_selected']} / {r['official_photos']}",r["manual_family"],"\u0414\u0430" if r["manual_exact_code_in_pdf"] else "\u041d\u0435\u0442",len(r["other_manuals_unverified"]),", ".join(r["blocking_gaps"]+r["advisory_gaps"])])
    razer_products=[p for p in batch['products'] if p['brand'].strip().casefold()=='razer']
    if razer_products:
        rr=book.create_sheet(_title('Готовность Razer',used));_headers(rr,['Товар','Код','Готовность','Модель','Конфигурация','Русская инструкция','Характеристики','Пробелы'])
        rc=book.create_sheet(_title('Кандидаты Razer',used));_headers(rc,['Товар','Раздел','Исходное поле','Значение','Причина'])
        rv=book.create_sheet(_title('Варианты Razer',used));_headers(rv,['Товар','Запрошенное поле','Значение','Подтверждение'])
        rd=book.create_sheet(_title('Документы Razer',used));_headers(rd,['Товар','Тип','Название','URL','Язык','Проверка'])
        raw=book.create_sheet(_title('Исходные факты Razer',used));_headers(raw,['Товар','Раздел','Исходное поле','Значение','URL'])
        from .razer_presentation import REASONS,CONFIG_LABELS
        for p in razer_products:
            ev=card_evidence.load(database,p['id'],'razer') or {};r=razer_pipeline.card_readiness(database,p['id']);i=r['identity'];rr.append([p['name'],p['search_code'],'Готова' if r['verdict']=='export_ready' else 'Не готова',i.get('model','Не подтверждена'),'Полный артикул' if i.get('exact_sku') else 'Не подтверждена',r['manual_status'],r['confirmed_specs'],'; '.join(razer_pipeline.GAPS[g] for g in r['gaps'])])
            for c in ev.get('rejected_specs',[]):rc.append([p['name'],c['section'],c['raw_label'],c['value'],REASONS.get(c['reason'],c['reason'])])
            for k,v in ev.get('requested_configuration',{}).items():rv.append([p['name'],CONFIG_LABELS.get(k,k),v,'Требует точного источника'])
            for c in ev.get('raw_specs',[]):raw.append([p['name'],c['section'],c['raw_label'],c['value'],ev.get('source_url','')])
            for d in ev.get('manuals',[]):rd.append([p['name'],d['type'],d['title'],d['url'],d['language'],'Проверена' if d.get('verified') else 'Не проверена'])
        for sheet in (rr,rc,rv,rd,raw):
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=38
    jbl_products = [p for p in batch["products"] if p["brand"].strip().upper() == "JBL"]
    if jbl_products:
        ready_jbl=book.create_sheet(_title("Готовность JBL",used))
        _headers(ready_jbl,["Товар","Полный артикул","Готовность","Пробелы","Русская инструкция","Связь модели","Связь варианта"])
        candidates_jbl=book.create_sheet(_title("Характеристики-кандидаты JBL",used))
        _headers(candidates_jbl,["Товар","Раздел","Исходное поле","Значение","Причина"])
        documents_jbl=book.create_sheet(_title("Документы-кандидаты JBL",used))
        _headers(documents_jbl,["Товар","Тип","Название","URL","Связь","Проверка"])
        for p in jbl_products:
            ev=card_evidence.load(database,p["id"],"jbl") or {};r=jbl_pipeline.card_readiness(database,p["id"])
            ready_jbl.append([p["name"],p["search_code"],jbl_pipeline.VERDICTS[r["verdict"]],"; ".join(jbl_pipeline.GAPS[g] for g in r["gaps"]),r["manual_status"],r["model_identity"],r["variant_identity"]])
            for c in ev.get("rejected_specs",[]):candidates_jbl.append([p["name"],c["section"],c["raw_label"],c["value"],c["reason"]])
            for d in ev.get("manuals",[]):documents_jbl.append([p["name"],d["type"],d["title"],d["url"],d["relation"],"Проверена" if d["verified"] else "Не проверена"])
    from . import playstation_pipeline
    from .playstation_identity import BRANDS as PLAYSTATION_BRANDS
    ps_products=[p for p in batch['products'] if p['brand'].strip().casefold() in PLAYSTATION_BRANDS]
    if ps_products:
        ps_ready=book.create_sheet(_title('Готовность PlayStation',used))
        _headers(ps_ready,['Товар','Артикул','Готовность','Модель','CFI','Конфигурация','Пользовательская инструкция','Пробелы','Область готовности','Розничный SKU','Confirmed specs','Quick Start'])
        ps_config=book.create_sheet(_title('Конфигурация PlayStation',used))
        _headers(ps_config,['Артикул','Поле','Значение','Официальный источник'])
        ps_candidates=book.create_sheet(_title('Кандидаты PlayStation',used))
        _headers(ps_candidates,['Артикул','Раздел','Исходное поле','Значение','Причина'])
        ps_documents=book.create_sheet(_title('Документы PlayStation',used))
        _headers(ps_documents,['Артикул','Тип','Название','URL','Связь','Проверка','Язык'])
        ps_links=book.create_sheet(_title('Связь SKU и CFI',used))
        _headers(ps_links,['Артикул','Розничный SKU','Опубликованный CFI','Связь','Официальный источник','Связь с запрошенным кодом','Применена к карточке'])
        for p in ps_products:
            ev=card_evidence.load(database,p['id'],'playstation') or {};r=playstation_pipeline.card_readiness(database,p['id']);i=r['identity']
            ps_ready.append([p['name'],p['search_code'],'Готова' if r['verdict']=='export_ready' else 'Не готова',i.get('model','unproven'),ev.get('hardware_model',''),i.get('configuration','unproven'),r['manual_status'],'; '.join(playstation_pipeline.GAPS[g] for g in r['gaps']),'Оборудование CFI' if ev.get('configuration_scope')=='hardware_only' else 'Розничный вариант',ev.get('retail_sku',''),r['confirmed_specs'],ev.get('quick_start_status','Не проверена')])
            for key,value in r['configuration_fields'].items():
                label={'storage':'Объём накопителя','color':'Цвет','disc':'Оптический привод','bundle':'Комплект','bundle_contents':'Комплектация','storefront':'Регион официального магазина'}.get(key,key)
                ps_config.append([p['search_code'],label,value,ev.get('exact_official_pdp','')])
            for c in ev.get('configuration_candidates',[]):ps_candidates.append([p['search_code'],c['section'],c['raw_label'],c['value'],c['reason']])
            for d in ev.get('manuals',[]):ps_documents.append([p['search_code'],d['type'],d['title'],d['url'],d['relation'],'Проверена' if d['verified'] else 'Не проверена',d.get('language','Не проверена')])
            for x in ev.get('sku_cfi_relations',[]):ps_links.append([p['search_code'],x.get('retail_sku',''),', '.join(x.get('published_cfi',[])),x.get('relation','unproven'),x.get('source_url',''),x.get('requested_relation','unproven'),'Да' if x.get('accepted_for_request') else 'Нет — candidate'])
            for c in ev.get('photo_candidates',[]):
                if photo_candidates is None:
                    photo_candidates=book.create_sheet(_title('Фото-кандидаты',used));_headers(photo_candidates,['Товар','Источник','Тип','URL','Выбрано для просмотра','Статус',*photo_columns])
                photo_candidates.append([p['name'],'PlayStation','Кандидат',c['url'],'Нет',c['reason'],'не определено','не определено','не определено','не определено'])
    xbox_products=[p for p in batch['products'] if p['brand'].strip().casefold() in XBOX_BRANDS]
    xbox_sheets=[]
    if xbox_products:
        for title,headers in [('Готовность Xbox',['Товар','Артикул','Готовность','Модель','Конфигурация','Оборудование','Пользовательская инструкция','Пробелы','Подтверждённых характеристик']),('Конфигурация Xbox',['Артикул','Поле','Значение','Официальный источник']),('Идентификаторы Xbox',['Артикул','Тип идентификатора','Значение','Регион','Источник']),('Кандидаты Xbox',['Артикул','Раздел','Исходное поле','Значение','Причина']),('Документы Xbox',['Артикул','Тип','Название','URL','Связь','Проверка','Язык']),('Связи Xbox',['Артикул','Product ID / catalog ID','SKU-кандидаты','Hardware number','Manufacturer part','Регион','Связь','Источник']),('Принятые факты Xbox',['Артикул','Раздел','Поле','Значение','Область','Источник'])]:
            sheet=book.create_sheet(_title(title,used));_headers(sheet,headers);xbox_sheets.append(sheet)
        xr,xc,xi,xk,xd,xrel,xa=xbox_sheets
        for p in xbox_products:
            ev=card_evidence.load(database,p['id'],'xbox') or {};r=xbox_pipeline.card_readiness(database,p['id']);i=r['identity']
            xr.append([p['name'],p['search_code'],'Готова' if r['verdict']=='export_ready' else 'Не готова',*[xbox_pipeline.SCOPES.get(i.get(k,'unproven'),'Не подтверждено') for k in ('model','configuration','hardware')],r['manual_status'],'; '.join(xbox_pipeline.GAPS[g] for g in r['gaps']),r['confirmed_specs']])
            for k,v in r['configuration_fields'].items():xc.append([p['search_code'],{'storage':'Встроенный накопитель','color':'Цвет','disc':'Оптический привод','bundle':'Название комплекта','bundle_contents':'Комплектация','included_game':'Игра в комплекте','game_pass':'Game Pass в комплекте','storefront':'Регион официального магазина'}.get(k,k),v,ev.get('exact_official_pdp','')])
            for k,v in ev.get('identifiers',{}).items():xi.append([p['search_code'],k,v or 'Не подтверждён',ev.get('requested_region',''),ev.get('exact_official_pdp','') or ev.get('support_url','')])
            for c in ev.get('configuration_candidates',[]):xk.append([p['search_code'],c['section'],c['raw_label'],c['value'],c['reason']])
            for d in ev.get('manuals',[]):xd.append([p['search_code'],d['type'],d['title'],d['url'],d['relation'],'Проверена' if d['verified'] else 'Не проверена',d.get('language','Не проверена')])
            for relation in ev.get('identity_relations',[]):xrel.append([p['search_code'],relation.get('product_id') or relation.get('value'),', '.join(relation.get('sku_ids',[])),', '.join(relation.get('hardware_model_numbers',[])),', '.join(relation.get('manufacturer_parts',[])),relation.get('region'),relation.get('relation'),relation.get('url') or relation.get('source_url')])
            for fact in ev.get('accepted_specs',[]):xa.append([p['search_code'],fact['section'],fact['label'],fact['value'],{'model':'Модель','hardware_configuration':'Оборудование / конфигурация','retail':'Retail вариант'}.get(fact['scope'],fact['scope']),fact['source_url']])
            for c in ev.get('photo_candidates',[]):
                if photo_candidates is None:
                    photo_candidates=book.create_sheet(_title('Фото-кандидаты',used));_headers(photo_candidates,['Товар','Источник','Тип','URL','Выбрано для просмотра','Статус',*photo_columns])
                if photo_candidates.cell(1,1).value=='SKU':photo_candidates.append([p['search_code'],p['name'],c['url'],c['reason'],'candidate','unproven',None,None,None,None])
                else:photo_candidates.append([p['name'],'Xbox / Microsoft','Кандидат',c['url'],'Нет',c['reason'],None,None,None,None])
    apple_products = [p for p in batch["products"] if p["brand"].strip().upper() == "APPLE"]
    if apple_products:
        from . import apple_pipeline
        ready_apple=book.create_sheet(_title("Готовность Apple",used))
        _headers(ready_apple,["Товар","Артикул","Готовность","Модель","Конфигурация","Вариант","Русская инструкция","Пробелы"])
        candidates_apple=book.create_sheet(_title("Варианты-кандидаты Apple",used))
        _headers(candidates_apple,["Товар","Раздел","Исходное поле","Значение","Причина"])
        for p in apple_products:
            ev=card_evidence.load(database,p["id"],"apple") or {};r=apple_pipeline.card_readiness(database,p["id"]);ident=r["identity"]
            ready_apple.append([p["name"],p["search_code"],"Готова" if r["verdict"]=="export_ready" else "Не готова",ident.get("model","unproven"),ident.get("configuration","unproven"),ident.get("variant","unproven"),r["manual_status"],"; ".join(r["blocking_gaps"])])
            for d in ev.get("manuals",[]):
                docs.append([p['name'],d.get('title',d['type']),d.get('language','Не проверена'),'','','Да' if d.get('verified') else 'Нет',ev.get('model_name',''),d['url'],d.get('chapter_url',d['url']),d.get('model_relation','unverified'),d['type'],'Проверена' if d.get('verified') else 'Не проверена'])
            for c in ev.get("configuration_candidates",[]):candidates_apple.append([p["name"],c["section"],c["raw_label"],c["value"],c["reason"]])
            for c in ev.get("photo_candidates",[]):
                if photo_candidates is None:
                    photo_candidates=book.create_sheet(_title("Фото-кандидаты",used))
                    _headers(photo_candidates,["SKU","Модель","URL","Причина candidate","Model relation","Color relation","Ширина, px","Высота, px","Размер файла, bytes","Формат"])
                if photo_candidates.cell(1,1).value == 'SKU':
                    photo_candidates.append([p["search_code"],ev.get("model_name",p["name"]),c["url"],c["reason"],c.get("model_relation","unproven"),c.get("color_relation","unproven"),c.get("width"),c.get("height"),c.get("file_size"),c.get("format")])
                else:
                    # Mixed-brand batches retain the existing candidate columns.
                    for col,label in enumerate(['SKU','Модель','Причина candidate','Model relation','Color relation'],11):
                        photo_candidates.cell(1,col,label).font=Font(bold=True)
                    photo_candidates.append([p['name'],'Apple','candidate',c['url'],'Нет',c['reason'],c.get('width'),c.get('height'),c.get('file_size'),c.get('format'),p['search_code'],ev.get('model_name',p['name']),c['reason'],c.get('model_relation','unproven'),c.get('color_relation','unproven')])

    for sheet in book.worksheets:
        for col in sheet.columns: sheet.column_dimensions[col[0].column_letter].width=min(60,max(12,max(len(str(c.value or "")) for c in col)+2))
    for sheet in xbox_sheets:
        from math import ceil
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for row in sheet.iter_rows():
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
            sheet.row_dimensions[row[0].row].height=min(180,max(24,16*max(ceil(len(str(c.value or ''))/max(8,(sheet.column_dimensions[c.column_letter].width or 12)-2)) for c in row)))
    if apple_products:
        from math import ceil
        for sheet in (ready_apple,candidates_apple,*([photo_candidates] if photo_candidates is not None else [])):
            for row in sheet.iter_rows():
                lines=1
                for cell in row:
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                    width=sheet.column_dimensions[cell.column_letter].width or 12
                    lines=max(lines,ceil(len(str(cell.value or ''))/max(8,width-2)))
                sheet.row_dimensions[row[0].row].height=max(22,16*lines)
    if lenovo_products:
        from math import ceil
        for sheet in (candidates, manual_candidates, ready):
            for row in sheet.iter_rows():
                lines=1
                for cell in row:
                    cell.alignment=Alignment(wrap_text=True,vertical="top")
                    width=sheet.column_dimensions[cell.column_letter].width or 12
                    lines=max(lines,ceil(len(str(cell.value or ""))/max(8,width-2)))
                sheet.row_dimensions[row[0].row].height=max(22,16*lines)
    if ps_products:
        from math import ceil
        for sheet in (ps_ready,ps_config,ps_candidates,ps_documents,ps_links):
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            for row in sheet.iter_rows():
                lines=1
                for cell in row:
                    cell.alignment=Alignment(wrap_text=True,vertical='top')
                    width=sheet.column_dimensions[cell.column_letter].width or 12
                    lines=max(lines,ceil(len(str(cell.value or ''))/max(8,width-2)))
                sheet.row_dimensions[row[0].row].height=min(180,max(24,16*lines))
    hx=[p for p in batch['products'] if p['brand'].strip().casefold()=='hyperx']
    if hx:
        from . import hyperx_audit,hyperx_presentation
        sheet=book.create_sheet(_title('Готовность HyperX',used));_headers(sheet,['Товар','Артикул','Готовность','Модель','Вариант','Инструкция','Характеристики','Пробелы'])
        variants=book.create_sheet(_title('Варианты HyperX',used));_headers(variants,['Товар','Тип идентификатора','Значение','Региональный суффикс','Опции','Связь с запросом','Источник'])
        rejected=book.create_sheet(_title('Кандидаты HyperX',used));_headers(rejected,['Товар','Раздел','Raw label','Значение','Причина','Источник'])
        manuals=book.create_sheet(_title('Документы HyperX',used));_headers(manuals,['Товар','Тип','Название','URL','Проверка'])
        photo_audit=book.create_sheet(_title('Фото HyperX',used));_headers(photo_audit,['Товар','Тип','URL','Выбрано','Связь с вариантом','Причина исключения'])
        for p in hx:
            ev=card_evidence.load(database,p['id'],'hyperx') or {};r=hyperx_audit.card_readiness(database,p['id'])
            sheet.append([p['name'],p['search_code'],'Готова' if r['verdict']=='export_ready' else 'Не готова',ev.get('model_name',''),ev.get('configuration_relation','unverified'),r['manual_status'],r['confirmed_specs'],'; '.join(hyperx_presentation.GAPS[g] for g in r['gaps'])])
            for rel in ev.get('identity_relations',[]):variants.append([p['name'],rel['type'],rel['value'],rel.get('regional_suffix',''),str(rel.get('options',{})),'Подтверждён' if rel.get('accepted') else 'Кандидат',rel.get('url','')])
            for f in ev.get('rejected_specs',[]):rejected.append([p['name'],f['section'],f['raw_label'],f['value'],f['reason'],f['url']])
            for d in ev.get('manuals',[]):manuals.append([p['name'],d['type'],d['title'],d['url'],'Проверена' if d['verified'] else 'Не проверена'])
            for photo in jobs.get_photo_candidates(database,p['id']):photo_audit.append([p['name'],photo['kind'],photo['url'],bool(photo['selected']),'Подтверждена' if photo['asset_key'] in ev.get('exact_photo_assets',[]) else 'Кандидат',photo['excluded_reason']])
    for family in ('xiaomi','dreame'):
        xiaomi_products=[p for p in batch['products'] if p['brand'].strip().casefold()=='dreame'] if family=='dreame' else [p for p in batch['products'] if p['brand'].strip().casefold() in XIAOMI_BRANDS]
        if xiaomi_products:
            from .xiaomi_presentation import label
            if family=='dreame':
                from .dreame_presentation import label
            photos.cell(1,5,'Область подтверждения')
            summary=book.create_sheet(_title(family.title()+' готовность',used));_headers(summary,['Товар','Поисковая модель','Модель','Рынок источника','Конфигурация','SKU','Инструкция','Характеристики','Готовность','Пробелы'])
            config=book.create_sheet(_title(family.title()+' конфигурация',used));_headers(config,['Товар','Поле','Запрошено','Статус'])
            raw=book.create_sheet(_title(family.title()+' кандидаты',used));_headers(raw,['Товар','Раздел','Исходное поле','Русское поле','Значение','Область','Причина'])
            gallery=book.create_sheet(_title(family.title()+' фото-кандидаты',used));_headers(gallery,['Товар','URL','Статус','Ширина','Высота','Размер','Формат'])
            model_codes=book.create_sheet(_title(family.title()+' номера моделей',used));_headers(model_codes,['Товар','Номер устройства','Тип','Источник','Связь с SKU'])
            components=None
            if family=='dreame':
                components=book.create_sheet(_title('Dreame компоненты',used));_headers(components,['Товар','Раздел источника','Компонент','Исходное поле','Исходное значение','Область','Решение','Источник'])
            for p in xiaomi_products:
                ev=card_evidence.load(database,p['id'],family) or {};r=xiaomi_pipeline.card_readiness(database,p['id'],family=family);i=r['identity'];summary.append([p['name'],p['search_code'],'Подтверждена' if r['exact_identity'] else 'Не подтверждена',i.get('source_region',''),'Данные входного файла; не блокируют готовность',i.get('retail_sku') or 'Не подтверждён',r['manual_status'],r['confirmed_specs'],'Готова' if r['verdict']=='export_ready' else 'Не готова','; '.join(xiaomi_pipeline.GAPS[g] for g in r['gaps'])])
                for k,v in r['input_configuration'].items():config.append([p['name'],{'RAM':'Оперативная память','storage':'Накопитель','color':'Цвет','region':'Региональная маркировка','supplier_part_number':'Артикул поставщика','bundle':'Комплект поставщика'}.get(k,k),v,'Данные входного файла; не официальный факт'])
                for f in ev.get('rejected_specs',[]):raw.append([p['name'],f.get('section',''),f['name'],label(f['name']),f['value'],f['scope'],f['reason']])
                if components is not None:
                    for f in ev.get('raw_specs',[]):components.append([p['name'],f.get('source_section',''),f.get('section',''),f['name'],f['value'],f.get('scope',''),'Кандидат' if f.get('reason') else 'Характеристика модели' if any(all(a.get(k)==f.get(k) for k in ('name','value','component')) for a in ev.get('accepted_specs',[])) else 'Исходные данные',ev.get('pdp_url','')])
                for photo in jobs.get_photo_candidates(database,p['id']):gallery.append([p['name'],photo['url'],'Фото модели; цвет поставщика не подтверждён' if xiaomi_pipeline.photo_verified(photo,ev) else 'Кандидат: product gallery не подтверждена',*_photo_cells(photo)])
                for c in ev.get('model_numbers',[]):model_codes.append([p['name'],c['value'],'Номер модели устройства',c['source_url'],'Не подтверждена'])
                if not jobs.get_documents(database,p['id']):docs.append([p['name'],'','','','','','','','','Официальный support inventory',family.title(),r['manual_status']])
            from math import ceil
            for sheet in (summary,config,raw,gallery,model_codes,photos,*([components] if components is not None else [])):
                for column in sheet.columns:
                    sheet.column_dimensions[column[0].column_letter].width=min(60,max(14,max(len(str(c.value or '')) for c in column)+2))
                for row in sheet.iter_rows():
                    lines=1
                    for cell in row:
                        cell.alignment=Alignment(wrap_text=True,vertical='top')
                        width=sheet.column_dimensions[cell.column_letter].width or 14
                        lines=max(lines,ceil(len(str(cell.value or ''))/max(8,width-2)))
                    sheet.row_dimensions[row[0].row].height=min(180,max(24,16*lines))
    electrolux_products=[p for p in batch['products'] if p['brand'].strip().casefold()=='electrolux']
    if electrolux_products:
        overview=book.create_sheet(_title('Electrolux готовность',used));_headers(overview,['Товар','Модель','Входной PNC','Product code / PNC','Revision','Регион','Связь продукта','Характеристики','Фото проверены','Инструкция','Готовность','Пробелы'])
        raw_sheet=book.create_sheet(_title('Electrolux исходные',used));_headers(raw_sheet,['Товар','Раздел','Исходное поле','Значение','Область','PNC','Регион','Результат','Причина'])
        docs_sheet=book.create_sheet(_title('Electrolux документы',used));_headers(docs_sheet,['Товар','Тип','Язык','Модель','PNC','Регион','URL','Связь проверена','Байты проверены','Стандарт источника','SHA256'])
        for p in electrolux_products:
            ev=card_evidence.load(database,p['id'],'electrolux') or {};r=electrolux_pipeline.card_readiness(database,p['id']);i=r['identity']
            overview.append([p['name'],i.get('model'),i.get('requested_pnc'),i.get('product_code'),i.get('revision'),i.get('region'),i.get('product_code_relation'),r['confirmed_specs'],len(ev.get('model_gallery_assets',[])),r['manual_status'],r['verdict'],'; '.join(electrolux_pipeline.GAPS[g] for g in r['gaps'])])
            for accepted,bucket in ((True,'accepted_facts'),(False,'rejected_facts')):
                for f in ev.get(bucket,[]):raw_sheet.append([p['name'],f['section'],f['label'],f['value'],f['scope'],f['product_code'],f['region'],'Принят' if accepted else 'Кандидат / отклонён',f.get('reason','')])
            for d in ev.get('document_inventory',[]):docs_sheet.append([p['name'],d['type'],d['language'],d['model'],d['product_code'],d['region'],d['url'],'Да' if d['identity_verified'] else 'Нет','Да' if d.get('bytes_verified') else 'Нет',d['standard_context'],d.get('sha256')])
        for sheet in (overview,raw_sheet,docs_sheet):
            for column in sheet.columns:sheet.column_dimensions[column[0].column_letter].width=min(60,max(14,max(len(str(c.value or '')) for c in column)+2))
            for row in sheet:
                for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    # Raw classification is technical evidence, never a source of product facts.
    # Keeping it separate prevents excluded help/marketing rows re-entering
    # product columns while preserving every audited input and its decision.
    audited=[(p,card_evidence.load(database,p['id'],'canonical_mapping_audit') or {}) for p in batch['products']]
    if any(a.get('rows') for p,a in audited):
        raw_mapping=book.create_sheet(_title('Исходные поля и mapping',used))
        _headers(raw_mapping,['Товар','Бренд','Категория','Source','Раздел','Подраздел','Исходное поле','Raw value','Регион','Компонент','Identity relation','Класс','Action','Причина','Canonical targets'])
        for p,audit in audited:
            for r in audit.get('rows',[]):
                value=str(r['raw_value'])
                if value.startswith('='):value="'"+value
                raw_mapping.append([p['name'],p['brand'],p['category'],r.get('resolved_source',r['source']),r['section'],r['subsection'],r['raw_label'],value,r['region'],r['component_scope'],json.dumps(r['identity_relation'],ensure_ascii=False),r['class'],r['action'],r['decision_reason'],'; '.join(r['canonical_targets'])])
    if any(p['brand'].strip().casefold() in {'jbl','samsung','lg','apple','playstation','sony','xbox','razer','hyperx','xiaomi','redmi','poco','dreame','electrolux'} for p in batch['products']):
        from math import ceil
        for sheet in book:
            sheet.sheet_view.showGridLines=False
            sheet.auto_filter.ref=sheet.dimensions
            for column in sheet.columns:
                label=str(column[0].value or '')
                width=min(54,max(14,max(min(52,len(str(c.value or ''))) for c in column)+2))
                if label in {'URL','URL фото','Значение','Причина','Доказательство'}:width=72
                sheet.column_dimensions[column[0].column_letter].width=width
            for row in sheet:
                lines=1
                for cell in row:
                    cell.font=Font(name='Arial',size=10,bold=cell.row==1,color='FFFFFF' if cell.row==1 else '172B4D')
                    cell.alignment=Alignment(wrap_text=True,vertical='top',horizontal='center' if cell.row==1 else 'right' if isinstance(cell.value,(int,float)) else 'left')
                    if cell.row==1:cell.fill=PatternFill('solid',fgColor='243B53')
                    width=sheet.column_dimensions[cell.column_letter].width
                    lines=max(lines,sum(max(1,ceil(len(part)/max(8,width-2))) for part in str(cell.value or '').split('\n')))
                sheet.row_dimensions[row[0].row].height=min(409,max(28,15*lines+8))
            if sheet.cell(1,1).value=='Строка':sheet.freeze_panes='F2'
    output=BytesIO(); book.save(output); book.close(); return output.getvalue()
