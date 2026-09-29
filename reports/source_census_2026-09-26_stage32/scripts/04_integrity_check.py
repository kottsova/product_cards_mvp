"""Stage 32 read-only integrity, selection, no-network gate and export checks."""
from __future__ import annotations
import gzip,hashlib,io,json,sys
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[2];STAGE=HERE.parent
sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
from openpyxl import load_workbook
import _pipeline_migration as mig
import _samsung_replay as R
from product_tool.coverage import classify as C,executor,planner
from product_tool.offline_guard import offline_only
from product_tool.adapters.policy_fetch import stopped_hosts_from_fetch_log
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def main():
 d=json.loads((STAGE/'raw/route_declaration.json').read_text(encoding='utf-8'))
 started=datetime.fromisoformat(d['declared_at']).timestamp()
 pins={p:mig.check_migrated_file(p,sha(ROOT/p))[0] for p in mig.ALL_AUTHORIZED_CHANGES}
 prior=[p for p in (ROOT/'reports').glob('source_census_*') if p!=STAGE]
 modified=sorted(p.relative_to(ROOT).as_posix() for directory in prior for p in directory.rglob('*') if p.is_file() and '__pycache__' not in p.parts and p.stat().st_mtime>started)
 records=[json.loads(x) for x in (STAGE/'route_check/responses/index.jsonl').read_text(encoding='utf-8').splitlines()]
 response_hashes={}
 for e in records:
  if e['saved_as']:
   with gzip.open(STAGE/'route_check/responses'/e['saved_as'],'rt',encoding='utf-8',newline='') as h:response_hashes[e['url']]=hashlib.sha256(h.read().encode('utf-8')).hexdigest()==e['sha256']
 route=json.loads((STAGE/'raw/route_check_result.json').read_text(encoding='utf-8'))
 replay=json.loads((STAGE/'raw/offline_replay_calls.json').read_text(encoding='utf-8'))
 cards=json.loads((STAGE/'raw/offline_two_result.json').read_text(encoding='utf-8'))['cards']
 with offline_only():
  units=[u for u in planner.build_plan()['units'] if u['family']=='samsung']
  selected=[u for u in units if u['status']==C.READY]
  targets=[u for u in units if u['seller_sku'] in d['products'].values()]
  with TemporaryDirectory() as tmp:
   session=R.Replay()
   checkpoint=executor.run_units(targets,workdir=Path(tmp),checkpoint_path=Path(tmp)/'checkpoint.json',factories=executor.RunFactories(session=session))
   gate={'target_outcomes':[v['outcome'] for v in checkpoint['units'].values()],'stopped_before_run_once':all(v['stopped_before_run_once'] for v in checkpoint['units'].values()),'request_calls':session.calls}
 book=load_workbook(io.BytesIO((STAGE/'export/samsung_two_export.xlsx').read_bytes()),read_only=True)
 log=STAGE/'route_check/workdir/samsung_fetch_log.json'
 entries=json.loads(log.read_text(encoding='utf-8')) if log.exists() else []
 result={'catalog_unchanged':sha(ROOT/'data/catalog_2026-09-21_filtered.xlsx')==d['catalog_sha256_before'],'source_registry_unchanged':sha(ROOT/'product_tool/config/source_catalog.v2.json')==d['registry_sha256_before'],'local_database_unchanged':sha(ROOT/'data/batches.sqlite3')==d['db_sha256_before'],'protected_pins_ok':all(pins.values()),'failed_pins':[p for p,ok in pins.items() if not ok],'prior_report_files_modified_since_declaration':modified,'recorded_response_hashes_ok':len(response_hashes)==2 and all(response_hashes.values()),'requests_spent':route['budget']['spent'],'stopped_hosts':sorted(stopped_hosts_from_fetch_log(x for x in entries if isinstance(x,dict))),'selected_samsung_products':len(selected),'selected_samsung_categories':len({u['category'] for u in selected}),'total_samsung_rows':len(units),'target_statuses':{u['seller_sku']:[u['status'],u['reason']] for u in targets},'target_gate':gate,'offline_replay_refused':replay['refused'],'offline_replay_dealer_calls':replay['dealer_calls'],'job_statuses':{c['article']:c['job_status'] for c in cards},'readiness':{c['article']:c['readiness']['verdict'] for c in cards},'export_sheets':book.sheetnames}
 (STAGE/'protected_hashes_check.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
 print(json.dumps(result,ensure_ascii=True,indent=2))
 assert result['catalog_unchanged'] and result['source_registry_unchanged'] and result['local_database_unchanged'] and result['protected_pins_ok'] and not modified
 assert result['recorded_response_hashes_ok'] and result['requests_spent']==2 and not result['stopped_hosts']
 assert (result['selected_samsung_products'],result['selected_samsung_categories'],result['total_samsung_rows'])==(20,20,1182)
 assert len(targets)==2 and all(u['status']==C.MANUAL and u['reason']=='not_selected_one_card_per_category' for u in targets)
 assert gate['stopped_before_run_once'] and not gate['request_calls'] and not replay['refused'] and not replay['dealer_calls']
 assert all(c['job_status']=='needs_review' and c['readiness']['verdict']=='not_ready' for c in cards)
 assert all(c['category'] in book.sheetnames for c in cards)
if __name__=='__main__':main()
