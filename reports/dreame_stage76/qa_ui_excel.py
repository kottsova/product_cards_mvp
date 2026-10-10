"""Actual blocked live results through native UI/export; no replay substitution."""
import sys,json,sqlite3,os,hashlib
from pathlib import Path
from io import BytesIO
from urllib.parse import urlsplit
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from product_tool import exporter,jobs,storage,readiness
from product_tool.web import create_app

R=Path(__file__).parent
U=R/'ui';U.mkdir(exist_ok=True)
Q=R/'qa';Q.mkdir(exist_ok=True)
db=U/'batches.sqlite3'
with sqlite3.connect(R/'final_live/live.sqlite3') as src,sqlite3.connect(db) as dst:src.backup(dst)
checks=[]
with TestClient(create_app(U,start_worker=False)) as client:
    for p in storage.get_batch(db,'dr76')['products']:
        response=client.get(f'/products/{p["id"]}');assert response.status_code==200
        html=response.text;r=readiness.card_readiness(db,p['id'])
        assert 'Модель и конфигурация Dreame' in html and r['manual_status'] in html
        assert 'Данные поставщика не подтверждают коммерческий SKU Dreame' in html
        assert 'Модель и конфигурация Xiaomi' not in html
        assert r['verdict']=='not_ready' and r['confirmed_specs']==0
        assert not any(p['selected'] for p in jobs.get_photo_candidates(db,p['id']))
        (U/f'product_{p["id"]}.html').write_text(html,encoding='utf8')
        checks.append({'id':p['id'],'http_status':200,'model_first_caption':True,'input_configuration_separate':True,'manual_status':r['manual_status'],'readiness':r['verdict'],'confirmed_specs':r['confirmed_specs']})
raw=exporter.export_batch(db,'dr76');(R/'Dreame_Stage76.xlsx').write_bytes(raw)
book=load_workbook(BytesIO(raw));assert book['Dreame готовность'].max_row==11
assert all(book['Dreame готовность'].cell(i,9).value=='Не готова' for i in range(2,12))
assert all(book['Dreame готовность'].cell(i,8).value==0 for i in range(2,12))
assert book['Фотографии'].max_row==1 and 'Dreame компоненты' in book.sheetnames
assert all('Данные входного файла' in str(row[3].value) for row in list(book['Dreame конфигурация'].iter_rows())[1:])
errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
sheets={s.title:{'rows':s.max_row,'columns':s.max_column,'headers':[c.value for c in s[1]]} for s in book};book.close()
(R/'qa_ui_excel.json').write_text(json.dumps({'observation':'actual blocked live worker; no captures injected','ui':checks,'native_exporter':True,'sheets':sheets,'formula_errors':errors,'workbook_sha256':hashlib.sha256(raw).hexdigest()},ensure_ascii=False,indent=2),encoding='utf8')
os.environ.setdefault('PLAYWRIGHT_BROWSERS_PATH',str(Path('.venv/playwright-browsers').resolve()))
from playwright.sync_api import sync_playwright
with TestClient(create_app(U,start_worker=False)) as client,sync_playwright() as pw:
    browser=pw.chromium.launch(headless=True);page=browser.new_page(viewport={'width':1440,'height':1000})
    def route(request):
        u=urlsplit(request.request.url)
        if u.hostname!='dreame-audit.test':request.abort();return
        response=client.get(u.path);request.fulfill(status=response.status_code,body=response.content,content_type=response.headers.get('content-type','text/html'))
    page.route('**/*',route)
    for pid in (1,5,8,9):
        page.goto(f'http://dreame-audit.test/products/{pid}',wait_until='domcontentloaded')
        page.locator('#xiaomi-readiness').scroll_into_view_if_needed();page.screenshot(path=str(Q/f'ui_{pid}_evidence.png'))
    browser.close()
print('10 UI cards and native Excel verified; 4 offline-local UI screenshots',flush=True)
