from pathlib import Path
from shutil import copyfile
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from product_tool import storage
from product_tool.web import create_app
root=Path('reports/bosch_stage58')
ui=root/'ui_batch'; ui.mkdir(exist_ok=True)
copyfile(root/'batch_pass.sqlite3',ui/'batches.sqlite3')
with storage._connection(ui/'batches.sqlite3') as db:
    products=[dict(r) for r in db.execute('select id,search_code from products order by id')]
with TestClient(create_app(ui,start_worker=False)) as client:
    for item in products:
        response=client.get(f'/products/{item["id"]}')
        html=response.text
        print('page',item['search_code'],response.status_code,'lightbox',('id="photo-lightbox"' in html),'main',('Основные характеристики' in html),'model',('Особенности модели' in html),'generic',('Дополнительная характеристика' in html),'supportcopy',('How can we help you' in html),'manual',('Русская инструкция:' in html))
        assert response.status_code==200
        assert 'Дополнительная характеристика' not in html
    response=client.get('/batches/batch_pass/export.xlsx')
    print('export',response.status_code,len(response.content))
    assert response.status_code==200
    book=load_workbook(BytesIO(response.content),read_only=True,data_only=True)
    for sheet in book:
        values=list(sheet.values)
        print('sheet',sheet.title,'rows',len(values),'columns',len(values[0]) if values else 0)
        if sheet.title.startswith(('Проверка источников','Инструкции','Фотографии')):
            print('head',sheet.title,values[0] if values else '')
    audit=next(sheet for sheet in book if sheet.title.startswith("Проверка источников"))
    assert "LG Казахстан" not in next(audit.values) and "LG Россия" not in next(audit.values)
    candidates=next(sheet for sheet in book if sheet.title.startswith("Фото-кандидаты"))
    assert candidates.max_row==45, candidates.max_row
    confirmed=next(sheet for sheet in book if sheet.title=="Фотографии")
    assert confirmed.max_row==48, confirmed.max_row
    for sheet in book:
        for row in sheet.values:
            for value in row:
                if isinstance(value,str):
                    assert 'Дополнительная характеристика' not in value
    book.close()
