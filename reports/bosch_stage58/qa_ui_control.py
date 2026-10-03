from pathlib import Path
from shutil import copyfile
from io import BytesIO
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from product_tool import jobs,storage
from product_tool.web import create_app
root=Path('reports/bosch_stage58')
ui=root/'ui_control';ui.mkdir(exist_ok=True)
copyfile(root/'control_pass.sqlite3',ui/'batches.sqlite3')
with storage._connection(ui/'batches.sqlite3') as db:
 pid=db.execute('select id from products').fetchone()[0]
photos=jobs.get_photo_candidates(ui/'batches.sqlite3',pid)
with TestClient(create_app(ui,start_worker=False)) as client:
 response=client.get(f'/products/{pid}');html=response.text
 print('page',response.status_code,'manual',('Русская инструкция: <strong>Проверена</strong>' in html),
       'lightbox',('id="photo-lightbox"' in html),'excluded',('Исключённые изображения' in html),
       'photo_count',len(photos),'selected',sum(p['selected'] for p in photos))
 first=next(p for p in photos if p['selected'])
 inspected=client.post(f'/products/{pid}/photos/{first["id"]}/inspect',follow_redirects=False)
 checked=jobs.get_photo_candidate(ui/'batches.sqlite3',pid,first['id'])
 print('inspect',inspected.status_code,checked['verified_width'],checked['verified_height'],checked['verified_bytes'],checked['verified_format'])
 exported=client.get('/batches/control_pass/export.xlsx')
 print('xlsx',exported.status_code,len(exported.content))
 book=load_workbook(BytesIO(exported.content),read_only=True,data_only=True)
 print('sheets',book.sheetnames)
 check=next(sheet for sheet in book if sheet.title.startswith('Проверка источников'))
 print('dealer_headers',[x for x in next(check.values) if x in ('Sulpak','DNS')])
 docs=next(sheet for sheet in book if sheet.title.startswith('Инструкции'))
 print('docs_rows',list(docs.values)[1:3])
 book.close()
