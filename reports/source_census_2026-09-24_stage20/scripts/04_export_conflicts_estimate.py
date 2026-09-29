"""Stage 20 step 4 -- OFFLINE, zero requests. Three things that need the pilot's database / saved responses:

  A. the ordinary Excel export of the pilot batch, counted per row (what actually reaches the sheets);
  B. the conflicts the resolver reported -- are they between sources, or inside one source? (a finding made AFTER the run:
     it is not in the taxonomy declared before, and is marked as such);
  C. how many of ALL 578 ready rows the current KZ lookup can reach, read from the KZ sitemap saved during the pilot (no request),
     under the current rule and under one corrected rule for the base-model suffix.

Output: raw/export_check.json, raw/conflicts_within_source.json, raw/estimate_578.json
"""
from __future__ import annotations

import gzip
import json
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
STAGE = HERE.parent
PILOT = STAGE / "pilot"
sys.path.insert(0, str(ROOT))

from product_tool.adapters.lg import _slug_key, lg_article_key, lg_base_model, lg_is_product_url, normalize_lg_sku  # noqa: E402
from product_tool.exporter import export_batch  # noqa: E402
from product_tool.offline_guard import offline_only  # noqa: E402


def export_check() -> dict:
    database = PILOT / "workdir/batches.sqlite3"
    book = load_workbook(BytesIO(export_batch(database, "lg-pilot")))
    sheets = {ws.title: ws for ws in book.worksheets}
    per_row = {}
    for title, ws in sheets.items():
        rows = list(ws.iter_rows(values_only=True))
        if not rows or rows[0][0] != "Строка" or title in ("Проверка источников", "Источники", "Инструкции", "Фотографии"):
            continue
        header = rows[0]
        for values in rows[1:]:
            article = values[3]
            filled = sum(1 for v in values[5:] if v not in (None, ""))
            per_row[article] = {"category_sheet": title, "attribute_columns": len(header) - 5, "attribute_cells_filled": filled}
    photos, docs = Counter(), Counter()
    product_names = {}
    connection = sqlite3.connect(f"file:{database}?mode=ro", uri=True)
    for pid, name, code in connection.execute("SELECT id, name, search_code FROM products"):
        product_names[name] = code
    connection.close()
    for values in list(sheets["Фотографии"].iter_rows(values_only=True))[1:]:
        photos[product_names.get(values[0], values[0])] += 1
    for values in list(sheets["Инструкции"].iter_rows(values_only=True))[1:]:
        docs[product_names.get(values[0], values[0])] += 1
    for article, info in per_row.items():
        info["photo_rows_in_export"] = photos.get(article, 0)
        info["instruction_rows_in_export"] = docs.get(article, 0)
    return {"sheets": sorted(sheets), "rows": per_row, "rows_with_attribute_cells": sum(1 for v in per_row.values() if v["attribute_cells_filled"]),
            "rows_with_photo_rows": sum(1 for v in per_row.values() if v["photo_rows_in_export"]), "rows_with_instruction_rows": sum(1 for v in per_row.values() if v["instruction_rows_in_export"]),
            "products_in_export": len(per_row)}


def conflicts_within_source() -> dict:
    connection = sqlite3.connect(f"file:{PILOT / 'workdir/batches.sqlite3'}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    out, within, between = [], 0, 0
    for conflict in connection.execute("SELECT p.search_code, r.normalized_name, r.status, r.product_id FROM resolved_attributes r JOIN products p ON p.id=r.product_id WHERE r.conflict=1 ORDER BY p.search_code, r.normalized_name"):
        facts = connection.execute("SELECT source_key, raw_name, raw_value FROM extracted_attribute_facts WHERE product_id=? AND normalized_name=?", (conflict["product_id"], conflict["normalized_name"])).fetchall()
        sources = {f["source_key"] for f in facts}
        raw_names = {f["raw_name"] for f in facts}
        single = len(sources) == 1
        within += single
        between += not single
        out.append({"seller_sku": conflict["search_code"], "normalized_name": conflict["normalized_name"], "status": conflict["status"], "sources": sorted(sources), "single_source": single,
                    "distinct_raw_names": len(raw_names), "facts": [[f["raw_name"], f["raw_value"][:50]] for f in facts]})
    connection.close()
    return {"note": "Finding made after the run; not in the taxonomy declared before it.", "conflicts": len(out), "within_one_source": within, "between_sources": between,
            "rows_with_conflicts": len({c["seller_sku"] for c in out}), "reading": "The resolver reports 'official_regions_conflict' when one normalized name has different values; NAME_RULES in normalization.py fold different measurements (weight with/without stand, indoor/outdoor unit, gross/net) into one name.",
            "items": out}


def estimate_578() -> dict:
    responses = PILOT / "responses"
    index = [json.loads(line) for line in (responses / "index.jsonl").read_text(encoding="utf-8").splitlines()]
    sitemap = next(e for e in index if e["url"].endswith("/kz/sitemap.xml"))
    with gzip.open(responses / sitemap["saved_as"], "rt", encoding="utf-8", newline="") as handle:
        text = handle.read()
    urls = [u for u in re.findall(r"<loc>\s*([^<\s]+)\s*</loc>", text) if lg_is_product_url(u)]
    slugs = defaultdict(list)
    for url in urls:
        slugs[_slug_key(url)].append(url)
    queue = [json.loads(line) for line in (ROOT / "reports/source_census_2026-09-24_stage19/queue/coverage_units.jsonl").read_text(encoding="utf-8").splitlines() if line]
    ready = [u for u in queue if u["family"] == "lg" and u["status"] == "ready_to_run"]

    def corrected_base(full: str) -> str:
        normalized = normalize_lg_sku(full)
        base, dot, suffix = normalized.rpartition(".")
        return base if dot and base and re.fullmatch(r"[A-Z]{3,}", suffix) else normalized

    rows, cats = [], defaultdict(Counter)
    for unit in ready:
        full = normalize_lg_sku(unit["seller_sku"])
        current_keys = {lg_article_key(full), lg_article_key(lg_base_model(full))}
        fixed_keys = current_keys | {lg_article_key(corrected_base(full))}
        current = any(slugs.get(k) for k in current_keys)
        fixed = any(slugs.get(k) for k in fixed_keys)
        rows.append((unit, current, fixed))
        cats[unit["category"]]["total"] += 1
        cats[unit["category"]]["found_now"] += current
        cats[unit["category"]]["found_with_suffix_fix"] += fixed
    dotted = [u for u, _, _ in rows if "." in u["seller_sku"]]
    suffix_len = Counter(len(u["seller_sku"].rsplit(".", 1)[1]) for u in dotted)
    return {
        "source": "the LG KZ sitemap saved during the pilot (no request)", "sitemap_product_urls": len(urls), "ready_rows": len(ready),
        "reachable_by_the_current_rule": sum(1 for _, c, _ in rows if c), "reachable_with_a_3plus_letter_suffix_rule": sum(1 for _, _, f in rows if f),
        "gain_of_the_suffix_correction": sum(1 for _, c, f in rows if f and not c), "not_in_the_kz_sitemap_at_all": sum(1 for _, _, f in rows if not f),
        "rows_with_a_dot_in_the_article": len(dotted), "suffix_letter_counts_among_dotted": dict(sorted(suffix_len.items())),
        "by_category": {c: dict(v) for c, v in sorted(cats.items())},
        "note": "'reachable' = a sitemap product URL whose slug equals the article or its base model; whether the page then shows the article (full_sku) is decided on the page.",
        "pilot_check": {"pilot_rows_reachable_now": sum(1 for u, c, _ in rows if c and u["seller_sku"] in {r["seller_sku"] for r in json.loads((STAGE / "raw/pilot_selection.json").read_text(encoding="utf-8"))["rows"]})},
        "pilot_rows_not_reached": {u["seller_sku"]: ("fixed_by_the_suffix_rule" if f else "absent_from_the_kz_sitemap")
                                   for u, c, f in rows if not c and u["seller_sku"] in {r["seller_sku"] for r in json.loads((STAGE / "raw/pilot_selection.json").read_text(encoding="utf-8"))["rows"]}},
        "unreachable_examples": [u["seller_sku"] for u, _, f in rows if not f][:15],
    }


def main() -> None:
    (STAGE / "raw/export_check.json").write_text(json.dumps(export_check(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (STAGE / "raw/conflicts_within_source.json").write_text(json.dumps(conflicts_within_source(), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    estimate = estimate_578()
    (STAGE / "raw/estimate_578.json").write_text(json.dumps(estimate, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    export = json.loads((STAGE / "raw/export_check.json").read_text(encoding="utf-8"))
    conflicts = json.loads((STAGE / "raw/conflicts_within_source.json").read_text(encoding="utf-8"))
    print("export:", {k: v for k, v in export.items() if k != "rows"})
    print("conflicts:", {k: v for k, v in conflicts.items() if k not in ("items", "note", "reading")})
    print("estimate:", {k: v for k, v in estimate.items() if k not in ("by_category", "note", "unreachable_examples")})
    for c, v in estimate["by_category"].items():
        print("   ", c, v)


if __name__ == "__main__":
    with offline_only():
        main()
