from __future__ import annotations
import hashlib,json,sqlite3
from contextlib import closing
from pathlib import Path
from io import BytesIO
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from product_tool import exporter
from product_tool.web import create_app
ROOT=Path('reports/lg_live_batch_2026-09-29_stage49')
before=ROOT/'batches_before_stage49.sqlite3'; after=Path('data/batches.sqlite3')
BATCH='ae3d2cb381744ba8a811e54231233254'
with closing(sqlite3.connect(before)) as a,closing(sqlite3.connect(after)) as b:
 a.row_factory=b.row_factory=sqlite3.Row
 tables=['batches','products','source_snapshots','fetch_attempts','product_documents','photo_candidates','extracted_attribute_facts','manual_attribute_decisions']
 unchanged={}
 for t in tables:
  old=[tuple(x) for x in a.execute(f'SELECT * FROM {t} ORDER BY rowid')]
  new=[tuple(x) for x in b.execute(f'SELECT * FROM {t} ORDER BY rowid')]
  unchanged[t]=len(old)
  assert old==new,(t,len(old),len(new))
 for t,pid in [('source_pages','product_id'),('resolved_attributes','product_id')]:
  old=[tuple(x) for x in a.execute(f'SELECT * FROM {t} WHERE {pid} NOT IN (10,15) ORDER BY rowid')]
  new=[tuple(x) for x in b.execute(f'SELECT * FROM {t} WHERE {pid} NOT IN (10,15) ORDER BY rowid')]
  assert old==new,t
 assert [tuple(x) for x in a.execute('SELECT * FROM search_jobs ORDER BY rowid')]==[tuple(x) for x in b.execute('SELECT * FROM search_jobs WHERE id NOT LIKE "stage49-offline-%" ORDER BY rowid')]
 assert [tuple(x) for x in a.execute('SELECT * FROM job_events ORDER BY rowid')]==[tuple(x) for x in b.execute('SELECT * FROM job_events WHERE job_id NOT LIKE "stage49-offline-%" ORDER BY rowid')]
 assert a.execute('SELECT count(*) FROM source_pages').fetchone()[0]+2==b.execute('SELECT count(*) FROM source_pages').fetchone()[0]
 assert a.execute('SELECT count(*) FROM search_jobs').fetchone()[0]+3==b.execute('SELECT count(*) FROM search_jobs').fetchone()[0]
 integrity=b.execute('PRAGMA integrity_check').fetchone()[0]
 fk=b.execute('PRAGMA foreign_key_check').fetchall()
 assert integrity=='ok' and not fk
 info={'before_sha256':hashlib.sha256(before.read_bytes()).hexdigest(),
       'after_sha256':hashlib.sha256(after.read_bytes()).hexdigest(),
       'unchanged_tables':unchanged,'outside_affected_source_and_resolution_unchanged':True,
       'old_jobs_and_events_unchanged':True,'new_offline_jobs':3,'new_support_pages':2,
       'integrity_check':integrity,'foreign_key_violations':len(fk),
       'catalog_sha256':hashlib.sha256(Path('data/catalog.xlsx').read_bytes()).hexdigest() if Path('data/catalog.xlsx').exists() else None}
ROOT.joinpath('integrity.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf8')
print('INTEGRITY',info)
blob=exporter.export_batch(after,BATCH)
xlsx=ROOT/'stage49.xlsx';xlsx.write_bytes(blob)
book=load_workbook(BytesIO(blob),read_only=True)
try:
 ready=next(s for s in book if s.title=='Готовность LG')
 rows=list(ready.iter_rows(values_only=True));print('READY_ROWS',len(rows)-1)
 assert len(rows)-1==12
 for row in rows[1:]:
  if row[1] in {'P12ED.NSAR + P12ED.USAR','ON77DKDRUSLLK','RNC9.DRUSLLK'}:
   print('READY',row[1],row[2],row[3],row[5])
 sources=next(s for s in book if s.title=='Источники')
 support_rows={(r[1],r[2]):r[4] for r in sources.iter_rows(min_row=2,values_only=True) if r[2] in {'LG RU support','LG Казахстан — поддержка'}}
 assert support_rows[('W4W8LVPKZHM.APBPCOM','LG RU support')]=='Связь инструкции с артикулом'
 assert support_rows[('ON77DKDRUSLLK','LG RU support')]=='Поддержка семейства/другого варианта'
 assert support_rows[('P12ED.NSAR + P12ED.USAR','LG Казахстан — поддержка')]=='Связь инструкции с артикулом'
 print('SUPPORT_LABELS',support_rows)
 docs=next(s for s in book if s.title=='Инструкции')
 for row in docs.iter_rows(min_row=2,values_only=True):
  if row[0] in ('ON77DKDRUSLLK','P12ED.NSAR + P12ED.USAR','RNC9.DRUSLLK'):
   print('DOC',row[0],row[-1])
 photos=next(s for s in book if s.title=='Фотографии')
 labels={}
 for row in photos.iter_rows(min_row=2,values_only=True):
  labels[(row[0],row[-1])]=labels.get((row[0],row[-1]),0)+1
 print('PHOTO_LABELS',labels)
finally:book.close()
with TestClient(create_app(Path('data'),start_worker=False)) as client:
 for pid in (4,15):
  response=client.get(f'/products/{pid}')
  assert response.status_code==200
  print('CARD',pid,response.status_code,'variant warning',('связь с вариантом не подтверждена' in response.text))
