"""Validate the native UI/export using the frozen live audit database."""
from pathlib import Path
import json,shutil,os
from io import BytesIO
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from product_tool import exporter,storage,jobs,playstation_pipeline,attribute_projection
from product_tool.web import create_app
r=Path('reports/playstation_stage67');u=r/'ui_batch';u.mkdir(exist_ok=True);shutil.copyfile(r/'replay_acceptance.sqlite3',u/'batches.sqlite3');p=u/'batches.sqlite3';out=[]
measurements=json.loads((r/'photo_inspection.json').read_text(encoding='utf-8')) if (r/'photo_inspection.json').exists() else []
for product in storage.get_batch(p,'ps67')['products']:
    for photo in jobs.get_photo_candidates(p,product['id']):
        m=next((m for m in measurements if m['url']==photo['url']),None)
        if m:jobs.save_photo_metadata(p,product['id'],photo['id'],photo['url'],width=m['width'],height=m['height'],size_bytes=m['size_bytes'],image_format=m['format'])
with TestClient(create_app(u,start_worker=False)) as client:
    for product in storage.get_batch(p,'ps67')['products']:
        pid=product['id'];response=client.get(f'/products/{pid}');assert response.status_code==200
        html=response.text;assert 'Модель и конфигурация PlayStation' in html;assert 'Фото-кандидаты' in html;assert 'Документы и руководства' in html
        (r/f'ui_{pid}.html').write_text(html,encoding='utf-8')
        out.append(dict(article=product['search_code'],http=response.status_code,readiness=playstation_pipeline.card_readiness(p,pid),labels=[x['display_name'] for x in attribute_projection.final_attribute_rows(p,pid)],descriptions_clean=all('How to' not in (s['description'] or '') for s in jobs.get_source_pages(p,pid))))
    data=exporter.export_batch(p,'ps67');(r/'export.xlsx').write_bytes(data);book=load_workbook(BytesIO(data))
    sheets={s.title:dict(rows=s.max_row,columns=s.max_column,headers=[c.value for c in s[1]]) for s in book}
    for name in ('Готовность PlayStation','Конфигурация PlayStation','Кандидаты PlayStation','Документы PlayStation','Связь SKU и CFI','Фото-кандидаты'):assert name in sheets
    assert 'Применена к карточке' in sheets['Связь SKU и CFI']['headers']
    errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
(r/'qa_ui_excel.json').write_text(json.dumps(dict(ui=out,sheets=sheets,formula_errors=errors,native_exporter=True),ensure_ascii=False,indent=2),encoding='utf-8')
print('10 UI HTTP200; native Excel export verified; no formula errors',flush=True)
os.environ.setdefault('PLAYWRIGHT_BROWSERS_PATH',str(Path('.venv/playwright-browsers').resolve()))
from playwright.sync_api import sync_playwright
from urllib.parse import urlsplit
with TestClient(create_app(u,start_worker=False)) as client,sync_playwright() as pw:
    browser=pw.chromium.launch(headless=True);page=browser.new_page(viewport=dict(width=1440,height=1000))
    def route(request):
        photo=next((m for m in measurements if m['url']==request.request.url),None)
        if photo:
            request.fulfill(status=200,body=(r/photo['file']).read_bytes(),content_type='image/jpeg' if photo['format']=='JPG' else 'image/'+photo['format'].lower());return
        url=urlsplit(request.request.url)
        if url.hostname!='playstation-audit.test':request.abort();return
        response=client.get(url.path);request.fulfill(status=response.status_code,body=response.content,content_type=response.headers.get('content-type','text/html'))
    page.route('**/*',route)
    for pid in (3,6,7,8,10):
        page.goto(f'http://playstation-audit.test/products/{pid}',wait_until='networkidle');page.screenshot(path=str(r/f'ui_{pid}_overview.png'))
        page.get_by_role('heading',name='Модель и конфигурация PlayStation',exact=True).scroll_into_view_if_needed();page.screenshot(path=str(r/f'ui_{pid}_evidence.png'))
        if pid==6:
            for img in page.locator('.photo-open img').all():
                img.scroll_into_view_if_needed();page.wait_for_function('(img)=>img.complete && img.naturalWidth>0',arg=img.element_handle())
            page.locator('.photo-open').first.click();assert page.locator('#photo-lightbox').is_visible();page.screenshot(path=str(r/'ui_6_lightbox.png'));page.get_by_role('button',name='Закрыть',exact=True).click()
    browser.close()
print('Native UI screenshots verified',flush=True)
