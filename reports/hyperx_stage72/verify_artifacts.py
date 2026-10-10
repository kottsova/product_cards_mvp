"""Read-back of actual stored cards, rendered UI, Excel and guarded lifecycle."""
import sys,json,sqlite3,tempfile,hashlib,time
from pathlib import Path
from io import BytesIO
from bs4 import BeautifulSoup
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from product_tool import jobs,exporter,discovery_trace
from product_tool.web import create_app
from fastapi.testclient import TestClient
from product_tool.adapters.hyperx import HyperXAdapter
from product_tool.adapters.access_stop import active_stops
from product_tool.adapters.policy_fetch import read_log
R=Path(__file__).parent;db=R/'live_release.sqlite3'
out={'provenance':'read-back of stored actual live release results; no new live success','ui':[],'traces':{}}
with tempfile.TemporaryDirectory(prefix='hyperx72_ui_') as temp:
    target=Path(temp)/'batches.sqlite3'
    source=sqlite3.connect(db);dest=sqlite3.connect(target)
    try:source.backup(dest)
    finally:source.close();dest.close()
    client=TestClient(create_app(Path(temp),start_worker=False))
    for pid in (1,4,6,10):
        response=client.get('/products/'+str(pid));assert response.status_code==200
        html=response.text;soup=BeautifulSoup(html,'html.parser');text=soup.get_text(' ',strip=True)
        (R/('ui_'+str(pid)+'.html')).write_text(html,encoding='utf8')
        out['ui'].append({'id':pid,'status':response.status_code,'hyperx_identity_panel':'Модель и конфигурация HyperX' in text,'manual_unchecked':'Не проверена' in text,'legacy_lg_counter_present':'Только у базовой модели LG' in text,'headings':[h.get_text(' ',strip=True) for h in soup.select('h2,h3')]})
    client.close()
data=exporter.export_batch(db,'hx72');(R/'HyperX_Stage72.xlsx').write_bytes(data)
book=load_workbook(BytesIO(data));out['excel']={'sheet_names':book.sheetnames,'headers':{s.title:[c.value for c in s[1]] for s in book.worksheets},'rows':{s.title:s.max_row for s in book.worksheets}};book.close()
for pid in range(1,11):
    jid=jobs.list_jobs(db,pid)[0]['id'];out['traces'][str(pid)]=discovery_trace.for_job(db,jid)
log=R/'hyperx_fetch_log.json';before=log.read_bytes();adapter=HyperXAdapter(discovery_enabled=True,fetch_log_path=log)
doc=adapter.find_source('6N0A7AA',name='HyperX Pulsefire Haste 2 Wired',deadline=time.monotonic()+20)
out['guarded_check']={'match_level':doc.match_level,'error':doc.error,'network_requests':adapter.request_count,'log_unchanged':log.read_bytes()==before,'active_stops':active_stops(read_log(log))}
assert doc.match_level=='blocked' and adapter.request_count==0 and log.read_bytes()==before
(R/'artifact_verification.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf8')
print(out['ui']);print(out['guarded_check']['match_level'],out['guarded_check']['network_requests']);print(out['excel']['sheet_names'])
