"""Actual-live model-first evidence QA and the existing application's Excel export."""
from pathlib import Path
import sys,json,hashlib,gzip
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from product_tool import jobs
from product_tool.exporter import export_batch
from product_tool.xiaomi_pipeline import photo_verified
from openpyxl import load_workbook
R=Path(__file__).parent;live=R.parent/'xiaomi_stage74/model_first_release_live'
data=json.loads((live/'results.json').read_text(encoding='utf8'));assert len(data['rows'])==10
assert data['dataset_sha256']==(R.parent/'xiaomi_stage74/dataset.sha256').read_text().strip()
index=[json.loads(x) for x in (live/'http/index.jsonl').read_text(encoding='utf8').splitlines()]
qa=R/'qa';qa.mkdir(exist_ok=True);rows=[]
for row in data['rows']:
    ev=row['evidence'];assert row['observation']=='actual_live_worker' and ev['live_success']
    assert row['job']['status']=='done' and ev['identity']['model_relation']=='model_confirmed'
    assert ev['identity']['retail_sku'] is None and not any(x['full_sku_confirmed'] for x in row['resolved'])
    assert not row['readiness']['conflicts'] and 'configuration_missing' not in row['readiness']['blocking_gaps'] and 'manual_unverified' not in row['readiness']['blocking_gaps']
    assert len([x for x in row['trace'] if x.get('observation')=='live_http' and x.get('source_type') in {'pdp','specs'}])==2
    assert len(row['photos'])==len({p['url'] for p in row['photos']})
    confirmed=[p for p in row['photos'] if photo_verified(p,ev)];assert all(p['selected'] for p in confirmed)
    for capture in ev['captures']:
        assert hashlib.sha256((live/'captures'/capture['file']).read_bytes()).hexdigest()==capture['sha256']
    for document in ev.get('manuals',[]):
        if document.get('file'):assert hashlib.sha256((live/'captures'/document['file']).read_bytes()).hexdigest()==document['sha256']
    for m in ev.get('image_measurements',[]):
        if not m.get('accepted_identity'):continue
        assert m['identity_scope']=='model' and m['supplier_color_relation']=='unproven'
        c=next(x for x in index if x['url']==m['url']);raw=gzip.decompress((live/'http'/c['saved_as']).read_bytes()).decode('utf8').encode('latin1');assert hashlib.sha256(raw).hexdigest()==m['sha256']
        (qa/f"model_{row['id']}_{m['sha256'][:12]}.{m['format'].lower()}").write_bytes(raw)
    rows.append({'id':row['id'],'model':row['input']['model'],'official_specs':row['readiness']['confirmed_specs'],'model_photos':len(confirmed),'manual':ev['manual_status'],'readiness':row['readiness']['verdict'],'blocking_gaps':row['readiness']['blocking_gaps'],'input_configuration':row['readiness']['input_configuration']})
file=R/'Xiaomi_Model_First.xlsx';file.write_bytes(export_batch(live/'live.sqlite3','xm74'));book=load_workbook(file,data_only=False)
assert book['Xiaomi готовность'].max_row==11
assert book['Фотографии'].max_row==1+sum(x['model_photos'] for x in rows)
assert all('Данные входного файла' in str(row[3].value) for row in list(book['Xiaomi конфигурация'].iter_rows())[1:])
assert not any(c.data_type=='e' for s in book for line in s for c in line);book.close()
result={'origin':'actual_live_worker; no replay rows substituted','dataset_sha256':data['dataset_sha256'],'rows':rows,'actual_live':10,'ready':sum(x['readiness']=='export_ready' for x in rows),'official_specs':sum(x['official_specs'] for x in rows),'confirmed_model_photos':sum(x['model_photos'] for x in rows),'confirmed_supplier_skus':0,'claimed_supplier_color_relations':0,'workbook_sha256':hashlib.sha256(file.read_bytes()).hexdigest()}
(R/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print({k:v for k,v in result.items() if k!='rows'})
