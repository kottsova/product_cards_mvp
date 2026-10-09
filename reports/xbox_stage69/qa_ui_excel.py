"""Native application UI/Excel QA; export from captured acceptance database."""
from pathlib import Path
from io import BytesIO
import json,shutil,os,re
import sqlite3
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from product_tool import exporter,storage,jobs,xbox_pipeline,attribute_projection
from product_tool.web import create_app
R=Path(__file__).parent;U=R/'ui';U.mkdir(exist_ok=True);db=U/'batches.sqlite3';out=[]
live=json.loads((R/'live_final_f.json').read_text(encoding='utf8'));assert len(live['rows'])==10 and all(r['readiness']['verdict']=='export_ready' for r in live['rows'])
with sqlite3.connect(R/'live_final_f'/'batches.sqlite3') as src,sqlite3.connect(db) as dst:src.backup(dst)
measurements=json.loads((R/'photo_inspection.json').read_text(encoding='utf8'))
for p in storage.get_batch(db,'xb69')['products']:
 for photo in jobs.get_photo_candidates(db,p['id']):
  m=next((m for m in measurements if m['url']==photo['url'] and not m.get('error')),None)
  if m:jobs.save_photo_metadata(db,p['id'],photo['id'],photo['url'],width=m['width'],height=m['height'],size_bytes=m['size_bytes'],image_format=m['format'])
with TestClient(create_app(U,start_worker=False)) as client:
 for p in storage.get_batch(db,'xb69')['products']:
  response=client.get(f'/products/{p["id"]}');assert response.status_code==200
  html=response.text;assert 'Модель и конфигурация Xbox' in html and 'Фото-кандидаты' in html and 'Документы и руководства' in html
  assert 'Только у базовой модели LG' not in html
  (U/f'product_{p["id"]}.html').write_text(html,encoding='utf8')
  descriptions=[s['description'] for s in jobs.get_source_pages(db,p['id'])]
  if p['row_number']==4:
   assert all('Carbon Black' not in d for d in descriptions)
   assert any('Robot White' in d for d in descriptions)
  assert not any(re.search(r'how to|reset your|repair|account setup|firmware update|drivers available',d,re.I) for d in descriptions)
  out.append(dict(article=p['search_code'],http=response.status_code,readiness=xbox_pipeline.card_readiness(db,p['id']),labels=[r['display_name'] for r in attribute_projection.final_attribute_rows(db,p['id'])],descriptions_clean=True))
 data=exporter.export_batch(db,'xb69');(R/'export.xlsx').write_bytes(data);book=load_workbook(BytesIO(data))
 sheets={s.title:dict(rows=s.max_row,columns=s.max_column,headers=[c.value for c in s[1]]) for s in book}
 for name in ('Готовность Xbox','Конфигурация Xbox','Идентификаторы Xbox','Кандидаты Xbox','Документы Xbox','Фото-кандидаты'):assert name in sheets
 errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
 assert book['Готовность Xbox'].max_row==11
 assert book['Связи Xbox'].max_row > 10
 assert book['Принятые факты Xbox'].max_row > 100
 book.close()
(R/'qa_ui_excel.json').write_text(json.dumps(dict(ui=out,sheets=sheets,formula_errors=errors,native_exporter=True),ensure_ascii=False,indent=2),encoding='utf8')
print('10 UI HTTP200, native Excel, scopes and clean descriptions verified',flush=True)
os.environ.setdefault('PLAYWRIGHT_BROWSERS_PATH',str(Path('.venv/playwright-browsers').resolve()))
from playwright.sync_api import sync_playwright
from urllib.parse import urlsplit
with TestClient(create_app(U,start_worker=False)) as client,sync_playwright() as pw:
 browser=pw.chromium.launch(headless=True);page=browser.new_page(viewport=dict(width=1440,height=1000))
 def route(request):
  m=next((m for m in measurements if m['url']==request.request.url and not m.get('error')),None)
  if m:request.fulfill(status=200,body=(R/m['file']).read_bytes(),content_type='image/jpeg');return
  u=urlsplit(request.request.url)
  if u.hostname!='xbox-audit.test':request.abort();return
  response=client.get(u.path);request.fulfill(status=response.status_code,body=response.content,content_type=response.headers.get('content-type','text/html'))
 page.route('**/*',route)
 for pid in (1,2,4,8,10):
  page.goto(f'http://xbox-audit.test/products/{pid}',wait_until='networkidle')
  page.locator('img[loading="lazy"]').evaluate_all('(nodes)=>nodes.forEach(n=>n.loading="eager")')
  page.wait_for_function('Array.from(document.querySelectorAll("img[loading]")).every(n=>n.complete && n.naturalWidth>0)',timeout=15000)
  page.screenshot(path=str(R/f'ui_{pid}_overview.png'))
  page.get_by_role('heading',name='Модель и конфигурация Xbox',exact=True).scroll_into_view_if_needed();page.screenshot(path=str(R/f'ui_{pid}_evidence.png'))
  if pid==4:
   page.locator('.photo-open').first.scroll_into_view_if_needed();page.locator('.photo-open').first.click();assert page.locator('#photo-lightbox').is_visible();page.screenshot(path=str(R/'ui_4_lightbox.png'));page.get_by_role('button',name='Закрыть',exact=True).click()
 browser.close()
print('Native UI screenshots saved',flush=True)
