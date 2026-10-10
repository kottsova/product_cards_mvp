"""Independent evidence checks and the application's existing Excel export."""
import sys,json,hashlib
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from product_tool import jobs,card_evidence,readiness
from product_tool.exporter import export_batch
from product_tool.xiaomi_page import parse_page
from product_tool.normalization import normalize_facts
from collections import defaultdict
from openpyxl import load_workbook
R=Path(__file__).parent;dataset=json.loads((R/'dataset.json').read_text(encoding='utf8'))
assert hashlib.sha256((R/'dataset.json').read_bytes()).hexdigest()==(R/'dataset.sha256').read_text().strip()
original=json.loads((R/'verified_live/results.json').read_text(encoding='utf8'))['rows'];repeat=json.loads((R/'stability/results.json').read_text(encoding='utf8'))['rows'];by_repeat={r['id']:r for r in repeat}
db=R/'verified_live/live.sqlite3';combined=[by_repeat.get(r['id'],r) for r in original];offline=[]
for r in combined:
    assert r['job']['status']=='done' and r['evidence']['live_success']
    assert r['evidence']['identity']['retail_sku'] is None
    assert not any(f['full_sku_confirmed'] for f in r['resolved'])
    assert not r['readiness']['conflicts']
    assert not any(p['selected'] for p in r['photos'])
    captures=R/('stability' if r['id'] in by_repeat else 'verified_live')/'captures'
    for c in r['evidence']['captures']:
        assert hashlib.sha256((captures/c['file']).read_bytes()).hexdigest()==c['sha256']
    source=next(s for s in r['sources'] if s['source_key']=='xiaomi_model');cap=next(c for c in r['evidence']['captures'] if c['url']==source['url'])
    doc,ev=parse_page((captures/cap['file']).read_text(encoding='utf8'),source['url'],r['input']['article'],r['input']['name'],r['input']['category']);group=defaultdict(set)
    for f in normalize_facts(doc.attributes):group[f.normalized_name].add(f.normalized_value)
    assert not any(len(v)>1 for v in group.values())
    offline.append({'id':r['id'],'observation':'offline_parser_only','raw':len(ev['raw_specs']),'accepted_raw_rows':len(ev['accepted_specs']),'rejected_raw_rows':len(ev['rejected_specs']),'normalized_model_fields':len(group),'model_numbers':ev['model_numbers']})
    # Preserve no model/regional source loss, duplicate or readiness instability.
    prior=next(x for x in original if x['id']==r['id'])
    assert {s['url'] for s in prior['sources']}=={s['url'] for s in r['sources']}
    assert prior['readiness']['verdict']==r['readiness']['verdict']
    assert prior['readiness']['confirmed_specs']==r['readiness']['confirmed_specs']
    keys=[(f['source_key'],f['normalized_name'],f['raw_value']) for f in jobs.get_facts(db,r['id'])]
    assert len(keys)==len(set(keys))
(R/'final_results.json').write_text(json.dumps({'origin':'actual live snapshots; five repeated live rows supersede earlier rows','rows':combined,'offline_parser':offline},ensure_ascii=False,indent=2),encoding='utf8')
file=R/'Xiaomi_Stage74.xlsx';file.write_bytes(export_batch(db,'xm74'));book=load_workbook(file,data_only=False)
assert book['Xiaomi готовность'].max_row==11
assert book['Фотографии'].max_row==1
assert book['Xiaomi фото-кандидаты'].max_row==1+sum(len(r['photos']) for r in combined)
assert all(row[5].value=='Не подтверждён' for row in list(book['Xiaomi готовность'].iter_rows())[1:])
errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
qa={'actual_live_models':len(combined),'actual_live_specs':sum(r['readiness']['confirmed_specs'] for r in combined),'confirmed_retail_skus':0,'ready':sum(r['readiness']['verdict']=='export_ready' for r in combined),'candidate_images':sum(len(r['photos']) for r in combined),'byte_verified_candidates':sum('width' in m for r in combined for m in r['evidence'].get('image_measurements',[])),'manual_checked':sum(r['evidence']['manual_status']!='Не проверена' for r in combined),'repeated_live_models':len(repeat),'duplicates':0,'false_confirmed_configuration':0,'export':{'file':file.name,'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'sheets':book.sheetnames,'formula_errors':errors},'frozen_input_metadata_issue':'PowerShell stdin replaced Cyrillic category labels with question marks during dataset creation. Frozen bytes were not changed; model identifiers/configuration and semantic product mix are intact.'}
book.close();(R/'verification.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf8');print({k:v for k,v in qa.items() if k not in {'export','frozen_input_metadata_issue'}})
