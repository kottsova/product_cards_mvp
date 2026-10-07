from pathlib import Path
import json,shutil
from product_tool.web import create_app
from product_tool import exporter,storage,jobs,apple_pipeline
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from io import BytesIO
r=Path('reports/apple_stage64');u=r/'ui_batch';u.mkdir(exist_ok=True);shutil.copyfile(r/'final_acceptance.sqlite3',u/'batches.sqlite3');p=u/'batches.sqlite3';pages=[]
with TestClient(create_app(u,start_worker=False)) as client:
 for product in storage.get_batch(p,'apple64')['products']:
  z=client.get(f"/products/{product['id']}");assert z.status_code==200;html=z.text;assert 'Конфигурация Apple' in html;assert 'Параметры и варианты-кандидаты' in html;assert 'Не проверена' in html
  (r/f"ui_{product['id']}.html").write_text(html,encoding='utf-8');pages.append({'article':product['search_code'],'http':z.status_code,'readiness':apple_pipeline.card_readiness(p,product['id']),'raw_help_not_description':all(not s['description'] for s in jobs.get_source_pages(p,product['id']))})
 data=exporter.export_batch(p,'apple64');(r/'export.xlsx').write_bytes(data);book=load_workbook(BytesIO(data));sheets={s.title:{'rows':s.max_row,'columns':s.max_column,'headers':[x.value for x in s[1]]} for s in book};assert 'Готовность Apple' in sheets;assert 'Варианты-кандидаты Apple' in sheets
 errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
(r/'qa_ui_excel.json').write_text(json.dumps({'ui':pages,'sheets':sheets,'formula_errors':errors,'note':'Native exporter; read-only Excel inspection. Common batch readiness has not yet been specialized for Apple, whereas the Apple audit panel reports its conservative configuration gate.'},ensure_ascii=False,indent=2),encoding='utf-8');print('11 UI pages HTTP200; native export verified; candidate rows',sheets['Варианты-кандидаты Apple']['rows'])
