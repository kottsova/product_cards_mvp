"""Read-only QA of the actual export route for the user's Apple batch."""
import sys,json,hashlib,urllib.request
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from product_tool import jobs,card_evidence
out=Path(__file__).parent;db=out.parents[1]/'data/batches.sqlite3'
url='http://127.0.0.1:8765/batches/6cb1a747ecd84b1ab9f079a14eeac52d/export.xlsx'
payload=urllib.request.urlopen(url,timeout=60).read();(out/'Apple_Acceptance.xlsx').write_bytes(payload)
book=load_workbook(BytesIO(payload),data_only=False)
sheet=book.worksheets[0];headers=[c.value for c in sheet[1]]
assert len(headers)==len(set(headers)), 'Duplicate product columns'
required=['Мощность беспроводной зарядки MagSafe, Вт','Мощность беспроводной зарядки Qi2, Вт','Идентификация аксессуаров NFC','Магнитометр','Быстрая зарядка USB-C','Быстрая зарядка MagSafe','Время быстрой зарядки USB-C, мин','Время быстрой зарядки MagSafe, мин','Максимальная типичная яркость экрана, нит','Минимальная яркость экрана, нит']
values={k:sheet.cell(2,headers.index(k)+1).value for k in required}
assert values['Мощность беспроводной зарядки MagSafe, Вт']=='До 25 Вт'
assert values['Время быстрой зарядки USB-C, мин']=='Около 15 мин'
assert values['Время быстрой зарядки MagSafe, мин']=='30 мин'
raw_sheet=book['Семантика и исходные значения'];raw=list(raw_sheet.iter_rows(min_row=2,values_only=True))
semantic=card_evidence.load(db,92,'semantic_decomposition')
expected=[(block['raw_value'],child['key']) for ev in semantic['sources'].values() for block in ev['decomposed'] for child in block['children']]
assert sorted((r[4],r[5]) for r in raw)==sorted(expected)
assert not any('MagSafe и беспроводная зарядка'==h for h in headers)
compat=[r for r in book['Атрибуты'].iter_rows(min_row=2,values_only=True) if r[4]=='Системные требования']
assert compat and all(r[2]=='Совместимость и системные требования' for r in compat)
errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e']
assert not errors
ranges=[dict(sheet=s.title,range='A1:'+get_column_letter(min(s.max_column,6))+str(min(s.max_row,4))) for s in book]
result=dict(url=url,sha256=hashlib.sha256(payload).hexdigest(),sheets=[dict(name=s.title,rows=s.max_row,columns=s.max_column) for s in book],duplicate_product_columns=0,formula_errors=errors,raw_children_preserved=len(raw),raw_blocks=sum(len(v['decomposed']) for v in semantic['sources'].values()),examples=values,all_product_values={h:sheet.cell(2,i+1).value for i,h in enumerate(headers)},render_ranges=ranges)
(out/'export_qa.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({k:v for k,v in result.items() if k not in {'all_product_values','render_ranges'}},ensure_ascii=False),flush=True)
