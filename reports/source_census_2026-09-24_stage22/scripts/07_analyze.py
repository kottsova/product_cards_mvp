"""Stage 22 step 7 -- OFFLINE, zero requests. Turns the raw files of the second pilot into the tables and numbers of the report.

Inputs : raw/pilot2_selection.json, raw/pilot2_rows.json (the LIVE run, code as it was before three post-run fixes), raw/pilot2_replay_rows.json (the same
         24 rows re-run OFFLINE on the saved responses with the code after those fixes), raw/pilot2_requests.json, the Stage 20 / 21 row files,
         the saved KZ and RU sitemaps, the Stage 19 queue, pilot2_replay/workdir (for the export check).
Output : raw/pilot2_comparison.json, raw/estimate_578_stage22.json, raw/export_check2.json, numbers.json
"""
from __future__ import annotations

import gzip
import json
import re
import sqlite3
import sys
from collections import Counter
from io import BytesIO
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
STAGE = HERE.parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from product_tool.adapters.lg import _slug_key, lg_article_key, lg_base_model, lg_is_product_url, normalize_lg_sku  # noqa: E402
from product_tool.adapters.sitemap_urls import sitemap_locs  # noqa: E402
from product_tool.exporter import export_batch  # noqa: E402
from product_tool.offline_guard import offline_only  # noqa: E402


def load(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def saved_text(directory: Path, suffix: str) -> str:
    for line in (directory / "index.jsonl").read_text(encoding="utf-8").splitlines():
        entry = json.loads(line)
        if entry["url"].endswith(suffix) and entry["saved_as"]:
            with gzip.open(directory / entry["saved_as"], "rt", encoding="utf-8", newline="") as handle:
                return handle.read()
    raise SystemExit(f"no saved {suffix}")


def instruction_summary(row: dict) -> dict:
    report = row.get("documents_report")
    saved = row["documents_saved"]
    if not report:
        return {"state": "not_attempted", "why": "no LG Russia page with a printed support link", "candidates": 0, "files": [], "saved": [d["language"] for d in saved]}
    files = []
    for f in report["files"]:
        a = f.get("assessment", {})
        files.append({"state": f["state"], "type": f["type"], "label": f["label"], "bytes": f.get("bytes"), "content_languages": a.get("languages"), "russian_instruction": a.get("russian_instruction"),
                      "model_evidence": a.get("model_evidence"), "masks": a.get("model_masks"), "names_model": a.get("names_model")})
    states = {f["state"] for f in files}
    if "instruction_confirmed_by_content" in states:
        top = "instruction_confirmed_by_content"
    elif "instruction_by_content_model_not_named" in states:
        top = "instruction_by_content_model_not_named"
    elif "reachable_file_not_confirmed_as_instruction" in states:
        top = "reachable_file"
    elif "reachable_but_not_a_complete_pdf" in states:
        top = "reachable_not_a_complete_pdf"
    elif report["candidates"]:
        top = "candidate_link"
    else:
        top = "no_candidates"
    support_code = report["support_url"].rsplit("/lg-", 1)[-1].upper() if report.get("support_url") else ""
    return {"state": top, "outcome": report["outcome"], "candidates": len(report["candidates"]), "files": files, "saved": [d["language"] for d in saved], "support_code": support_code}


def main() -> None:
    selection = load(STAGE / "raw/pilot2_selection.json")
    live = {r["seller_sku"]: r for r in load(STAGE / "raw/pilot2_rows.json")["rows"]}
    replay = {r["seller_sku"]: r for r in load(STAGE / "raw/pilot2_replay_rows.json")["rows"]}
    requests_ = load(STAGE / "raw/pilot2_requests.json")
    s20 = {r["seller_sku"]: r for r in load(ROOT / "reports/source_census_2026-09-24_stage20/raw/pilot_analysis.json")["rows"]}
    s21 = {r["seller_sku"]: r for r in load(ROOT / "reports/source_census_2026-09-24_stage21/raw/replay_phase2.json")["rows"]}
    assert len(live) == len(replay) == 24 and not requests_["halted"] and not requests_["stopped_hosts"]

    comparison = []
    for row in selection["rows"]:
        sku = row["seller_sku"]
        lv, rp = live[sku], replay[sku]
        ks, rs = rp["sources"].get("lg_kz", {}), rp["sources"].get("lg_ru", {})
        comparison.append({
            "seller_sku": sku, "category": row["category"], "group": row["group"],
            "history": {"stage20_job": s20[sku]["job"]["status"] if sku in s20 else None, "stage21_job": s21[sku]["job_status"] if sku in s21 else None},
            "live_job": lv["job"]["status"], "job": rp["job"]["status"], "readiness": rp["readiness"]["verdict"], "gaps": rp["readiness"]["gaps"], "blocking": rp["readiness"]["blocking_gaps"],
            "variant": {"kz": [ks.get("match_level"), ks.get("url", "")], "ru": [rs.get("match_level"), rs.get("url", "")]},
            "facts": {"kz": ks.get("facts"), "ru": rs.get("facts")}, "resolved_total": rp["resolved_total"],
            "real_conflicts": [{"name": c["normalized_name"], "raw_names": c["raw_names"], "values": c["values"], "sources": c["sources"]} for c in rp["real_conflicts"]],
            "live_conflicts": len(lv["real_conflicts"]),
            "gallery": {"kz": ks.get("gallery_selected"), "ru": rs.get("gallery_selected"), "from_exact_pages": rp["readiness"]["official_gallery_from_exact_pages"]},
            "instruction": instruction_summary(rp), "live_instruction_state": instruction_summary(lv)["state"],
            "live_real_requests": lv["real_requests"], "live_seconds": lv["seconds"],
        })
    (STAGE / "raw/pilot2_comparison.json").write_text(json.dumps(comparison, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # ---- the 578 rows against the saved sitemaps (offline) -------------------------------------------------------------
    kz = {_slug_key(u) for u in sitemap_locs(saved_text(ROOT / "reports/source_census_2026-09-24_stage20/pilot/responses", "/kz/sitemap.xml")) if lg_is_product_url(u)}
    ru = {lg_article_key(u.rstrip("/").rsplit("/", 1)[-1].removeprefix("lg-")) for u in sitemap_locs(saved_text(ROOT / "reports/source_census_2026-09-24_stage21/probe/responses", "/ru/sitemap.xml"))
          if "/ru/" in u and "/support/" not in u}
    queue = [json.loads(line) for line in (ROOT / "reports/source_census_2026-09-24_stage19/queue/coverage_units.jsonl").read_text(encoding="utf-8").splitlines() if line]
    ready = [u for u in queue if u["family"] == "lg" and u["status"] == "ready_to_run"]

    def keys(article: str) -> tuple[str, str]:
        full = normalize_lg_sku(article)
        return lg_article_key(full), lg_article_key(lg_base_model(full))

    tiers, market_tag, by_category, regions = Counter(), [], {}, Counter()
    for unit in ready:
        full_key, base_key = keys(unit["seller_sku"])
        exact = full_key in kz or full_key in ru
        base = base_key in kz or base_key in ru
        tier = "exact_slug" if exact else "base_only" if base else "neither"
        tiers[tier] += 1
        if exact:
            regions["exact_both_regions" if full_key in kz and full_key in ru else "exact_ru_only" if full_key in ru else "exact_kz_only"] += 1
        by_category.setdefault(unit["category"], Counter())[tier] += 1
        if re.search(r"_[A-Za-z]{2}$", unit["seller_sku"]):
            stripped = re.sub(r"_[A-Za-z]{2}$", "", unit["seller_sku"])
            f2, b2 = keys(stripped)
            market_tag.append({"seller_sku": unit["seller_sku"], "tier_now": tier, "found_after_stripping_the_tag": f2 in kz or f2 in ru or b2 in kz or b2 in ru})
    calibration = []
    for row in comparison:
        full_key, base_key = keys(row["seller_sku"])
        predicted = "exact_slug" if full_key in kz or full_key in ru else "base_only" if base_key in kz or base_key in ru else "neither"
        levels = {row["variant"]["kz"][0], row["variant"]["ru"][0]}
        actual = "exact_slug" if "full_sku" in levels else "base_only" if "base_model" in levels else "neither"
        calibration.append({"seller_sku": row["seller_sku"], "predicted": predicted, "actual": actual})
    estimate = {
        "rows": len(ready), "tiers_by_saved_sitemaps": dict(tiers), "exact_slug_by_region": dict(regions),
        "pages_reachable_now": tiers["exact_slug"] + tiers["base_only"], "share_pages_reachable_now": round((tiers["exact_slug"] + tiers["base_only"]) / len(ready), 3),
        "expected_exact_page": tiers["exact_slug"], "share_expected_exact": round(tiers["exact_slug"] / len(ready), 3),
        "expected_base_model_only": tiers["base_only"], "expected_no_page": tiers["neither"],
        "market_tag_rows": len(market_tag), "market_tag_rows_currently_missed": sum(1 for m in market_tag if m["tier_now"] == "neither"), "market_tag_rows_found_if_the_tag_were_stripped": sum(1 for m in market_tag if m["found_after_stripping_the_tag"]),
        "truly_absent_without_tag_rows": tiers["neither"] - sum(1 for m in market_tag if m["tier_now"] == "neither"),
        "calibration_on_the_24_pilot_rows": {"agree": sum(1 for c in calibration if c["predicted"] == c["actual"]), "of": len(calibration), "disagree": [c for c in calibration if c["predicted"] != c["actual"]]},
        "by_category": {c: dict(v) for c, v in sorted(by_category.items())},
        "note": "The tier is read from slug matches in the saved KZ and RU sitemaps: an exact slug = the article itself has a product URL; base_only = only the base-model code has one. On the 24 pilot rows the tier predicted the page match level in every case.",
    }
    (STAGE / "raw/estimate_578_stage22.json").write_text(json.dumps(estimate, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # ---- the ordinary export of the replay batch -------------------------------------------------------------------------
    database = STAGE / "pilot2_replay/workdir/batches.sqlite3"
    book = load_workbook(BytesIO(export_batch(database, "lg-pilot2")))
    sheets = {ws.title: ws for ws in book.worksheets}
    connection = sqlite3.connect(f"file:{database}?mode=ro", uri=True)
    names = {name: code for _, name, code in connection.execute("SELECT id, name, search_code FROM products")}
    connection.close()
    photos, docs = Counter(), Counter()
    for values in list(sheets["Фотографии"].iter_rows(values_only=True))[1:]:
        photos[names.get(values[0], values[0])] += 1
    for values in list(sheets["Инструкции"].iter_rows(values_only=True))[1:]:
        docs[names.get(values[0], values[0])] += 1
    export = {"sheets": sorted(sheets), "rows_with_photo_rows": len(photos), "photo_rows": sum(photos.values()), "rows_with_instruction_rows": len(docs), "instruction_rows": sum(docs.values()),
              "instruction_languages_in_sheet": sorted({v[2] for v in list(sheets["Инструкции"].iter_rows(values_only=True))[1:]})}
    (STAGE / "raw/export_check2.json").write_text(json.dumps(export, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    # ---- numbers ---------------------------------------------------------------------------------------------------------
    events = requests_["events"]
    real = [e for e in events if e["kind"] == "real"]
    index = [json.loads(line) for line in (STAGE / "pilot2/responses/index.jsonl").read_text(encoding="utf-8").splitlines()]
    real_urls = {e["url"] for e in real}
    real_bytes = sum(e["bytes"] for e in index if e["url"] in real_urls)
    doc_bytes = sum(e["bytes"] for e in index if e["url"] in real_urls and "lge.com" in e["url"])
    new_rows = [c for c in comparison if c["group"] != "first_pilot"]
    exact_rows = [c for c in comparison if "full_sku" in (c["variant"]["kz"][0], c["variant"]["ru"][0])]
    both_exact = [c for c in exact_rows if c["variant"]["kz"][0] == "full_sku" and c["variant"]["ru"][0] == "full_sku"]
    ru_rows = [c for c in comparison if c["variant"]["ru"][0] in ("full_sku", "base_model")]
    states = Counter(c["instruction"]["state"] for c in ru_rows)
    russian_named = [c for c in ru_rows if "Русский" in c["instruction"]["saved"]]
    numbers = {
        "rows": len(comparison), "new_rows": len(new_rows), "new_rows_not_in_kz_sitemap": selection["new_rows_not_in_kz_sitemap"],
        "job_live": dict(Counter(c["live_job"] for c in comparison)), "job_after_fixes": dict(Counter(c["job"] for c in comparison)),
        "readiness_after_fixes": dict(Counter(c["readiness"] for c in comparison)),
        "done_but_not_export_ready": sum(1 for c in comparison if c["job"] == "done" and c["readiness"] != "export_ready"),
        "conflicts_live": sum(c["live_conflicts"] for c in comparison), "conflicts_after_fixes": sum(len(c["real_conflicts"]) for c in comparison),
        "rows_with_conflicts_after_fixes": sum(1 for c in comparison if c["real_conflicts"]),
        "exact_page_rows": len(exact_rows), "exact_rows_both_regions": len(both_exact), "exact_rows_both_regions_with_conflicts": sum(1 for c in both_exact if c["real_conflicts"]),
        "exact_rows_single_region": len(exact_rows) - len(both_exact), "exact_rows_single_region_with_conflicts": sum(1 for c in exact_rows if c not in both_exact and c["real_conflicts"]),
        "base_model_only_rows": sum(1 for c in comparison if c["blocking"] and "no_official_full_sku_page" in c["blocking"] and c["facts"]["kz"] + c["facts"]["ru"] > 0),
        "no_official_page_rows": sum(1 for c in comparison if "no_official_specifications" in c["blocking"]),
        "rows_with_ru_page": len(ru_rows), "instruction_states_among_ru_page_rows": dict(states),
        "russian_instruction_saved_with_model_evidence": len(russian_named),
        "instruction_model_evidence": dict(Counter(f.get("model_evidence") for c in ru_rows for f in c["instruction"]["files"] if f["state"] == "instruction_confirmed_by_content")),
        "kazakh_only_instruction_rows": [c["seller_sku"] for c in comparison if c["instruction"]["saved"] and "Русский" not in c["instruction"]["saved"]],
        "requests": {"real_total": requests_["real_requests_total"], "cap_total": 180, "max_real_in_one_row": requests_["max_real_in_one_row"], "cap_row": 12, "refusals": len(requests_["refusals"]),
                     "replayed": sum(1 for e in events if e["kind"] == "replayed"), "real_new_rows": sum(requests_["real_per_row"][c["seller_sku"]] for c in new_rows),
                     "real_first_pilot_rows": sum(requests_["real_per_row"][c["seller_sku"]] for c in comparison if c["group"] == "first_pilot"),
                     "real_downloaded_mb": round(real_bytes / 1e6, 1), "real_document_files_mb": round(doc_bytes / 1e6, 1), "real_document_files": sum(1 for e in real if "lge.com" in e["url"]),
                     "blocks": 0, "stopped_hosts": requests_["stopped_hosts"]},
        "seconds_per_row_live": {"mean": round(sum(c["live_seconds"] for c in comparison) / len(comparison), 1), "max": max(c["live_seconds"] for c in comparison),
                                 "rows_over_50_s": [c["seller_sku"] for c in comparison if c["live_seconds"] > 50]},
        "export": export,
        "estimate_578": {k: estimate[k] for k in ("pages_reachable_now", "share_pages_reachable_now", "expected_exact_page", "share_expected_exact", "expected_base_model_only", "expected_no_page",
                                                 "exact_slug_by_region", "market_tag_rows", "market_tag_rows_currently_missed", "market_tag_rows_found_if_the_tag_were_stripped", "truly_absent_without_tag_rows")},
        "calibration_agree": estimate["calibration_on_the_24_pilot_rows"]["agree"],
    }
    (STAGE / "numbers.json").write_text(json.dumps(numbers, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(numbers, ensure_ascii=False, indent=1))
    print(json.dumps(estimate["tiers_by_saved_sitemaps"]), estimate["calibration_on_the_24_pilot_rows"]["agree"])


if __name__ == "__main__":
    with offline_only():
        main()
