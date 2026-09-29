"""Offline Stage 45 replay of the existing product row on a SQLite backup."""
from __future__ import annotations

import gc
import gzip
import io
import json
import sqlite3
import tempfile
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from product_tool import exporter, jobs, worker
from product_tool.readiness import card_readiness
from product_tool.adapters.common import SourceDocument
from product_tool.adapters.lg_support import LGSupportAdapter

OUT = Path(__file__).resolve().parent
DB = ROOT / "data/batches.sqlite3"
BATCH = "ae3d2cb381744ba8a811e54231233254"
PRODUCT = 4
URL = "https://www.lg.com/kz/support/product-support/cs-P12ED.USAR/"

class EmptyOfficial:
    def __init__(self, key):
        self.source_key = self.site_name = key
    def find_source(self, *args, **kwargs):
        return SourceDocument(self.source_key, self.site_name, "", match_level="mismatch", evidence="Saved official result replayed offline")

class StoppedDealer:
    source_key = site_name = "sulpak"
    def find_source(self, *args, **kwargs):
        return SourceDocument(self.source_key, self.site_name,
            "https://www.sulpak.kz/g/kondicioneriy_split_sistemiy_lg_p12ednsar___p12edusar",
            match_level="unknown", error="policy_host_stopped",
            evidence="Host stop followed a different Sulpak page; this candidate was not requested")

class EmptyDns:
    source_key = site_name = "dns"
    document_urls = {}
    def find_source(self, *args, **kwargs):
        return SourceDocument(self.source_key, self.site_name, "", match_level="dealer_url_needed")
    def find_documents(self, *args, **kwargs):
        return [], "No verified exact URL"

class NoHttp:
    calls = 0
    def get(self, *args, **kwargs):
        self.calls += 1
        raise AssertionError("saved support snapshot must not trigger HTTP")

with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as directory:
    copy = Path(directory) / "batch.sqlite3"
    source = sqlite3.connect(DB)
    target = sqlite3.connect(copy)
    source.backup(target)
    source.close(); target.close()
    before = jobs.get_source_pages(copy, PRODUCT)
    http = NoHttp()
    job_id = jobs.enqueue(copy, PRODUCT, [1,2,3,4,6])
    assert worker.run_once(copy,
        adapter_factory=lambda: (EmptyOfficial("lg_kz"), EmptyOfficial("lg_ru"), StoppedDealer()),
        dns_adapter_factory=EmptyDns,
        lg_support_adapter_factory=lambda: LGSupportAdapter(http=http))
    connection = sqlite3.connect(copy)
    job = connection.execute("SELECT status,message FROM search_jobs WHERE id=?", (job_id,)).fetchone()
    product_count = connection.execute("SELECT count(*) FROM products WHERE batch_id=?", (BATCH,)).fetchone()[0]
    integrity = connection.execute("PRAGMA integrity_check").fetchone()[0]
    foreign_keys = len(connection.execute("PRAGMA foreign_key_check").fetchall())
    connection.close()
    after = jobs.get_source_pages(copy, PRODUCT)
    support = next(x for x in after if x["source_key"] == "lg_kz_support")
    workbook = load_workbook(io.BytesIO(exporter.export_batch(copy, BATCH)), read_only=True, data_only=True)
    export_urls = [str(cell) for sheet in workbook for row in sheet.iter_rows(values_only=True)
                   for cell in row if cell is not None and URL in str(cell)]
    workbook.close()
    result = {
        "product_id": PRODUCT, "batch_rows": product_count, "job_status": job[0],
        "support_match_level": support["match_level"], "support_url": support["url"],
        "support_http_calls": http.calls, "sources_before": len(before), "sources_after": len(after),
        "card_readiness": card_readiness(copy, PRODUCT)["verdict"],
        "attributes": len(jobs.get_facts(copy, PRODUCT)),
        "photos": len(jobs.get_photo_candidates(copy, PRODUCT, include_excluded=False)),
        "documents": len(jobs.get_documents(copy, PRODUCT)),
        "excel_support_url_hits": len(export_urls), "integrity": integrity,
        "foreign_key_errors": foreign_keys,
    }
    assert result["batch_rows"] == 12 and result["job_status"] == "needs_review"
    assert result["support_match_level"] == "support_candidate" and result["support_http_calls"] == 0
    assert result["excel_support_url_hits"] > 0 and result["integrity"] == "ok"
    (OUT / "offline_replay.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result))
    gc.collect()
