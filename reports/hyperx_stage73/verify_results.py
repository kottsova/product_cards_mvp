"""Compare actual runs, replay parsers explicitly offline, and read UI/Excel."""
import sys,json,hashlib,sqlite3,tempfile
from pathlib import Path
from io import BytesIO
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from product_tool import jobs,exporter
from product_tool.fetch_history import latest_source_snapshot
from product_tool.adapters.hyperx import HyperXAdapter
from product_tool.hyperx_page import refine_page
from product_tool.web import create_app
from fastapi.testclient import TestClient
from openpyxl import load_workbook
R=Path(__file__).parent;db=R/'stability_actual.sqlite3'
a=json.loads((R/'release_actual.json').read_text(encoding='utf8'));b=json.loads((R/'stability_actual.json').read_text(encoding='utf8'))
assert len(a['rows'])==10 and len(b['rows'])==5 and a['contexts']==b['contexts']==1
def source(r):return next(s for s in r['sources'] if s['source_key']=='hyperx')
def facts(r):return sorted((f['raw_name'],f['raw_value'],f['section']) for f in r['facts'] if f['source_key']=='hyperx')
def photos(r):return sorted((p['asset_key'],p['url'],p['verified_width'],p['verified_height'],p['verified_bytes'],p['verified_format']) for p in r['photos'] if p['selected'] and not p['excluded_reason'] and p['kind']=='product_gallery')
stability=[];offline=[]
for y in b['rows']:
 x=next(x for x in a['rows'] if x['id']==y['id'])
 s=source(y);fresh=[e for e in y['trace'] if e.get('provider')=='official_live_browser' and e.get('url')==s['url'] and e.get('accepted')]
 assert fresh and s['match_level'] in {'exact_variant','model_confirmed'} and not s['error']
 row={'id':y['id'],'fresh_pdp':True,'same_source':source(x)['url']==s['url'],'same_identity':x['readiness']['identity']==y['readiness']['identity'],'same_facts':facts(x)==facts(y),'same_readiness':x['readiness']['verdict']==y['readiness']['verdict'],'same_verified_photos':photos(x)==photos(y),'manuals_retained':{d['direct_url'] for d in x['documents']}<={d['direct_url'] for d in y['documents']},'no_photo_duplicates':len(y['photos'])==len({p['asset_key'] for p in y['photos']}),'no_fact_duplicates':len(facts(y))==len(set(facts(y)))}
 assert all(v for k,v in row.items() if k!='id'),row
 stability.append(row)
for y in a['rows']:
 s=source(y)
 assert not s['error'] and s['match_level'] in {'exact_variant','model_confirmed'}
 assert any(e.get('provider')=='official_live_browser' and e.get('source_type')=='fresh_dom' and e.get('url')==s['url'] and e.get('http_status')==200 and e.get('accepted') for e in y['trace'])
 assert all(all(value is not None and value>0 for value in photo[2:5]) for photo in photos(y))
 snapshot=latest_source_snapshot(R/'release_actual.sqlite3',y['product_id'],'hyperx');html=snapshot['content'];p=y['input']
 doc,ev=refine_page(HyperXAdapter().parse_page(html,s['url'],catalog_code=p['article']),html,s['url'],p['article'],p['name'])
 same=ev['accepted_specs']==y['evidence']['accepted_specs'] and ev['rejected_specs']==y['evidence']['rejected_specs']
 assert same and len(ev['raw_specs'])==len(ev['accepted_specs'])+len(ev['rejected_specs'])
 offline.append({'id':y['id'],'mode':'offline parser on stored actual DOM; no new live request','sha256':hashlib.sha256(html.encode()).hexdigest(),'same_accepted_rejected':same,'raw':len(ev['raw_specs']),'accepted':len(ev['accepted_specs']),'rejected':len(ev['rejected_specs']),'attribute_count':len(doc.attributes)})
assert sum(len(photos(r)) for r in a['rows'])==56
assert sum(r['readiness']['verdict']=='export_ready' for r in a['rows'])==9
assert all(r['evidence']['manual_status']!='Не проверена' for r in a['rows'])
(R/'stability.json').write_text(json.dumps({'actual_live':5,'categories':5,'contexts_per_batch':1,'distinct_sessions':a['session_id']!=b['session_id'],'same_persistent_profile':True,'rows':stability},ensure_ascii=False,indent=2),encoding='utf8')
(R/'offline_parser.json').write_text(json.dumps(offline,ensure_ascii=False,indent=2),encoding='utf8')
ui=[]
with tempfile.TemporaryDirectory(prefix='hyperx73_ui_') as temp:
 target=Path(temp)/'batches.sqlite3';src=sqlite3.connect(db);dst=sqlite3.connect(target)
 try:src.backup(dst)
 finally:src.close();dst.close()
 with TestClient(create_app(Path(temp),start_worker=False)) as client:
  for pid in (1,2,8,10):
   response=client.get('/products/'+str(pid));assert response.status_code==200
   (R/f'ui_{pid}.html').write_text(response.text,encoding='utf8');ui.append({'id':pid,'status':200})
data=exporter.export_batch(db,'hx73');(R/'HyperX_Stage73.xlsx').write_bytes(data)
book=load_workbook(BytesIO(data));excel={'sheets':book.sheetnames,'rows':{s.title:s.max_row for s in book.worksheets}};book.close()
(R/'artifact_verification.json').write_text(json.dumps({'origin':'readback; no new live success','ui':ui,'excel':excel},ensure_ascii=False,indent=2),encoding='utf8')
print('Actual 10/10; 9 ready; 56 byte-verified images; manuals checked 10/10; stable 5/5 types; parser matches; UI/Excel readback OK')
