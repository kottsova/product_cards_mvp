"""Real localhost UI: original uploads, default worker, JSON polling and downloads."""
import hashlib,io,json,sys,time,zipfile,urllib.request
from pathlib import Path
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[2];sys.path.insert(0,str(root));out=Path(__file__).parent
from product_tool import jobs,storage,discovery_trace
base='http://127.0.0.1:8765';phase=sys.argv[1];filename='LG_12_original.xlsx' if phase=='lg12' else 'Mixed_original.xlsx';file=out/filename
dest=out/phase;dest.mkdir(exist_ok=True);(dest/'runtime.json').write_bytes(urllib.request.urlopen(base+'/diagnostics/runtime',timeout=5).read())
proof=dict(phase=phase,file_sha256=hashlib.sha256(file.read_bytes()).hexdigest(),flow='actual UI upload and queue, default production factories',polls=[],page_errors=[])
with sync_playwright() as pw:
 browser=pw.chromium.launch(headless=True);page=browser.new_page(viewport={'width':1900,'height':1080},accept_downloads=True)
 page.on('pageerror',lambda e:proof['page_errors'].append(str(e)))
 for attempt in range(2):
  try:page.goto(base+'/',wait_until='domcontentloaded',timeout=15000);break
  except Exception:
   if attempt:raise
 page.locator('form[action="/upload"] input[type=file]').set_input_files(str(file));page.locator('form[action="/upload"] button[type=submit]').click();page.wait_for_url('**/drafts/**')
 draft=page.url.rsplit('/',1)[1];proof['batch_id']=draft
 page.screenshot(path=str(dest/'preview.png'),full_page=True);(dest/'preview.html').write_text(page.content(),encoding='utf8')
 page.locator(f'form[action="/drafts/{draft}/confirm"] button[type=submit]').click();page.wait_for_url('**/batches/**')
 batch=storage.get_batch(root/'data/batches.sqlite3',draft);ids=[p['id'] for p in batch['products']];proof['ids']=ids;proof['brands']=sorted({p['brand'] for p in batch['products']})
 assert len(ids)==(12 if phase=='lg12' else 8)
 assert hashlib.sha256((root/'data/uploads'/f'{draft}.xlsx').read_bytes()).hexdigest()==proof['file_sha256']
 assert page.locator('table').count()==1
 page.locator('#select-none').click();assert page.locator('[name="product_ids"]:checked').count()==0
 page.locator('#select-all').click();assert page.locator('[name="product_ids"]:checked').count()==len(ids)
 page.screenshot(path=str(dest/'before_queue.png'),full_page=True)
 page.get_by_role('button',name='Поставить выбранные в очередь',exact=True).click();page.wait_for_load_state('domcontentloaded')
 token='stage85-'+phase;page.evaluate('(value)=>window.stage85Document=value',token)
 checkbox=page.locator('[name="product_ids"]').last;checkbox.uncheck()
 def collect(response):
  if response.url==base+f'/batches/{draft}/progress' and response.status==200:
   try:proof['polls'].append(dict(observed_at=time.time(),state=response.json()))
   except Exception as e:proof.setdefault('poll_errors',[]).append(str(e))
 page.on('response',collect)
 start=time.monotonic();running_screenshot=False
 while time.monotonic()-start<2400:
  page.wait_for_timeout(1000)
  assert page.evaluate('window.stage85Document')==token,'full page reload'
  assert not checkbox.is_checked(),'poll overwrote selection'
  if proof['polls']:
   state=proof['polls'][-1]['state']
   if not running_screenshot and state['summary']['running']:
    page.screenshot(path=str(dest/'running.png'),full_page=True);running_screenshot=True
   if state['summary']['completed']==len(ids):break
  if int(time.monotonic()-start)%30==0:(dest/'progress.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2),encoding='utf8')
 assert proof['polls'] and proof['polls'][-1]['state']['summary']['completed']==len(ids),'batch deadline'
 proof['seconds']=time.monotonic()-start;proof['final']=proof['polls'][-1]['state']
 assert page.locator('#batch-progress').get_attribute('value')==str(len(ids))
 assert page.locator('#batch-heading').inner_text()=='Результаты обработки'
 page.screenshot(path=str(dest/'final.png'),full_page=True);(dest/'final.html').write_text(page.content(),encoding='utf8')
 with page.expect_download() as downloaded:page.get_by_role('link',name='Скачать Excel',exact=True).click()
 downloaded.value.save_as(str(dest/'result.xlsx'));book=load_workbook(dest/'result.xlsx',read_only=True);proof['excel_sheets']=book.sheetnames;book.close()
 results=[]
 for pid in ids:
  history=jobs.list_jobs(root/'data/batches.sqlite3',pid);assert len(history)==1
  job=history[0]
  with storage._connection(root/'data/batches.sqlite3') as c:events=[dict(x) for x in c.execute('SELECT phase,created_at FROM job_phase_events WHERE job_id=? ORDER BY id',(job['id'],))]
  result=dict(product=jobs.get_product(root/'data/batches.sqlite3',pid),job=job,phase_events=events,sources=jobs.get_source_pages(root/'data/batches.sqlite3',pid),trace=discovery_trace.for_job(root/'data/batches.sqlite3',job['id']))
  (dest/f'{pid}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');results.append(result)
 photo_row=max(proof['final']['rows'],key=lambda r:r['photos_count']);pid=photo_row['id'];proof['photo_product_id']=pid
 page.locator(f'[data-product-row="{pid}"] a',has_text='Подробнее').click();page.wait_for_url(f'**/products/{pid}');proof['card_opened']=True
 page.goto(base+f'/products/{pid}/photos',wait_until='domcontentloaded');assert page.locator('#photos').count()==1
 # Keep the ZIP acceptance small and observable: save exactly two actual assets.
 boxes=page.locator('[name="asset_keys"]');assert boxes.count()>0
 for box in boxes.all():box.uncheck()
 for box in boxes.all()[:2]:box.check()
 page.get_by_role('button',name='Сохранить выбор',exact=True).click();page.wait_for_load_state('domcontentloaded')
 tile=page.locator('.photo-open').first;tile.click();assert page.locator('#photo-lightbox').is_visible();page.get_by_role('button',name='Закрыть',exact=True).click()
 page.screenshot(path=str(dest/'photos.png'),full_page=True)
 with page.expect_download(timeout=120000) as downloaded:page.get_by_role('link',name='Скачать выбранные фото',exact=True).click()
 downloaded.value.save_as(str(dest/'selected_photos.zip'));proof['zip_filename']=downloaded.value.suggested_filename
 from PIL import Image
 with zipfile.ZipFile(dest/'selected_photos.zip') as z:
  proof['zip_entries']=[];hashes=set()
  for name in z.namelist():
   raw=z.read(name);digest=hashlib.sha256(raw).hexdigest();assert digest not in hashes;hashes.add(digest)
   with Image.open(io.BytesIO(raw)) as image:proof['zip_entries'].append(dict(name=name,width=image.width,height=image.height,format=image.format,size=len(raw),sha256=digest))
  assert len(proof['zip_entries'])>=1
 proof['no_reload']=True;proof['selection_preserved']=True;assert not proof['page_errors']
 browser.close()
(dest/'acceptance.json').write_text(json.dumps(proof,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(dict(phase=phase,summary=proof['final']['summary'],seconds=proof['seconds'],zip_entries=len(proof['zip_entries'])),ensure_ascii=True))
