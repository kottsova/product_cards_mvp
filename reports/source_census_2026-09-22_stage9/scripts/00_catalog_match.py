import json
from pathlib import Path
from openpyxl import load_workbook

OUT = Path(r'A:\work\dev\product_cards_mvp\reports\source_census_2026-09-22_stage9')

VERIFIED_PAGE_URL = "https://direct.playstation.com/en-us/buy-accessories/dualsense-wireless-controller-cosmic-red-for-ps5-pc-mac-mobile"

wb = load_workbook(r'A:\work\dev\product_cards_mvp\data\catalog_2026-09-21_filtered.xlsx', read_only=True, data_only=True)
goods = wb['Товары']


def norm(s):
    return (s or '').strip().lower()


candidates = []
for r in goods.iter_rows(min_row=2, values_only=True):
    brand, cat, art, wbart, name = r[0], r[1], r[2], r[3], r[4]
    if cat and 'геймпад' in norm(cat) and 'dualsense' in norm(name) and 'cosmic red' in norm(name):
        candidates.append({'brand': brand, 'category': cat, 'seller_article': art, 'wb_articles': wbart, 'name': name})

result = {
    'schema_version': 'stage9_catalog_match.v1',
    'verified_page_url': VERIFIED_PAGE_URL,
    'verified_page_provenance': 'Already confirmed official Product page from Stage 8.2 (product_page_fixture_result=verified, catalog_identity_result=exact_model at that time)',
    'search_scope': "Товары sheet, Категория contains 'Геймпады', Наименование contains both 'DualSense' and 'Cosmic Red' (case-insensitive)",
    'candidates_found': candidates,
    'candidate_count': len(candidates),
}

json.dump(result, open(OUT / 'catalog_match.json', 'w', encoding='utf-8'), indent=2, ensure_ascii=False)
print(json.dumps(result, indent=2, ensure_ascii=False))
