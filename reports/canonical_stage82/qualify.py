"""Read-only native UI/Excel verification and observed 38-category schema union."""
import sys,json,hashlib,sqlite3
from contextlib import closing
from collections import Counter,defaultdict
from pathlib import Path
from io import BytesIO
sys.path.insert(0,str(Path(__file__).resolve().parents[2]));sys.stdout.reconfigure(encoding='utf8')
from product_tool import jobs,attribute_projection,card_presentation,exporter,photo_inventory
from product_tool.web import create_app
from fastapi.testclient import TestClient
from bs4 import BeautifulSoup
from openpyxl import load_workbook
R=Path(__file__).parent;S=R.parent/'cross_brand_stage81';db=R/'offline/batch.sqlite3';ui=R/'native_ui';ui.mkdir(exist_ok=True)
with closing(sqlite3.connect(db)) as source,closing(sqlite3.connect(ui/'batches.sqlite3')) as target:source.backup(target)
inputs=json.loads((R/'dataset.json').read_text(encoding='utf8'))['rows'];classification=json.loads((R/'classification.json').read_text(encoding='utf8'))
inventory=defaultdict(dict);checks=[];additions=[];removed=[];gallery_before=[];gallery_after=[]
with TestClient(create_app(ui,start_worker=False),base_url='http://stage82-audit.test') as client:
 for inp in inputs:
  pid=inp['id'];facts=jobs.get_facts(db,pid);sources=jobs.get_source_pages(db,pid)
  rows=attribute_projection.final_attribute_rows(db,pid);before=attribute_projection.final_attribute_rows(S/'qualified_live/batch.sqlite3',pid)
  keys=[r['normalized_name'] for r in rows];assert len(keys)==len(set(keys)),(pid,'duplicate canonical keys')
  old_keys={r['normalized_name'] for r in before}
  additions.extend({'id':pid,'brand':inp['brand'],'category':inp['category'],'key':r['normalized_name'],'label':r['display_name'],'group':r['parent_group']} for r in rows if r['normalized_name'] not in old_keys)
  removed.extend({'id':pid,'brand':inp['brand'],'key':r['normalized_name'],'label':r['display_name']} for r in before if r['normalized_name'] not in set(keys))
  for d in classification['rows']:
   if d['product_id']==pid and d['action'] in {'mapped','feature','parser_fixed'}:assert d['canonical_targets'] and set(d['canonical_targets'])<=set(keys),(d['row_id'],'no canonical target')
  presented=card_presentation.present_card_rows(rows,facts,sources,inp['category'])
  rendered=client.get('/products/'+str(pid));assert rendered.status_code==200
  soup=BeautifulSoup(rendered.text,'html.parser');assert len(soup.select('.card-attributes tbody>tr'))==presented['count']==len(rows)
  assert soup.select_one('#attribute-search') and soup.select_one('#attribute-role')
  (ui/f'product_{pid}.html').write_text(rendered.text,encoding='utf8')
  counts=Counter()
  for role,groups in presented['groups'].items():
   for group,members in groups.items():
    for r in members:
     counts[role]+=1;key=(group,r['display_name'])
     record=inventory[inp['category']].setdefault(key,{'attribute':r['display_name'],'parent_group':group,'roles':set(),'canonical_keys':set(),'brands':set(),'models':set()})
     for k,v in [('roles',role),('canonical_keys',r['normalized_name']),('brands',inp['brand']),('models',pid)]:record[k].add(v)
  photos=photo_inventory.enriched_photos(db,pid)
  gallery_before.extend((pid,p['asset_key'],p.get('sha256'),p['url']) for p in photo_inventory.enriched_photos(S/'qualified_live/batch.sqlite3',pid) if p['identity_confirmed'] and p['kind']!='excluded')
  gallery_after.extend((pid,p['asset_key'],p.get('sha256'),p['url']) for p in photos if p['identity_confirmed'] and p['kind']!='excluded')
  checks.append({'id':pid,'brand':inp['brand'],'before':len(before),'after':len(rows),'core':counts['core'],'feature':counts['feature'],'UI':presented['count'],'conflicts':sum(bool((r.get('resolved') or {}).get('conflict')) for r in rows),'photos':sum(p['identity_confirmed'] and p['kind']!='excluded' for p in photos)})
assert len(inventory)==38 and len(checks)==103
assert sorted(gallery_before)==sorted(gallery_after),'Stage 82 changed confirmed gallery evidence'
book_bytes=exporter.export_batch(db,'cb81');(R/'Canonical_Stage82.xlsx').write_bytes(book_bytes)
book=load_workbook(BytesIO(book_bytes));assert book['Атрибуты'].max_row-1==sum(c['after'] for c in checks)
assert book['Исходные поля и mapping'].max_row-1==1643
assert book['Фотографии'].max_row-1==len(gallery_after)==1984
errors=[(s.title,c.coordinate) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
schema={category:[{**e,**{k:sorted(e[k]) for k in ('roles','canonical_keys','brands','models')}} for e in sorted(values.values(),key=lambda e:(e['parent_group'],e['attribute']))] for category,values in sorted(inventory.items())}
def save(name,data):(R/name).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
save('category_inventory_38.json',schema);save('native_output_checks.json',{'mode':'Native UI/Excel over offline parser qualification; not actual live success','inputs':103,'categories':38,'rows':checks,'duplicate_canonical_keys':0,'canonical_UI_Excel_rows':sum(c['after'] for c in checks),'raw_classification_Excel_rows':1643,'confirmed_gallery_unchanged':True,'confirmed_photos':len(gallery_after),'workbook_sha256':hashlib.sha256(book_bytes).hexdigest(),'formula_errors':errors,'sheet_names':book.sheetnames})
save('new_canonical_fields.json',{'added_occurrences':len(additions),'new_normalized_keys':len({a['key'] for a in additions}),'added':additions,'removed':removed})
book.close();print({'canonical':sum(c['after'] for c in checks),'added_occurrences':len(additions),'new_keys':len({a['key'] for a in additions}),'removed_occurrences':len(removed),'photo_evidence_unchanged':True},flush=True)
