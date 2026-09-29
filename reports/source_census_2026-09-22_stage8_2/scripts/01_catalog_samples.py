import json
from openpyxl import load_workbook

wb = load_workbook('data/catalog_2026-09-21_filtered.xlsx', read_only=True, data_only=True)
bc = wb['Бренд_Категория']
goods = wb['Товары']

# brand label (as it appears in the catalog) -> (source_family for Stage 8 profiles, 2 target categories)
PLAN = {
    'Samsung': {'family': 'samsung', 'categories': ['Телевизоры', 'Смартфоны']},
    'Apple': {'family': 'apple', 'categories': ['Смартфоны', 'Чехлы для телефонов']},
    'LG': {'family': 'lg', 'categories': ['Телевизоры', 'Стиральные машины']},
    'JBL': {'family': 'jbl', 'categories': ['Колонки', 'Наушники беспроводные']},
    'Playstation': {'family': 'playstation', 'categories': ['Диски с играми', 'Геймпады']},
    'Razer': {'family': 'razer', 'categories': ['Мыши', 'Наушники игровые']},
    'HIPER': {'family': 'hiper', 'categories': ['Наушники беспроводные', 'Внешние аккумуляторы']},
    'Microsoft': {'family': 'microsoft_xbox_check', 'categories': ['Геймпады', 'Игровые консоли']},
}

bc_rows = list(bc.iter_rows(min_row=2, values_only=True))
bc_index = {(r[0], r[1]): r for r in bc_rows}

# Build a lookup from seller article -> full 'Товары' row
goods_header = next(goods.iter_rows(min_row=1, max_row=1, values_only=True))
print('Товары header:', goods_header)
article_col = goods_header.index('Артикул продавца')
name_col = goods_header.index('Наименование')
cat_col = goods_header.index('Категория')
alt_col = goods_header.index('Альтернативные наименования') if 'Альтернативные наименования' in goods_header else None
wb_col = goods_header.index('Артикулы WB') if 'Артикулы WB' in goods_header else None

wanted_articles = {}
for brand, cfg in PLAN.items():
    for cat in cfg['categories']:
        row = bc_index.get((brand, cat))
        if not row:
            print('MISSING bc row for', brand, cat)
            continue
        examples = [x.strip() for x in (row[4] or '').split(';') if x.strip()]
        for art in examples[:3]:
            wanted_articles[art] = (brand, cat)

print('looking for', len(wanted_articles), 'articles')
found = {}
for r in goods.iter_rows(min_row=2, values_only=True):
    art = r[article_col]
    if art in wanted_articles and art not in found:
        found[art] = r

print('found', len(found), 'of', len(wanted_articles))
for art, (brand, cat) in wanted_articles.items():
    if art not in found:
        print('  NOT FOUND IN GOODS SHEET:', brand, cat, art)

out = []
for art, r in found.items():
    brand, cat = wanted_articles[art]
    out.append({
        'catalog_brand_label': brand,
        'catalog_category': cat,
        'seller_article': art,
        'wb_article': r[wb_col] if wb_col is not None else None,
        'name': r[name_col],
        'alt_name': r[alt_col] if alt_col is not None else None,
    })

with open('reports/source_census_2026-09-22_stage8_2/catalog_samples_raw.json', 'w', encoding='utf-8') as f:
    json.dump(out, f, indent=2, ensure_ascii=False)
print('wrote', len(out), 'rows')
for o in out:
    print(o['catalog_brand_label'], '|', o['catalog_category'], '|', o['seller_article'], '|', o['name'])
