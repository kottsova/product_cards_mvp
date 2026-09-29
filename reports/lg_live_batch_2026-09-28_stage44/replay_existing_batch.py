"""Offline replay of the existing P12ED row through the ordinary worker path."""
from __future__ import annotations

import gzip
import io
import json
import sqlite3
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from product_tool import exporter, jobs, worker
from product_tool.adapters.common import SourceDocument
from product_tool.adapters.lg_support import LGSupportAdapter

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
DB = ROOT / "data/batches.sqlite3"
COPY = OUT / "batch_replay.sqlite3"
PRODUCT_ID = 4
BATCH_ID = "ae3d2cb381744ba8a811e54231233254"
SAVED = next((OUT / "official_support").glob("001_*.gz"))


class Response:
    status_code = 200
    url = "https://www.lg.com/kz/support/product-support/cs-P12ED.USAR/"

    def __init__(self):
        with gzip.open(SAVED, "rt", encoding="utf-8") as handle:
            self.text = handle.read()
        self.content = self.text.encode()

    def raise_for_status(self):
        pass


class ReplayHttp:
    calls = 0

    def get(self, url, **_):
        self.calls += 1
        if url != Response.url:
            raise AssertionError(f"Unrecorded URL: {url}")
        return Response()


class EmptyOfficial:
    def __init__(self, key):
        self.source_key, self.site_name = key, key

    def find_source(self, *_args, **_kwargs):
        return SourceDocument(self.source_key, self.site_name, "", match_level="mismatch",
                              evidence="Existing product-page search result replayed offline.")


class EmptyDealer:
    source_key, site_name = "sulpak", "Sulpak"

    def find_source(self, *_args, **_kwargs):
        return SourceDocument(self.source_key, self.site_name,
                              "https://www.sulpak.kz/g/kondicioneriy_split_sistemiy_lg_p12ednsar___p12edusar",
                              match_level="unknown", error="policy_host_stopped",
                              evidence="Host stop was caused by another Sulpak page; this candidate was not fetched.")


class EmptyDns:
    source_key, site_name = "dns", "DNS"
    document_urls = {}

    def find_source(self, *_args, **_kwargs):
        return SourceDocument(self.source_key, self.site_name, "", match_level="dealer_url_needed")

    def find_documents(self, *_args, **_kwargs):
        return [], "No verified exact URL."


def main():
    with sqlite3.connect(DB) as source, sqlite3.connect(COPY) as target:
        source.backup(target)
    before = jobs.get_product(COPY, PRODUCT_ID)
    http = ReplayHttp()
    job_id = jobs.enqueue(COPY, PRODUCT_ID, [1, 2, 3, 4, 6])
    worker.run_once(
        COPY,
        adapter_factory=lambda: (EmptyOfficial("lg_kz"), EmptyOfficial("lg_ru"), EmptyDealer()),
        dns_adapter_factory=EmptyDns,
        lg_support_adapter_factory=lambda: LGSupportAdapter(http=http),
    )
    with sqlite3.connect(COPY) as connection:
        job = connection.execute("SELECT status,message FROM search_jobs WHERE id=?", (job_id,)).fetchone()
        product_count = connection.execute("SELECT count(*) FROM products WHERE batch_id=?", (BATCH_ID,)).fetchone()[0]
    sources = jobs.get_source_pages(COPY, PRODUCT_ID)
    binary = exporter.export_batch(COPY, BATCH_ID)
    workbook = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
    export_hits = {}
    for sheet in workbook:
        hits = []
        for row in sheet.iter_rows(values_only=True):
            line = " | ".join(str(value) for value in row if value is not None)
            if "P12ED" in line:
                hits.append(line[:1500])
        if hits:
            export_hits[sheet.title] = hits[:5]
    result = {
        "existing_product_id": before["id"],
        "products_in_batch": product_count,
        "replayed_job_id": job_id,
        "job_status": job[0],
        "job_message": job[1],
        "support_http_calls": http.calls,
        "support_source": next((x for x in sources if x["source_key"] == "lg_kz_support"), None),
        "documents": jobs.get_documents(COPY, PRODUCT_ID),
        "export_hits": export_hits,
    }
    result["support_source"].pop("id", None)
    (OUT / "replay_result.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in result.items() if key != "export_hits"}, ensure_ascii=False, indent=2)[:5000])


if __name__ == "__main__":
    main()
