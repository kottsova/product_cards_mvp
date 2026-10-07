from pathlib import Path
from types import SimpleNamespace
from shutil import copyfile
from io import BytesIO
import json,re
from fastapi.testclient import TestClient
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from product_tool.web import create_app
from product_tool import jobs,storage,card_evidence,jbl_pipeline,attribute_projection,photo_metadata
r=Path('reports/jbl_stage63');ui=r/'ui_batch';ui.mkdir(exist_ok=True);p=ui/'batches.sqlite3'
if not p.exists():copyfile(r/'release_pass.sqlite3',p)
checks=[];measurements=[]
# Inspect saved actual image URLs through the normal app route; blocked images stay unmeasured.
with TestClient(create_app(ui,start_worker=False)) as client:
 for product in storage.get_batch(p,'stage62')['products']:
  pid=product['id'];response=client.get(f'/products/{pid}');html=response.text;rows=attribute_projection.final_attribute_rows(p,pid);ev=card_evidence.load(p,pid,'jbl') or {};facts=jobs.get_facts(p,pid)
  assert response.status_code==200 and 'Проверка JBL' in html
  assert 'Характеристики-кандидаты' in html and 'Документы по типам' in html
  assert all(f['section'] and f['raw_name'] and f['raw_value'] for f in facts)
  assert not any(re.search(r'FAQ|firmware|troubleshoot|warranty|reset instructions',f['raw_name'],re.I) for f in facts)
  labels=[row['display_name'] for row in rows]
  if product['search_code']=='JBLLIVE770NCBLK':assert any('без ANC' in x for x in labels) and any('с ANC' in x for x in labels)
  if product['search_code']=='JBLTBUDSBLK':assert 'Вес зарядного кейса, г' in labels and 'Вес наушников (один или пара не уточнено), г' in labels
  checks.append({'article':product['search_code'],'http':response.status_code,'labels':labels,'readiness':jbl_pipeline.card_readiness(p,pid),'raw_preserved':True,'support_fields_leaked':False})
  if product['search_code'] in {'JBLXTREME4BLUEP','JBLT520BTBLKEU','JBLBAR500PROBLKEP','JBLT520BTWHTEU'}:
   photos=jobs.get_photo_candidates(p,pid)
   if photos:
    photo=photos[0];z=SimpleNamespace(status_code=200) if photo.get('verified_width') else client.post(f"/products/{pid}/photos/{photo['id']}/inspect");saved=jobs.get_photo_candidates(p,pid)[0];measurements.append({'article':product['search_code'],'url':photo['url'],'http':z.status_code,'width':saved.get('verified_width'),'height':saved.get('verified_height'),'bytes':saved.get('verified_bytes'),'format':saved.get('verified_format')})
 z=client.get('/batches/stage62/export.xlsx');assert z.status_code==200;(r/'export.xlsx').write_bytes(z.content);book=load_workbook(BytesIO(z.content),read_only=True,data_only=True)
 sheets={s.title:{'rows':s.max_row,'columns':s.max_column,'headers':list(next(s.values))} for s in book}
 assert book['Готовность JBL'].max_row==12
 readiness=list(book['Готовность JBL'].values)[1:];assert sum(x[2]=='Готова' for x in readiness)==11;assert sum(x[2]=='Готова с пробелами' for x in readiness)==0;assert sum(x[2]=='Не готова' for x in readiness)==0
 errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
 result={'ui':checks,'photos_measured':measurements,'xlsx_bytes':len(z.content),'sheets':sheets,'formula_errors':errors,'scope':'Production exporter; static values. UI HTTP inspection plus actual saved-photo metadata route; independent renders follow.'};(r/'qa_ui_excel.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps({'rows':len(checks),'sheets':{k:v['rows'] for k,v in sheets.items()},'photos':measurements},ensure_ascii=True),flush=True);book.close()
