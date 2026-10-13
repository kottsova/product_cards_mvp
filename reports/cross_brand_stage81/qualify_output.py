"""Read-only audit of actual observations and the current native output layers."""
import sys,json,re,hashlib,sqlite3
from pathlib import Path
from collections import Counter,defaultdict
from contextlib import closing
from io import BytesIO
sys.path.insert(0,str(Path(__file__).resolve().parents[2]));sys.stdout.reconfigure(encoding='utf8')
from product_tool import jobs,storage,attribute_projection,card_presentation,card_evidence,canonical_attributes,photo_inventory,readiness,exporter
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from product_tool.web import create_app
R=Path(__file__).parent;phase=sys.argv[1] if len(sys.argv)>1 else 'final_live';root=R/phase
Q=R/(phase+'_audit');Q.mkdir(exist_ok=True)
U=Q/'ui';U.mkdir(exist_ok=True);db=U/'batches.sqlite3'
with closing(sqlite3.connect(root/'batch.sqlite3')) as source,closing(sqlite3.connect(db)) as target:source.backup(target)
dataset=json.loads((R/'dataset.json').read_text(encoding='utf8'))['rows']
observations={int(p.parent.name):json.loads(p.read_text(encoding='utf8')) for p in root.glob('*/result.json')}
def norm(v):return re.sub(r'\s+',' ',str(v or '')).strip().casefold()
def label(row):return row.get('raw_label',row.get('label',row.get('name','')))
def value(row):return row.get('raw_value',row.get('value',''))
def classification(reason):
 for pattern,kind in ((r'duplicat|alias|equivalent','duplicate'),(r'malformed|invalid|empty','malformed'),(r'metadata|internal|logistic|heading','logistics/internal'),(r'not.*map|unsupported|unmapped|requires.canonical.mapping|lacks.safe.canonical.mapping','unsupported canonical mapping'),(r'marketing|promotional|prose','marketing'),(r'variant|conditional|region|configuration|option|case size|material','variant-sensitive'),(r'unproven|unresolved|not proven|identity|not confirmed','unknown identity')):
  if re.search(pattern,reason,re.I):return kind
 return 'other'
def lists(node,where='evidence'):
 if isinstance(node,dict):
  for key,items in node.items():
   if isinstance(items,list) and key in {'raw_specs','raw_facts','accepted_specs','accepted_facts','admitted','rejected_specs','rejected_facts','candidates','configuration_candidates'}:yield where+'.'+key,items
   elif isinstance(items,dict):yield from lists(items,where+'.'+key)
def imperial_duplicate(f,facts):
 units={'in':(25.4,{'mm','cm'}),'lb':(0.45359237,{'kg','kgs','g'}),'lbs':(0.45359237,{'kg','kgs','g'}),'oz':(28.349523125,{'g','kg'}),'ft':(0.3048,{'m','cm'})}
 match=re.search(r'\((in|lbs?|oz|ft)\)$',f['raw_name'],re.I)
 if not match:return None
 base=norm(re.sub(r'\([^()]+\)$','',f['raw_name']));factor,metric_units=units[match[1].lower()]
 nums=lambda s:[float(n.replace(',','.')) for n in re.findall(r'\d+(?:[.,]\d+)?',s)]
 a=nums(f['raw_value'])
 for other in facts:
  suffix=re.search(r'\((mm|cm|kgs?|g|m)\)$',other['raw_name'],re.I)
  if not suffix or suffix[1].lower() not in metric_units or norm(re.sub(r'\([^()]+\)$','',other['raw_name']))!=base:continue
  b=nums(other['raw_value']);unit=suffix[1].lower();scale=factor
  if match[1].lower()=='in' and unit=='cm':scale/=10
  if match[1].lower() in {'lb','lbs'} and unit=='g':scale*=1000
  if match[1].lower()=='oz' and unit=='kg':scale/=1000
  if match[1].lower()=='ft' and unit=='cm':scale*=100
  if len(a)==len(b) and a and all(abs(x*scale-y)<=max(abs(y)*.035,.2) for x,y in zip(a,b)):
   return {'counterpart_fact_id':other['id'],'counterpart_label':other['raw_name'],'counterpart_value':other['raw_value'],'conversion_factor':scale,'tolerance':'published rounded rendition, 3.5% or 0.2 metric unit'}
 return None
rows=[];details={};schema=defaultdict(dict);gallery=[];colors=[];ui_checks=[]
with TestClient(create_app(U,start_worker=False)) as client:
 for inp in dataset:
  pid=inp['id']
  if pid not in observations:continue
  obs=observations[pid];facts=jobs.get_facts(db,pid);sources=jobs.get_source_pages(db,pid);product=jobs.get_product(db,pid)
  ev=card_evidence.load(db,pid,photo_inventory.evidence_kind(product)) or {}
  canonical=attribute_projection.final_attribute_rows(db,pid);presented=card_presentation.present_card_rows(canonical,facts,sources,inp['category'])
  ui_rows=[r for groups in presented['groups'].values() for section in groups.values() for r in section]
  lineage=defaultdict(list)
  for r in canonical:
   for key in {r['normalized_name'],r.get('derived_from'),r.get('refined_from'),*r.get('derived_from_names',[])}-{None}:lineage[key].append(r['normalized_name'])
  missing=[]
  for f in facts:
   if lineage.get(f['normalized_name']):continue
   duplicate=imperial_duplicate(f,facts)
   metadata=bool(re.fullmatch(r'safety and product resources|bar.?code|ean|штрих.?код',f['raw_name'],re.I))
   unknown=f['normalized_value']=='unknown' or f['unit']=='unknown'
   reason='duplicate' if duplicate else 'logistics/internal' if metadata else 'resolver rejection' if unknown else 'parser loss'
   missing.append({**f,'classification':reason,'explanation':'Numerically checked metric counterpart' if duplicate else 'Support/identity metadata retained in source audit' if metadata else 'Missing physical value, not boolean absence' if unknown else 'Accepted fact has no output lineage','counterpart':duplicate})
  raw=[];admitted=[];rejected=[]
  for location,items in lists(ev):
   for item in items:
    if not isinstance(item,dict) or not label(item):continue
    record={**item,'location':location,'label':label(item),'value':value(item),'source':item.get('source_url') or item.get('url') or ev.get('pdp_url') or ev.get('source_url') or ev.get('model_evidence',{}).get('url') or next((s['url'] for s in sources if not s['error']),'')}
    if location.startswith('evidence.model_evidence.'):record['source']=ev['model_evidence']['url']
    if location.endswith(('.raw_specs','.raw_facts')):raw.append(record)
    elif location.endswith(('.accepted_specs','.accepted_facts','.admitted')):admitted.append(record)
    elif 'reason' in record:rejected.append(record)
  if not raw:
   for file in (root/str(pid)/'raw_sources').glob('*.json'):
    source=json.loads(file.read_text(encoding='utf8'))
    raw.extend({**r,'source':source['url'],'location':'returned live SourceDocument'} for r in source['raw_specs'])
  unique={}
  for r in raw:unique.setdefault((r['source'],norm(r.get('source_section',r.get('section',''))),norm(r.get('subsection',r.get('component',''))),norm(r['label']),norm(r['value'])),r)
  classified=[]
  for r in unique.values():
   match=lambda z:norm(z['label'])==norm(r['label']) and norm(z['value'])==norm(r['value'])
   a=[z for z in admitted if match(z)];reject=[z for z in rejected if match(z)]
   related=[f for f in facts if norm(f['raw_name'])==norm(r['label']) and norm(f['raw_value'])==norm(r['value'])]
   if inp['brand']=='Apple':
    source_key='apple_model' if r['location'].startswith('evidence.model_evidence.') else 'apple'
    # The source parser deliberately aggregates paragraphs in their own section.
    # A raw paragraph retained verbatim in that admitted section is not lost.
    related.extend(f for f in facts if f['source_key']==source_key and norm(f['section'])==norm(r.get('section')) and len(norm(r['value']))>=4 and norm(r['value']) in norm(f['raw_value']))
   for item in a:
    labels={item.get('label',''),item.get('name',''),item.get('accepted_label',''),*item.get('canonical_fields',[])}
    related.extend(f for f in facts if f['raw_name'] in labels and norm(f['raw_value'])==norm(item['value']))
    for derived in item.get('derived_fields',[]):related.extend(f for f in facts if norm(f['raw_name'])==norm(derived.get('name')) and norm(f['raw_value'])==norm(derived.get('value')))
   output=sorted({n for f in related for n in lineage.get(f['normalized_name'],[])})
   if output:state='accepted';reason='Source fact is present or explicitly decomposed/derived'
   elif reject:
    reason='; '.join(dict.fromkeys(z.get('reason','candidate') for z in reject));kind=classification(reason)
    state='candidate' if kind in {'variant-sensitive','unknown identity','unsupported canonical mapping'} else 'rejected'
   elif related:
    reason='; '.join(dict.fromkeys(m['classification'] for m in missing if m['id'] in {f['id'] for f in related}));state='rejected';kind=reason or 'resolver rejection'
   elif re.fullmatch(r'(?:product )?(?:name|model(?: number)?)|station model|omni station model',r['label'],re.I):state='rejected';reason='Model/hardware identity, preserved in identity inventory';kind='logistics/internal'
   else:state='candidate';reason='Raw field retained; mapping lineage needs further qualification';kind='unsupported canonical mapping'
   classified.append({**r,'classification':'accepted' if state=='accepted' else kind,'state':state,'reason':reason,'fact_ids':sorted({f['id'] for f in related}),'canonical_names':output,'partial_candidates':[z for z in reject if z.get('admitted_model_fields')]})
  photos=photo_inventory.enriched_photos(db,pid);valid=[p for p in photos if p['kind']!='excluded'];confirmed=[p for p in valid if p['identity_confirmed']]
  color=canonical_attributes.color_result(product,canonical,sources,ev);colors.append({'id':pid,'brand':inp['brand'],'model':inp['model'],'input_color':inp['input_color'],**color})
  for role,groups in presented['groups'].items():
   for group,items in groups.items():
    for r in items:
     key=(group,r['display_name']);entry=schema[inp['category']].setdefault(key,{'parent_group':group,'attribute':r['display_name'],'roles':set(),'canonical_keys':set(),'brands':set(),'models':set()})
     entry['roles'].add(role);entry['canonical_keys'].add(r['normalized_name']);entry['brands'].add(inp['brand']);entry['models'].add(pid)
  response=client.get('/products/'+str(pid));assert response.status_code==200
  soup=BeautifulSoup(response.text,'html.parser');visible=soup.select('.card-attributes tbody > tr');tiles=soup.select('.photo-tile')
  assert len(visible)==len(ui_rows),(pid,'UI attribute loss',len(visible),len(ui_rows))
  assert len(tiles)==len(valid),(pid,'UI gallery loss',len(tiles),len(valid))
  assert soup.select_one('#attribute-search') and soup.select_one('#attribute-role')
  (U/f'product_{pid}.html').write_text(response.text,encoding='utf8')
  counts=Counter(r['state'] for r in classified)
  row={'id':pid,'brand':inp['brand'],'category':inp['category'],'model':inp['model'],'color':color['value'],'live_official_data':any(s['url'] and not s['error'] and s['source_key'] not in {'dns','sulpak'} for s in sources) and bool(facts),'source_observation':obs.get('observed_at'),'runtime_seconds':obs['seconds'],'raw_spec_rows':len(raw),'unique_raw_rows':len(classified),'accepted_raw':counts['accepted'],'candidate_raw':counts['candidate'],'rejected_raw':counts['rejected'],'stored_facts':len(facts),'canonical_rows':len(canonical),'ui_rows':len(ui_rows),'useful_projection_loss':sum(m['classification']=='parser loss' for m in missing),'duplicate_output_keys':[key for key,n in Counter(r['normalized_name'] for r in canonical).items() if n>1],'all_photos':len(valid),'confirmed_photos':len(confirmed),'selected_photos':sum(bool(p['selected']) for p in confirmed),'byte_verified_photos':sum(p['byte_verified'] for p in confirmed),'readiness':readiness.card_readiness(db,pid,manuals_required=False)}
  assert len(classified)==sum(counts.values());rows.append(row)
  details[str(pid)]={'raw_rows':classified,'accepted_source_fact_omissions':missing,'rejected_records':[{**r,'classification':classification(r.get('reason',''))} for r in rejected],'fact_output_lineage':[{'fact_id':f['id'],'canonical_name':f['normalized_name'],'output':lineage.get(f['normalized_name'],[])} for f in facts]}
  gallery.append({'id':pid,'brand':inp['brand'],'model':inp['model'],'native_media_inventory':ev.get('native_media_inventory',ev.get('gallery_audit',[])),'raw_saved_rows':len(jobs.get_photo_candidates(db,pid,include_excluded=True)),'assets':photos,'completion':(card_evidence.load(db,pid,'photo_inventory') or {}).get('last_completion')})
  ui_checks.append({'id':pid,'attributes':len(visible),'gallery_tiles':len(tiles),'status':response.status_code})
inventory={category:[{**entry,**{k:sorted(entry[k]) for k in ('roles','canonical_keys','brands','models')}} for entry in sorted(entries.values(),key=lambda e:(e['parent_group'],e['attribute']))] for category,entries in sorted(schema.items())}
def save(name,payload):(Q/name).write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
save('loss_report.json',{'observed_inputs':len(rows),'rows':rows,'details':details,'method':'Source rows deduplicated by provenance and raw content. Mixed parent rows retain candidate parts; canonical decomposition is counted separately. No replay is counted as live.'})
save('category_inventory.json',inventory);save('gallery_report.json',gallery);save('color_resolution.json',colors)
if len(rows)==103:
 workbook=exporter.export_batch(db,'cb81');(R/'CrossBrand_Stage81.xlsx').write_bytes(workbook)
 book=load_workbook(BytesIO(workbook));errors=[(s.title,c.coordinate) for s in book for row in s for c in row if c.data_type=='e'];assert not errors
 attrs=book['Атрибуты'];assert attrs.max_row-1==sum(r['ui_rows'] for r in rows)
 photo_sheet=book['Фотографии'];assert photo_sheet.max_row-1==sum(r['confirmed_photos'] for r in rows)
 save('native_ui_excel_qa.json',{'ui':ui_checks,'sheets':{s.title:{'rows':s.max_row,'columns':s.max_column,'headers':[c.value for c in s[1]]} for s in book},'workbook_sha256':hashlib.sha256(workbook).hexdigest(),'formula_errors':errors,'attribute_rows':attrs.max_row-1,'confirmed_photo_rows':photo_sheet.max_row-1})
 book.close()
print({'observed':len(rows),'categories':len(inventory),'live_official_data':sum(r['live_official_data'] for r in rows),'raw_unique':sum(r['unique_raw_rows'] for r in rows),'canonical':sum(r['canonical_rows'] for r in rows),'projection_loss':sum(r['useful_projection_loss'] for r in rows),'photos':sum(r['confirmed_photos'] for r in rows),'byte_verified':sum(r['byte_verified_photos'] for r in rows)},flush=True)
