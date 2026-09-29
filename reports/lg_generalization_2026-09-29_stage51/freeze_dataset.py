"""Freeze Stage 51 LG catalog rows before inspecting pilot outcomes."""

from __future__ import annotations

import hashlib
import json
import sqlite3
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
CATALOG = ROOT / "data" / "catalog_2026-09-21.xlsx"
DATABASE = ROOT / "data" / "batches.sqlite3"
OUTPUT = Path(__file__).with_name("dataset.json")
CURRENT_BATCH = "ae3d2cb381744ba8a811e54231233254"

# Deliberately spans easy and difficult formats; no observed outcome informed selection.
SELECTION = [
    ("F2Y1HS5W.AGWPCOM", "Стиральная машина: полный региональный код с точкой."),
    ("W4W8LVPK4HM", "WashTower: базовый код без суффикса, иной вариант, чем в текущей партии."),
    ("DC90V9V9W.ABWPCOM_KZ", "Сушильная машина: полный код с каталожным суффиксом _KZ."),
    ("DC90V5V0W", "Сушильная машина: базовый код для сравнения с полным артикулом."),
    ("GC-B399SQCL.ASWQCIS", "Холодильник: полный региональный вариант."),
    ("GC-B509MLWM.ADSQCIS", "Холодильник: похожая семейная линейка с иным размером и точным цветовым вариантом."),
    ("32LQ63006LA.ARUG", "Телевизор: полный RU артикул с точкой."),
    ("43NANO81A6A.ARUG", "Телевизор: соседний размер к прежней модели; проверка запрета переноса фото."),
    ("GXDCISLLK", "Саундбар: слитный каталожный код против формата модели и суффикса."),
    ("CL87.DCISLLK", "Аудиосистема: полный код с точкой и региональным вариантом."),
    ("VC5316NNTS.APSQCIS", "Пылесос: полный код с точкой."),
    ("A9K_MAX1", "Вертикальный пылесос: подчёркивание в артикуле, возможное отличие формата сайта."),
    ("MS2044V.BSSQCIS", "Микроволновая печь: полный код с точкой."),
    ("AC09BK", "Сплит-система: базовый код; проверка отличия от комплекта P12ED."),
]


def main() -> None:
    if OUTPUT.exists():
        raise SystemExit(f"Dataset already frozen: {OUTPUT}")
    sheet = load_workbook(CATALOG, read_only=True, data_only=True)["Товары"]
    catalog_rows = {str(row[2]).strip(): row for row in sheet.iter_rows(min_row=2, values_only=True)
                    if row[0] == "LG" and row[2]}
    with sqlite3.connect(DATABASE) as connection:
        old = {row[0] for row in connection.execute(
            "SELECT search_code FROM products WHERE batch_id=?", (CURRENT_BATCH,))}
    rows = []
    for code, reason in SELECTION:
        if code in old:
            raise ValueError(f"Already in current batch: {code}")
        catalog = catalog_rows.get(code)
        if catalog is None:
            raise ValueError(f"Not in catalog: {code}")
        rows.append({"category": catalog[1], "article": code,
                     "catalog_name": catalog[4], "reason": reason})
    if len(rows) != len({r["article"] for r in rows}):
        raise ValueError("Duplicate articles")
    payload = {
        "stage": 51,
        "frozen_at_utc": datetime.now(timezone.utc).isoformat(),
        "catalog_sha256": hashlib.sha256(CATALOG.read_bytes()).hexdigest(),
        "excluded_current_batch": CURRENT_BATCH,
        "sampling_note": "Only one kit is in the LG catalog, and it belongs to the excluded current batch.",
        "rows": rows,
    }
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Frozen {len(rows)} rows in {OUTPUT}")


if __name__ == "__main__":
    main()
