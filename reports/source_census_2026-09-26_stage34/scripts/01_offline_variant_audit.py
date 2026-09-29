"""Offline exact-code and regional-link audit before any Stage 34 network call."""
from __future__ import annotations
import gzip,hashlib,json,re,sqlite3,sys
from datetime import datetime,timezone
from pathlib import Path
from bs4 import BeautifulSoup
from openpyxl import load_workbook
HERE=Path(__file__).resolve().parent;ROOT=HERE.parents[2];STAGE=HERE.parent
sys.path.insert(0,str(ROOT))
from product_tool.adapters.structured_page import extract_json_ld_product
from product_tool.offline_guard import offline_only
S33=ROOT/'reports/source_census_2026-09-26_stage33'
TARGETS={'MMB2111M':3051,'TWK7203':3702}
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
 with offline_only():
  wb=load_workbook(ROOT/'data/catalog_2026-09-21_filtered.xlsx',read_only=True,data_only=True)
  sheet=wb['Товары'];catalog={}
  for rownum in TARGETS.values():
   row=next(sheet.iter_rows(min_row=rownum,max_row=rownum,values_only=True))
   catalog[row[2]]={'catalog_row':rownum,'brand':row[0],'category':row[1],'seller_article':row[2],'marketplace_articles':row[3],'title':row[4],'alternate_title':row[5],'extra_region_or_enr_fields':False}
  wb.close()
  prev=json.loads((S33/'raw/bosch_route_declaration.json').read_text(encoding='utf-8'))
  index=[json.loads(x) for x in (S33/'bosch_route/responses/index.jsonl').read_text(encoding='utf-8').splitlines()]
  pages={}
  for sku,route in prev['first_batch'].items():
   entry=next(x for x in index if x['url']==route['url'])
   with gzip.open(S33/'bosch_route/responses'/entry['saved_as'],'rt',encoding='utf-8') as h:html=h.read()
   decoded=html.replace('\\"','"')
   fields=extract_json_ld_product(html,route['url'])
   identity={f.name:f.value for f in fields if f.name in ('name','mpn','gtin','sku','productID')}
   current_product_ids=[m.group(1) for m in re.finditer(r'"tracking":\{"productId":"([^"]+)"',decoded)]
   revision_codes=sorted(set(re.findall(re.escape(sku)+r'/\d{2}(?!\d)',decoded,re.I)))
   enumber_ui_labels=re.findall(r'"eNumber":"([^"]+)"',decoded)
   country_switch=re.findall(r'"countrySwitchLink":\{[^}]*?"href":"([^"]+)"',decoded)
   pages[sku]={'page_url':route['url'],'page_market':route['market'],'catalog':catalog[sku],'identity':identity,'current_product_ids':current_product_ids,'product_specific_enr_revisions':revision_codes,'enumber_ui_labels':enumber_ui_labels,'enumber_examples_are_generic':sorted(set(re.findall(r'WM\d{7}/\d{2}',decoded))),'country_switch_links':country_switch,'model_code_present':identity.get('mpn')==sku,'gtin_in_catalog':False,'variant_roadblock':'catalog lacks market/GTIN/E-Nr and other-market page supplies no product-specific E-Nr revision or KZ identity'}
   assert identity.get('mpn')==sku and not revision_codes and all(label in ('E-Number','E-Nummer') for label in enumber_ui_labels)
  c=sqlite3.connect(ROOT/'reports/source_census_2026-09-22_stage6/source_snapshots.sqlite3')
  html=c.execute("select content from source_snapshots where source_url=?",('https://www.bosch-home.com/',)).fetchone()[0];c.close()
  soup=BeautifulSoup(html,'html.parser')
  kz=[{'url':a['href'],'label':a.get_text(' ',strip=True)} for a in soup.select('a[href]') if a['href']=='https://www.bosch-home.com/kz/']
  assert len(kz)==1 and kz[0]['label']=='Kazakhstan'
  observed={'source':'reports/source_census_2026-09-22_stage6/source_snapshots.sqlite3 source_snapshots id=1','root':'https://www.bosch-home.com/','kz_country_link':kz[0],'de_product_country_switch':pages['TWK7203']['country_switch_links'],'mt_product_country_switch':pages['MMB2111M']['country_switch_links']}
  out={'products':pages,'regional_evidence':observed,'instruction_family':'Stage 33: full 76-page multilingual PDF, Russian operating section pages 64-67, TWK720. printed, TWK7203 not printed'}
  (STAGE/'raw/offline_variant_audit.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
  declaration={'declared_at':datetime.now(timezone.utc).isoformat(timespec='seconds'),'purpose':'Two Bosch Home catalog models only; observed KZ country link, then at most two newly observed regional navigation URLs; no guessed product URL','observed_start_url':kz[0]['url'],'allowed_hosts':['www.bosch-home.com'],'budget':{'max_real_requests_total':3,'max_per_host':3,'pacing_seconds':1.5,'page_cap_bytes':25000000,'dealer_requests':0,'documents_requests':0,'images_requests':0,'stop':'401/403/429 or confirmed challenge stops host and remaining requests'},'catalog_sha256_before':sha(ROOT/'data/catalog_2026-09-21_filtered.xlsx'),'registry_sha256_before':sha(ROOT/'product_tool/config/source_catalog.v2.json'),'db_sha256_before':sha(ROOT/'data/batches.sqlite3'),'planner_sha256_before':sha(ROOT/'product_tool/config/coverage_planner.v1.json')}
  (STAGE/'raw/regional_declaration.json').write_text(json.dumps(declaration,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
  print(json.dumps({'codes':{k:{'catalog':v['catalog'],'identity':v['identity'],'tracking_product_ids':v['current_product_ids'],'enr_revision':v['product_specific_enr_revisions'],'e_number_ui_labels':v['enumber_ui_labels'],'generic_enr_examples':v['enumber_examples_are_generic'],'country_switch':v['country_switch_links']} for k,v in pages.items()},'observed_kz':kz,'declaration':declaration['budget']},ensure_ascii=True,indent=2))
if __name__=='__main__':main()
