"""Build a two-row LG trial book from the unchanged catalog."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parents[1] / "lg_first_test.xlsx"
CATALOG = ROOT / "data/catalog_2026-09-21_filtered.xlsx"
WANTED = {"XL7S", "MS2032GAS"}


def main():
    source = load_workbook(CATALOG, read_only=True, data_only=True)
    try:
        rows = iter(source["Товары"].values)
        header = next(rows)
        names = {name: index for index, name in enumerate(header)}
        selected = []
        for row in rows:
            code = str(row[names["Артикул продавца"]] or "").strip().upper()
            if code in WANTED and str(row[names["Бренд"]]).strip().upper() == "LG":
                selected.append((row[names["Категория"]], row[names["Бренд"]],
                                 row[names["Наименование"]], row[names["Артикул продавца"]]))
        assert len(selected) == 2 and {item[3] for item in selected} == WANTED
    finally:
        source.close()
    book = Workbook()
    sheet = book.active
    sheet.title = "Товары"
    sheet.append(("Категория", "Бренд", "Наименование", "Артикул продавца"))
    for row in selected:
        sheet.append(row)
    book.save(OUT)
    book.close()
    print(OUT)


if __name__ == "__main__":
    main()
