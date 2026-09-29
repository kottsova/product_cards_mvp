"""Read-only Stage 50 audit of the existing LG batch and Stage 49 workbook."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from product_tool import attribute_projection, jobs, readiness, storage
from product_tool.lg_identity import document_tied_to_article, photo_tied_to_article
from product_tool.web import create_app


ROOT = Path(__file__).resolve().parents[2]
DATABASE = ROOT / "data" / "batches.sqlite3"
BATCH_ID = "ae3d2cb381744ba8a811e54231233254"
WORKBOOK = ROOT / "reports" / "lg_live_batch_2026-09-29_stage49" / "stage49.xlsx"


def main() -> None:
    batch = storage.get_batch(DATABASE, BATCH_ID)
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    product_sheets = [sheet for sheet in book if sheet.title not in {
        "Проверка источников", "Источники", "Инструкции", "Фотографии", "Готовность LG",
    }]
    xlsx_rows = {}
    for sheet in product_sheets:
        values = sheet.values
        header = next(values)
        for row in values:
            xlsx_rows[str(row[3])] = dict(zip(header, row))
    ready_rows = list(book["Готовность LG"].values)
    photo_rows = list(book["Фотографии"].values)
    app = create_app(ROOT / "data", start_worker=False)
    result = {"batch_file": batch["filename"], "products": [], "workbook": {
        "sheets": book.sheetnames,
        "product_rows": len(xlsx_rows),
        "ready_rows": len(ready_rows) - 1,
        "photo_statuses": dict(Counter(str(row[-1]) for row in photo_rows[1:])),
    }}
    with TestClient(app) as client:
        for product in batch["products"]:
            code = product["search_code"]
            if product["brand"].strip().upper() != "LG":
                continue
            pid = product["id"]
            sources = jobs.get_source_pages(DATABASE, pid)
            read = readiness.card_readiness(DATABASE, pid)
            resolved = jobs.get_resolved(DATABASE, pid)
            docs = jobs.get_documents(DATABASE, pid)
            photos = jobs.get_photo_candidates(DATABASE, pid, include_excluded=False)
            latest = jobs.list_jobs(DATABASE, pid)[0]
            html = client.get(f"/products/{pid}")
            soup = BeautifulSoup(html.text, "html.parser")
            card = soup.select_one("#lg-readiness")
            user_reasons = [item.get_text(" ", strip=True) for item in card.select("li")]
            xlsx = xlsx_rows.get(code)
            final = attribute_projection.final_attribute_rows(DATABASE, pid)
            conflicts = [{"name": r["normalized_name"], "reason": r["reason"], "value": r["selected_value"]}
                         for r in resolved if r["conflict"]]
            item = {
                "id": pid, "article": code, "name": product["name"], "category": product["category"],
                "job": latest["status"], "verdict": read["verdict"],
                "confirmed_specs": sum(bool(r["selected_value"] and r["full_sku_confirmed"] and not r["conflict"])
                                       for r in resolved),
                "confirmed_photos": sum(bool(p["selected"] and photo_tied_to_article(p, sources)) for p in photos),
                "selected_unconfirmed_photos": sum(bool(p["selected"] and not photo_tied_to_article(p, sources)) for p in photos),
                "manuals": [{"title": d["title"], "language": d["language"],
                             "tied": document_tied_to_article(code, d, sources)} for d in docs],
                "gaps": read["blocking_gaps"], "conflicts": conflicts,
                "sources": [{"key": s["source_key"], "match": s["match_level"], "url": s["url"],
                             "error": s["error"]} for s in sources],
                "ui_status": html.status_code, "ui_has_article": code in soup.get_text(" ", strip=True),
                "ui_reasons": user_reasons,
                "ui_candidate_count": len(soup.select("#photos span")) and soup.select_one("#photos").get_text(" ", strip=True).count("кандидат:"),
                "xlsx_found": xlsx is not None,
                "xlsx_nonempty_specs": sum(bool(v) for k, v in (xlsx or {}).items() if k not in
                                           {"Строка", "Название", "Бренд", "Полный артикул", "Базовая модель"}),
                "xlsx_marker_values": [str(v) for v in (xlsx or {}).values() if str(v).strip() in {"●", "•", "○", "-", "—", "O"}],
                "final_names_with_value": [r["display_name"] for r in final if (r.get("resolved") or {}).get("display_value")],
            }
            result["products"].append(item)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
