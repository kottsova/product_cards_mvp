"""Read-only UI and Excel audit for the frozen Stage 56 batch."""
from __future__ import annotations

import json
import re
import shutil
import sys
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT))

from product_tool import attribute_projection, card_presentation, exporter, jobs, manual_status, photo_metadata
from product_tool.web import create_app

HERE=Path(__file__).resolve().parent
DB=ROOT/"data/stage56/stage56.sqlite3"
OUT=HERE/(sys.argv[1] if len(sys.argv)>1 else "ui_excel_audit.json")

def main():
    if OUT.exists(): raise SystemExit("UI audit already exists")
    first=json.loads((HERE/"first_pass.json").read_text(encoding="utf-8"))
    records=[]
    with TemporaryDirectory() as temp:
        copy=Path(temp)/"batches.sqlite3"
        shutil.copy2(DB,copy)
        with TestClient(create_app(temp,start_worker=False)) as client:
            for item in first["rows"]:
                pid=item["product_id"]
                product=jobs.get_product(DB,pid)
                sources=jobs.get_source_pages(DB,pid)
                facts=jobs.get_facts(DB,pid)
                rows=attribute_projection.final_attribute_rows(DB,pid)
                presented=card_presentation.present_card_rows(rows,facts,sources,product["category"])
                flattened=[r for sections in presented["groups"].values() for group in sections.values() for r in group]
                response=client.get(f"/products/{pid}")
                soup=BeautifulSoup(response.text,"html.parser")
                docs=jobs.get_documents(DB,pid)
                photos=jobs.get_photo_candidates(DB,pid,include_excluded=False)
                tiles=soup.select(".photo-open")
                tile_mismatches=[]
                by_url={x["url"]:x for x in photos}
                for tile in tiles:
                    photo=by_url.get(tile.get("data-url"))
                    if photo is None:
                        tile_mismatches.append({"url":tile.get("data-url"),"reason":"not in DB"})
                        continue
                    expected=(f'{photo["verified_width"]} × {photo["verified_height"]} px'
                              if photo["verified_width"] and photo["verified_height"] else "Не определено")
                    if tile.get("data-resolution")!=expected:
                        tile_mismatches.append({"url":photo["url"],"reason":"resolution mismatch"})
                    if tile.get("data-size")!=photo_metadata.format_file_size(photo.get("verified_bytes")):
                        tile_mismatches.append({"url":photo["url"],"reason":"size mismatch"})
                description_text=" ".join(x.get_text(" ",strip=True) for x in soup.select(".description"))
                suspicious=re.findall(r"(?i)(product support|customer service|download manual|find a service center|support home)",description_text)
                conflict_rows=[r for r in flattened if (r.get("resolved") or {}).get("conflict")]
                columns=[k for k,_ in presented["columns"]]
                dealer_columns=[k for k in columns if k in {"dns","sulpak"}]
                records.append({
                    "article":item["article"],"status_code":response.status_code,
                    "manual_status_expected":manual_status.russian_status(DB,pid,product["search_code"]),
                    "manual_status_rendered":manual_status.russian_status(DB,pid,product["search_code"]) in response.text,
                    "attribute_count":len(flattened),
                    "has_core_heading":"Основные характеристики" in response.text if presented["groups"]["core"] else True,
                    "has_feature_heading":"Особенности модели" in response.text if presented["groups"]["feature"] else True,
                    "sections":{role:list(groups) for role,groups in presented["groups"].items()},
                    "generic_names":[r["display_name"] for r in flattened if r["display_name"]=="Дополнительная характеристика"],
                    "noncompact_statuses":[r["compact_status"] for r in flattened if len(r["compact_status"])>48 or "\n" in r["compact_status"]],
                    "conflicts":len(conflict_rows),
                    "conflicts_with_both_sides":sum(len(r["conflict_sides"])>=2 for r in conflict_rows),
                    "dealer_columns":dealer_columns,
                    "dealer_columns_without_facts":[k for k in dealer_columns if not any(f["source_key"]==k for f in facts)],
                    "photo_tiles":len(tiles),"photo_candidates":len(photos),
                    "photo_metadata_mismatches":tile_mismatches,
                    "lightbox_present":bool(soup.find("dialog",id="photo-lightbox")),
                    "lightbox_script_present":bool(soup.find("script",src="/static/photo_lightbox.js")),
                    "suspicious_description_terms":suspicious,
                    "documents":len(docs)
                })
    batch=first["batch_id"]
    book=load_workbook(BytesIO(exporter.export_batch(DB,batch)),read_only=True)
    try:
        instruction=list(book["Инструкции"].values)
        check=list(book["Проверка источников"].values)
        dealer_col=check[0].index("Sulpak") if "Sulpak" in check[0] else None
        excel={
            "sheet_names":book.sheetnames,
            "manual_status_column_last":instruction[0][-1]=="Статус проверки",
            "instruction_rows":len(instruction)-1,
            "generic_headers":[{"sheet":sheet.title,"header":cell} for sheet in book
                               if sheet.title not in ("Источники","Инструкции","Фотографии","Проверка источников")
                               for cell in next(sheet.values)
                               if cell=="Дополнительная характеристика"],
            "sulpak_check_column_present":dealer_col is not None,
            "sulpak_check_column_all_empty":dealer_col is not None and all(not row[dealer_col] for row in check[1:])
        }
    finally:book.close()
    payload={"rows":records,"excel":excel}
    OUT.write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print("UI_EXCEL_AUDIT",len(records),"issues",sum(bool(r["generic_names"] or r["noncompact_statuses"] or r["photo_metadata_mismatches"] or r["suspicious_description_terms"] or r["dealer_columns_without_facts"]) for r in records))

if __name__=="__main__":main()
