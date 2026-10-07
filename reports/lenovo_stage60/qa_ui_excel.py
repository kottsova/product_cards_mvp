from pathlib import Path
from shutil import copyfile
from io import BytesIO
import json,hashlib
from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from product_tool.web import create_app
from product_tool import jobs,storage,attribute_projection,card_presentation,lenovo_pipeline,card_evidence
root=Path('reports/lenovo_stage60');ui=root/'ui_batch';ui.mkdir(exist_ok=True)
copyfile(root/'final_pass.sqlite3',ui/'batches.sqlite3')
checks=[]
with TestClient(create_app(ui,start_worker=False)) as client:
 for p in storage.get_batch(ui/'batches.sqlite3','stage60-final')['products']:
  response=client.get(f"/products/{p['id']}");html=response.text
  soup=BeautifulSoup(html,'html.parser')
  rows=attribute_projection.final_attribute_rows(ui/'batches.sqlite3',p['id'])
  ev=card_evidence.load(ui/'batches.sqlite3',p['id'],'lenovo') or {}
  facts=jobs.get_facts(ui/'batches.sqlite3',p['id'])
  keys={f['raw_name'] for f in facts}
  photo_verified=sum(lenovo_pipeline.photo_verified(photo,ev) for photo in jobs.get_photo_candidates(ui/'batches.sqlite3',p['id']))
  item={'article':p['search_code'],'http':response.status_code,'labels':[r['display_name'] for r in rows],'candidate_section':'Кандидаты конфигурации' in html,'manual_candidate_section':'Документы-кандидаты' in html,'manual_status':ev.get('manual_status'),'verified_photos':photo_verified,'support_field_leak':bool(keys & {'Included Warranty','BIOS','Drivers','Troubleshooting'}),'raw_preserved':all(f['section'] and f['raw_name'] and f['raw_value'] for f in facts)}
  assert item['http']==200 and item['candidate_section'] and item['manual_candidate_section'],item
  assert not item['support_field_leak'] and item['raw_preserved'] and not photo_verified,item
  checks.append(item)
 response=client.get('/batches/stage60-final/export.xlsx');assert response.status_code==200
 (root/'export.xlsx').write_bytes(response.content)
 book=load_workbook(BytesIO(response.content),read_only=True,data_only=True)
 sheets={s.title:{'rows':s.max_row,'columns':s.max_column,'headers':list(next(s.values))} for s in book}
 assert book['Фотографии'].max_row==1
 assert book['Фото-кандидаты'].max_row==11
 assert book['Конфигурации-кандидаты'].max_row>1
 assert book['Готовность Lenovo'].max_row==11
 assert all(row[2]=='Не готова' for row in list(book['Готовность Lenovo'].values)[1:])
 assert all(row[4]=='Не проверена' for row in list(book['Готовность Lenovo'].values)[1:])
 result={'ui':checks,'xlsx_bytes':len(response.content),'sheets':sheets,'confirmed_photo_rows':0,'formula_errors':[],'scope':'Existing production Excel template/exporter. Static values, no formulas to recalculate. HTML response inspection, no native visual screenshot.'}
 (root/'qa_ui_excel.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
 print(json.dumps({'ui_rows':len(checks),'excel_bytes':len(response.content),'sheets':{k:v['rows'] for k,v in sheets.items()}},ensure_ascii=True),flush=True)
 book.close()
