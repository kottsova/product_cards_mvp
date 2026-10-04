"""Read-only Stage 59 UI and generated Excel audit of the final isolated batch."""
from __future__ import annotations
from io import BytesIO
from pathlib import Path
from shutil import copyfile
import json, sqlite3
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from product_tool.web import create_app

root = Path(__file__).resolve().parent
ui = root / 'ui_batch'
ui.mkdir(exist_ok=True)
copyfile(root / 'post_fix.sqlite3', ui / 'batches.sqlite3')
with sqlite3.connect(ui / 'batches.sqlite3') as db:
    products = list(db.execute('select id, search_code from products order by id'))
checks = []
with TestClient(create_app(ui, start_worker=False)) as client:
    for pid, article in products:
        response = client.get(f'/products/{pid}')
        html = response.text
        item = {'article': article, 'http': response.status_code,
                'main_group': '\u041e\u0441\u043d\u043e\u0432\u043d\u044b\u0435 \u0445\u0430\u0440\u0430\u043a\u0442\u0435\u0440\u0438\u0441\u0442\u0438\u043a\u0438' in html,
                'features_group': '\u041e\u0441\u043e\u0431\u0435\u043d\u043d\u043e\u0441\u0442\u0438 \u043c\u043e\u0434\u0435\u043b\u0438' in html,
                'lightbox': 'id="photo-lightbox"' in html,
                'generic_label': '\u0414\u043e\u043f\u043e\u043b\u043d\u0438\u0442\u0435\u043b\u044c\u043d\u0430\u044f \u0445\u0430\u0440\u0430\u043a\u0442\u0435\u0440\u0438\u0441\u0442\u0438\u043a\u0430' in html,
                'support_copy': 'How can we help you' in html or 'Troubleshooting' in html,
                'photo_resolution_label': '\u0420\u0430\u0437\u0440\u0435\u0448\u0435\u043d\u0438\u0435:' in html,
                'manual_status_label': '\u0420\u0443\u0441\u0441\u043a\u0430\u044f \u0438\u043d\u0441\u0442\u0440\u0443\u043a\u0446\u0438\u044f:' in html}
        checks.append(item)
        assert response.status_code == 200, item
        assert not item['generic_label'] and not item['support_copy'], item
        assert item['main_group'] and item['features_group'] and item['lightbox'] and item['manual_status_label'], item
    response = client.get('/batches/post_fix/export.xlsx')
    assert response.status_code == 200
    book = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
    sheets = {sheet.title: {'rows': sheet.max_row, 'cols': sheet.max_column, 'head': list(next(sheet.values))} for sheet in book}
    all_text = [cell for sheet in book for row in sheet.values for cell in row if isinstance(cell, str)]
    assert not any('\u0420\u0430\u0441\u0445\u043e\u0436\u0434\u0435\u043d\u0438\u044f \u0441 \u0434\u0438\u043b\u0435\u0440\u043e\u043c' in str(value) or 'Sulpak' in str(value) for sheet in sheets.values() for value in sheet['head'])
    assert not any('\u0414\u043e\u043f\u043e\u043b\u043d\u0438\u0442\u0435\u043b\u044c\u043d\u0430\u044f \u0445\u0430\u0440\u0430\u043a\u0442\u0435\u0440\u0438\u0441\u0442\u0438\u043a\u0430' in value for value in all_text)
    candidates = next((sheet for sheet in book if sheet.title.startswith('\u0424\u043e\u0442\u043e-\u043a\u0430\u043d\u0434\u0438\u0434\u0430\u0442\u044b')), None)
    assert candidates is not None and candidates.max_row > 1
    assert list(next(candidates.values))[:6] == ["\u0422\u043e\u0432\u0430\u0440", "\u0418\u0441\u0442\u043e\u0447\u043d\u0438\u043a", "\u0422\u0438\u043f", "URL", "\u0412\u044b\u0431\u0440\u0430\u043d\u043e \u0434\u043b\u044f \u043f\u0440\u043e\u0441\u043c\u043e\u0442\u0440\u0430", "\u0421\u0442\u0430\u0442\u0443\u0441"]
    docs = next(sheet for sheet in book if sheet.title.startswith('\u0418\u043d\u0441\u0442\u0440\u0443\u043a\u0446\u0438\u0438'))
    statuses = sorted({str(row[-1]) for row in list(docs.values)[1:] if row[-1]})
    result = {'ui': checks, 'workbook_bytes': len(response.content), 'sheets': sheets,
              'manual_statuses': statuses, 'photo_candidate_rows': candidates.max_row - 1}
    (root / 'qa_ui_excel.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({'ui_rows': len(checks), 'excel_bytes': len(response.content), 'sheets': {k: v['rows'] for k,v in sheets.items()},
                      'photo_candidate_rows': result['photo_candidate_rows'], 'manual_statuses': statuses}, ensure_ascii=True))
    book.close()
