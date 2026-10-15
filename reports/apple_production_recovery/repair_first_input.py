"""One explicit owner-authorized first-card repair from the preserved workbook row."""
import json,sqlite3,sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from product_tool import storage,jobs,discovery_trace
from product_tool.identity import ProductIdentity
from openpyxl import load_workbook
root=Path(__file__).resolve().parents[2];db=root/'data/batches.sqlite3';report=Path(__file__).parent
p=jobs.get_product(db,84)
assert p['batch_id']=='70c329c089c2432fbe042651944c6acd' and p['search_code']=='A3714'
w=load_workbook(root/'data/uploads'/f"{p['batch_id']}.xlsx",read_only=True,data_only=True)
header=next(w.active.iter_rows(min_row=1,max_row=1,values_only=True))
row=next(w.active.iter_rows(min_row=p['row_number'],max_row=p['row_number'],values_only=True));w.close()
assert header[3].strip()=='Наименование на сайте' and header[4]=='Цвет'
assert row[2]==p['search_code'] and row[3]=='iPhone 18 Pro 256GB'
before=dict(product=p,jobs=jobs.list_jobs(db,84),trace=discovery_trace.for_job(db,jobs.list_jobs(db,84)[0]['id']),events=jobs.list_events(db,jobs.list_jobs(db,84)[0]['id']))
(report/'original_failure.json').write_text(json.dumps(before,ensure_ascii=False,indent=2),encoding='utf8')
backup=root/'data/backups/apple_first_card_before_repair.sqlite3';backup.parent.mkdir(exist_ok=True)
if not backup.exists():
 with sqlite3.connect(db) as source,sqlite3.connect(backup) as target:source.backup(target)
new={**p,'name':row[3],'alternate_code':'','variant_attributes':{'color':row[4],'storage':'256GB'}}
identity=ProductIdentity.from_product(new).to_dict()
identity['input_provenance']={'name_column':4,'color_column':5,'hardware_column':3,'source':'original uploaded row; input configuration only'}
with storage._connection(db) as c:
 assert not c.execute("SELECT 1 FROM search_jobs WHERE status IN ('queued','running')").fetchone()
 c.execute('UPDATE products SET name=?,alternate_code=?,identity_json=? WHERE id=?',(new['name'],'',json.dumps(identity,ensure_ascii=False),84))
(report/'input_repair.json').write_text(json.dumps(dict(product_id=84,header=header,original_row=row,old_name=p['name'],new_name=new['name'],old_alternate=p['alternate_code'],new_alternate='',identity=identity,backup=str(backup),other_rows_modified=False),ensure_ascii=False,indent=2),encoding='utf8')
print('Repaired only product 84 from original cells; backup:',backup)
