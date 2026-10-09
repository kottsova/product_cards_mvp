from pathlib import Path
from io import BytesIO
import json,sqlite3,os,re
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from product_tool import exporter,storage,jobs,razer_pipeline,attribute_projection
from product_tool.web import create_app
R=Path(__file__).parent;U=R/'ui';U.mkdir(exist_ok=True);db=U/'batches.sqlite3'
with sqlite3.connect(R/'corrected_replay_release.sqlite3') as src,sqlite3.connect(db) as dst:src.backup(dst)
out=[]
with TestClient(create_app(U,start_worker=False)) as client:
 for p in storage.get_batch(db,'rz70')['products']:
  response=client.get(f'/products/{p["id"]}');assert response.status_code==200;html=response.text;assert 'Модель и конфигурация Razer' in html and 'Характеристики-кандидаты' in html and 'Документы по типам' in html
  (U/f'product_{p["id"]}.html').write_text(html,encoding='utf8');rows=attribute_projection.final_attribute_rows(db,p['id']);r=razer_pipeline.card_readiness(db,p['id']);assert r['manual_status'] in html
  assert not any(photo['selected'] for photo in jobs.get_photo_candidates(db,p['id']))
  assert not any(re.search(r'firmware update|how to reset|warranty procedure|Synapse setup',s['description'],re.I) for s in jobs.get_source_pages(db,p['id']))
  out.append({'id':p['row_number'],'http':response.status_code,'labels':[x['display_name'] for x in rows],'readiness':r,'description_clean':True,'observation':'verified capture replay UI, not live readiness'})
 data=exporter.export_batch(db,'rz70');(R/'export.xlsx').write_bytes(data);book=load_workbook(BytesIO(data));sheets={s.title:{'rows':s.max_row,'columns':s.max_column,'headers':[c.value for c in s[1]]} for s in book};assert book['Готовность Razer'].max_row==11
 for name in ('Варианты Razer','Кандидаты Razer','Документы Razer','Исходные факты Razer','Фото-кандидаты'):assert name in sheets
 errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors;assert all(book['Готовность Razer'].cell(i,3).value=='Не готова' for i in range(2,12));book.close()
(R/'qa_ui_excel.json').write_text(json.dumps({'ui':out,'sheets':sheets,'formula_errors':errors,'native_exporter':True},ensure_ascii=False,indent=2),encoding='utf8');print('10 UI HTTP 200 and native Excel verified',flush=True)
os.environ.setdefault('PLAYWRIGHT_BROWSERS_PATH',str(Path('.venv/playwright-browsers').resolve()))
from playwright.sync_api import sync_playwright
from urllib.parse import urlsplit
with TestClient(create_app(U,start_worker=False)) as client,sync_playwright() as pw:
 browser=pw.chromium.launch(headless=True);page=browser.new_page(viewport={'width':1440,'height':1000})
 def route(request):
  u=urlsplit(request.request.url)
  if u.hostname!='razer-audit.test':request.abort();return
  response=client.get(u.path);request.fulfill(status=response.status_code,body=response.content,content_type=response.headers.get('content-type','text/html'))
 page.route('**/*',route)
 for pid in (1,3,7,10):
  page.goto(f'http://razer-audit.test/products/{pid}',wait_until='domcontentloaded');page.locator('#razer-readiness').scroll_into_view_if_needed();page.screenshot(path=str(R/f'ui_{pid}_evidence.png'));page.get_by_role('heading',name='Основные характеристики',exact=True).first.scroll_into_view_if_needed();page.screenshot(path=str(R/f'ui_{pid}_specs.png'))
 browser.close()
